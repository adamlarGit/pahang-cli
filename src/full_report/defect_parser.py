"""CBM Defect Detail Header Parsing & D37 Naming for Full Report (Ticket 004 / T1.3b)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

import docx


@dataclass(frozen=True)
class CbmDefectSliceMetadata:
    """Metadata extracted from Table 1 of a sliced CBM defect page."""

    equipment_category: str  # "swg", "tx", "fp", "battery"
    equipment_instance: str  # e.g. "swg1", "tx1", "fp1"
    sequence: str  # e.g. "p04", "f02", "s01", "p00"
    equipment_id: str  # e.g. "CKN01309", "TX1", "FP1"
    defect_area: str  # e.g. "FUSE_COMPARTMENT", "HV_BUSHING"
    severity: str = "DEFECT"  # "DEFECT", "NORMAL"
    index: int = 1  # 1-based defect index for D37 naming
    substation: str = ""
    manufacturer: str = ""
    model: str = ""
    filename: str = ""
    slice_path: Path | None = None


def normalize_defect_area(area_str: str) -> str:
    """Normalize defect area string to uppercase snake_case."""
    raw = (area_str or "").strip()
    if not raw or raw == "-":
        return "OVERVIEW"
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", raw)
    normalized = cleaned.strip("_").upper()
    return normalized or "OVERVIEW"


def normalize_equipment_id(id_str: str) -> str:
    """Normalize equipment or panel ID string to uppercase snake_case."""
    raw = (id_str or "").strip()
    if not raw or raw == "-":
        return ""
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", raw)
    return cleaned.strip("_").upper()


def build_d37_defect_filename(metadata: CbmDefectSliceMetadata) -> str:
    """Format filename adhering strictly to D37 grammar: {eq_instance}_{seq}_{id}_{area}_{idx}.docx."""
    eq_inst = metadata.equipment_instance.strip().lower()
    seq = metadata.sequence.strip().lower()
    eq_id = metadata.equipment_id.strip().upper()
    area = metadata.defect_area.strip().upper()
    idx = f"{metadata.index:02d}"
    return f"{eq_inst}_{seq}_{eq_id}_{area}_{idx}.docx"


class CbmDefectHeaderParser:
    """Parser extracting Table 1 metadata and constructing D37 compliant filenames."""

    def parse(
        self,
        target: Any,
        sequence_hint: str = "",
        index: int = 1,
        slice_path: Path | None = None,
    ) -> CbmDefectSliceMetadata:
        """Parse Table 1 from a docx file path, Document, Table, or 2D string grid."""
        grid = self._extract_grid(target)
        if not grid:
            raise ValueError("No table rows found to parse CBM defect header.")

        fields = self._extract_key_values(grid)

        # 1. Equipment Category & Instance
        raw_eq = fields.get("equipment", "")
        category = self._infer_equipment_category(raw_eq)
        raw_tx_no = fields.get("tx no.", fields.get("tx no", fields.get("transformer no", "")))
        instance = self._infer_equipment_instance(category, raw_eq, raw_tx_no)

        # 2. Defect Area
        raw_area = fields.get("area", "")
        defect_area = normalize_defect_area(raw_area)

        # 3. Sequence Token (e.g. p04, f02, s01, p00)
        panel_no_val = fields.get(
            "panel no.",
            fields.get(
                "panel no",
                fields.get(
                    "feeder no.",
                    fields.get(
                        "feeder no",
                        fields.get("tx no.", fields.get("battery no.", "")),
                    ),
                ),
            ),
        )
        sequence = self._infer_sequence(
            category=category,
            panel_no_text=panel_no_val,
            area_text=raw_area,
            sequence_hint=sequence_hint,
        )

        # 4. Equipment ID (e.g. CKN01309, TX1, FP1)
        panel_name_val = fields.get(
            "panel name",
            fields.get("feeder name", fields.get("location", "")),
        )
        equipment_id = self._infer_equipment_id(
            category=category,
            panel_name=panel_name_val,
            panel_no=panel_no_val,
            eq_instance=instance,
        )

        # 5. Severity
        severity = self._extract_severity(fields, grid)

        metadata = CbmDefectSliceMetadata(
            equipment_category=category,
            equipment_instance=instance,
            sequence=sequence,
            equipment_id=equipment_id,
            defect_area=defect_area,
            severity=severity,
            index=index,
            substation=fields.get("substation", ""),
            manufacturer=fields.get("manufacturer", ""),
            model=fields.get("model", ""),
            slice_path=slice_path,
        )

        filename = build_d37_defect_filename(metadata)
        # return new instance with filename populated
        return CbmDefectSliceMetadata(
            equipment_category=metadata.equipment_category,
            equipment_instance=metadata.equipment_instance,
            sequence=metadata.sequence,
            equipment_id=metadata.equipment_id,
            defect_area=metadata.defect_area,
            severity=metadata.severity,
            index=metadata.index,
            substation=metadata.substation,
            manufacturer=metadata.manufacturer,
            model=metadata.model,
            filename=filename,
            slice_path=metadata.slice_path,
        )

    def parse_batch(
        self,
        targets: Any,
        sequence_hints: Any | None = None,
    ) -> list[CbmDefectSliceMetadata]:
        """Parse multiple defect tables, automatically deduplicating and indexing D37 filenames."""
        counts: dict[tuple[str, str, str, str], int] = {}
        results: list[CbmDefectSliceMetadata] = []

        target_list = list(targets)
        hints = list(sequence_hints) if sequence_hints is not None else [""] * len(target_list)

        for target, hint in zip(target_list, hints):
            # Parse with index=1 first to get base keys
            base = self.parse(target, sequence_hint=hint, index=1)
            key = (
                base.equipment_instance,
                base.sequence,
                base.equipment_id,
                base.defect_area,
            )
            idx = counts.get(key, 0) + 1
            counts[key] = idx

            if idx == 1:
                results.append(base)
            else:
                updated = self.parse(target, sequence_hint=hint, index=idx)
                results.append(updated)

        return results

    def _extract_grid(self, target: Any) -> list[list[str]]:
        """Extract a 2D text matrix from various inputs."""
        if isinstance(target, (str, Path)):
            p = Path(target)
            if not p.exists():
                raise FileNotFoundError(f"File not found: {p}")
            doc = docx.Document(p)
            if not doc.tables:
                raise ValueError(f"No tables found in docx: {p}")
            return self._table_to_grid(doc.tables[0])

        if hasattr(target, "tables"):
            if not target.tables:
                raise ValueError("No tables found in Document.")
            return self._table_to_grid(target.tables[0])

        if hasattr(target, "rows"):
            return self._table_to_grid(target)

        if isinstance(target, (list, tuple)):
            grid: list[list[str]] = []
            for row in target:
                if isinstance(row, (list, tuple)):
                    grid.append([str(c).strip() for c in row])
                elif hasattr(row, "cells"):
                    grid.append([c.text.strip() for c in row.cells])
            return grid

        raise TypeError(f"Unsupported target type for table extraction: {type(target)}")

    def _table_to_grid(self, table: Any) -> list[list[str]]:
        """Convert a docx Table into a clean 2D string grid without consecutive duplicates."""
        grid: list[list[str]] = []
        for row in table.rows:
            row_texts: list[str] = []
            for cell in row.cells:
                txt = cell.text.strip()
                # de-duplicate consecutive identical merged cells
                if not row_texts or row_texts[-1] != txt:
                    row_texts.append(txt)
            grid.append(row_texts)
        return grid

    def _extract_key_values(self, grid: list[list[str]]) -> dict[str, str]:
        """Scan table grid rows to extract canonical field values by label."""
        known_labels = {
            "substation",
            "date:",
            "equipment",
            "manufacturer",
            "time:",
            "model",
            "rating",
            "humidity:",
            "area",
            "panel no.",
            "panel no",
            "tx no.",
            "tx no",
            "transformer no",
            "feeder no.",
            "feeder no",
            "battery no.",
            "battery no",
            "panel name",
            "location",
            "feeder name",
            "severity",
            "severity:",
        }

        fields: dict[str, str] = {}

        prefix_pattern = re.compile(
            r"^(substation|equipment|manufacturer|model|rating|area|panel\s*no\.?|tx\s*no\.?|transformer\s*no\.?|feeder\s*no\.?|battery\s*no\.?|panel\s*name|location|feeder\s*name|severity)\s*[:.]?\s*(.*)$",
            re.IGNORECASE,
        )

        for row in grid:
            non_empty = [c for c in row if c]
            i = 0
            while i < len(non_empty):
                cell = non_empty[i]
                norm_label = cell.lower().strip()

                # Check for inline "Label: Value" or "Label Value" (e.g. "Panel No. 2", "Date: 28-Aug-2026")
                colon_match = re.match(r"^([A-Za-z0-9\s.]+):\s*(.+)$", cell)
                if colon_match:
                    k = colon_match.group(1).lower().strip()
                    v = colon_match.group(2).strip()
                    fields[k] = v
                    i += 1
                    continue

                pref_match = prefix_pattern.match(cell)
                if pref_match:
                    raw_k = pref_match.group(1).lower().strip().rstrip(":")
                    inline_v = pref_match.group(2).strip()
                    clean_k = "panel no." if "panel" in raw_k and "no" in raw_k else (
                        "tx no." if "tx" in raw_k or "transformer" in raw_k else (
                            "feeder no." if "feeder" in raw_k and "no" in raw_k else (
                                "battery no." if "battery" in raw_k and "no" in raw_k else raw_k
                            )
                        )
                    )
                    if inline_v:
                        fields[clean_k] = inline_v
                        i += 1
                        continue

                    val = ""
                    if i + 1 < len(non_empty):
                        next_cell = non_empty[i + 1]
                        if not prefix_pattern.match(next_cell) and not re.match(r"^([A-Za-z0-9\s.]+):", next_cell):
                            val = next_cell
                            i += 1
                    fields[clean_k] = val
                    i += 1
                    continue

                if norm_label in known_labels:
                    clean_k = norm_label.rstrip(":")
                    val = ""
                    if i + 1 < len(non_empty):
                        next_cell = non_empty[i + 1]
                        next_norm = next_cell.lower().strip()
                        if next_norm not in known_labels and not re.match(r"^([A-Za-z0-9\s.]+):", next_cell):
                            val = next_cell
                            i += 1
                    fields[clean_k] = val
                i += 1

        return fields

    def _infer_equipment_category(self, raw_eq: str) -> str:
        """Infer equipment category ('swg', 'tx', 'fp', 'battery') from equipment text."""
        eq = raw_eq.upper()
        if any(k in eq for k in ("RMU", "VCB", "SWG", "SWITCHGEAR", "AIS", "GIS")):
            return "swg"
        if any(k in eq for k in ("TRANSFORMER", "TX")):
            return "tx"
        if any(k in eq for k in ("FEEDER PILLAR", "FP", "LVDB")):
            return "fp"
        if "BATTERY" in eq:
            return "battery"
        return "swg"

    def _infer_equipment_instance(self, category: str, raw_eq: str, raw_tx_no: str) -> str:
        """Determine equipment instance string (e.g. swg1, tx1, fp1)."""
        if category == "tx" and raw_tx_no:
            match = re.search(r"(\d+)", raw_tx_no)
            if match:
                return f"tx{match.group(1)}"

        # Check for trailing or embedded digit in equipment string (e.g. "SWG 2", "FP 2")
        match = re.search(r"(?:SWG|TX|FP|LVDB|BATTERY)\s*(\d+)", raw_eq.upper())
        if match:
            return f"{category}{match.group(1)}"

        return f"{category}1"

    def _infer_sequence(
        self,
        category: str,
        panel_no_text: str,
        area_text: str,
        sequence_hint: str = "",
    ) -> str:
        """Infer physical sequence token (p04, f02, s01, p00)."""
        if sequence_hint:
            return sequence_hint.strip().lower()

        area_norm = normalize_defect_area(area_text)

        if category == "swg":
            if area_norm == "OVERVIEW" or panel_no_text in ("-", ""):
                return "p00"
            match = re.search(r"(\d+)", panel_no_text)
            if match:
                return f"p{int(match.group(1)):02d}"
            p_match = re.search(r"p(\d+)", panel_no_text, re.IGNORECASE)
            if p_match:
                return f"p{int(p_match.group(1)):02d}"
            return "p01"

        if category == "fp":
            if area_norm == "OVERVIEW" or panel_no_text in ("-", ""):
                return "f00"
            p_upper = panel_no_text.upper()
            if "IN" in p_upper:
                m = re.search(r"(\d+)", p_upper)
                return f"in{int(m.group(1)):02d}" if m else "in01"
            if "OT" in p_upper:
                m = re.search(r"(\d+)", p_upper)
                return f"ot{int(m.group(1)):02d}" if m else "ot01"
            m = re.search(r"(\d+)", panel_no_text)
            if m:
                return f"f{int(m.group(1)):02d}"
            return "f01"

        if category == "tx":
            tx_area_map = {
                "OVERVIEW": "s00",
                "OVERVIEW_TOP": "s01",
                "HV_BUSHING": "s02",
                "HV_CABLE": "s03",
                "HV_CABLE_SPLIT": "s04",
                "LV_BUSHING": "s05",
                "LV_CABLE": "s06",
            }
            if area_norm in tx_area_map:
                return tx_area_map[area_norm]
            m = re.search(r"(\d+)", panel_no_text)
            if m:
                return f"s{int(m.group(1)):02d}"
            return "s01"

        if category == "battery":
            return "b00"

        return "s01"

    def _infer_equipment_id(
        self,
        category: str,
        panel_name: str,
        panel_no: str,
        eq_instance: str,
    ) -> str:
        """Infer canonical alphanumeric equipment identifier (e.g. CKN01309)."""
        # Look for standard TNB plant alphanumeric codes like CKN01309, CRA00223
        for val in (panel_name, panel_no):
            if val:
                match = re.search(r"\b([A-Z]{3,4}\d{4,6})\b", val.upper())
                if match:
                    return match.group(1)

        # For Transformer or Battery Bank, use the equipment instance (e.g. TX1, BATTERY1)
        # rather than room location (e.g. "TX ROOM") or raw digit numbers ("1")
        if category in ("tx", "battery"):
            return eq_instance.upper()

        # If panel_name exists and is meaningful (and not a generic room or bare digits)
        clean_name = normalize_equipment_id(panel_name)
        if clean_name and clean_name != "-" and not clean_name.isdigit() and "ROOM" not in clean_name:
            return clean_name

        # If panel_no exists and is meaningful (and not just bare digits)
        clean_no = normalize_equipment_id(panel_no)
        if clean_no and clean_no != "-" and not clean_no.isdigit():
            return clean_no

        return eq_instance.upper()

    def _extract_severity(self, fields: dict[str, str], grid: list[list[str]]) -> str:
        """Extract defect severity classification from fields or table grid."""
        raw_sev = fields.get("severity", "").upper()
        if "DEFECT" in raw_sev or "CRITICAL" in raw_sev:
            return "DEFECT"
        if "NORMAL" in raw_sev:
            return "NORMAL"

        # Check for Analysis rows in grid
        for row in grid:
            for cell in row:
                txt = cell.lower()
                if "please refer to the following page for details defect" in txt:
                    return "DEFECT"
                if "no anomaly" in txt:
                    return "NORMAL"

        # Default for CBM defect detail pages
        return "DEFECT"
