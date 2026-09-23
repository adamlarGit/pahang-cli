"""Executive Summary Equipment Census Builder with Group Vertical Merge (Ticket #28 / T3.2)."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Sequence

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docxtpl import DocxTemplate

from src.core.contract import normalize_swg_compartment
from src.core.normalizers import normalize_tx_id
from src.core.topology import (
    BayRole,
    SwitchgearArchetype,
    SwitchgearTopologyEngine,
    VoltageClass,
    classify_bay_role,
)
from src.full_report.models import (
    FullReportScanPackage,
    SwitchgearScanSpec,
    TRANSFORMER_STANDARD_COMPONENTS,
    has_hv_cable_split,
)
from src.quick_report.cbm_render import _build_jinja_env
from src.quick_report.cbm_summary import (
    format_db_reading,
    format_summary_equipment,
    format_temperature_reading,
    prepare_tech_summary_rows,
)
from src.quick_report.defects import CbmDefectRecord
from src.quick_report.utils import clear_cell_text, set_cell_shading
from src.testsheet.models import (
    BatteryBankSpec,
    LVDBSpec,
    SubstationEquipmentPackage,
    SwitchgearPanelSpec,
    SwitchgearSpec,
    TransformerSpec,
)


# Standard 7 components for transformer census table per ADR 0004
TRANSFORMER_CENSUS_COMPONENTS: tuple[str, ...] = (
    "OVERVIEW",
    "OVERVIEW TOP",
    "HV BUSHING",
    "HV CABLES",
    "CABLE SPLIT",
    "LV BUSHING",
    "LV CABLES",
)


# ==============================================================================
# Domain Dataclasses
# ==============================================================================

@dataclass
class CensusRowItem:
    """Row representation for Executive Summary Census (Table 2)."""

    no: str = ""
    equipment: str = ""
    defect_area: str = ""
    ir_abs: str = "-"
    us_dB: str = "-"
    tev_dB: str = "-"
    severity: str = "NORMAL"
    group_no: int = 1
    is_overview: bool = False
    is_defect: bool = False
    eq_group_key: str = ""

    def to_dict(self) -> dict[str, Any]:
        """Serialize row to dictionary matching template placeholders."""
        return {
            "no": self.no,
            "equipment": self.equipment,
            "defect_area": self.defect_area,
            "ir_abs": self.ir_abs,
            "us_dB": self.us_dB,
            "tev_dB": self.tev_dB,
            "severity": self.severity,
        }

    def __getitem__(self, key: str) -> Any:
        return getattr(self, key)


@dataclass
class ExecutiveSummaryCensusContext:
    """Jinja render context for Executive Summary Census template."""

    census_items: list[CensusRowItem] = field(default_factory=list)
    substation_number: int | str = ""
    station_name: str = ""

    def to_dict(self) -> dict[str, Any]:
        """Serialize context to dictionary for DocxTemplate rendering."""
        return {
            "census_items": [item.to_dict() for item in self.census_items],
            "substation_number": self.substation_number,
            "station_name": self.station_name,
        }

    def __getitem__(self, key: str) -> Any:
        if key == "census_items":
            return self.census_items
        return getattr(self, key)


@dataclass
class ExecutiveSummaryCensusResult:
    """Outcome of Executive Summary Census generation."""

    docx_path: Path | None = None
    items: list[CensusRowItem] = field(default_factory=list)
    group_count: int = 0
    total_rows: int = 0
    has_defects: bool = False


# ==============================================================================
# Defect Cross-Referencing Helpers
# ==============================================================================

def _format_defect_readings(
    matched_defects: Sequence[CbmDefectRecord],
) -> tuple[str, str, str]:
    """Format combined (ir_abs, us_dB, tev_dB) readings from matched defect records."""
    ir_str = "-"
    us_str = "-"
    tev_str = "-"

    for d in matched_defects:
        raw_ir = d.ir_reading or (d.raw_measurement if d.technology == "IR" else "")
        if raw_ir and raw_ir != "-":
            formatted = format_temperature_reading(raw_ir)
            if formatted != "-":
                ir_str = formatted

        raw_us = d.us_reading or (d.raw_measurement if d.technology == "US" else "")
        if raw_us and raw_us != "-":
            formatted = format_db_reading(raw_us)
            if formatted != "-":
                us_str = formatted

        raw_tev = d.tev_reading or (d.raw_measurement if d.technology == "TEV" else "")
        if raw_tev and raw_tev != "-":
            formatted = format_db_reading(raw_tev)
            if formatted != "-":
                tev_str = formatted

    return ir_str, us_str, tev_str


def _match_defects_for_swg_panel(
    swg: SwitchgearSpec,
    panel: SwitchgearPanelSpec,
    compartment: str,
    defects: Sequence[CbmDefectRecord],
) -> list[CbmDefectRecord]:
    """Find CbmDefectRecords matching a specific switchgear panel and compartment."""
    matched: list[CbmDefectRecord] = []
    comp_upper = compartment.upper()

    for d in defects:
        eq_upper = d.equipment.upper()
        eq_id_upper = d.equipment_id.upper()
        area_upper = f"{d.defect_area} {d.additional_remarks}".upper()

        # Check if defect targets switchgear
        is_swg_related = any(
            k in eq_upper for k in ("RMU", "VCB", "SWG", "SWITCHGEAR", "MRMU", "GIS", "EARTHING")
        ) or any(
            k in eq_id_upper for k in ("PANEL", "CKN", "CRB", "F1", "F2", "F3", "F4", "BAY")
        )

        if not is_swg_related:
            continue

        # Check if defect targets this specific panel
        feeder_no = (panel.panel_feeder_no or "").strip().upper()
        panel_name = (panel.name or "").strip().upper()
        panel_num = panel.panel_no

        panel_match = False
        if eq_id_upper in ("ALL", "ALL PANELS", "ALL PANEL", "SWITCHGEAR"):
            panel_match = True
        elif feeder_no and (feeder_no in eq_id_upper or feeder_no in eq_upper):
            panel_match = True
        elif panel_name and len(panel_name) >= 3 and (panel_name in eq_id_upper or panel_name in eq_upper):
            panel_match = True
        elif re.search(rf"\b(?:PANEL|BAY|P|F)\s*0?{panel_num}\b", eq_id_upper) or re.search(rf"\bPANEL\s*0?{panel_num}\b", eq_upper):
            panel_match = True
        elif len(swg.panels) == 1:
            panel_match = True

        if not panel_match:
            continue

        # Check if defect targets this specific compartment
        norm_comp = normalize_swg_compartment(compartment)
        norm_area = normalize_swg_compartment(d.defect_area)
        norm_remarks = normalize_swg_compartment(area_upper)

        comp_match = False
        if norm_area in (
            "BREAKER COMPARTMENT",
            "CABLE COMPARTMENT",
            "BUSBAR COMPARTMENT",
            "PT COMPARTMENT",
            "SECONDARY COMPARTMENT",
            "FRONT COMPARTMENT",
            "REAR COMPARTMENT",
            "CABLE ENTRY",
            "FUSE COMPARTMENT",
        ):
            comp_match = (norm_comp == norm_area)
        elif norm_remarks in (
            "BREAKER COMPARTMENT",
            "CABLE COMPARTMENT",
            "BUSBAR COMPARTMENT",
            "PT COMPARTMENT",
            "SECONDARY COMPARTMENT",
            "FRONT COMPARTMENT",
            "REAR COMPARTMENT",
            "CABLE ENTRY",
            "FUSE COMPARTMENT",
        ):
            comp_match = (norm_comp == norm_remarks)
        elif comp_upper == "FUSE COMPARTMENT":
            comp_match = "FUSE" in area_upper
        elif comp_upper == "CABLE ENTRY":
            comp_match = "ENTRY" in area_upper
        elif comp_upper == "CABLE COMPARTMENT":
            comp_match = ("CABLE" in area_upper or "BOX" in area_upper or "TERMINATION" in area_upper) and not ("ENTRY" in area_upper or "FUSE" in area_upper)
            if not comp_match and not any(k in area_upper for k in ("FUSE", "ENTRY", "OVERVIEW", "BREAKER", "BUSBAR", "PT", "SECONDARY", "FRONT", "BACK", "REAR")):
                comp_match = True
        elif comp_upper in area_upper:
            comp_match = True
        elif not any(k in area_upper for k in ("FUSE", "ENTRY", "OVERVIEW", "BREAKER", "BUSBAR", "PT", "SECONDARY", "FRONT", "BACK", "REAR", "CABLE")):
            comp_match = True

        if comp_match:
            matched.append(d)

    return matched


def _match_defects_for_tx_component(
    tx: TransformerSpec,
    component: str,
    defects: Sequence[CbmDefectRecord],
    total_transformers: int = 1,
) -> list[CbmDefectRecord]:
    """Find CbmDefectRecords matching a specific transformer component."""
    matched: list[CbmDefectRecord] = []
    comp_upper = component.upper()

    for d in defects:
        eq_upper = d.equipment.upper()
        eq_id_upper = d.equipment_id.upper()
        area_upper = f"{d.defect_area} {d.additional_remarks} {d.equipment_id}".upper()

        is_tx_related = any(
            k in eq_upper for k in ("TX", "TRANSFORMER", "CABLE TX", "ALATUBAH")
        ) or any(
            k in eq_id_upper for k in ("TX", "TRANSFORMER")
        )

        if not is_tx_related:
            continue

        if total_transformers > 1:
            norm_id = normalize_tx_id(tx.tx_id)
            target = f"{eq_upper} {eq_id_upper}".replace(" ", "")
            m_target = re.search(r"(?:TX|TRANSFORMER)\s*[-_]?\s*(\d+)", f"{eq_upper} {eq_id_upper}")
            if m_target:
                target_tx_id = f"TX{m_target.group(1)}"
                if norm_id != target_tx_id:
                    continue
            elif norm_id not in target and (tx.tx_id or "").upper().replace(" ", "") not in target:
                continue

        comp_match = False
        if comp_upper in ("CABLE SPLIT", "HV CABLE SPLIT"):
            comp_match = "SPLIT" in area_upper
        elif comp_upper == "HV BUSHING":
            comp_match = ("HV" in area_upper and "BUSHING" in area_upper) or ("BUSHING" in area_upper and d.hv_lv == "HV") or ("BUSHING" in area_upper and "LV" not in area_upper)
        elif comp_upper in ("HV CABLE", "HV CABLES"):
            comp_match = (("HV" in area_upper and "CABLE" in area_upper) or ("CABLE" in area_upper and d.hv_lv == "HV")) and "SPLIT" not in area_upper
        elif comp_upper == "LV BUSHING":
            comp_match = ("LV" in area_upper and "BUSHING" in area_upper) or ("BUSHING" in area_upper and d.hv_lv == "LV")
        elif comp_upper in ("LV CABLE", "LV CABLES"):
            comp_match = ("LV" in area_upper and "CABLE" in area_upper) or ("CABLE" in area_upper and d.hv_lv == "LV")
        elif comp_upper == "OVERVIEW TOP":
            comp_match = "TOP" in area_upper
        elif comp_upper == "OVERVIEW":
            comp_match = "BODY" in area_upper or "OVERVIEW" in area_upper

        if comp_match:
            matched.append(d)

    return matched


def _match_defects_for_lvdb(
    lvdb: LVDBSpec,
    defects: Sequence[CbmDefectRecord],
    total_lvdbs: int = 1,
) -> list[CbmDefectRecord]:
    """Find CbmDefectRecords matching a specific Feeder Pillar / LVDB."""
    matched: list[CbmDefectRecord] = []
    lvdb_name_upper = (lvdb.name or "").strip().upper()

    for d in defects:
        eq_upper = d.equipment.upper()
        eq_id_upper = d.equipment_id.upper()
        combined = f"{eq_upper} {eq_id_upper}"

        is_lvdb = any(k in combined for k in ("FP", "LVDB", "PILLAR", "FEEDER PILLAR", "WAY")) or (
            lvdb_name_upper and lvdb_name_upper in combined
        )

        if not is_lvdb:
            continue

        if total_lvdbs > 1:
            # Check if this defect targets another specific FP by TX source (e.g. FP TX1 vs FP TX2)
            if lvdb.source and any(k in lvdb.source.upper() for k in ("TX", "TRANSFORMER")):
                self_tx = normalize_tx_id(lvdb.source)
                m_tx = re.search(r"\b(?:TX|TRANSFORMER)\s*[-_]?\s*(\d+)\b", combined)
                if m_tx:
                    target_tx = f"TX{m_tx.group(1)}"
                    if self_tx != target_tx:
                        continue

            if lvdb_name_upper:
                # Check if this defect targets another specific FP (e.g. FP 1 vs FP 2)
                m_self = re.search(r"\b(?:FP|LVDB)\s*(\d+)\b", lvdb_name_upper)
                m_target = re.search(r"\b(?:FP|LVDB)\s*(\d+)\b", combined)
                if m_self and m_target and m_self.group(1) != m_target.group(1):
                    continue

        matched.append(d)

    return matched


def _match_defects_for_battery(
    bb: BatteryBankSpec,
    defects: Sequence[CbmDefectRecord],
) -> list[CbmDefectRecord]:
    """Find CbmDefectRecords matching a DC Battery Bank."""
    matched: list[CbmDefectRecord] = []
    bb_name = (bb.name or "").strip().upper()
    for d in defects:
        eq_upper = d.equipment.upper()
        eq_id_upper = d.equipment_id.upper()
        area_upper = f"{d.defect_area} {d.additional_remarks}".upper()
        combined = f"{eq_upper} {eq_id_upper} {area_upper}"
        if "BATTERY" in combined or (bb_name and bb_name in combined):
            matched.append(d)
    return matched


def _resolve_lvdb_defect_area_name(record: CbmDefectRecord) -> str:
    """Format standardized defect area name for Feeder Pillar defect row."""
    raw_area = record.defect_area.strip()
    raw_id = record.equipment_id.strip()
    remarks = record.additional_remarks.strip()

    combined = f"{raw_id} {raw_area} {remarks}"
    match = re.search(r"\b(?:WAY|F|OT|IN)\s*(\d+)\b", combined, re.I)
    if match:
        way_num = match.group(1)
        sub = raw_area or remarks or "HOTSPOT"
        sub_upper = sub.upper()
        if sub_upper.startswith(f"WAY {way_num}") or sub_upper.startswith(f"F{way_num}"):
            return sub_upper
        return f"WAY {way_num} - {sub_upper}"

    return raw_area.upper() if raw_area else "DEFECT"


# ==============================================================================
# Post-Render OpenXML DOM Manipulation Helpers (D21 & D24)
# ==============================================================================

def apply_column_vertical_merge(
    table: Any,
    spans_or_items: Sequence[tuple[int, int]] | Sequence[CensusRowItem],
    group_numbers: Sequence[int] | None = None,
) -> None:
    """Apply OpenXML <w:vMerge> on Column 0 (NO.) with group numbering per D21.

    Restart cell receives <w:vMerge w:val="restart"/> and integer group numbering 'X.'.
    Continuing cells receive <w:vMerge/> and text is cleared.
    """
    if not spans_or_items:
        return

    groups: list[tuple[int, int, int]] = []  # (start_row, end_row, group_num)

    if isinstance(spans_or_items[0], CensusRowItem):
        items = list(spans_or_items)
        curr_group = items[0].group_no
        start_idx = 0
        for i, item in enumerate(items):
            if item.group_no != curr_group:
                groups.append((start_idx + 1, i, curr_group))
                curr_group = item.group_no
                start_idx = i
        groups.append((start_idx + 1, len(items), curr_group))
    else:
        for idx, span in enumerate(spans_or_items):
            g_num = group_numbers[idx] if group_numbers and idx < len(group_numbers) else (idx + 1)
            groups.append((span[0], span[1], g_num))

    for start_row, end_row, g_num in groups:
        if start_row >= len(table.rows):
            continue

        # 1. Restart cell
        start_tc = table.rows[start_row]._tr.tc_lst[0]
        start_cell = _Cell(start_tc, table)
        if start_cell.paragraphs:
            p = start_cell.paragraphs[0]
            if p.runs:
                p.runs[0].text = f"{g_num}."
                for r_extra in p.runs[1:]:
                    r_extra.text = ""
            else:
                p.text = f"{g_num}."
        else:
            p = start_cell.add_paragraph(f"{g_num}.")

        # Preserve / enforce center alignment and line spacing
        pPr = p._p.get_or_add_pPr()
        jc = pPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            pPr.append(jc)
        jc.set(qn("w:val"), "center")

        spacing = pPr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            pPr.append(spacing)
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")

        # Strip any indentation (w:ind) or list numbering (w:numPr) to prevent horizontal offset
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            pPr.remove(ind)
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)

        tcPr_start = start_tc.get_or_add_tcPr()
        v_align = tcPr_start.find(qn("w:vAlign"))
        if v_align is None:
            v_align = OxmlElement("w:vAlign")
            tcPr_start.append(v_align)
        v_align.set(qn("w:val"), "center")

        existing_start = tcPr_start.find(qn("w:vMerge"))
        if existing_start is not None:
            tcPr_start.remove(existing_start)
        v_merge_restart = OxmlElement("w:vMerge")
        v_merge_restart.set(qn("w:val"), "restart")
        tcPr_start.append(v_merge_restart)

        # 2. Continuing cells
        max_r = min(end_row, len(table.rows) - 1)
        for r in range(start_row + 1, max_r + 1):
            cont_tc = table.rows[r]._tr.tc_lst[0]
            cont_cell = _Cell(cont_tc, table)
            clear_cell_text(cont_cell)
            tcPr_cont = cont_tc.get_or_add_tcPr()
            v_align_c = tcPr_cont.find(qn("w:vAlign"))
            if v_align_c is None:
                v_align_c = OxmlElement("w:vAlign")
                tcPr_cont.append(v_align_c)
            v_align_c.set(qn("w:val"), "center")
            existing_cont = tcPr_cont.find(qn("w:vMerge"))
            if existing_cont is not None:
                tcPr_cont.remove(existing_cont)
            v_merge_cont = OxmlElement("w:vMerge")
            tcPr_cont.append(v_merge_cont)


def apply_equipment_vertical_merge(
    table: Any,
    items: Sequence[CensusRowItem],
) -> None:
    """Apply dynamic OpenXML <w:vMerge> on Column 1 (EQUIPMENT).

    Dynamic grouping rules:
    - Switchgear Overview: merged across overview rows of that switchgear.
    - Switchgear Panels: merged across compartment rows belonging to each panel.
    - Transformers: merged across all component rows belonging to that transformer group.
    - Single-row items (e.g. INDKOM 1-comp panel, FP overview/defects) receive no merge.

    First row of a multi-row group receives <w:vMerge w:val="restart"/> (keeping existing rendered text/formatting).
    Subsequent rows receive <w:vMerge/> and text is cleared.
    """
    if not items or len(table.rows) <= 1:
        return

    groups: list[tuple[int, int]] = []
    start_idx = 0
    curr_key = getattr(items[0], "eq_group_key", None) or f"{items[0].group_no}_{items[0].equipment}"

    for i, item in enumerate(items):
        key = getattr(item, "eq_group_key", None) or f"{item.group_no}_{item.equipment}"
        if key != curr_key:
            groups.append((start_idx, i - 1))
            curr_key = key
            start_idx = i
    groups.append((start_idx, len(items) - 1))

    for start_i, end_i in groups:
        span = end_i - start_i + 1
        if span <= 1:
            continue

        start_row = 1 + start_i
        end_row = 1 + end_i

        if start_row >= len(table.rows):
            continue

        # 1. Restart cell
        start_tc = table.rows[start_row]._tr.tc_lst[1]
        tcPr_start = start_tc.get_or_add_tcPr()
        v_align = tcPr_start.find(qn("w:vAlign"))
        if v_align is None:
            v_align = OxmlElement("w:vAlign")
            tcPr_start.append(v_align)
        v_align.set(qn("w:val"), "center")

        existing_start = tcPr_start.find(qn("w:vMerge"))
        if existing_start is not None:
            tcPr_start.remove(existing_start)
        v_merge_restart = OxmlElement("w:vMerge")
        v_merge_restart.set(qn("w:val"), "restart")
        tcPr_start.append(v_merge_restart)

        # 2. Continuing cells
        max_r = min(end_row, len(table.rows) - 1)
        for r in range(start_row + 1, max_r + 1):
            cont_tc = table.rows[r]._tr.tc_lst[1]
            cont_cell = _Cell(cont_tc, table)
            clear_cell_text(cont_cell)
            tcPr_cont = cont_tc.get_or_add_tcPr()
            v_align_c = tcPr_cont.find(qn("w:vAlign"))
            if v_align_c is None:
                v_align_c = OxmlElement("w:vAlign")
                tcPr_cont.append(v_align_c)
            v_align_c.set(qn("w:val"), "center")
            existing_cont = tcPr_cont.find(qn("w:vMerge"))
            if existing_cont is not None:
                tcPr_cont.remove(existing_cont)
            v_merge_cont = OxmlElement("w:vMerge")
            tcPr_cont.append(v_merge_cont)


def apply_severity_shading(
    table: Any,
    items: Sequence[CensusRowItem],
) -> None:
    """Apply dynamic Green (00B050) / Red (EE0000) Severity cell shading per D24.

    - Overview rows: text '-' and unshaded background.
    - Healthy component rows: shaded Green '00B050', text cleared.
    - Defective component rows: shaded Red 'EE0000', text cleared.
    """
    for i, item in enumerate(items):
        row_idx = 1 + i
        if row_idx >= len(table.rows):
            break
        cell_tc = table.rows[row_idx]._tr.tc_lst[6]
        cell = _Cell(cell_tc, table)
        tcPr = cell_tc.get_or_add_tcPr()
        existing_shd = tcPr.find(qn("w:shd"))

        if item.is_overview or item.severity == "-":
            # Overview row: unshaded background, text '-'
            if existing_shd is not None:
                tcPr.remove(existing_shd)
            if not cell.text.strip():
                cell.text = "-"
        elif item.severity == "DEFECT" or item.is_defect:
            # Defect row: Red EE0000, clear text
            clear_cell_text(cell)
            set_cell_shading(cell, "EE0000")
        else:
            # Normal row: Green 00B050, clear text
            clear_cell_text(cell)
            set_cell_shading(cell, "00B050")


def apply_post_render_dom(
    doc: Any,
    items: Sequence[CensusRowItem],
    table_index: int = 0,
) -> Any:
    """Execute complete post-render DOM pass (vertical merge + severity shading) on census table."""
    document = doc.docx if hasattr(doc, "docx") else doc
    if not document.tables or table_index >= len(document.tables):
        return doc
    table = document.tables[table_index]
    apply_column_vertical_merge(table, items)
    apply_equipment_vertical_merge(table, items)
    apply_severity_shading(table, items)
    return doc


# ==============================================================================
# Deep Module: ExecutiveSummaryCensusBuilder
# ==============================================================================

class ExecutiveSummaryCensusBuilder:
    """Traverses SubstationEquipmentPackage and builds Executive Summary Equipment Census table."""

    def __init__(self, template_path: str | Path | None = None) -> None:
        self.template_path = (
            Path(template_path)
            if template_path
            else Path("templates/FULL REPORT/executive_summary_census.docx")
        )

    def build_census_rows(
        self,
        package: SubstationEquipmentPackage | FullReportScanPackage,
        defects: Sequence[CbmDefectRecord] = (),
    ) -> list[CensusRowItem]:
        """Assemble exhaustive census row items across all physical equipment groups."""
        swgs = package.switchgears
        txs = package.transformers
        lvdbs = package.lvdb_specs
        bbs = package.battery_banks

        rows: list[CensusRowItem] = []
        current_group = 1

        # ----------------------------------------------------------------------
        # 1. Switchgear lineup(s)
        # ----------------------------------------------------------------------
        for swg_idx, swg in enumerate(swgs, 1):
            group_num = current_group
            current_group += 1

            # Phase 1: Board-level resolution via SwitchgearTopologyEngine
            if isinstance(swg, SwitchgearScanSpec):
                archetype = swg.archetype
                overview_compartments = swg.overview_compartments
            else:
                board = SwitchgearTopologyEngine.classify_board(
                    switchgear_type=getattr(swg, "switchgear_type", ""),
                    manufacturer=getattr(swg, "manufacturer", ""),
                    model=getattr(swg, "model", ""),
                    rating=getattr(swg, "rating", ""),
                    swg=swg,
                )
                archetype = getattr(swg, "archetype", None) or board.archetype
                overview_compartments = (
                    getattr(swg, "overview_compartments", None)
                    or board.overview_compartments
                )
                # Backward compatibility for legacy tests specifying INDKOM with blank model
                mfg_upper = (getattr(swg, "manufacturer", "") or "").strip().upper()
                model_str = (getattr(swg, "model", "") or "").strip()
                if mfg_upper == "INDKOM" and not model_str and not getattr(swg, "archetype", None):
                    overview_compartments = ("OVERVIEW",)

            mfg = (swg.manufacturer or "").strip()
            swg_type = (swg.switchgear_type or "RMU SF6").strip()
            if mfg and mfg.upper() not in swg_type.upper():
                swg_overview_title = f"{swg_type}, {mfg}"
            else:
                swg_overview_title = swg_type

            ov_group_key = f"swg_{swg_idx}_overview"

            # Switchgear Overview Row(s)
            for comp in overview_compartments:
                rows.append(
                    CensusRowItem(
                        no=f"{group_num}.",
                        equipment=swg_overview_title,
                        defect_area=comp,
                        ir_abs="-",
                        us_dB="-",
                        tev_dB="-",
                        severity="-",
                        group_no=group_num,
                        is_overview=True,
                        is_defect=False,
                        eq_group_key=ov_group_key,
                    )
                )

            # Switchgear Panels
            for panel in swg.panels:
                p_feeder = (panel.panel_feeder_no or "").strip()
                p_name = (panel.name or "").strip()

                clean_feeder = re.sub(r"^PANEL\s*", "", p_feeder, flags=re.I).strip()
                clean_name = p_name

                # Strip redundant feeder prefix from clean_name if present
                if clean_feeder:
                    clean_name = re.sub(
                        rf"^(?:PANEL\s+)?{re.escape(clean_feeder)}\s*[-–:]?\s*",
                        "",
                        clean_name,
                        flags=re.I,
                    ).strip()
                    clean_name = re.sub(r"^PANEL\s*", "", clean_name, flags=re.I).strip()
                elif not clean_feeder and clean_name:
                    m_fn = re.match(r"^(?:PANEL\s+)?([A-Za-z0-9]+)\s*[-–:]\s*(.+)$", clean_name, flags=re.I)
                    if m_fn:
                        clean_feeder = m_fn.group(1).strip()
                        clean_name = m_fn.group(2).strip()

                if clean_feeder and clean_name:
                    panel_eq = f"PANEL {clean_feeder}\n{clean_name}"
                elif clean_feeder:
                    panel_eq = f"PANEL {clean_feeder}"
                elif clean_name:
                    if clean_name.upper().startswith("PANEL"):
                        panel_eq = clean_name
                    else:
                        panel_eq = f"PANEL {panel.panel_no}\n{clean_name}"
                else:
                    panel_eq = f"PANEL {panel.panel_no}"

                panel_group_key = f"swg_{swg_idx}_p{panel.panel_no}"

                # Phase 2: Dynamic compartment resolution via SwitchgearTopologyEngine
                if hasattr(panel, "compartments") and panel.compartments:
                    compartments = panel.compartments
                else:
                    mfg_upper = (getattr(swg, "manufacturer", "") or "").strip().upper()
                    model_str = (getattr(swg, "model", "") or "").strip()
                    bay_role = classify_bay_role(
                        name=getattr(panel, "name", ""),
                        panel_feeder_no=getattr(panel, "panel_feeder_no", ""),
                    )
                    if mfg_upper == "INDKOM" and not model_str and not getattr(swg, "archetype", None) and bay_role == BayRole.TRANSFORMER:
                        compartments = ("FUSE COMPARTMENT",)
                    else:
                        compartments = SwitchgearTopologyEngine.resolve_panel_compartments(
                            archetype=archetype,
                            panel=panel,
                        )

                for comp in compartments:
                    matched = _match_defects_for_swg_panel(swg, panel, comp, defects)
                    if matched:
                        ir_s, us_s, tev_s = _format_defect_readings(matched)
                        rows.append(
                            CensusRowItem(
                                no=f"{group_num}.",
                                equipment=panel_eq,
                                defect_area=comp,
                                ir_abs=ir_s,
                                us_dB=us_s,
                                tev_dB=tev_s,
                                severity="DEFECT",
                                group_no=group_num,
                                is_overview=False,
                                is_defect=True,
                                eq_group_key=panel_group_key,
                            )
                        )
                    else:
                        rows.append(
                            CensusRowItem(
                                no=f"{group_num}.",
                                equipment=panel_eq,
                                defect_area=comp,
                                ir_abs="-",
                                us_dB="-",
                                tev_dB="-",
                                severity="NORMAL",
                                group_no=group_num,
                                is_overview=False,
                                is_defect=False,
                                eq_group_key=panel_group_key,
                            )
                        )

        # ----------------------------------------------------------------------
        # 2. Transformer(s) - Dynamic vertical merge across rows per ADR 0004 & D27
        # ----------------------------------------------------------------------
        for idx, tx in enumerate(txs):
            group_num = current_group
            current_group += 1

            tx_label = normalize_tx_id(tx.tx_id, default_idx=idx + 1)
            mfg = (tx.manufacturer or "").strip()
            rating_raw = (tx.rating_kva or "").strip()
            if rating_raw.upper() in ("NOT ACCESSIBLE", "-", "N/A", "NONE", "NAN", ""):
                rating_str = ""
            elif not rating_raw.upper().endswith("KVA"):
                rating_str = f"{rating_raw}kVA"
            else:
                rating_str = rating_raw

            spec_str = f"{mfg} {rating_str}".strip() if (mfg and rating_str) else (mfg or rating_str)
            tx_title = f"{tx_label} - {spec_str}".strip() if spec_str else tx_label

            components = TRANSFORMER_CENSUS_COMPONENTS
            if not has_hv_cable_split(tx):
                components = tuple(c for c in components if c != "CABLE SPLIT")

            tx_group_key = f"tx_{idx + 1}"

            for comp in components:
                is_ov = comp in ("OVERVIEW", "OVERVIEW TOP")
                matched = _match_defects_for_tx_component(
                    tx, comp, defects, total_transformers=len(txs)
                )
                if matched:
                    ir_s, us_s, tev_s = _format_defect_readings(matched)
                    rows.append(
                        CensusRowItem(
                            no=f"{group_num}.",
                            equipment=tx_title,
                            defect_area=comp,
                            ir_abs=ir_s,
                            us_dB=us_s,
                            tev_dB=tev_s,
                            severity="DEFECT",
                            group_no=group_num,
                            is_overview=False,
                            is_defect=True,
                            eq_group_key=tx_group_key,
                        )
                    )
                elif is_ov:
                    rows.append(
                        CensusRowItem(
                            no=f"{group_num}.",
                            equipment=tx_title,
                            defect_area=comp,
                            ir_abs="-",
                            us_dB="-",
                            tev_dB="-",
                            severity="-",
                            group_no=group_num,
                            is_overview=True,
                            is_defect=False,
                            eq_group_key=tx_group_key,
                        )
                    )
                else:
                    rows.append(
                        CensusRowItem(
                            no=f"{group_num}.",
                            equipment=tx_title,
                            defect_area=comp,
                            ir_abs="-",
                            us_dB="-",
                            tev_dB="-",
                            severity="NORMAL",
                            group_no=group_num,
                            is_overview=False,
                            is_defect=False,
                            eq_group_key=tx_group_key,
                        )
                    )

        # ----------------------------------------------------------------------
        # 3. Feeder Pillar / LVDB (D22 - 1 Overview always, defects appended)
        # ----------------------------------------------------------------------
        for fp_idx, lvdb in enumerate(lvdbs, 1):
            group_num = current_group
            current_group += 1
            fp_title = lvdb.name.strip() or "FEEDER PILLAR"

            # Always emit 1 OVERVIEW row
            rows.append(
                CensusRowItem(
                    no=f"{group_num}.",
                    equipment=fp_title,
                    defect_area="OVERVIEW",
                    ir_abs="-",
                    us_dB="-",
                    tev_dB="-",
                    severity="-",
                    group_no=group_num,
                    is_overview=True,
                    is_defect=False,
                    eq_group_key=f"fp_{fp_idx}_ov",
                )
            )

            # Append active defect feeder rows
            lvdb_defects = _match_defects_for_lvdb(lvdb, defects, total_lvdbs=len(lvdbs))
            if lvdb_defects:
                pe_info = {"equipment_package": package, "equipment_specs": package}
                qr_rows = prepare_tech_summary_rows(lvdb_defects, pe_info=pe_info)
                for def_idx, qr_r in enumerate(qr_rows, 1):
                    eq_name = (qr_r.equipment or "").strip() or fp_title
                    if eq_name.upper() in (fp_title.upper(), "FEEDER PILLAR", "FP", "LVDB"):
                        eq_name = fp_title
                    elif not eq_name.upper().startswith(fp_title.upper()):
                        cleaned_sub = re.sub(
                            r"^(?:FEEDER\s*PILLAR|FP|LVDB)(?:\s*\d+)?(?:\s*[-–]\s*|\s+)",
                            "",
                            eq_name,
                            flags=re.IGNORECASE,
                        ).strip()
                        eq_name = f"{fp_title} - {cleaned_sub}" if cleaned_sub else fp_title
                    rows.append(
                        CensusRowItem(
                            no=f"{group_num}.",
                            equipment=eq_name,
                            defect_area=qr_r.defect_area or "DEFECT",
                            ir_abs=qr_r.ir_reading or "-",
                            us_dB=qr_r.us_reading or "-",
                            tev_dB=qr_r.tev_reading or "-",
                            severity="DEFECT",
                            group_no=group_num,
                            is_overview=False,
                            is_defect=True,
                            eq_group_key=f"fp_{fp_idx}_def_{def_idx}",
                        )
                    )

        # ----------------------------------------------------------------------
        # 4. Battery Bank(s)
        # ----------------------------------------------------------------------
        for bb_idx, bb in enumerate(bbs, 1):
            group_num = current_group
            current_group += 1
            bb_title = bb.name.strip() or "BATTERY BANK 1"
            matched = _match_defects_for_battery(bb, defects)

            if matched:
                ir_s, us_s, tev_s = _format_defect_readings(matched)
                rows.append(
                    CensusRowItem(
                        no=f"{group_num}.",
                        equipment=bb_title,
                        defect_area="BATTERY SYSTEM",
                        ir_abs=ir_s,
                        us_dB=us_s,
                        tev_dB=tev_s,
                        severity="DEFECT",
                        group_no=group_num,
                        is_overview=False,
                        is_defect=True,
                        eq_group_key=f"bb_{bb_idx}_def",
                    )
                )
            else:
                rows.append(
                    CensusRowItem(
                        no=f"{group_num}.",
                        equipment=bb_title,
                        defect_area="OVERVIEW",
                        ir_abs="-",
                        us_dB="-",
                        tev_dB="-",
                        severity="-",
                        group_no=group_num,
                        is_overview=True,
                        is_defect=False,
                        eq_group_key=f"bb_{bb_idx}_ov",
                    )
                )

        return rows

    def build_context(
        self,
        package: SubstationEquipmentPackage | FullReportScanPackage,
        defects: Sequence[CbmDefectRecord] = (),
        substation_number: int | str = "",
        station_name: str = "",
    ) -> ExecutiveSummaryCensusContext:
        """Build render context containing assembled census items and metadata."""
        items = self.build_census_rows(package, defects=defects)
        return ExecutiveSummaryCensusContext(
            census_items=items,
            substation_number=substation_number,
            station_name=station_name,
        )

    def render(
        self,
        package_or_context: SubstationEquipmentPackage | FullReportScanPackage | ExecutiveSummaryCensusContext,
        defects: Sequence[CbmDefectRecord] = (),
        output_path: str | Path | None = None,
        substation_number: int | str = "",
        station_name: str = "",
    ) -> ExecutiveSummaryCensusResult:
        """Render Executive Summary Census table into Word document template."""
        if not self.template_path.exists():
            raise FileNotFoundError(f"Census template not found at {self.template_path}")

        if isinstance(package_or_context, ExecutiveSummaryCensusContext):
            context = package_or_context
            items = context.census_items
        else:
            items = self.build_census_rows(package_or_context, defects=defects)
            context = self.build_context(
                package_or_context,
                defects=defects,
                substation_number=substation_number,
                station_name=station_name,
            )

        doc = DocxTemplate(str(self.template_path))
        doc.render(context.to_dict(), jinja_env=_build_jinja_env(), autoescape=True)

        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)
        reloaded_doc = docx.Document(bio)
        apply_post_render_dom(reloaded_doc, items)

        out_p = Path(output_path) if output_path else None
        if out_p:
            out_p.parent.mkdir(parents=True, exist_ok=True)
            reloaded_doc.save(str(out_p))

        groups = {item.group_no for item in items}
        has_defects = any(item.is_defect or item.severity == "DEFECT" for item in items)

        return ExecutiveSummaryCensusResult(
            docx_path=out_p,
            items=items,
            group_count=len(groups),
            total_rows=len(items),
            has_defects=has_defects,
        )


__all__ = [
    "TRANSFORMER_CENSUS_COMPONENTS",
    "CensusRowItem",
    "ExecutiveSummaryCensusBuilder",
    "ExecutiveSummaryCensusContext",
    "ExecutiveSummaryCensusResult",
    "apply_column_vertical_merge",
    "apply_equipment_vertical_merge",
    "apply_post_render_dom",
    "apply_severity_shading",
]
