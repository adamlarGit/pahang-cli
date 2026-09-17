"""Unified cell shading and post-render styling engine.

Single source of truth for:
- Canonical color hex codes (COLOR_DEFECT, COLOR_HEALTHY, COLOR_NORMAL, COLOR_WHITE)
- Technology severity placeholder markers (__SEVERITY_IR__, __SEVERITY_US__, __SEVERITY_TEV__)
- Technology normalization supporting composite '+', ',', '/' delimiters
- Cell technology detection and banner text identification
- OpenXML table cell shading, border, and text clearing operations
- Dynamic technology severity shading and analysis/recommendation banner shading
- Dash measurement unit cleanup ('-dB', '-°C', etc.)
- apply_scan_post_processing() pipeline for Quick Report and Full Report
"""

from __future__ import annotations

import re
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ─── Canonical Color Hex Codes ───────────────────────────────────────────────
COLOR_DEFECT: str = "EE0000"   # Red
COLOR_HEALTHY: str = "00B050"  # Green
COLOR_NORMAL: str = "00B050"   # Alias for healthy
COLOR_WHITE: str = "FFFFFF"    # White

# ─── Canonical Banner Texts ──────────────────────────────────────────────────
BANNER_HEALTHY_ANALYSIS: str = "No Anomaly."
BANNER_HEALTHY_RECOMMENDATION: str = "-"
BANNER_DEFECT_FORWARDING: str = "Please refer to the following page for details defect."
BANNER_DEFECT_FORWARDING_ANALYSIS: str = "Please refer to the following page for details defect."
BANNER_DEFECT_FORWARDING_RECOMMENDATION: str = "Please refer to the following page for details defect."

# ─── Sentinels for Jinja Severity Placeholders ────────────────────────────────
SEVERITY_MARKER_IR: str = "__SEVERITY_IR__"
SEVERITY_MARKER_US: str = "__SEVERITY_US__"
SEVERITY_MARKER_TEV: str = "__SEVERITY_TEV__"

from src.core.contract import (
    TEV_BLANKED_SWG_COMPARTMENTS,
    TEV_ELIGIBLE_SWG_COMPARTMENTS,
    US_BLANKED_SWG_COMPARTMENTS,
    is_swg_compartment_tev_eligible,
    is_swg_compartment_us_eligible,
    is_swg_tev_active,
    is_swg_us_active,
    is_tev_contract_awarded,
    is_us_contract_awarded,
    normalize_swg_compartment,
    normalize_technologies,
)

# Backward-compatible aliases
_normalize_technologies = normalize_technologies



def clear_cell_text(cell: Any) -> None:
    """Clear paragraph text, runs, and drawing objects in a python-docx table cell or raw <w:tc> element."""
    for paragraph in getattr(cell, "paragraphs", ()):
        for run in getattr(paragraph, "runs", ()):
            run.text = ""
        paragraph.text = ""
    tc = _resolve_tc(cell)
    if tc is not None and hasattr(tc, "iter"):
        for parent in list(tc.iter()):
            for child in list(parent):
                tag = getattr(child, "tag", "")
                if isinstance(tag, str) and tag.endswith(("}drawing", "}pict", "}object", "drawing", "pict", "object")):
                    parent.remove(child)
        for node in tc.iter():
            if isinstance(node.tag, str) and node.tag.endswith("}t"):
                node.text = ""


def _resolve_tc(target: Any) -> Any:
    """Extract <w:tc> oxml element from a Cell or raw tc element, or return None."""
    if hasattr(target, "_tc"):
        return target._tc
    tag = getattr(target, "tag", None)
    if isinstance(tag, str) and (tag.endswith("}tc") or tag in ("w:tc", "tc")):
        return target
    return None


def set_cell_shading(cell: Any, hex_color: str) -> None:
    """Set background fill color on a python-docx table cell (e.g. 'EE0000' or '00B050')."""
    tc = _resolve_tc(cell)
    if tc is None:
        return
    tcPr = tc.get_or_add_tcPr() if hasattr(tc, "get_or_add_tcPr") else tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr) if hasattr(tc, "insert") else tc.append(tcPr)
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#").upper())


def get_cell_shading(cell: Any) -> str | None:
    """Read w:fill hex color attribute from a table cell's tcPr/w:shd XML element."""
    tc = _resolve_tc(cell)
    if tc is None:
        return None
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    return shd.attrib.get(qn("w:fill"))


def clear_cell_shading(cell: Any) -> None:
    """Remove background shading (<w:shd>) from a table cell or raw <w:tc> element."""
    tc = _resolve_tc(cell)
    if tc is None:
        return
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is not None:
        shd = tcPr.find(qn("w:shd"))
        if shd is not None:
            tcPr.remove(shd)


def _set_tc_border_nil(cell: Any, border_name: str) -> None:
    """Set a specific border on a cell or raw <w:tc> element to nil."""
    tc = _resolve_tc(cell)
    if tc is None:
        return
    tcPr = tc.get_or_add_tcPr() if hasattr(tc, "get_or_add_tcPr") else tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr) if hasattr(tc, "insert") else tc.append(tcPr)
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    border = tcBorders.find(qn(f"w:{border_name}"))
    if border is None:
        border = OxmlElement(f"w:{border_name}")
        tcBorders.append(border)
    border.set(qn("w:val"), "nil")


def set_cell_no_borders(cell: Any) -> None:
    """Set all borders on a python-docx table cell or raw <w:tc> element to nil (invisible)."""
    tc = _resolve_tc(cell)
    if tc is None:
        return
    tcPr = tc.get_or_add_tcPr() if hasattr(tc, "get_or_add_tcPr") else tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr) if hasattr(tc, "insert") else tc.append(tcPr)
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV", "tl2br", "tr2bl"):
        border = tcBorders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            tcBorders.append(border)
        border.set(qn("w:val"), "nil")


def detect_cell_technology(cell: Any) -> str | None:
    """Detect if a table cell represents a technology severity cell (IR, US, TEV).

    Returns 'IR', 'US', 'TEV', or None.
    """
    text = cell.text.strip() if hasattr(cell, "text") else ""
    if not text:
        return None

    text_upper = text.upper()

    # 1. Exact or marker substring matches
    if SEVERITY_MARKER_IR in text_upper or "{{ IR.SEVERITY }}" in text_upper or "{{IR.SEVERITY}}" in text_upper:
        return "IR"
    if SEVERITY_MARKER_US in text_upper or "{{ US.SEVERITY }}" in text_upper or "{{US.SEVERITY}}" in text_upper:
        return "US"
    if SEVERITY_MARKER_TEV in text_upper or "{{ TEV.SEVERITY }}" in text_upper or "{{TEV.SEVERITY}}" in text_upper:
        return "TEV"

    # 2. Jinja variable pattern
    if re.search(r"\{\{\s*ir\.severity\s*\}\}", text, re.I):
        return "IR"
    if re.search(r"\{\{\s*us\.severity\s*\}\}", text, re.I):
        return "US"
    if re.search(r"\{\{\s*tev\.severity\s*\}\}", text, re.I):
        return "TEV"

    # 3. Simple token patterns
    if text_upper in ("IR", "IR_SEVERITY", "[IR_SEVERITY]", "IR.SEVERITY", "IR SEVERITY"):
        return "IR"
    if text_upper in ("US", "US_SEVERITY", "[US_SEVERITY]", "US.SEVERITY", "US SEVERITY", "U/S", "U/S_SEVERITY", "U/S.SEVERITY", "U/S SEVERITY"):
        return "US"
    if text_upper in ("TEV", "TEV_SEVERITY", "[TEV_SEVERITY]", "TEV.SEVERITY", "TEV SEVERITY"):
        return "TEV"

    # 4. Heading + technology patterns
    tokens = set(re.findall(r"\b[A-Z0-9/]+\b", text_upper))
    if "SEVERITY" in tokens:
        if "IR" in tokens:
            return "IR"
        if "US" in tokens or "U/S" in tokens:
            return "US"
        if "TEV" in tokens:
            return "TEV"

    return None


def is_defect_forwarding_text(text: str) -> bool:
    """Check if cell text corresponds to downstream defect forwarding prose."""
    lower = text.lower()
    return (
        "please refer to the following page for detail" in lower
        or "please refer to the following page for details" in lower
        or "refer to the following page" in lower
        or "following page for details defect" in lower
        or "following page for detail defect" in lower
    )


def is_healthy_banner_text(text: str) -> bool:
    """Check if cell text corresponds to healthy Analysis / Recommendation prose."""
    lower = text.lower().strip()
    if not lower:
        return False
    if is_defect_forwarding_text(lower):
        return False

    # Check for explicit English phrases
    if "no anomaly" in lower or "no defect" in lower:
        return True

    # Require explicit prefix: 'analysis:' or 'recommendation:'
    m = re.match(r"^(?:analysis|recommendation)\s*:\s*(.*)$", lower)
    if m:
        clean = m.group(1).strip()
        if not clean or clean in ("-", "--", "none", "n/a", "na", "nil", "normal"):
            return True
        if "no anomaly" in clean or "no defect" in clean:
            return True

    return False


def _get_cell_tc(cell: Any) -> Any:
    """Extract underlying XML _tc element from a cell or wrapper."""
    if hasattr(cell, "_tc"):
        return cell._tc
    return cell


def _is_cell(target: Any) -> bool:
    """Check if target represents a single table cell."""
    return hasattr(target, "paragraphs") and hasattr(target, "_tc") and not hasattr(target, "rows")


def _iter_unique_cells(target: Any) -> list[Any]:
    """Extract a deduplicated list of table cells from various containers.

    Accepts a single _Cell, _Row, Table, Document, DocxTemplate, or sequence thereof.
    Ensures merged cells sharing the same underlying XML _tc element are only yielded once.
    """
    cells: list[Any] = []
    seen_tcs: set[Any] = set()

    def _collect(obj: Any) -> None:
        if obj is None:
            return
        if _is_cell(obj):
            tc = _get_cell_tc(obj)
            if tc not in seen_tcs:
                seen_tcs.add(tc)
                cells.append(obj)
            return
        if hasattr(obj, "cells"):  # Table row (_Row)
            for c in obj.cells:
                _collect(c)
            return
        if hasattr(obj, "rows"):  # Table
            for r in obj.rows:
                for c in r.cells:
                    _collect(c)
            return
        if hasattr(obj, "docx") and hasattr(obj.docx, "tables"):  # DocxTemplate
            for t in obj.docx.tables:
                _collect(t)
            return
        if hasattr(obj, "tables"):  # docx.Document
            for t in obj.tables:
                _collect(t)
            return
        if isinstance(obj, (list, tuple, set)):
            for item in obj:
                _collect(item)

    _collect(target)
    return cells


def apply_technology_severity_shading(
    target: Any,
    *,
    defective_technologies: set[str] | list[str] | tuple[str, ...] | str | None = None,
    technology: str | None = None,
    is_defective: bool | None = None,
    is_overview: bool = False,
) -> Any:
    """Dynamically shade technology severity cells (IR, US, TEV) Green or Red, clearing text.

    - Healthy cell: Green '00B050', text cleared.
    - Defective cell: Red 'EE0000', text cleared.
    - On overview pages (is_overview=True): cell text is cleared and set to '-'.

    Can operate on a single _Cell, a _Row, a Table, a Document, a DocxTemplate, or an iterable thereof.
    """
    def_techs = _normalize_technologies(defective_technologies)
    cells = _iter_unique_cells(target)
    single_cell_mode = len(cells) == 1 and _is_cell(target)

    for cell in cells:
        detected_tech = detect_cell_technology(cell)
        tech = (technology or detected_tech or "").upper()

        if not tech:
            if single_cell_mode and is_defective is not None:
                cell_defective = is_defective
            else:
                continue
        else:
            if is_defective is not None:
                cell_defective = is_defective
            else:
                cell_defective = tech in def_techs

        clear_cell_text(cell)
        if is_overview:
            if getattr(cell, "paragraphs", None):
                cell.paragraphs[0].text = "-"
        else:
            set_cell_shading(cell, COLOR_DEFECT if cell_defective else COLOR_HEALTHY)

    return target


def _sanitize_healthy_banner_cell(cell: Any) -> None:
    """Sanitize leftover defect forwarding prose in a healthy banner cell."""
    text = cell.text.strip()
    lower = text.lower()
    if lower.startswith("analysis:"):
        clear_cell_text(cell)
        p = cell.paragraphs[0] if getattr(cell, "paragraphs", None) else cell.add_paragraph()
        r0 = p.add_run("Analysis: ")
        r0.bold = True
        r0.italic = True
        p.add_run(" ")
        p.add_run("No Anomaly.")
    elif lower.startswith("recommendation:"):
        clear_cell_text(cell)
        p = cell.paragraphs[0] if getattr(cell, "paragraphs", None) else cell.add_paragraph()
        r0 = p.add_run("Recommendation: ")
        r0.bold = True
        r0.italic = True
        p.add_run(" ")
        p.add_run("-")
    elif is_defect_forwarding_text(text):
        clear_cell_text(cell)
        p = cell.paragraphs[0] if getattr(cell, "paragraphs", None) else cell.add_paragraph()
        p.add_run("No Anomaly.")


def apply_banner_shading(
    target: Any,
    *,
    is_defective: bool | None = None,
) -> Any:
    """Dynamically shade Analysis & Recommendation banner cells.

    - Healthy ('No Anomaly.'): Green '00B050'. If cell text has leftover defect text,
      sanitizes it to 'Analysis:  No Anomaly.' / 'Recommendation:  -'.
    - Defect Forwarding prose or is_defective=True: Red 'EE0000'.

    Can operate on a single _Cell, a _Row, a Table, a Document, a DocxTemplate, or an iterable thereof.
    """
    cells = _iter_unique_cells(target)
    single_cell_mode = len(cells) == 1 and _is_cell(target)

    for cell in cells:
        text = cell.text.strip()
        lower = text.lower()
        is_header = lower.rstrip(":").strip() in ("analysis & recommendations", "analysis & recommendation")
        if is_header:
            continue

        if not single_cell_mode:
            is_banner_candidate = (
                lower.startswith("analysis:")
                or lower.startswith("recommendation:")
                or is_defect_forwarding_text(text)
            )
            if not is_banner_candidate:
                continue

        if is_defective is True:
            if (
                single_cell_mode
                or is_defect_forwarding_text(text)
                or lower.startswith("analysis:")
                or lower.startswith("recommendation:")
            ):
                set_cell_shading(cell, COLOR_DEFECT)
        elif is_defective is False:
            if (
                single_cell_mode
                or is_healthy_banner_text(text)
                or lower.startswith("analysis:")
                or lower.startswith("recommendation:")
                or is_defect_forwarding_text(text)
            ):
                if is_defect_forwarding_text(text) or "defect" in lower:
                    _sanitize_healthy_banner_cell(cell)
                set_cell_shading(cell, COLOR_HEALTHY)
        else:  # is_defective is None
            if is_defect_forwarding_text(text):
                set_cell_shading(cell, COLOR_DEFECT)
            elif is_healthy_banner_text(text):
                set_cell_shading(cell, COLOR_HEALTHY)

    return target


def cleanup_dash_measurement_units(target: Any) -> Any:
    """Clean up '-dB', '- dB', '-°C', '- °C', '-%' artifacts produced when a measurement is '-'."""
    cells = _iter_unique_cells(target)
    for cell in cells:
        for p in getattr(cell, "paragraphs", ()):
            for run in getattr(p, "runs", ()):
                if run.text:
                    for bad in ("-dB", "- dB", "-°C", "- °C", "-%"):
                        if bad in run.text:
                            run.text = run.text.replace(bad, "-")
    return target





def _collect_tables(obj: Any) -> list[Any]:
    collected = []
    if hasattr(obj, "docx") and hasattr(obj.docx, "tables"):
        collected.extend(obj.docx.tables)
    elif hasattr(obj, "tables"):
        collected.extend(obj.tables)
    elif hasattr(obj, "rows") and hasattr(obj, "columns"):
        collected.append(obj)
    elif isinstance(obj, (list, tuple, set)):
        for item in obj:
            collected.extend(_collect_tables(item))
    return collected


def blank_swg_tev_cells(target: Any) -> Any:
    """Blank TEV measurement cells in swg-panel.docx 37x24 table.

    Operates on rows 21–32 (indices 21..32 inclusive) and columns 11–22 (indices 11..22 inclusive).
    Directly iterates over raw XML <w:tc> elements in each row's <w:tr> to bypass python-docx
    vMerge masking (where table.cell() on vertically merged continuation cells resolves back
    to the anchor cell, leaving continuation row borders visible around the {{ tev.prpd }} quadrant).
    - Clears cell text
    - Removes background cell shading (<w:shd>)
    - Sets all borders to <w:val="nil"/>
    Also clears bordering spacer borders facing the TEV block (col 10 right border,
    col 23 left border, row 20 bottom border, row 33 top border) to eliminate ghost borders.
    """
    tables = _collect_tables(target)

    for table in tables:
        # Match swg-panel.docx 37x24 table layout (tightened guard: exactly 24 columns)
        if len(table.rows) >= 33 and len(table.columns) == 24:
            # 1. Process rows 21..32 at the raw XML <w:tc> level
            for r_idx in range(21, 33):
                row = table.rows[r_idx]
                col_idx = 0
                for tc in row._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span

                    # If this cell falls entirely within TEV columns (cols 11..22)
                    if col_start >= 11 and col_end <= 22:
                        clear_cell_text(tc)
                        clear_cell_shading(tc)
                        set_cell_no_borders(tc)

                    # If this cell is the adjacent left spacer (ending at col 10)
                    elif col_end == 10:
                        _set_tc_border_nil(tc, "right")

                    # If this cell is the adjacent right spacer (starting at col 23)
                    elif col_start == 23:
                        _set_tc_border_nil(tc, "left")

            # 2. Clear bottom borders of adjacent top spacer cells (Row 20) facing TEV columns
            if len(table.rows) > 20:
                row_20 = table.rows[20]
                col_idx = 0
                for tc in row_20._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span
                    if max(col_start, 11) <= min(col_end, 22):
                        _set_tc_border_nil(tc, "bottom")

            # 3. Clear top borders of adjacent bottom spacer cells (Row 33) facing TEV columns
            if len(table.rows) > 33:
                row_33 = table.rows[33]
                col_idx = 0
                for tc in row_33._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span
                    if max(col_start, 11) <= min(col_end, 22):
                        _set_tc_border_nil(tc, "top")

    return target


def blank_swg_us_cells(target: Any) -> Any:
    """Blank Ultrasound (US) measurement cells in swg-panel.docx 37x24 table.

    Operates on rows 21–32 (indices 21..32 inclusive) and columns 1–8 (indices 1..8 inclusive).
    Directly iterates over raw XML <w:tc> elements in each row's <w:tr> to bypass python-docx
    vMerge masking (where table.cell() on vertically merged continuation cells resolves back
    to the anchor cell, leaving continuation row borders visible around the {{ us.prpd }} quadrant).
    - Clears cell text
    - Removes background cell shading (<w:shd>)
    - Sets all borders to <w:val="nil"/>
    Also clears bordering spacer borders facing the US block (col 0 right border,
    col 9 left border, row 20 bottom border, row 33 top border) to eliminate ghost borders.
    """
    tables = _collect_tables(target)

    for table in tables:
        # Match swg-panel.docx 37x24 table layout (tightened guard: exactly 24 columns)
        if len(table.rows) >= 33 and len(table.columns) == 24:
            # 1. Process rows 21..32 at the raw XML <w:tc> level
            for r_idx in range(21, 33):
                row = table.rows[r_idx]
                col_idx = 0
                for tc in row._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span

                    # If this cell falls entirely within US columns (cols 1..8)
                    if col_start >= 1 and col_end <= 8:
                        clear_cell_text(tc)
                        clear_cell_shading(tc)
                        set_cell_no_borders(tc)

                    # If this cell is the adjacent left spacer (ending at col 0)
                    elif col_end == 0:
                        _set_tc_border_nil(tc, "right")

                    # If this cell is the adjacent right spacer (starting at col 9)
                    elif col_start == 9:
                        _set_tc_border_nil(tc, "left")

            # 2. Clear bottom borders of adjacent top spacer cells (Row 20) facing US columns
            if len(table.rows) > 20:
                row_20 = table.rows[20]
                col_idx = 0
                for tc in row_20._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span
                    if max(col_start, 1) <= min(col_end, 8):
                        _set_tc_border_nil(tc, "bottom")

            # 3. Clear top borders of adjacent bottom spacer cells (Row 33) facing US columns
            if len(table.rows) > 33:
                row_33 = table.rows[33]
                col_idx = 0
                for tc in row_33._tr.findall(qn("w:tc")):
                    tcPr = tc.find(qn("w:tcPr"))
                    gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                    span = 1
                    if gridSpan_elem is not None:
                        try:
                            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1))
                        except (ValueError, TypeError):
                            span = 1
                    col_start = col_idx
                    col_end = col_idx + span - 1
                    col_idx += span
                    if max(col_start, 1) <= min(col_end, 8):
                        _set_tc_border_nil(tc, "top")

    return target


def apply_scan_post_processing(
    target: Any,
    *,
    defective_technologies: set[str] | list[str] | tuple[str, ...] | str | None = None,
    is_defective: bool | None = None,
    is_overview: bool = False,
    blank_tev: bool = False,
    blank_us: bool = False,
) -> Any:
    """Unified post-render processing engine for scanning pages and CBM defect pages.

    1. If blank_tev is True, clears text, shading, and sets nil borders for TEV cells on swg-panel.docx.
    2. If blank_us is True, clears text, shading, and sets nil borders for US cells on swg-panel.docx.
    3. Applies technology severity shading (IR, US, TEV cells) Green/Red, clearing text.
       (If is_overview is True, clears text and sets to '-' if applicable).
    4. Applies banner shading (Analysis / Recommendation) Green ('00B050') or Red ('EE0000').
    5. Cleans up dash measurement unit artifacts ('-dB', '-°C', etc.).
    6. If blank_tev is True, re-verifies TEV cells remain strictly blanked, unshaded, and borderless.
    7. If blank_us is True, re-verifies US cells remain strictly blanked, unshaded, and borderless.
    """
    if blank_tev:
        blank_swg_tev_cells(target)
    if blank_us:
        blank_swg_us_cells(target)

    def_techs = _normalize_technologies(defective_technologies)
    if is_defective is True and not def_techs and not is_overview:
        def_techs.add("IR")

    if blank_tev and "TEV" in def_techs:
        def_techs.remove("TEV")
    if blank_us and "US" in def_techs:
        def_techs.remove("US")

    apply_technology_severity_shading(
        target,
        defective_technologies=def_techs,
        is_overview=is_overview,
    )
    apply_banner_shading(target, is_defective=is_defective)
    cleanup_dash_measurement_units(target)

    if blank_tev:
        blank_swg_tev_cells(target)
    if blank_us:
        blank_swg_us_cells(target)

    return target


__all__ = [
    "COLOR_DEFECT",
    "COLOR_HEALTHY",
    "COLOR_NORMAL",
    "COLOR_WHITE",
    "BANNER_HEALTHY_ANALYSIS",
    "BANNER_HEALTHY_RECOMMENDATION",
    "BANNER_DEFECT_FORWARDING",
    "BANNER_DEFECT_FORWARDING_ANALYSIS",
    "BANNER_DEFECT_FORWARDING_RECOMMENDATION",
    "SEVERITY_MARKER_IR",
    "SEVERITY_MARKER_US",
    "SEVERITY_MARKER_TEV",
    "TEV_ELIGIBLE_SWG_COMPARTMENTS",
    "TEV_BLANKED_SWG_COMPARTMENTS",
    "US_BLANKED_SWG_COMPARTMENTS",
    "clear_cell_text",
    "set_cell_shading",
    "get_cell_shading",
    "clear_cell_shading",
    "set_cell_no_borders",
    "detect_cell_technology",
    "is_defect_forwarding_text",
    "is_healthy_banner_text",
    "normalize_swg_compartment",
    "is_swg_compartment_tev_eligible",
    "is_swg_compartment_us_eligible",
    "is_tev_contract_awarded",
    "is_us_contract_awarded",
    "is_swg_tev_active",
    "is_swg_us_active",
    "normalize_technologies",
    "_normalize_technologies",
    "blank_swg_tev_cells",
    "blank_swg_us_cells",
    "apply_technology_severity_shading",
    "apply_banner_shading",
    "cleanup_dash_measurement_units",
    "apply_scan_post_processing",
]
