"""33kV PPU testsheet parser implementing TestsheetParserStrategy."""

from __future__ import annotations

from datetime import datetime
import logging
from pathlib import Path
import re
from typing import Any, ClassVar, Mapping, Sequence
import warnings
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from src.core.normalizers import (
    format_humidity_str,
    format_load_amp,
    format_testsheet_time,
    parse_background_temp,
)
from src.testsheet.models import (
    LVDBSpec,
    PhotoRange,
    PowerTransformerSpec,
    RawPhotoRanges,
    SubstationEquipmentPackage,
    SwitchgearPanelSpec,
    SwitchgearSpec,
    TestsheetData,
    ThermalReadingSpec,
    TransformerSpec,
    VisualInspectionFinding,
)
from src.testsheet.parsers.base import TestsheetParserStrategy
from src.testsheet.parsers.common import (
    clean_val,
    find_next_val,
    is_date_folder_name,
    is_marked,
    normalize_fl_erms,
    parse_int_safe,
    parse_photo_numbers,
    parse_range_val,
    to_excel_date,
)

logger = logging.getLogger(__name__)

_GPS_COORD_RE = re.compile(r"\d+\.\d+.*,\s*\d+\.\d+")


def classify_transformer_rating(
    rating: object,
    tx_id: str = "",
    unit_hint: str = "",
) -> bool:
    """Classify whether a transformer is a PowerTransformerSpec (True) or TransformerSpec (False).

    Returns:
        True if rating >= 5 MVA (PowerTransformerSpec).
        False if rating in 100-1000 kVA, labeled LOCAL TX, or out-of-range / ambiguous (standard TransformerSpec).
    """
    if "LOCAL" in tx_id.upper():
        return False

    raw_str = str(rating).strip() if rating is not None else ""
    if not raw_str or raw_str.upper() in ("-", "NONE", "N/A", "NAN"):
        logger.warning(
            "Empty or missing rating for transformer '%s'; defaulting to standard TransformerSpec",
            tx_id or "unnamed",
        )
        return False

    s_upper = raw_str.upper()

    m = re.search(r"(\d+(?:\.\d+)?)", s_upper)
    if not m:
        logger.warning(
            "Unrecognized transformer rating '%s' for '%s'; defaulting to standard TransformerSpec",
            raw_str,
            tx_id or "unnamed",
        )
        return False

    val = float(m.group(1))

    has_mva = "MVA" in s_upper or ("KVA" not in s_upper and unit_hint.upper() == "MVA")
    has_kva = "KVA" in s_upper or ("MVA" not in s_upper and unit_hint.upper() in ("KVA", "VA"))

    if not has_mva and not has_kva:
        if 5.0 <= val < 100.0:
            has_mva = True
        elif 100.0 <= val <= 1000.0:
            has_kva = True
        elif val >= 5000.0:
            return True
        else:
            logger.warning(
                "Ambiguous transformer rating '%s' without unit for '%s'; defaulting to standard TransformerSpec",
                raw_str,
                tx_id or "unnamed",
            )
            return False

    if has_mva:
        if val >= 5.0:
            return True
        else:
            logger.warning(
                "Out-of-range Power Transformer rating '%s' (< 5 MVA) for '%s'; defaulting to standard TransformerSpec",
                raw_str,
                tx_id or "unnamed",
            )
            return False

    if has_kva:
        if 100.0 <= val <= 1000.0:
            return False
        elif val >= 5000.0:
            return True
        else:
            logger.warning(
                "Out-of-range transformer rating '%s' (not in 100-1000 kVA or >= 5 MVA) for '%s'; defaulting to standard TransformerSpec",
                raw_str,
                tx_id or "unnamed",
            )
            return False

    return False


class PpuTestsheetParser(TestsheetParserStrategy):
    """Parser strategy for 33kV PPU/SSU/PMU multi-sheet workbooks."""

    __test__ = False
    HANDLED_PREFIXES: ClassVar[tuple[str, ...]] = (
        "33kV (Main)",
        "33kV (Extra)",
        "11kV (Main)",
        "11kV (Extra)",
    )

    def can_handle_sheets(self, sheet_names: Sequence[str]) -> bool:
        """Return True if sheets contain 33kV or 11kV PPU sheets."""
        for s in sheet_names:
            s_clean = s.strip()
            if any(s_clean.startswith(prefix) for prefix in self.HANDLED_PREFIXES):
                return True
        return False

    def extract_photo_ranges(self, workbook_path: Path | str) -> RawPhotoRanges:
        """Extract IR and DG photo range bounds from 33kV testsheet RAW DATA sheet."""
        path = Path(workbook_path)
        if not path.exists():
            raise FileNotFoundError(f"Testsheet workbook not found: {path}")

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            try:
                if "RAW DATA" in wb.sheetnames:
                    return self._parse_raw_data_sheet(wb["RAW DATA"])
            finally:
                wb.close()

        return RawPhotoRanges()

    def extract_testsheet_metadata(
        self,
        workbook_path: Path | str,
        station_hint: str = "",
        date_hint: str = "",
    ) -> TestsheetData:
        """Fast read-only extraction of header metadata required for TOTAL PE."""
        path = Path(workbook_path)
        if not path.exists():
            raise FileNotFoundError(f"Testsheet workbook not found: {path}")

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            try:
                sheet_names = wb.sheetnames
                ws_33_main = wb["33kV (Main)"] if "33kV (Main)" in sheet_names else None
                ws_11_main = wb["11kV (Main)"] if "11kV (Main)" in sheet_names else None

                # Substation number resolution hierarchy
                substation_number = self._resolve_substation_number(path, ws_33_main, ws_11_main)

                # Header metadata priority: 33kV (Main) -> 11kV (Main)
                ws_header = ws_33_main if ws_33_main is not None else ws_11_main
                meta = self._extract_header_metadata(ws_header, station_hint, date_hint)

                # Substation type
                substation_type = self._resolve_substation_type(
                    meta["substation_name_site"], path.name, meta["fl_erms"]
                )

                # GPS from Visual Inspec. sheet
                ws_vi = self._find_vi_sheet(wb)
                sub_name = station_hint or meta.get("substation_name_site") or meta.get("substation_name_erms") or ""
                gps_coordinate = self._extract_gps_coordinate(ws_vi, substation_name=sub_name)

                equipment = SubstationEquipmentPackage(
                    switchgears=(meta["swg_meta"],) if meta["swg_meta"] is not None else ()
                )

            finally:
                wb.close()

        return TestsheetData(
            substation_number=substation_number,
            station_name=station_hint or meta["substation_name_erms"],
            date_str=meta["date_str"],
            fl_erms=meta["fl_erms"],
            fl_site=meta["fl_site"],
            substation_name_erms=meta["substation_name_erms"],
            substation_name_site=meta["substation_name_site"],
            gps_coordinate=gps_coordinate,
            substation_type=substation_type,
            ambient=meta["ambient"],
            humidity=meta["humidity"],
            time=meta["time_str"],
            tev_background=meta["tev_background"],
            cycle_1=meta["cycle_1"],
            equipment=equipment,
        )

    def extract_testsheet_data(
        self,
        workbook_path: Path | str,
        station_hint: str = "",
        date_hint: str = "",
    ) -> TestsheetData:
        """Parse 33kV testsheet workbook."""
        path = Path(workbook_path)
        if not path.exists():
            raise FileNotFoundError(f"Testsheet workbook not found: {path}")

        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet_names = wb.sheetnames
            ws_33_main = wb["33kV (Main)"] if "33kV (Main)" in sheet_names else None
            ws_11_main = wb["11kV (Main)"] if "11kV (Main)" in sheet_names else None

            # 1. Substation number resolution hierarchy
            substation_number = self._resolve_substation_number(path, ws_33_main, ws_11_main)

            # 2. Header metadata extraction (33kV -> 11kV fallback)
            ws_header = ws_33_main if ws_33_main is not None else ws_11_main
            meta = self._extract_header_metadata(ws_header, station_hint, date_hint)

            # Substation type
            substation_type = self._resolve_substation_type(
                meta["substation_name_site"], path.name, meta["fl_erms"]
            )

            # GPS from Visual Inspec. sheet
            ws_vi = self._find_vi_sheet(wb)
            sub_name = station_hint or meta.get("substation_name_site") or meta.get("substation_name_erms") or ""
            gps_coordinate = self._extract_gps_coordinate(ws_vi, substation_name=sub_name)

            # 3. Switchgear lineups assembly
            switchgears: list[SwitchgearSpec] = []

            # 33kV Lineup
            if ws_33_main is not None:
                extra_33 = self._find_extra_sheets(wb, "33kV (Extra)")
                swg_33 = self._parse_switchgear_lineup(ws_33_main, extra_sheets=extra_33)
                if swg_33 is not None:
                    switchgears.append(swg_33)

            # 11kV Lineup
            if ws_11_main is not None:
                extra_11 = self._find_extra_sheets(wb, "11kV (Extra)")
                swg_11 = self._parse_switchgear_lineup(ws_11_main, extra_sheets=extra_11)
                if swg_11 is not None:
                    switchgears.append(swg_11)

            if not switchgears:
                raise ValueError(
                    f"No populated 33kV or 11kV switchgear lineups found in {path.name}"
                )

            # 4. Transformers and LVDB extraction (TX sheet)
            transformers: list[TransformerSpec] = []
            lvdb_specs: list[LVDBSpec] = []
            if "TX" in wb.sheetnames:
                transformers, lvdb_specs = self._extract_tx_sheet(wb["TX"])

            # 5. Visual Inspection findings
            visual_findings: list[VisualInspectionFinding] = []
            if ws_vi is not None:
                visual_findings = self._parse_visual_inspection_findings(ws_vi)

            # 6. Photo range bounds
            photo_ranges = RawPhotoRanges()
            if "RAW DATA" in wb.sheetnames:
                photo_ranges = self._parse_raw_data_sheet(wb["RAW DATA"])

        finally:
            wb.close()

        has_cc = any(
            any("CELLAR" in str(comp).upper() for comp in getattr(s, "overview_compartments", ()))
            for s in switchgears
        )
        equipment = SubstationEquipmentPackage(
            switchgears=tuple(switchgears),
            transformers=tuple(transformers),
            lvdb_specs=tuple(lvdb_specs),
            has_cable_cellar=has_cc,
        )

        return TestsheetData(
            substation_number=substation_number,
            station_name=station_hint or meta["substation_name_erms"],
            date_str=meta["date_str"],
            fl_erms=meta["fl_erms"],
            fl_site=meta["fl_site"],
            photo_ranges=photo_ranges,
            substation_name_erms=meta["substation_name_erms"],
            substation_name_site=meta["substation_name_site"],
            gps_coordinate=gps_coordinate,
            substation_type=substation_type,
            ambient=meta["ambient"],
            humidity=meta["humidity"],
            time=meta["time_str"],
            tev_background=meta["tev_background"],
            cycle_1=meta["cycle_1"],
            equipment=equipment,
            visual_findings=tuple(visual_findings),
        )

    def _resolve_substation_number(
        self,
        path: Path,
        ws_33: Worksheet | None,
        ws_11: Worksheet | None,
    ) -> int:
        """Resolve substation number in strict hierarchy: V2 33kV -> V2 11kV -> filename prefix -> ValueError."""
        # 1. Cell V2 on 33kV (Main)
        if ws_33 is not None:
            num = self._parse_positive_int(ws_33["V2"].value)
            if num is not None:
                return num

        # 2. Cell V2 on 11kV (Main)
        if ws_11 is not None:
            num = self._parse_positive_int(ws_11["V2"].value)
            if num is not None:
                return num

        # 3. Filename or parent folder leading numeric prefix
        m = re.match(r"^(\d+)", path.name)
        if not m and path.parent and not is_date_folder_name(path.parent.name):
            m = re.match(r"^(\d+)", path.parent.name)
        if m:
            num = int(m.group(1))
            if num > 0:
                return num

        # 4. Fail fast
        raise ValueError(f"Unable to determine substation number for {path.name}")

    def _extract_header_metadata(
        self,
        ws_header: Worksheet | None,
        station_hint: str = "",
        date_hint: str = "",
    ) -> dict[str, Any]:
        """Extract header metadata cells (D6..D9, U6..U8, L6..L11) from main switchgear sheet."""
        fl_site = ""
        fl_erms = ""
        substation_name_site = ""
        substation_name_erms = ""
        date_str = date_hint
        cycle_1: datetime | None = None
        time_str = "-"
        ambient = "-"
        humidity = "-"
        tev_background = "-"
        swg_meta: SwitchgearSpec | None = None

        if ws_header is not None:
            raw_name = clean_val(ws_header["D6"].value) or ""
            substation_name_site = raw_name
            substation_name_erms = raw_name or station_hint

            raw_fl = ws_header["D7"].value
            fl_site = clean_val(raw_fl) or ""
            fl_erms = normalize_fl_erms(raw_fl)

            cycle_1 = to_excel_date(ws_header["D8"].value)
            if cycle_1 is not None:
                if not date_str:
                    date_str = cycle_1.strftime("%d-%m-%Y")
            elif not date_str and ws_header["D8"].value is not None:
                raw_d8 = str(ws_header["D8"].value).strip()
                if raw_d8 and raw_d8.lower() not in ("none", "-", "n/a"):
                    date_str = raw_d8

            time_str = format_testsheet_time(ws_header["D9"].value)
            ambient = parse_background_temp(ws_header["U6"].value)
            humidity = format_humidity_str(ws_header["U7"].value)
            tev_bg_val = clean_val(ws_header["U8"].value)
            if tev_bg_val is not None:
                tev_background = str(tev_bg_val)

            # Switchgear Details (L6..L11, M6)
            manufacturer = clean_val(ws_header["L6"].value) or ""
            type_from_m6 = clean_val(ws_header["M6"].value) or ""
            year = clean_val(ws_header["L7"].value) or ""
            model = clean_val(ws_header["L8"].value) or ""
            serial_no = clean_val(ws_header["L9"].value) or ""
            busbar_type = clean_val(ws_header["L10"].value) or ""
            rating = clean_val(ws_header["L11"].value) or ""
            swg_type = type_from_m6 or busbar_type or "VCB"
            if any((manufacturer, year, model, serial_no, swg_type, rating)):
                swg_meta = SwitchgearSpec(
                    switchgear_type=swg_type,
                    manufacturer=manufacturer,
                    model=model,
                    manufactured_year=year,
                    rating=rating,
                    serial_no=serial_no,
                )

        return {
            "fl_site": fl_site,
            "fl_erms": fl_erms,
            "substation_name_site": substation_name_site,
            "substation_name_erms": substation_name_erms,
            "date_str": date_str,
            "cycle_1": cycle_1,
            "time_str": time_str,
            "ambient": ambient,
            "humidity": humidity,
            "tev_background": tev_background,
            "swg_meta": swg_meta,
        }

    def _parse_positive_int(self, val: object) -> int | None:
        """Convert a cell value into a positive integer, returning None on invalid/non-positive."""
        if val is None or val == "":
            return None
        val_str = str(val).strip()
        if val_str in ("", "-", "None", "N/A"):
            return None
        try:
            num = int(float(val_str))
            if num > 0:
                return num
        except (ValueError, TypeError):
            pass
        return None

    def _resolve_substation_type(
        self, name: str, filename: str, fl_erms: str = ""
    ) -> str:
        """Extract substation type (PPU, SSU, PMU, MPPU) from station name, filename, or FL token."""
        for candidate in (name, filename, fl_erms):
            if candidate:
                cand_str = str(candidate).strip()
                if re.search(r"\bMINI[\s_-]+PPU\b", cand_str, re.IGNORECASE) or re.search(
                    r"\b(MPPU|MPU)\b", cand_str, re.IGNORECASE
                ):
                    return "MPPU"
                m = re.search(r"\b(PPU|SSU|PMU)\b", cand_str, re.IGNORECASE)
                if m:
                    return m.group(1).upper()
        return ""

    def _find_vi_sheet(self, wb: openpyxl.Workbook) -> Worksheet | None:
        """Locate Visual Inspection worksheet by known name patterns."""
        for sname in wb.sheetnames:
            s_clean = sname.strip().lower()
            if s_clean in ("visual inspec.", "visual inspection", "visual inspec"):
                return wb[sname]
            if s_clean.startswith("visual inspec"):
                return wb[sname]
        return None

    def _extract_gps_coordinate(self, ws_vi: Worksheet | None, substation_name: str = "") -> str:
        """Extract GPS coordinates from Visual Inspection sheet.

        Locates coordinates from cell M3 (or adjacent to "GPS:" label).
        Fallback to cell D3 is retained only when D3 matches coordinate regex
        and does not match the substation name.
        """
        if ws_vi is None:
            return ""

        # 1. Look for cell adjacent to "GPS:" label in header rows (rows 1..5)
        for r in range(1, 6):
            for c in range(1, 25):
                val = clean_val(ws_vi.cell(r, c).value)
                if val and "gps" in val.lower():
                    m = _GPS_COORD_RE.search(val)
                    if m:
                        return m.group(0).strip()
                    adj_val = clean_val(ws_vi.cell(r, c + 1).value)
                    if adj_val:
                        return adj_val

        # 2. Check cell M3 directly
        m3_val = clean_val(ws_vi["M3"].value)
        if m3_val:
            return m3_val

        # 3. Fallback to D3 only when D3 matches coordinate regex pattern
        #    and does not match the substation name
        d3_val = clean_val(ws_vi["D3"].value)
        if d3_val and _GPS_COORD_RE.search(d3_val):
            sub_norm = substation_name.strip().lower()
            d3_norm = d3_val.strip().lower()
            if not sub_norm or (sub_norm != d3_norm and sub_norm not in d3_norm):
                return d3_val

        return ""

    def _find_extra_sheets(self, wb: openpyxl.Workbook, prefix: str) -> list[Worksheet]:
        """Find and sequentially sort extra switchgear sheets (e.g. 33kV (Extra))."""
        matched: list[Worksheet] = []
        for sname in wb.sheetnames:
            s_clean = sname.strip()
            if s_clean.startswith(prefix):
                matched.append(wb[sname])

        def sort_key(ws: Worksheet) -> int:
            m = re.search(r"\((\d+)\)", ws.title)
            if m:
                return int(m.group(1))
            return 1

        return sorted(matched, key=sort_key)

    def _parse_switchgear_lineup(
        self,
        ws_main: Worksheet,
        extra_sheets: list[Worksheet],
    ) -> SwitchgearSpec | None:
        """Parse switchgear lineup across main and extra sheets."""
        # Details from L6..L11, M6
        manufacturer = clean_val(ws_main["L6"].value) or ""
        type_from_m6 = clean_val(ws_main["M6"].value) or ""
        manufactured_year = clean_val(ws_main["L7"].value) or ""
        model = clean_val(ws_main["L8"].value) or ""
        serial_no = clean_val(ws_main["L9"].value) or ""
        busbar_type = clean_val(ws_main["L10"].value) or ""
        rating = clean_val(ws_main["L11"].value) or ""
        switchgear_type = type_from_m6 or busbar_type or "VCB"

        # Panels extraction across main sheet and extra sheets
        panels: list[SwitchgearPanelSpec] = []
        panel_counter = 1

        for ws in [ws_main] + extra_sheets:
            sheet_panels = self._parse_sheet_panels(ws, start_index=panel_counter, switchgear_type=switchgear_type)
            panels.extend(sheet_panels)
            panel_counter += len(sheet_panels)

        # Overview Bays (Rows 63-74 on main sheet)
        # Desired canonical order: LHS pairs before RHS pairs
        # Row mapping:
        # RHS: 63 FRONT, 64 BACK (REAR), 65 TOP, 66 CABLE CELLAR, 67 MBB, 68 RBB
        # LHS: 69 FRONT, 70 BACK (REAR), 71 TOP, 72 CABLE CELLAR, 73 MBB, 74 RBB
        row_map = {
            63: ("OVERVIEW FRONT (RHS)", parse_photo_numbers(ws_main["P63"].value)),
            64: ("OVERVIEW REAR (RHS)", parse_photo_numbers(ws_main["P64"].value)),
            65: ("OVERVIEW TOP (RHS)", parse_photo_numbers(ws_main["P65"].value)),
            66: ("OVERVIEW CABLE CELLAR (RHS)", parse_photo_numbers(ws_main["P66"].value)),
            67: ("OVERVIEW MAIN BUSBAR (RHS)", parse_photo_numbers(ws_main["P67"].value)),
            68: ("OVERVIEW RESERVE BUSBAR (RHS)", parse_photo_numbers(ws_main["P68"].value)),
            69: ("OVERVIEW FRONT (LHS)", parse_photo_numbers(ws_main["P69"].value)),
            70: ("OVERVIEW REAR (LHS)", parse_photo_numbers(ws_main["P70"].value)),
            71: ("OVERVIEW TOP (LHS)", parse_photo_numbers(ws_main["P71"].value)),
            72: ("OVERVIEW CABLE CELLAR (LHS)", parse_photo_numbers(ws_main["P72"].value)),
            73: ("OVERVIEW MAIN BUSBAR (LHS)", parse_photo_numbers(ws_main["P73"].value)),
            74: ("OVERVIEW RESERVE BUSBAR (LHS)", parse_photo_numbers(ws_main["P74"].value)),
        }

        order = [
            (69, 63),  # FRONT
            (70, 64),  # REAR
            (71, 65),  # TOP
            (72, 66),  # CABLE CELLAR (optional)
            (73, 67),  # MAIN BUSBAR
            (74, 68),  # RESERVE BUSBAR (optional)
        ]
        ov_comps: list[str] = []
        ov_photos: list[int] = []
        has_any_overview_photo = any(bool(p) for _, p in row_map.values())
        if has_any_overview_photo:
            for lhs_r, rhs_r in order:
                lhs_name, lhs_p = row_map[lhs_r]
                rhs_name, rhs_p = row_map[rhs_r]
                if lhs_p or rhs_p or "CELLAR" not in lhs_name:
                    if lhs_p or "RESERVE" not in lhs_name:
                        ov_comps.append(lhs_name)
                        ov_photos.append(lhs_p[0] if lhs_p else 0)
                    if rhs_p or "RESERVE" not in rhs_name:
                        ov_comps.append(rhs_name)
                        ov_photos.append(rhs_p[0] if rhs_p else 0)

        overview_rhs = self._extract_overview_photos(ws_main, start_row=63, end_row=68)
        overview_lhs = self._extract_overview_photos(ws_main, start_row=69, end_row=74)
        all_overview_photos: list[int] = []
        for n in ov_photos:
            if n > 0 and n not in all_overview_photos:
                all_overview_photos.append(n)
        for n in overview_rhs + overview_lhs:
            if n not in all_overview_photos:
                all_overview_photos.append(n)

        # If no physical panels found, return None (overview photos alone do not constitute an active lineup)
        if not panels:
            return None

        # Synthetic panels removed: overview compartments are stored cleanly on SwitchgearSpec
        return SwitchgearSpec(
            switchgear_type=switchgear_type,
            manufacturer=manufacturer,
            model=model,
            manufactured_year=manufactured_year,
            rating=rating,
            serial_no=serial_no,
            panels=tuple(panels),
            photo_numbers=tuple(ov_photos if any(p > 0 for p in ov_photos) else all_overview_photos),
            overview_compartments=tuple(ov_comps),
        )

    def _parse_sheet_panels(
        self,
        ws: Worksheet,
        start_index: int = 1,
        switchgear_type: str = "",
    ) -> list[SwitchgearPanelSpec]:
        """Parse 6-row stride switchgear panels from Rows 15 to 62."""
        panels: list[SwitchgearPanelSpec] = []
        panel_idx = start_index

        for r in range(15, 63, 6):
            if r > ws.max_row:
                break

            b_val = ws[f"B{r}"].value
            b_str = str(b_val).strip() if b_val is not None else ""
            c_raw = clean_val(ws[f"C{r}"].value) or ""
            d_raw = clean_val(ws[f"D{r}"].value) or ""

            # Detect layout: live testsheet layout vs testsheet factory layout
            # In live layout: Col B is panel_no (e.g. 1, 2, "1", "2" -> pure digits)
            # In factory layout: Col B is feeder_no (e.g. "3P1", "1K1" -> feeder token)
            b_is_digits = b_str.isdigit()

            if b_is_digits:
                # Live testsheet layout: Col B is panel_no, Col C is feeder_code, Col D is feeder_name
                panel_feeder_no = c_raw
                feeder_name = d_raw
                is_factory_layout = False
            elif b_str:
                # Factory layout: Col B is feeder_no, Col C is name, Col D is status
                panel_feeder_no = b_str
                feeder_name = c_raw
                is_factory_layout = True
            else:
                # Col B is empty
                if d_raw.upper() in ("CLOSE", "OPEN", "TRIP", "SERVICE", "TEST", "RACK IN", "RACK OUT"):
                    panel_feeder_no = ""
                    feeder_name = c_raw
                    is_factory_layout = True
                else:
                    panel_feeder_no = c_raw
                    feeder_name = d_raw
                    is_factory_layout = False

            # Unpopulated template bays (blank panel no and feeder name) are skipped
            if not panel_feeder_no and not feeder_name:
                continue

            if is_factory_layout:
                status = d_raw or clean_val(ws[f"E{r}"].value) or ""
                load_raw = clean_val(ws[f"F{r}"].value) or ""
                cable_type = clean_val(ws[f"G{r}"].value) or clean_val(ws[f"E{r}"].value) or ""
                heater_amp = clean_val(ws[f"H{r}"].value) or ""
                panel_type = clean_val(ws[f"J{r}"].value) or switchgear_type or "VCB"
                busbar_position = ""
                breaker_serial = clean_val(ws[f"V{r}"].value) or clean_val(ws[f"I{r}"].value) or ""
            else:
                # Cable Type: Col E
                cable_type = clean_val(ws[f"E{r}"].value) or ""

                # Breaker Position & Status: Col F & Col G
                breaker_pos = clean_val(ws[f"F{r}"].value) or ""
                status_val = clean_val(ws[f"G{r}"].value) or ""
                if any(t in switchgear_type.upper() for t in ("VCB", "GIS")):
                    if breaker_pos and status_val and breaker_pos.upper() != status_val.upper() and "/" not in status_val:
                        status = f"{breaker_pos}/{status_val}"
                    else:
                        status = status_val or breaker_pos or ""
                else:
                    status = status_val or breaker_pos or ""

                # Busbar Position: Col H
                busbar_position = clean_val(ws[f"H{r}"].value) or ""

                # Load (A): Col I
                load_raw = clean_val(ws[f"I{r}"].value) or ""

                # Heater Amp: Col J across 6-row stride
                heater_amp = ""
                for sub_r in range(r, r + 6):
                    if sub_r <= ws.max_row:
                        j_val = clean_val(ws[f"J{sub_r}"].value)
                        if j_val is not None:
                            j_str = str(j_val).strip()
                            m_num = re.search(r"(\d+(?:\.\d+)?)", j_str)
                            if m_num and float(m_num.group(1)) > 0:
                                heater_amp = j_str
                                break
                            elif not heater_amp and j_str and j_str.upper() not in ("-", "NONE", "N/A"):
                                heater_amp = j_str

                panel_type = switchgear_type or "VCB"
                breaker_serial = clean_val(ws[f"V{r}"].value) or ""

            norm_load = format_load_amp(load_raw)
            load_amp = "" if norm_load == "-" else norm_load

            if breaker_serial and "SN:" in breaker_serial.upper():
                m_sn = re.search(r"SN\s*:\s*([^\n\r]+)", breaker_serial, re.I)
                if m_sn:
                    breaker_serial = m_sn.group(1).strip()
            elif not breaker_serial:
                for sub_r in range(r, r + 6):
                    if sub_r <= ws.max_row:
                        val = clean_val(ws[f"V{sub_r}"].value)
                        if val and "SN:" in val.upper():
                            m_sn = re.search(r"SN\s*:\s*([^\n\r]+)", val, re.I)
                            if m_sn:
                                breaker_serial = m_sn.group(1).strip()
                                break

            # Ultrasound (Col T) and TEV (Col R / Col S) for Breaker Compartment (Row r)
            # Col U: Type of Characteristic (for TEV / US; defaults to "NORMAL" if empty)
            # Col V: Remarks / Other Defects (e.g. SN: breaker serial, GIS gas pressure)
            us_reading = clean_val(ws[f"T{r}"].value) or ""
            tev_reading = clean_val(ws[f"R{r}"].value) or ""
            tev_ppc = clean_val(ws[f"S{r}"].value) or ""

            # Cable Compartment TEV and Ultrasound readings (Row r + 1)
            cable_tev_reading = ""
            cable_tev_ppc = ""
            cable_us_reading = ""
            if r + 1 <= ws.max_row:
                cable_tev_reading = clean_val(ws[f"R{r + 1}"].value) or ""
                cable_tev_ppc = clean_val(ws[f"S{r + 1}"].value) or ""
                cable_us_reading = clean_val(ws[f"T{r + 1}"].value) or ""

            # Characteristic: Column U (empty if no defect -> defaults to "NORMAL")
            u_char = clean_val(ws[f"U{r}"].value) or ""
            if not u_char:
                for sub_r in range(r, r + 6):
                    if sub_r <= ws.max_row:
                        val = clean_val(ws[f"U{sub_r}"].value)
                        if val is not None:
                            u_char = val
                            break

            us_char = u_char or "NORMAL"
            tev_char = u_char or "NORMAL"

            # 6-Row Stride photo parsing (Col P)
            # Row 1 (r): Breaker
            breaker_photos = parse_photo_numbers(ws[f"P{r}"].value)
            breaker_photo = breaker_photos[0] if breaker_photos else None

            # Row 2 (r+1): Cable
            cable_photos = parse_photo_numbers(ws[f"P{r + 1}"].value)
            cable_photo = cable_photos[0] if cable_photos else None

            # Row 3 (r+2): Busbar (Top)
            busbar_photos = parse_photo_numbers(ws[f"P{r + 2}"].value)
            busbar_photo = busbar_photos[0] if busbar_photos else None

            # Row 4 (r+3): PT Compartment
            pt_photos = parse_photo_numbers(ws[f"P{r + 3}"].value)
            pt_photo = pt_photos[0] if pt_photos else None

            # Row 5 (r+4): Cable Entry
            cable_entry_photos = parse_photo_numbers(ws[f"P{r + 4}"].value)
            cable_entry_photo = cable_entry_photos[0] if cable_entry_photos else None

            # Row 6 (r+5): Link Box
            link_box_photos = parse_photo_numbers(ws[f"P{r + 5}"].value)
            link_box_photo = link_box_photos[0] if link_box_photos else None

            # Secondary Compartment: Inspect Column Q across panel rows for regex SP: <photo>
            secondary_photo: int | None = None
            for sub_r in range(r, r + 6):
                if sub_r <= ws.max_row:
                    q_val = str(ws[f"Q{sub_r}"].value or "").strip()
                    if q_val:
                        m_sp = re.search(r"(?:\bSP:|\bSP\s*:)\s*(\d+)", q_val, re.IGNORECASE)
                        if m_sp:
                            secondary_photo = int(m_sp.group(1))
                            break

            # PT measurement detection in Row 4 (r+3)
            has_pt = False
            pt_row = r + 3
            if pt_row <= ws.max_row:
                for col in ("K", "L", "M", "N", "R", "S", "T", "U", "V"):
                    if clean_val(ws[f"{col}{pt_row}"].value) is not None:
                        has_pt = True
                        break

            # Aggregate all distinct photo integers for this panel in order
            discovered: list[int] = []
            for n_list in (
                breaker_photos,
                cable_photos,
                busbar_photos,
                pt_photos,
                cable_entry_photos,
                link_box_photos,
            ):
                for n in n_list:
                    if n not in discovered:
                        discovered.append(n)
            if secondary_photo is not None and secondary_photo not in discovered:
                discovered.append(secondary_photo)

            panel = SwitchgearPanelSpec(
                panel_no=panel_idx,
                panel_feeder_no=panel_feeder_no,
                name=feeder_name,
                panel_type=panel_type,
                serial_no=breaker_serial,
                status=status,
                load_amp=load_amp,
                cable_type=cable_type,
                busbar_position=busbar_position,
                heater_amp=heater_amp,
                us_reading=us_reading,
                us_char=us_char,
                tev_reading=tev_reading,
                tev_ppc=tev_ppc,
                tev_char=tev_char,
                cable_tev_reading=cable_tev_reading,
                cable_tev_ppc=cable_tev_ppc,
                cable_us_reading=cable_us_reading,
                breaker_photo=breaker_photo,
                cable_photo=cable_photo,
                busbar_photo=busbar_photo,
                pt_photo=pt_photo,
                cable_entry_photo=cable_entry_photo,
                link_box_photo=link_box_photo,
                secondary_photo=secondary_photo,
                has_pt_measurement=has_pt,
                photo_numbers=tuple(discovered),
            )
            panels.append(panel)
            panel_idx += 1

        return panels

    def _extract_overview_photos(
        self,
        ws: Worksheet,
        start_row: int,
        end_row: int,
    ) -> tuple[int, ...]:
        """Extract overview photo numbers from Column P within row range."""
        photos: list[int] = []
        for r in range(start_row, end_row + 1):
            if r <= ws.max_row:
                nums = parse_photo_numbers(ws[f"P{r}"].value)
                for n in nums:
                    if n not in photos:
                        photos.append(n)
        return tuple(photos)

    def _parse_visual_inspection_findings(
        self,
        ws: Worksheet,
    ) -> list[VisualInspectionFinding]:
        """Parse Visual Inspec. sheet checklist defects (Rows 8-76) and Category 8 (Rows 78-80)."""
        findings: list[VisualInspectionFinding] = []
        current_category = ""

        # Checklist Rows 8-76
        for r in range(8, 77):
            if r > ws.max_row:
                break

            b_val = clean_val(ws[f"B{r}"].value)
            c_val = clean_val(ws[f"C{r}"].value)

            if not b_val and not c_val:
                continue

            # Section header (e.g. "1. CIVIL & STRUCTURAL")
            if b_val and re.match(r"^\d+\.\s+[A-Za-z]", b_val):
                current_category = b_val
                continue

            # Item number (e.g. "1.1", "2.1")
            is_item = b_val and re.match(r"^\d+\.\d+", b_val)
            item_no = b_val if is_item else ""
            description = c_val or ""

            # Check status
            status = "YES"
            d_val = ws[f"D{r}"].value
            e_val = ws[f"E{r}"].value

            if clean_val(d_val) == "NO" or clean_val(e_val) == "NO":
                status = "NO"
            elif is_marked(e_val) and not is_marked(d_val):
                status = "NO"
            elif is_marked(d_val):
                status = "YES"
            else:
                status = ""

            remarks = clean_val(ws[f"F{r}"].value) or ""
            if not remarks:
                # Check for remarks in neighboring columns
                for col in ("G", "H", "I"):
                    val = clean_val(ws[f"{col}{r}"].value)
                    if val:
                        remarks = val
                        break

            # Record items marked 'NO' or having non-empty remarks
            if status == "NO" or bool(remarks):
                findings.append(
                    VisualInspectionFinding(
                        category=current_category,
                        item_no=item_no,
                        description=description,
                        status=status,
                        remarks=remarks,
                    )
                )

        # Category 8 "Other Findings" (Rows 78-80)
        for r in range(78, 81):
            if r > ws.max_row:
                break
            desc_val = clean_val(ws[f"C{r}"].value) or ""
            remark_val = clean_val(ws[f"F{r}"].value) or ""

            if desc_val.upper() in ("OTHER FINDINGS", "8.", "8", "NO", "YES"):
                desc_val = ""
            if remark_val.upper() in ("OTHER FINDINGS", "8.", "8", "NO", "YES"):
                remark_val = ""

            if not desc_val and not remark_val:
                for col in ("B", "D", "E"):
                    val = clean_val(ws[f"{col}{r}"].value)
                    if val and val.upper() not in ("OTHER FINDINGS", "8.", "8", "NO", "YES") and not re.match(r"^8\.\d+$", val):
                        desc_val = val
                        break

            if desc_val or remark_val:
                item_no = clean_val(ws[f"B{r}"].value) or f"8.{r - 77}"
                findings.append(
                    VisualInspectionFinding(
                        category="Other Findings",
                        item_no=item_no,
                        description=desc_val or remark_val,
                        status="NO",
                        remarks=remark_val or desc_val,
                    )
                )

        return findings

    def _parse_raw_data_sheet(self, ws_raw: Worksheet) -> RawPhotoRanges:
        """Extract IR and DG photo range bounds from RAW DATA sheet."""
        ir_start: int | None = None
        ir_end: int | None = None
        dg_start: int | None = None
        dg_end: int | None = None

        for row in ws_raw.iter_rows(values_only=True):
            row_cells = [str(c).strip() if c is not None else "" for c in row]
            if not row_cells:
                continue

            # 1. Grid table schema (e.g. ['IR', 49, 66] / ['DG', 1715, 1739])
            first_cell = row_cells[0].upper()
            if first_cell in ("IR", "FLIR") and len(row_cells) >= 3:
                if ir_start is None and row_cells[1]:
                    ir_start = parse_int_safe(row_cells[1])
                if ir_end is None and row_cells[2]:
                    ir_end = parse_int_safe(row_cells[2])
            elif first_cell in ("DG", "IMG") and len(row_cells) >= 3:
                if dg_start is None and row_cells[1]:
                    dg_start = parse_int_safe(row_cells[1])
                if dg_end is None and row_cells[2]:
                    dg_end = parse_int_safe(row_cells[2])

            # 2. Label-based schema
            for idx, text in enumerate(row_cells):
                text_upper = text.upper()
                if "IR START" in text_upper or "FLIR START" in text_upper:
                    val = find_next_val(row_cells, idx)
                    parsed = parse_int_safe(val)
                    if parsed is not None:
                        ir_start = parsed
                elif "IR END" in text_upper or "FLIR END" in text_upper:
                    val = find_next_val(row_cells, idx)
                    parsed = parse_int_safe(val)
                    if parsed is not None:
                        ir_end = parsed
                elif "DG START" in text_upper or "IMG START" in text_upper:
                    val = find_next_val(row_cells, idx)
                    parsed = parse_int_safe(val)
                    if parsed is not None:
                        dg_start = parsed
                elif "DG END" in text_upper or "IMG END" in text_upper:
                    val = find_next_val(row_cells, idx)
                    parsed = parse_int_safe(val)
                    if parsed is not None:
                        dg_end = parsed
                elif "IR RANGE" in text_upper or "FLIR RANGE" in text_upper:
                    val = find_next_val(row_cells, idx)
                    s, e = parse_range_val(val)
                    if ir_start is None:
                        ir_start = s
                    if ir_end is None:
                        ir_end = e
                elif "DG RANGE" in text_upper or "IMG RANGE" in text_upper:
                    val = find_next_val(row_cells, idx)
                    s, e = parse_range_val(val)
                    if dg_start is None:
                        dg_start = s
                    if dg_end is None:
                        dg_end = e

        return RawPhotoRanges(
            ir=PhotoRange(start_num=ir_start, end_num=ir_end),
            dg=PhotoRange(start_num=dg_start, end_num=dg_end),
        )

    def _extract_tx_sheet(
        self,
        ws: Worksheet,
    ) -> tuple[list[TransformerSpec], list[LVDBSpec]]:
        """Parse TX sheet into unified PowerTransformerSpec / TransformerSpec and LVDBSpec."""
        slots = [
            {"slot_id": 1, "default_tx_id": "TX 1", "cols": ("G", "H", "I", "J"), "part1_row": 14, "part2_row": 33, "unit_hint": "MVA"},
            {"slot_id": 2, "default_tx_id": "TX 2", "cols": ("K", "L", "M"), "part1_row": 17, "part2_row": 36, "unit_hint": "MVA"},
            {"slot_id": 3, "default_tx_id": "TX 3", "cols": ("N", "O", "P"), "part1_row": 20, "part2_row": 39, "unit_hint": "MVA"},
            {"slot_id": 4, "default_tx_id": "TX 4", "cols": ("Q", "R", "S"), "part1_row": 23, "part2_row": 42, "unit_hint": "MVA"},
            {"slot_id": 5, "default_tx_id": "LOCAL TX", "cols": ("T", "U", "V"), "part1_row": 26, "part2_row": 45, "unit_hint": "KVA"},
        ]

        transformers: list[TransformerSpec] = []
        for slot in slots:
            tx = self._extract_tx_slot(ws, slot)
            if tx is not None:
                transformers.append(tx)

        lvdb_specs = self._extract_station_service_lvdb(ws)
        return transformers, lvdb_specs

    def _extract_tx_slot(
        self,
        ws: Worksheet,
        slot: Mapping[str, Any],
    ) -> TransformerSpec | None:
        """Extract a single transformer slot into PowerTransformerSpec or TransformerSpec."""
        default_id = slot["default_tx_id"]
        slot_id = slot.get("slot_id", 1)
        cols = slot["cols"]
        unit_hint = slot["unit_hint"]

        part1_row, part2_row = self._resolve_tx_slot_rows(
            ws, slot_id, default_id, slot["part1_row"], slot["part2_row"]
        )

        # 1. Header metadata (Rows 6-9)
        mfg = self._get_tx_meta_val(ws, 7, cols, r"(?:2\.1\s*)?Manufacturer\s*[:.]*")
        rating = self._get_tx_meta_val(ws, 8, cols, r"(?:2\.2\s*)?Rating\s*(?:\([^\)]*\))?\s*[:.]*")
        sn = self._get_tx_meta_val(ws, 9, cols, r"(?:2\.3\s*)?Serial\s*No\.?\s*[:.]*")
        model = self._get_tx_model(ws, cols)

        p1_name = clean_val(ws[f"C{part1_row}"].value) if part1_row <= ws.max_row else None
        p2_name = clean_val(ws[f"C{part2_row}"].value) if part2_row <= ws.max_row else None
        header_name = clean_val(ws[f"{cols[0]}6"].value)
        tx_name = p1_name or p2_name or (header_name if header_name and "TX" in header_name.upper() else default_id)

        # 2. Part 1: Primary terminations, load, cable, ultrasound, cable thermals
        load_amp = ""
        cable_type = ""
        hv_ultrasound = ""
        hv_ultrasound_char = ""
        hv_ultrasound_file = ""
        lv_ultrasound = ""
        lv_ultrasound_char = ""
        lv_ultrasound_file = ""
        hv_cable_photo: int | None = None
        lv_cable_photo: int | None = None
        hv_cable_thermal = ThermalReadingSpec()
        lv_cable_thermal = ThermalReadingSpec()
        all_unit_photos: list[int] = []

        if part1_row <= ws.max_row:
            load_amp = clean_val(ws[f"F{part1_row}"].value) or ""
            cable_type = clean_val(ws[f"H{part1_row}"].value) or ""
            hv_ultrasound = clean_val(ws[f"J{part1_row}"].value) or ""
            hv_ultrasound_char = clean_val(ws[f"K{part1_row}"].value) or ""
            hv_ultrasound_file = clean_val(ws[f"M{part1_row}"].value) or ""
            lv_ultrasound = clean_val(ws[f"N{part1_row}"].value) or ""
            lv_ultrasound_char = clean_val(ws[f"O{part1_row}"].value) or ""
            lv_ultrasound_file = clean_val(ws[f"Q{part1_row}"].value) or ""

            hv_cable_thermal = self._extract_th_reading(ws, ("R", "S", "T", "U"), part1_row)
            lv_cable_thermal = self._extract_th_reading(ws, ("W", "X", "Y", "Z"), part1_row)

            for col in ("V", "AA"):
                photos = parse_photo_numbers(ws[f"{col}{part1_row}"].value)
                if photos:
                    if col == "V" and hv_cable_photo is None:
                        hv_cable_photo = photos[0]
                    elif col == "AA" and lv_cable_photo is None:
                        lv_cable_photo = photos[0]
                    for p in photos:
                        if p not in all_unit_photos:
                            all_unit_photos.append(p)

        # 3. Part 2: Auxiliary component thermals
        mk_thermal = ThermalReadingSpec()
        ct_thermal = ThermalReadingSpec()
        body_thermal = ThermalReadingSpec()
        cf_thermal = ThermalReadingSpec()
        remarks = ""
        mk_photo: int | None = None
        ct_photo: int | None = None
        body_photo: int | None = None
        cf_photo: int | None = None

        if part2_row <= ws.max_row:
            mk_thermal = self._extract_th_reading(ws, ("F", "G", "H", "I"), part2_row)
            ct_thermal = self._extract_th_reading(ws, ("L", "M", "N", "O"), part2_row)
            body_thermal = self._extract_th_reading(ws, ("R", "S", "T", "U"), part2_row)
            cf_thermal = self._extract_th_reading(ws, ("X", "Y", "Z", "AA"), part2_row)
            remarks = clean_val(ws[f"AC{part2_row}"].value) or ""

            for col in ("J", "P", "V", "AB"):
                photos = parse_photo_numbers(ws[f"{col}{part2_row}"].value)
                if photos:
                    if col == "J" and mk_photo is None:
                        mk_photo = photos[0]
                    elif col == "P" and ct_photo is None:
                        ct_photo = photos[0]
                    elif col == "V" and body_photo is None:
                        body_photo = photos[0]
                    elif col == "AB" and cf_photo is None:
                        cf_photo = photos[0]
                    for p in photos:
                        if p not in all_unit_photos:
                            all_unit_photos.append(p)

        has_electrical_or_thermal_data = bool(
            all_unit_photos
            or load_amp
            or cable_type
            or hv_ultrasound
            or lv_ultrasound
            or hv_cable_thermal != ThermalReadingSpec()
            or lv_cable_thermal != ThermalReadingSpec()
            or body_thermal != ThermalReadingSpec()
            or mk_thermal != ThermalReadingSpec()
            or ct_thermal != ThermalReadingSpec()
            or cf_thermal != ThermalReadingSpec()
            or remarks
        )
        has_nameplate_data = bool(rating or sn)

        if not (has_electrical_or_thermal_data or has_nameplate_data):
            return None

        is_power_tx = classify_transformer_rating(rating, tx_id=tx_name, unit_hint=unit_hint)

        if is_power_tx:
            return PowerTransformerSpec(
                tx_id=tx_name,
                rating_kva=rating,
                manufacturer=mfg,
                model=model,
                serial_no=sn,
                hv_cable_type=cable_type,
                lv_cable_type=cable_type,
                load_amp=load_amp,
                remarks=remarks,
                hv_ultrasound=hv_ultrasound,
                hv_ultrasound_char=hv_ultrasound_char,
                hv_ultrasound_file=hv_ultrasound_file,
                lv_ultrasound=lv_ultrasound,
                lv_ultrasound_char=lv_ultrasound_char,
                lv_ultrasound_file=lv_ultrasound_file,
                hv_cable_thermal=hv_cable_thermal,
                lv_cable_thermal=lv_cable_thermal,
                marshalling_kiosk_temp=mk_thermal,
                conservator_temp=ct_thermal,
                body_temp=body_thermal,
                cooling_fin_temp=cf_thermal,
                marshalling_kiosk_photo=mk_photo,
                conservator_photo=ct_photo,
                body_photo=body_photo,
                cooling_fin_photo=cf_photo,
                hv_cable_photo=hv_cable_photo,
                lv_cable_photo=lv_cable_photo,
                photo_numbers=tuple(all_unit_photos),
            )
        else:
            return TransformerSpec(
                tx_id=tx_name,
                rating_kva=rating,
                manufacturer=mfg,
                model=model,
                serial_no=sn,
                hv_cable_type=cable_type,
                lv_cable_type=cable_type,
                load_amp=load_amp,
                remarks=remarks,
                us_reading=hv_ultrasound or lv_ultrasound,
                us_char=hv_ultrasound_char or lv_ultrasound_char,
                hv_cable_thermal=hv_cable_thermal,
                lv_cable_thermal=lv_cable_thermal,
                body_thermal=body_thermal,
                photo_numbers=tuple(all_unit_photos),
            )

    def _resolve_tx_slot_rows(
        self,
        ws: Worksheet,
        slot_id: int,
        default_id: str,
        default_p1: int,
        default_p2: int,
    ) -> tuple[int, int]:
        """Dynamically resolve Part 1 and Part 2 row indices for a transformer slot."""
        norm_target = re.sub(r"[\s_]+", "", default_id).upper()
        p1_row = default_p1
        p2_row = default_p2

        # 1. Resolve Part 1 row (search rows 12..30)
        for r in range(12, min(31, ws.max_row + 1)):
            c_val = clean_val(ws[f"C{r}"].value) or ""
            b_val = clean_val(ws[f"B{r}"].value) or ""
            norm_c = re.sub(r"[\s_]+", "", c_val).upper()
            if norm_c == norm_target:
                p1_row = r
                break
            b_num = self._parse_positive_int(b_val)
            if b_num == slot_id and ("TX" in norm_c or "LOCAL" in norm_c):
                p1_row = r
                break

        # 2. Resolve Part 2 row (search rows 30..50)
        for r in range(30, min(51, ws.max_row + 1)):
            c_val = clean_val(ws[f"C{r}"].value) or ""
            b_val = clean_val(ws[f"B{r}"].value) or ""
            norm_c = re.sub(r"[\s_]+", "", c_val).upper()
            if norm_c == norm_target:
                p2_row = r
                break
            b_num = self._parse_positive_int(b_val)
            if b_num == slot_id and ("TX" in norm_c or "LOCAL" in norm_c):
                p2_row = r
                break

        return p1_row, p2_row

    @staticmethod
    def _extract_station_service_lvdb(ws: Worksheet) -> list[LVDBSpec]:
        """Extract station service LVDB dynamically by finding label rows."""
        lvdb_specs: list[LVDBSpec] = []
        if ws.max_row < 40:
            return lvdb_specs

        rating_row = None
        sn_row = None
        model_row = None
        mfg_row = None

        for r in range(48, min(65, ws.max_row + 1)):
            b_val = clean_val(ws[f"B{r}"].value) or ""
            if not b_val:
                continue
            b_lower = b_val.lower()
            if "rating" in b_lower and rating_row is None:
                rating_row = r
            elif "serial" in b_lower and sn_row is None:
                sn_row = r
            elif ("type" in b_lower or "model" in b_lower) and model_row is None:
                model_row = r
            elif "manufacturer" in b_lower and mfg_row is None:
                mfg_row = r

        r_rating = rating_row or 51
        r_sn = sn_row or 52
        r_model = model_row or 53
        r_mfg = mfg_row or 54

        lvdb_rating = clean_val(ws[f"D{r_rating}"].value) or "" if r_rating <= ws.max_row else ""
        lvdb_photos: Sequence[int] = ()
        if r_rating <= ws.max_row:
            for p_col in ("G", "E", "F", "H"):
                photos = parse_photo_numbers(ws[f"{p_col}{r_rating}"].value)
                if photos:
                    lvdb_photos = photos
                    break

        lvdb_sn = clean_val(ws[f"D{r_sn}"].value) or "" if r_sn <= ws.max_row else ""
        lvdb_model = clean_val(ws[f"D{r_model}"].value) or "" if r_model <= ws.max_row else ""
        lvdb_mfg = clean_val(ws[f"D{r_mfg}"].value) or "" if r_mfg <= ws.max_row else ""

        if lvdb_rating or lvdb_photos or lvdb_sn or lvdb_model or lvdb_mfg:
            lvdb_specs.append(
                LVDBSpec(
                    name="LVDB 1",
                    label="LVDB",
                    source="LOCAL TX",
                    manufacturer=lvdb_mfg,
                    model=lvdb_model,
                    serial_no=lvdb_sn,
                    rating=lvdb_rating,
                    photo_numbers=tuple(lvdb_photos),
                )
            )
        return lvdb_specs

    @staticmethod
    def _extract_th_reading(
        ws: Worksheet,
        cols: tuple[str, str, str, str],
        row: int,
    ) -> ThermalReadingSpec:
        """Extract ThermalReadingSpec from 4 temperature column coordinates (ref, max, delta, avg)."""
        ref_col, max_col, del_col, avg_col = cols

        def _clean_temp(val: object) -> str:
            if val is None:
                return ""
            if isinstance(val, float):
                return f"{round(val, 2):g}"
            s = str(val).replace("\t", "").strip()
            return "" if s in ("-", "None", "NONE", "N/A", "#REF!", "nan") else s

        return ThermalReadingSpec(
            tmin=_clean_temp(ws[f"{ref_col}{row}"].value),
            tmax=_clean_temp(ws[f"{max_col}{row}"].value),
            delta_t=_clean_temp(ws[f"{del_col}{row}"].value),
            avg=_clean_temp(ws[f"{avg_col}{row}"].value),
        )

    @staticmethod
    def _get_tx_meta_val(
        ws: Worksheet,
        row_idx: int,
        cols: Sequence[str],
        label_pat: str,
    ) -> str:
        """Extract metadata value within a column block matching a label pattern across candidate rows."""
        is_mfg = bool(re.search(r"manufacturer", label_pat, re.IGNORECASE))
        is_rating = bool(re.search(r"rating", label_pat, re.IGNORECASE))
        is_sn = bool(re.search(r"serial", label_pat, re.IGNORECASE))

        candidate_rows = [row_idx]
        for offset in (1, -1, 2, -2):
            r = row_idx + offset
            if 7 <= r <= 12 and r not in candidate_rows:
                candidate_rows.append(r)
        for r in range(7, 13):
            if r not in candidate_rows:
                candidate_rows.append(r)

        # 1. Search for matching label in candidate rows (Rows 7..12)
        for r in candidate_rows:
            if r > ws.max_row:
                continue
            for idx, col in enumerate(cols):
                val = clean_val(ws[f"{col}{r}"].value)
                if val and re.search(label_pat, val, re.IGNORECASE):
                    sub = re.sub(label_pat, "", val, flags=re.IGNORECASE).strip(" :-\t.")
                    if sub:
                        return sub
                    for next_col in cols[idx + 1:]:
                        n_val = clean_val(ws[f"{next_col}{r}"].value)
                        if n_val and not re.search(r"^\d\.\d", n_val) and not re.search(label_pat, n_val, re.IGNORECASE):
                            return n_val

        # 2. Rating-specific heuristic: if label was omitted/overwritten, check for rating pattern
        if is_rating:
            for r in candidate_rows:
                if r > ws.max_row:
                    continue
                row_text = " ".join(str(ws[f"{c}{r}"].value or "") for c in cols)
                if re.search(r"(?:manufacturer|serial)", row_text, re.IGNORECASE):
                    continue
                for col in reversed(cols):
                    val = clean_val(ws[f"{col}{r}"].value)
                    if val and re.search(r"\d+(?:\.\d+)?\s*(?:MVA|KVA)", val, re.IGNORECASE):
                        return val

        # 3. Fallback on expected row_idx if it has a non-label, non-header value and no conflicting label
        if 7 <= row_idx <= ws.max_row:
            row_text = " ".join(str(ws[f"{c}{row_idx}"].value or "") for c in cols)
            conflict_pat = r"serial" if is_mfg else (r"manufacturer" if is_sn else r"(?:manufacturer|serial)")
            if not re.search(conflict_pat, row_text, re.IGNORECASE):
                for col in reversed(cols):
                    val = clean_val(ws[f"{col}{row_idx}"].value)
                    if (
                        val
                        and not re.search(r"^(?:\d\.\d|model|tx|local)\b", val, re.IGNORECASE)
                        and not re.search(label_pat, val, re.IGNORECASE)
                        and not re.search(r"(?:manufacturer|rating|serial)", val, re.IGNORECASE)
                    ):
                        return val

        return ""

    @staticmethod
    def _get_tx_model(ws: Worksheet, cols: Sequence[str]) -> str:
        """Extract Model specification within a column block across candidate rows."""
        for r in (6, 7):
            if r > ws.max_row:
                continue
            for idx, col in enumerate(cols):
                val = clean_val(ws[f"{col}{r}"].value)
                if val and re.search(r"model", val, re.IGNORECASE):
                    sub = re.sub(r"model\s*[:.]*", "", val, flags=re.IGNORECASE).strip()
                    if sub:
                        return sub
                    for next_col in cols[idx + 1:]:
                        n_val = clean_val(ws[f"{next_col}{r}"].value)
                        if n_val and not re.search(r"^model\s*[:.]*$", n_val, re.IGNORECASE):
                            return n_val
        for col in reversed(cols):
            val = clean_val(ws[f"{col}6"].value)
            if val and not re.search(r"^(?:model|tx|local)\b", val, re.IGNORECASE):
                return val
        return ""
