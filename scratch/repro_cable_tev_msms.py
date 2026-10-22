"""Reproduction script for Ticket #114:
Demonstrating extraction and formatting of cable compartment TEV and pulse readings in workflow 12.
"""

from __future__ import annotations

import csv
from pathlib import Path
import sys
import tempfile

# Ensure workspace root is in sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from src.project.environment import ProjectEnvironment
from src.project.models import ProjectMetadata
from src.project.storage import LocalWorkspaceStorage
from src.workflows.models import PopulateDataMsmsRequest, PopulateMode
from src.workflows.populate_data_msms import PopulateDataMsmsWorkflow


def test_cable_compartment_tev_msms() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        base_path = Path(tmpdir)
        meta = ProjectMetadata(
            key="test_project",
            name="Test Project",
            po_number="PO123",
            state="PAHANG",
            voltage_type="33kV",
            year=2026,
            cycle="Cycle 1",
            technologies=("TEV", "US", "IR"),
            base_path=str(base_path),
        )
        storage = LocalWorkspaceStorage(base_path)
        env = ProjectEnvironment(metadata=meta, storage=storage)

        date_str = "07-10-2026"
        ts_dir = base_path / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / date_str
        ts_dir.mkdir(parents=True)
        (ts_dir / "UNSORTED RAW DATA").mkdir()

        # Create 33kV testsheet workbook with PPU APMM structure
        wb_path = ts_dir / "001. PPU APMM 33-11kV.xlsx"
        wb = openpyxl.Workbook()

        # Sheet 1: Visual Inspec.
        ws_vi = wb.active
        ws_vi.title = "Visual Inspec."
        ws_vi["D3"] = "PPU APMM 33-11kV"
        ws_vi["D4"] = "CGBG/PPUAPMM"
        ws_vi["D5"] = date_str

        # Sheet 2: 33kV (Main)
        ws_33 = wb.create_sheet("33kV (Main)")
        ws_33["C13"] = "1"
        ws_33["D13"] = "INCOMING"
        ws_33["E13"] = "3P1"
        # Row 15: Breaker compartment (FP)
        ws_33["R15"] = "5 dB"
        ws_33["S15"] = 0
        # Row 16: Cable compartment (BP)
        ws_33["R16"] = "7 dB"    # Cable TEV with unit text
        ws_33["S16"] = 0         # Cable Pulse integer

        # Sheet 3: 11kV (Main)
        ws_11 = wb.create_sheet("11kV (Main)")
        ws_11["C13"] = "1"
        ws_11["D13"] = "TX 1"
        ws_11["E13"] = "1P1"
        # Row 16: Cable compartment (BP)
        ws_11["R16"] = "10dB"    # Cable TEV attached unit text
        ws_11["S16"] = 15.0      # Cable Pulse float

        wb.save(wb_path)
        wb.close()

        # Setup TOTAL PE.xlsx
        total_pe_path = env.storage.get_total_pe_path()
        env.storage.ensure_directory(total_pe_path.parent)
        wb_pe = openpyxl.Workbook()
        ws_pe = wb_pe.active
        ws_pe.title = "DataCycle1"
        ws_pe.append(["PE NO", "FL NUMBER", "SUBSTATION NAME", "DATE", "TYPE", "WO", "SCOPE"])
        ws_pe.append([1, "CGBG/PPUAPMM", "PPU APMM 33-11kV", date_str, "PPU", "W1616898", "CBM"])
        wb_pe.save(total_pe_path)
        wb_pe.close()

        # Setup MSMS CSV in TO BE FILLED/
        filled_dir = env.storage.get_msms_to_be_filled_dir()
        completed_dir = env.storage.get_msms_completed_dir()
        env.storage.ensure_directory(filled_dir)
        env.storage.ensure_directory(completed_dir)

        csv_path = filled_dir / "07-10-2026_001.csv"
        csv_rows = [
            "ACTFINISH,ACTSTART,JPNUM,WONUM,TNBLOCATION,METERNAME,TNBNEWREADING,TNBNEWREADINGDATE,TNBCOMMENTS,METERTYPE\n",
            # Cable Compartment 33kV TEV and Pulse
            ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_BP_P31VV,,,,GAUGE\n",
            ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_BP_PUL_P31VV,,,,GAUGE\n",
            # Cable Compartment 11kV TEV and Pulse
            ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,TEV_S11_BP_P31VV,,,,GAUGE\n",
            ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,TEV_S11_BP_PUL_P31VV,,,,GAUGE\n",
            # Breaker Compartment 33kV TEV
            ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_FP_P31VV,,,,GAUGE\n",
        ]
        csv_path.write_text("".join(csv_rows), encoding="utf-8")

        # Execute PopulateDataMsmsWorkflow
        workflow = PopulateDataMsmsWorkflow()
        request = PopulateDataMsmsRequest(
            mode=PopulateMode.AUTO,
            overwrite=True,
            move_to_completed=True,
        )
        result = workflow.execute(env, request)

        assert result.csv_files_processed == 1, f"Expected 1 processed CSV, got {result.csv_files_processed}"
        assert result.rows_populated == 5, f"Expected 5 populated rows, got {result.rows_populated}"

        completed_csv = completed_dir / "07-10-2026_001.csv"
        assert completed_csv.exists(), "Completed CSV was not generated in COMPLETED/"

        with open(completed_csv, mode="r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            records = {r["METERNAME"]: r for r in reader}

        # Assert clean integer readings
        assert records["TEV_S33_BP_P31VV"]["TNBNEWREADING"] == "7", (
            f"Expected '7' for TEV_S33_BP_P31VV, got {records['TEV_S33_BP_P31VV']['TNBNEWREADING']!r}"
        )
        assert records["TEV_S33_BP_PUL_P31VV"]["TNBNEWREADING"] == "0", (
            f"Expected '0' for TEV_S33_BP_PUL_P31VV, got {records['TEV_S33_BP_PUL_P31VV']['TNBNEWREADING']!r}"
        )
        assert records["TEV_S11_BP_P31VV"]["TNBNEWREADING"] == "10", (
            f"Expected '10' for TEV_S11_BP_P31VV, got {records['TEV_S11_BP_P31VV']['TNBNEWREADING']!r}"
        )
        assert records["TEV_S11_BP_PUL_P31VV"]["TNBNEWREADING"] == "15", (
            f"Expected '15' for TEV_S11_BP_PUL_P31VV, got {records['TEV_S11_BP_PUL_P31VV']['TNBNEWREADING']!r}"
        )
        assert records["TEV_S33_FP_P31VV"]["TNBNEWREADING"] == "5", (
            f"Expected '5' for TEV_S33_FP_P31VV, got {records['TEV_S33_FP_P31VV']['TNBNEWREADING']!r}"
        )

        print("All cable and breaker compartment TEV/pulse checks PASSED successfully!")


if __name__ == "__main__":
    test_cable_compartment_tev_msms()
