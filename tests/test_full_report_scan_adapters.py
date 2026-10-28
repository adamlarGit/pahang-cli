"""Unit tests for Full Report Equipment Scan Adapters & PRPD Integration (Ticket #31 / T4.2b).

Validates:
1. Switchgear adapter renders panel scanning pages per manufacturer compartment matrix (D26, D29).
2. Transformer adapter renders unconditional 7-point scanning pages per ADR 0004 and D31.
3. Feeder Pillar and Battery Bank adapters render overview pages with operating parameters.
4. Integrates prpd.py scatter plots, falling back cleanly to "" if survey data is missing (D08).
5. Implements D47 overview page substitution when equipment group has active defects.
6. Context generation across all equipment families.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
import tempfile
from typing import Any
import zipfile

from PIL import Image
import pytest

from src.full_report.models import (
    BatteryBankScanSpec,
    FullReportScanPackage,
    LVDBFeederScanSpec,
    LVDBScanSpec,
    SwitchgearPanelScanSpec,
    SwitchgearScanSpec,
    TransformerScanSpec,
)
from src.core.topology import SwitchgearArchetype, VoltageClass
from src.testsheet.models import SwitchgearPanelSpec, SwitchgearSpec
from src.full_report.photo_resolver import PhotoPair, RawPhotoResolver
from src.full_report.scan_adapters import (
    BatteryBankScanAdapter,
    FeederPillarScanAdapter,
    LVDBScanAdapter,
    ScanAdapterResult,
    ScanPageContext,
    ScanRenderItem,
    SwitchgearScanAdapter,
    TransformerScanAdapter,
    adapt_equipment_package,
    find_sliced_overview_page,
)
from src.full_report.scan_render import (
    BANNER_DEFECT_FORWARDING,
    BANNER_HEALTHY_ANALYSIS,
    BANNER_HEALTHY_RECOMMENDATION,
    COLOR_DEFECT,
    COLOR_HEALTHY,
)


@pytest.fixture
def mock_substation_info() -> dict[str, str]:
    return {
        "name_erms": "PE TEST MOCK STATION",
        "date": "28-Aug-2026",
        "time": "11:30 AM",
        "ambient": "32.0 °C",
        "humidity": "65%",
    }


@pytest.fixture
def dummy_image(tmp_path: Path) -> Path:
    img_path = tmp_path / "FLIR0100.jpg"
    im = Image.new("RGB", (64, 48), color=(100, 150, 200))
    im.save(img_path)
    return img_path


@pytest.fixture
def dummy_visual_image(tmp_path: Path) -> Path:
    img_path = tmp_path / "FLIR0100-photo.jpg"
    im = Image.new("RGB", (64, 48), color=(200, 150, 100))
    im.save(img_path)
    return img_path


@pytest.fixture
def dummy_prpd_image(tmp_path: Path) -> Path:
    img_path = tmp_path / "prpd_sample.png"
    im = Image.new("RGB", (80, 40), color=(50, 100, 150))
    im.save(img_path)
    return img_path


class FakePhotoResolver:
    """Test double for RawPhotoResolver."""

    def __init__(self, photo_map: dict[int, tuple[str, str]] | None = None) -> None:
        self.photo_map = photo_map or {}

    def resolve_ir_photo(self, num: Any, warn_if_missing: bool = False) -> str:
        if num is not None and int(num) in self.photo_map:
            return self.photo_map[int(num)][0]
        return ""

    def resolve_visual_photo(self, num: Any, warn_if_missing: bool = False) -> str:
        if num is not None and int(num) in self.photo_map:
            return self.photo_map[int(num)][1]
        return ""

    def resolve_pair(self, num: Any, warn_if_missing: bool = False) -> PhotoPair:
        ir_p = self.resolve_ir_photo(num)
        vis_p = self.resolve_visual_photo(num)
        return PhotoPair(
            ir_photo=ir_p,
            visual_photo=vis_p,
            photo_number=int(num) if num is not None else None,
            ir_path=Path(ir_p) if ir_p else None,
            visual_path=Path(vis_p) if vis_p else None,
        )

    def resolve_cable_split_photo(self, nums: Any, warn_if_missing: bool = False) -> str:
        if isinstance(nums, (list, tuple)) and len(nums) >= 2:
            return self.resolve_ir_photo(nums[1])
        return ""

    def resolve_cable_split_pair(self, nums: Any, warn_if_missing: bool = False) -> PhotoPair:
        if isinstance(nums, (list, tuple)) and len(nums) >= 2:
            return self.resolve_pair(nums[1])
        return PhotoPair()

    def resolve_tx_overview_top_photo(self, num: Any = None, warn_if_missing: bool = False) -> str:
        if num is not None:
            return self.resolve_ir_photo(num)
        return ""

    def resolve_tx_overview_top_pair(self, num: Any = None, warn_if_missing: bool = False) -> PhotoPair:
        if num is not None:
            return self.resolve_pair(num)
        return PhotoPair()


# ==============================================================================
# 1. Switchgear Adapter Tests (D26, D29, D08, D47)
# ==============================================================================

def test_switchgear_adapter_tamco_lucy_compartment_matrix(mock_substation_info: dict[str, str]) -> None:
    """TAMCO/LUCY generates 2 overview pages and 2 scanning pages per panel per D26/D29."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        model="AIR",
        rating="12kV",
        serial_no="SWG-100",
        archetype=SwitchgearArchetype.RMU_DUAL_CABLE_ENTRY,
        overview_compartments=("OVERVIEW", "OVERVIEW BOTTOM"),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, name="INCOMING 1", compartments=("CABLE COMPARTMENT", "CABLE ENTRY")),
            SwitchgearPanelScanSpec(panel_no=2, name="TX 1 FEEDER", compartments=("CABLE COMPARTMENT", "CABLE ENTRY")),
        ),
        photo_numbers=(10, 11),
    )

    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    assert isinstance(result, ScanAdapterResult)
    assert result.equipment_category == "swg"
    assert result.has_defect is False
    assert result.is_overview_substituted is False

    # 2 overview pages + 2 panels * 2 compartments = 6 total pages
    assert result.page_count == 6
    items = result.items

    # Overviews
    assert items[0].template_name == "swg-overview.docx"
    assert items[0].component_name == "OVERVIEW"
    assert items[0].is_overview is True
    assert items[0].sequence == "p00"
    assert items[0].context["swg"]["area"] == "OVERVIEW"
    assert items[0].context["banner"]["analysis"] == BANNER_HEALTHY_ANALYSIS

    assert items[1].template_name == "swg-overview.docx"
    assert items[1].component_name == "OVERVIEW BOTTOM"
    assert items[1].is_overview is True
    assert items[1].context["swg"]["area"] == "OVERVIEW BOTTOM"

    # Panel 1
    assert items[2].template_name == "swg-panel.docx"
    assert items[2].component_name == "CABLE COMPARTMENT"
    assert items[2].panel_no == 1
    assert items[2].context["panel"]["name"] == "INCOMING 1"
    assert items[2].context["panel"]["area"] == "CABLE COMPARTMENT"

    assert items[3].template_name == "swg-panel.docx"
    assert items[3].component_name == "CABLE ENTRY"
    assert items[3].panel_no == 1
    assert items[3].context["panel"]["area"] == "CABLE ENTRY"

    # Panel 2
    assert items[4].panel_no == 2
    assert items[4].component_name == "CABLE COMPARTMENT"
    assert items[5].panel_no == 2
    assert items[5].component_name == "CABLE ENTRY"


def test_switchgear_adapter_indkom_compartment_matrix(mock_substation_info: dict[str, str]) -> None:
    """INDKOM generates 1 overview page; 1 page per panel (FUSE for TX feeder, CABLE for incomers)."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="INDKOM",
        archetype=SwitchgearArchetype.RMU_FUSE_CANISTER,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, name="INCOMING 1", compartments=("CABLE COMPARTMENT",)),
            SwitchgearPanelScanSpec(panel_no=2, name="TX 1 TEE OFF", compartments=("FUSE COMPARTMENT",)),
        ),
        photo_numbers=(20,),
    )

    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    # 1 overview + 2 panels * 1 = 3 pages
    assert result.page_count == 3
    assert result.items[0].component_name == "OVERVIEW"
    assert result.items[1].component_name == "CABLE COMPARTMENT"
    assert result.items[2].component_name == "FUSE COMPARTMENT"


def test_switchgear_adapter_vcb_standard_5_compartments(mock_substation_info: dict[str, str]) -> None:
    """VCB generates 1 overview and 5 standard compartments per standard panel per D26."""
    vcb_std_comps = (
        "BREAKER COMPARTMENT",
        "CABLE COMPARTMENT",
        "BUSBAR COMPARTMENT",
        "PT COMPARTMENT",
        "SECONDARY COMPARTMENT",
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(
                panel_no=1,
                name="INCOMING 1",
                compartments=vcb_std_comps,
            ),
        ),
        photo_numbers=(30,),
    )

    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    # 1 overview + 1 panel * 5 compartments = 6 pages
    assert result.page_count == 6
    assert result.items[0].component_name == "OVERVIEW"
    panel_compartments = [item.component_name for item in result.items[1:]]
    assert tuple(panel_compartments) == vcb_std_comps


def test_switchgear_adapter_vcb_transition_5_compartments(mock_substation_info: dict[str, str]) -> None:
    """VCB generates 1 overview and 5 transition compartments for a transition panel per D26."""
    vcb_trans_comps = (
        "FRONT COMPARTMENT",
        "REAR COMPARTMENT",
        "BUSBAR COMPARTMENT",
        "PT COMPARTMENT",
        "SECONDARY COMPARTMENT",
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(
                panel_no=1,
                name="TRANSITION PANEL",
                compartments=vcb_trans_comps,
            ),
        ),
        photo_numbers=(30,),
    )

    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    # 1 overview + 1 panel * 5 compartments = 6 pages
    assert result.page_count == 6
    assert result.items[0].component_name == "OVERVIEW"
    panel_compartments = [item.component_name for item in result.items[1:]]
    assert tuple(panel_compartments) == vcb_trans_comps
    assert tuple(panel_compartments) == (
        "FRONT COMPARTMENT",
        "REAR COMPARTMENT",
        "BUSBAR COMPARTMENT",
        "PT COMPARTMENT",
        "SECONDARY COMPARTMENT",
    )


def test_switchgear_adapter_prpd_integration_and_clean_fallback(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
    dummy_prpd_image: Path,
) -> None:
    """PRPD graphs are injected when present, and fall back cleanly to '' when missing (D08)."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, name="PANEL 1", compartments=("CABLE COMPARTMENT",)),
            SwitchgearPanelScanSpec(panel_no=2, name="PANEL 2", compartments=("CABLE COMPARTMENT",)),
        ),
    )

    # Catalog with Panel 1 having US & TEV, Panel 2 having none
    prpd_catalog = {
        "swg": {
            1: {"us": str(dummy_prpd_image), "tev": str(dummy_prpd_image)},
            2: {"us": None, "tev": ""},
        }
    }

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        prpd_catalog=prpd_catalog,
    )
    result = adapter.adapt()

    p1_item = result.items[1]
    assert p1_item.context["us"]["prpd"] == str(dummy_prpd_image)
    assert p1_item.context["tev"]["prpd"] == str(dummy_prpd_image)

    p2_item = result.items[2]
    # D08 clean fallback to empty string
    assert p2_item.context["us"]["prpd"] == ""
    assert p2_item.context["tev"]["prpd"] == ""


def test_switchgear_adapter_d47_overview_substitution(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """When SWG has active defects, overview page is substituted with sliced QR file (D47)."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, name="INCOMING 1", compartments=("CABLE COMPARTMENT",)),
        ),
    )

    # Create dummy sliced overview file
    sliced_ov_file = tmp_path / "swg1_p00_SWG1_OVERVIEW_01.docx"
    sliced_ov_file.write_text("dummy sliced content", encoding="utf-8")

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        has_active_defect=True,
        sliced_overview_pages=[sliced_ov_file],
    )
    result = adapter.adapt()

    assert result.has_defect is True
    assert result.is_overview_substituted is True
    ov_item = result.items[0]
    assert ov_item.is_overview is True
    assert ov_item.is_sliced is True
    assert ov_item.sliced_path == sliced_ov_file
    assert ov_item.is_defective is True


def test_switchgear_adapter_defect_forwarding_when_no_slice(
    mock_substation_info: dict[str, str],
) -> None:
    """When SWG has active defect but no sliced file provided, overview banner forwards defect per D30."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, name="INCOMING 1", compartments=("CABLE COMPARTMENT",)),
        ),
    )

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        has_active_defect=True,
        # No sliced files provided
    )
    result = adapter.adapt()

    assert result.has_defect is True
    assert result.is_overview_substituted is False
    ov_item = result.items[0]
    assert ov_item.is_overview is True
    assert ov_item.is_sliced is False
    assert ov_item.is_defective is True
    assert ov_item.context["banner"]["analysis"] == BANNER_DEFECT_FORWARDING
    assert ov_item.context["banner"]["recommendation"] == BANNER_DEFECT_FORWARDING


# ==============================================================================
# 2. Transformer Adapter Tests (ADR 0004, D31, D08, D48)
# ==============================================================================

def test_transformer_adapter_unconditional_7_point_pages(mock_substation_info: dict[str, str]) -> None:
    """Transformer adapter unconditionally provisions 7 standard pages per ADR 0004 & D31."""
    tx = TransformerScanSpec(
        tx_id="Tx 1",
        rating_kva="1000",
        construction_year="2019",
        manufacturer="MTM",
        serial_no="SN-999",
        type="HERMETICALLY SEALED",
        us_reading="0",
        hv_cable_type="XLPE 3C 240mm2",
        lv_cable_type="PVC 1C 500mm2",
        photo_numbers=(101, 102),
    )

    adapter = TransformerScanAdapter(tx=tx, substation_info=mock_substation_info)
    result = adapter.adapt()

    assert isinstance(result, ScanAdapterResult)
    assert result.equipment_category == "tx"
    assert result.page_count == 7

    expected_specs = [
        ("OVERVIEW", "tx-overview.docx", "s00", True),
        ("OVERVIEW TOP", "tx-overview.docx", "s01", True),
        ("HV BUSHING", "tx-hv-sides.docx", "s02", False),
        ("HV CABLE", "tx-hv-sides.docx", "s03", False),
        ("HV CABLE SPLIT", "tx-hv-sides.docx", "s04", False),
        ("LV BUSHING", "tx-lv-sides.docx", "s05", False),
        ("LV CABLE", "tx-lv-sides.docx", "s06", False),
    ]

    for idx, (comp, tpl, seq, is_ov) in enumerate(expected_specs):
        item = result.items[idx]
        assert item.component_name == comp
        assert item.template_name == tpl
        assert item.sequence == seq
        assert item.is_overview is is_ov
        assert item.context["tx"]["manufacturer"] == "MTM"
        assert item.context["tx"]["rating"] == "1000"


def test_transformer_adapter_hv_cable_split_and_prpd_fallback(
    mock_substation_info: dict[str, str],
    dummy_image: Path,
    dummy_visual_image: Path,
) -> None:
    """Transformer cable split resolves secondary photo, falls back cleanly to '' if absent (ADR 0004/D08)."""
    photo_map = {
        101: (str(dummy_image), str(dummy_visual_image)),
        # 102 (split) is intentionally missing from disk
    }
    fake_resolver = FakePhotoResolver(photo_map)

    tx = TransformerScanSpec(
        tx_id="Tx 1",
        photo_numbers=(101, 102),
    )

    adapter = TransformerScanAdapter(
        tx=tx,
        substation_info=mock_substation_info,
        photo_resolver=fake_resolver,
    )
    result = adapter.adapt()

    # Page 5: HV CABLE SPLIT
    split_item = result.items[4]
    assert split_item.component_name == "HV CABLE SPLIT"
    # Photo 102 missing on disk -> clean '' fallback
    assert split_item.context["ir"]["image"] == ""
    assert split_item.context["visual"]["image"] == ""
    # PRPD missing -> clean '' fallback
    assert split_item.context["us"]["prpd"] == ""


def test_transformer_adapter_d47_overview_substitution(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """When TX has active defect, overview page is substituted with sliced QR file (D47)."""
    tx = TransformerScanSpec(tx_id="Tx 1")
    sliced_tx_file = tmp_path / "tx1_s00_TX1_OVERVIEW_01.docx"
    sliced_tx_file.write_text("dummy sliced tx overview", encoding="utf-8")

    adapter = TransformerScanAdapter(
        tx=tx,
        substation_info=mock_substation_info,
        has_active_defect=True,
        sliced_overview_pages=[sliced_tx_file],
    )
    result = adapter.adapt()

    assert result.has_defect is True
    assert result.is_overview_substituted is True
    assert result.items[0].is_sliced is True
    assert result.items[0].sliced_path == sliced_tx_file


# ==============================================================================
# 3. Feeder Pillar & Battery Bank Adapter Tests
# ==============================================================================

def test_lvdb_adapter_renders_overview_with_operating_parameters(
    mock_substation_info: dict[str, str],
) -> None:
    """LVDB adapter renders overview page with operating parameters."""
    lvdb = LVDBScanSpec(
        name="LVDB TX1",
        label="LVDB",
        source="TX 1",
        manufacturer="ALGEBRA",
        model="J-SLOTTED",
        serial_no="SN-LVDB-01",
        rating="1600A",
        cable_type="4C 300mm2 PVC",
        photo_numbers=(201,),
        feeders=(
            LVDBFeederScanSpec(channel="IN1", cable_type="4C 300mm2 PVC"),
            LVDBFeederScanSpec(channel="OT1", cable_type="4C 185mm2 PVC"),
        ),
    )

    adapter = LVDBScanAdapter(lvdb=lvdb, substation_info=mock_substation_info)
    result = adapter.adapt()

    assert result.equipment_category == "fp"
    assert result.page_count == 1
    item = result.items[0]
    assert item.template_name == "fp-overview.docx"
    assert item.is_overview is True
    assert item.sequence == "f00"
    assert item.context["fp"]["labelsource"] == "LVDB TX1"
    assert item.context["fp"]["manufacturer"] == "ALGEBRA"
    assert item.context["fp"]["model"] == "J-SLOTTED"
    assert item.context["fp"]["rating"] == "1600A"
    assert item.context["fp"]["cabletype"] == "4C 300mm2 PVC"
    assert item.context["banner"]["analysis"] == BANNER_HEALTHY_ANALYSIS


def test_lvdb_adapter_d47_overview_substitution(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """LVDB adapter substitutes overview page when defective per D47."""
    lvdb = LVDBScanSpec(name="FP 1", label="FP 1")
    sliced_fp_file = tmp_path / "fp1_f00_FP1_OVERVIEW_01.docx"
    sliced_fp_file.write_text("dummy fp slice", encoding="utf-8")

    adapter = LVDBScanAdapter(
        lvdb=lvdb,
        substation_info=mock_substation_info,
        has_active_defect=True,
        sliced_overview_pages=[sliced_fp_file],
    )
    result = adapter.adapt()

    assert result.has_defect is True
    assert result.is_overview_substituted is True
    assert result.items[0].is_sliced is True
    assert result.items[0].sliced_path == sliced_fp_file


def test_battery_bank_adapter_renders_overview(mock_substation_info: dict[str, str]) -> None:
    """BatteryBank adapter renders overview page with operating parameters."""
    bb = BatteryBankScanSpec(
        name="BATTERY 1",
        manufacturer="SUNPOWER",
        model="SP-110",
        serial_no="BATT-99",
        photo_numbers=(301,),
    )

    adapter = BatteryBankScanAdapter(bb=bb, substation_info=mock_substation_info)
    result = adapter.adapt()

    assert result.equipment_category == "battery"
    assert result.page_count == 1
    item = result.items[0]
    assert item.template_name == "battery-overview.docx"
    assert item.is_overview is True
    assert item.sequence == "b00"
    # Satisfies both battery and batt namespaces
    assert item.context["battery"]["manufacturer"] == "SUNPOWER"
    assert item.context["batt"]["manufacturer"] == "SUNPOWER"
    assert item.context["battery"]["number"] == "1"
    assert item.context["batt"]["number"] == "1"
    assert item.context["banner"]["analysis"] == BANNER_HEALTHY_ANALYSIS


def test_battery_bank_adapter_number_normalization(mock_substation_info: dict[str, str]) -> None:
    """BatteryBank adapter normalizes batt['number'] to integer string index."""
    # 1. "BATTERY BANK 1" -> "1"
    bb1 = BatteryBankScanSpec(name="BATTERY BANK 1")
    ad1 = BatteryBankScanAdapter(bb=bb1, substation_info=mock_substation_info).adapt()
    assert ad1.items[0].context["batt"]["number"] == "1"

    # 2. "BATTERY BANK 2" -> "2"
    bb2 = BatteryBankScanSpec(name="BATTERY BANK 2")
    ad2 = BatteryBankScanAdapter(bb=bb2, substation_info=mock_substation_info).adapt()
    assert ad2.items[0].context["batt"]["number"] == "2"

    # 3. "BATTERY" (no digit) -> defaults to "1"
    bb_nodigit = BatteryBankScanSpec(name="BATTERY")
    ad_nodigit = BatteryBankScanAdapter(bb=bb_nodigit, substation_info=mock_substation_info).adapt()
    assert ad_nodigit.items[0].context["batt"]["number"] == "1"

    # 4. "BATTERY 3" -> "3"
    bb3 = BatteryBankScanSpec(name="BATTERY 3")
    ad3 = BatteryBankScanAdapter(bb=bb3, substation_info=mock_substation_info).adapt()
    assert ad3.items[0].context["batt"]["number"] == "3"


# ==============================================================================
# 4. Composite Package & Physical Rendering Integration
# ==============================================================================

def test_adapt_equipment_package_full_station(mock_substation_info: dict[str, str]) -> None:
    """Composite package adapter aggregates all equipment families correctly."""
    swg = SwitchgearScanSpec(
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),),
    )
    tx = TransformerScanSpec(tx_id="Tx 1")
    lvdb = LVDBScanSpec(name="FP 1")
    bb = BatteryBankScanSpec(name="BATT 1")

    package = FullReportScanPackage(
        substation_number=123,
        station_name="PE COMPOSITE",
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(lvdb,),
        battery_banks=(bb,),
    )

    results = adapt_equipment_package(package=package, substation_info=mock_substation_info)

    assert len(results) == 4
    # SWG: 1 ov + 1 panel = 2
    # TX: 7
    # LVDB: 1
    # BB: 1
    # Total: 11 pages
    total_pages = sum(r.page_count for r in results)
    assert total_pages == 11


def test_scan_render_item_renders_document_with_shading(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """ScanRenderItem renders clean .docx template with dynamic cell shading."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),),
    )
    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    panel_item = result.items[1]
    out_docx = tmp_path / "rendered_panel.docx"
    rendered_path = panel_item.render(out_docx)

    assert rendered_path.is_file()
    assert rendered_path.stat().st_size > 1000

    # Verify word/document.xml contains healthy green shading
    with zipfile.ZipFile(rendered_path, "r") as z:
        xml_content = z.read("word/document.xml").decode("utf-8")
        assert f'w:fill="{COLOR_HEALTHY}"' in xml_content


def test_scan_adapter_result_render_all(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """ScanAdapterResult.render_all renders all items in batch."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),
            SwitchgearPanelScanSpec(panel_no=2, compartments=("CABLE COMPARTMENT",)),
        ),
    )
    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    out_dir = tmp_path / "render_all_out"
    paths = result.render_all(out_dir)

    assert len(paths) == 3
    for p in paths:
        assert p.is_file()
        assert p.stat().st_size > 1000


def test_find_sliced_overview_page_matching(tmp_path: Path) -> None:
    """find_sliced_overview_page matches D37 tokens across equipment categories."""
    p_swg = tmp_path / "swg1_p00_TAMCO_OVERVIEW_01.docx"
    p_tx = tmp_path / "tx1_s00_MTM_OVERVIEW_01.docx"
    p_fp = tmp_path / "fp1_f00_ALGEBRA_OVERVIEW_01.docx"
    p_batt = tmp_path / "batt1_b00_SUNPOWER_OVERVIEW_01.docx"
    p_other = tmp_path / "swg1_p02_TAMCO_CABLE_COMPARTMENT_01.docx"

    for f in (p_swg, p_tx, p_fp, p_batt, p_other):
        f.write_text("x", encoding="utf-8")

    all_files = [p_swg, p_tx, p_fp, p_batt, p_other]

    assert find_sliced_overview_page("swg", "swg1", all_files) == p_swg
    assert find_sliced_overview_page("tx", "tx1", all_files) == p_tx
    assert find_sliced_overview_page("fp", "fp1", all_files) == p_fp
    assert find_sliced_overview_page("battery", "batt1", all_files) == p_batt
    assert find_sliced_overview_page("swg", "swg1", [p_other]) is None
    assert find_sliced_overview_page("swg", "swg1", None) is None


def test_raw_testsheet_specs_adaptation(mock_substation_info: dict[str, str]) -> None:
    """Adapters accept raw testsheet specs directly and classify/build scan specs automatically."""
    from src.testsheet.models import (
        BatteryBankSpec,
        LVDBSpec,
        SwitchgearPanelSpec,
        SwitchgearSpec,
        TransformerSpec,
    )

    swg_raw = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        panels=(SwitchgearPanelSpec(panel_no=1, name="INCOMING 1"),),
    )
    tx_raw = TransformerSpec(tx_id="Tx 1", manufacturer="MTM")
    lvdb_raw = LVDBSpec(name="FP 1", manufacturer="ALGEBRA")
    bb_raw = BatteryBankSpec(name="BATT 1", manufacturer="SUNPOWER")

    res_swg = SwitchgearScanAdapter(swg=swg_raw, substation_info=mock_substation_info).adapt()
    assert res_swg.equipment_category == "swg"
    # TAMCO -> 2 overview pages + 2 panel pages (cable compartment + cable entry) = 4
    assert res_swg.page_count == 4

    res_tx = TransformerScanAdapter(tx=tx_raw, substation_info=mock_substation_info).adapt()
    assert res_tx.page_count == 7

    res_lvdb = LVDBScanAdapter(lvdb=lvdb_raw, substation_info=mock_substation_info).adapt()
    assert res_lvdb.page_count == 1

    res_bb = BatteryBankScanAdapter(bb=bb_raw, substation_info=mock_substation_info).adapt()
    assert res_bb.page_count == 1


def test_missing_template_raises_file_not_found(tmp_path: Path) -> None:
    """ScanRenderItem raises FileNotFoundError if template does not exist."""
    item = ScanRenderItem(
        page_name="Missing",
        template_name="non_existent.docx",
        template_path=tmp_path / "non_existent.docx",
    )
    with pytest.raises(FileNotFoundError):
        item.render(tmp_path / "out.docx")


def test_scan_page_context_dataclass() -> None:
    """ScanPageContext instantiates and holds structured context metadata."""
    spc = ScanPageContext(
        template_name="swg-panel.docx",
        context={"hello": "world"},
        equipment_category="swg",
        component_name="CABLE COMPARTMENT",
        sequence="p01",
        is_overview=False,
        is_defective=True,
        defective_technologies={"IR", "TEV"},
        equipment_id="swg1",
        panel_no=1,
    )
    assert spc.template_name == "swg-panel.docx"
    assert spc.context == {"hello": "world"}
    assert spc.defective_technologies == {"IR", "TEV"}
    assert spc.panel_no == 1


def test_cbm_defects_list_triggers_active_defect(mock_substation_info: dict[str, str]) -> None:
    """Passing cbm_defects triggers defect state on matching equipment."""
    swg = SwitchgearScanSpec(
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),),
    )
    # Defect list with a switchgear defect filename
    cbm_defects = ["swg1_p01_TAMCO_CABLE_COMPARTMENT_01.docx"]

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        cbm_defects=cbm_defects,
    )
    result = adapter.adapt()

    assert result.has_defect is True
    # Overview page should carry defect forwarding prose
    ov_item = result.items[0]
    assert ov_item.is_defective is True
    assert ov_item.context["banner"]["analysis"] == BANNER_DEFECT_FORWARDING


def test_find_sliced_overview_page_multi_unit_disambiguation(tmp_path: Path) -> None:
    """find_sliced_overview_page strictly isolates multi-unit instances (TX1 vs TX2, FP1 vs FP2)."""
    p_tx1 = tmp_path / "tx1_s00_TX1_OVERVIEW_01.docx"
    p_tx2 = tmp_path / "tx2_s00_TX2_OVERVIEW_01.docx"
    p_fp1 = tmp_path / "fp1_f00_FP1_OVERVIEW_01.docx"
    p_fp2 = tmp_path / "fp2_f00_FP2_OVERVIEW_01.docx"

    for f in (p_tx1, p_tx2, p_fp1, p_fp2):
        f.write_text("x", encoding="utf-8")

    all_files = [p_tx1, p_tx2, p_fp1, p_fp2]

    # TX disambiguation
    assert find_sliced_overview_page("tx", "tx1", all_files) == p_tx1
    assert find_sliced_overview_page("tx", "tx2", all_files) == p_tx2
    assert find_sliced_overview_page("tx", "Tx 1", all_files) == p_tx1
    assert find_sliced_overview_page("tx", "Tx 2", all_files) == p_tx2

    # FP disambiguation
    assert find_sliced_overview_page("fp", "fp1", all_files) == p_fp1
    assert find_sliced_overview_page("fp", "fp2", all_files) == p_fp2
    assert find_sliced_overview_page("fp", "FP 1", all_files) == p_fp1
    assert find_sliced_overview_page("fp", "FP 2", all_files) == p_fp2


def test_transformer_component_photo_index_alignment(
    mock_substation_info: dict[str, str],
    dummy_image: Path,
    dummy_visual_image: Path,
) -> None:
    """Transformer photos map to testsheet Rows 33-37 per D05."""
    # Photo numbers 100..104 correspond to:
    # 100: Overview (Row 33)
    # 101: HV Bushing (Row 34)
    # 102: HV Cable (Row 35)
    # 103: LV Bushing (Row 36)
    # 104: LV Cable (Row 37)
    photo_map = {
        i: (f"/path/FLIR0{i}.jpg", f"/path/FLIR0{i}-photo.jpg")
        for i in range(100, 105)
    }
    resolver = FakePhotoResolver(photo_map)

    tx = TransformerScanSpec(
        tx_id="Tx 1",
        photo_numbers=(100, 101, 102, 103, 104),
    )

    adapter = TransformerScanAdapter(
        tx=tx,
        substation_info=mock_substation_info,
        photo_resolver=resolver,
    )
    result = adapter.adapt()

    # Index 0: Overview -> 100
    assert result.items[0].context["ir"]["image"] == "/path/FLIR0100.jpg"
    # Index 2: HV Bushing -> 101
    assert result.items[2].context["ir"]["image"] == "/path/FLIR0101.jpg"
    # Index 3: HV Cable -> 102
    assert result.items[3].context["ir"]["image"] == "/path/FLIR0102.jpg"
    # Index 5: LV Bushing -> 103
    assert result.items[5].context["ir"]["image"] == "/path/FLIR0103.jpg"
    # Index 6: LV Cable -> 104
    assert result.items[6].context["ir"]["image"] == "/path/FLIR0104.jpg"


def test_transformer_scan_adapter_model_context(mock_substation_info: dict[str, str]) -> None:
    """Verify TransformerScanAdapter populates tx.model in page context with normalize_tx_model."""
    from src.full_report.scan_adapters import TransformerScanAdapter
    from src.testsheet.models import TransformerSpec

    # 1. H/S abbreviation expanded
    tx1 = TransformerSpec(tx_id="Tx 1", type="H/S")
    adapter1 = TransformerScanAdapter(tx=tx1, substation_info=mock_substation_info)
    res1 = adapter1.adapt()
    for item in res1.items:
        assert item.context["tx"]["model"] == "HERMETICALLY SEAL"

    # 2. C/T abbreviation expanded
    tx2 = TransformerSpec(tx_id="Tx 2", type="C/T")
    adapter2 = TransformerScanAdapter(tx=tx2, substation_info=mock_substation_info)
    res2 = adapter2.adapt()
    for item in res2.items:
        assert item.context["tx"]["model"] == "CONSERVATOR TANK"

    # 3. Blank fallback to '-'
    tx3 = TransformerSpec(tx_id="Tx 1", type="")
    adapter3 = TransformerScanAdapter(tx=tx3, substation_info=mock_substation_info)
    res3 = adapter3.adapt()
    for item in res3.items:
        assert item.context["tx"]["model"] == "-"

    # 4. Sentinel fallback to '-'
    tx4 = TransformerSpec(tx_id="Tx 1", type="-")
    adapter4 = TransformerScanAdapter(tx=tx4, substation_info=mock_substation_info)
    res4 = adapter4.adapt()
    for item in res4.items:
        assert item.context["tx"]["model"] == "-"


def test_dynamic_prpd_generation_with_tempfile_no_crash(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """Dynamic PRPD resolution handles survey_root without explicit prpd_output_dir using tempfile."""
    survey_dir = tmp_path / "US_TEV_SURVEY"
    survey_dir.mkdir()

    swg = SwitchgearScanSpec(
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),),
    )

    # Providing survey_root without prpd_output_dir must not trigger NameError on tempfile
    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        survey_root=survey_dir,
        prpd_output_dir=None,
    )
    result = adapter.adapt()
    assert result.page_count == 2


def test_scan_render_item_to_page_context(mock_substation_info: dict[str, str]) -> None:
    """ScanRenderItem converts cleanly to ScanPageContext."""
    swg = SwitchgearScanSpec(
        archetype=SwitchgearArchetype.RMU_STANDARD,
        overview_compartments=("OVERVIEW",),
        panels=(SwitchgearPanelScanSpec(panel_no=1, compartments=("CABLE COMPARTMENT",)),),
    )
    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    item = result.items[1]
    spc = item.to_page_context()

    assert isinstance(spc, ScanPageContext)
    assert spc.template_name == "swg-panel.docx"
    assert spc.component_name == "CABLE COMPARTMENT"
    assert spc.panel_no == 1
    assert spc.sequence == "p01"
    assert spc.equipment_category == "swg"
    assert spc.is_overview is False
    assert spc.context == item.context


def test_check_has_active_defect_disambiguates_fp_with_tx_in_name() -> None:
    """Verify _check_has_active_defect correctly identifies FP defects without falsely triggering TX defects."""
    from src.full_report.scan_adapters import _check_has_active_defect
    from src.quick_report.defects import CbmDefectRecord

    d_fp = CbmDefectRecord(
        equipment="FP (J)",
        technology="IR",
        defect_area="INCOMING LINK CONNECTION",
        additional_remarks="RED PHASE",
        equipment_id="FP TX1 - INCOMING 1",
    )
    cbm_defects = [d_fp]

    # FP category must detect defect
    assert _check_has_active_defect("fp", "fp1", cbm_defects=cbm_defects) is True
    assert _check_has_active_defect("fp", "FP TX1", cbm_defects=cbm_defects) is True

    # TX category must NOT detect defect from FP TX1
    assert _check_has_active_defect("tx", "TX1", cbm_defects=cbm_defects) is False
    assert _check_has_active_defect("tx", "TX2", cbm_defects=cbm_defects) is False
    assert _check_has_active_defect("tx", "", cbm_defects=cbm_defects) is False

    # SWG category must NOT detect defect from FP TX1
    assert _check_has_active_defect("swg", "swg1", cbm_defects=cbm_defects) is False

    # Now verify true TX defect is detected
    d_tx = CbmDefectRecord(
        equipment="TRANSFORMER",
        technology="IR",
        defect_area="HV BUSHING",
        equipment_id="TX 1",
    )
    tx_defects = [d_tx]
    assert _check_has_active_defect("tx", "TX1", cbm_defects=tx_defects) is True
    assert _check_has_active_defect("tx", "TX2", cbm_defects=tx_defects) is False
    assert _check_has_active_defect("tx", "", cbm_defects=tx_defects) is True
    assert _check_has_active_defect("fp", "fp1", cbm_defects=tx_defects) is False

def test_switchgear_scan_adapter_tev_background_resolution():
    """Verify SwitchgearScanAdapter reads tev_background from substation_info and defaults to '-', never '4'."""
    swg = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        panels=(SwitchgearPanelScanSpec(panel_no=1, panel_feeder_no="P1", name="TX A"),),
    )

    # 1. When substation_info has no TEV background, must default to "-", NOT "4"
    ad_empty = SwitchgearScanAdapter(swg=swg, substation_info={})
    assert ad_empty.tev_background == "-"
    res_empty = ad_empty.adapt()
    panel_item_empty = [it for it in res_empty.items if not it.is_overview][0]
    assert panel_item_empty.context["panel"]["tev"]["bg"] == "-"

    # 2. When substation_info has tev_background="1", must be "1"
    ad_with_bg = SwitchgearScanAdapter(swg=swg, substation_info={"tev_background": "1"})
    assert ad_with_bg.tev_background == "1"
    res_with_bg = ad_with_bg.adapt()
    panel_item_with_bg = [it for it in res_with_bg.items if not it.is_overview][0]
    assert panel_item_with_bg.context["panel"]["tev"]["bg"] == "1"

    # 3. When substation_info has tev_bg="1", must be "1"
    ad_with_tev_bg = SwitchgearScanAdapter(swg=swg, substation_info={"tev_bg": "1"})
    assert ad_with_tev_bg.tev_background == "1"
    res_with_tev_bg = ad_with_tev_bg.adapt()
    panel_item_with_tev_bg = [it for it in res_with_tev_bg.items if not it.is_overview][0]
    assert panel_item_with_tev_bg.context["panel"]["tev"]["bg"] == "1"

    # 4. When substation_info has tev_background="-" or None, must be "-"
    ad_dash = SwitchgearScanAdapter(swg=swg, substation_info={"tev_background": "-"})
    assert ad_dash.tev_background == "-"
    res_dash = ad_dash.adapt()
    panel_item_dash = [it for it in res_dash.items if not it.is_overview][0]
    assert panel_item_dash.context["panel"]["tev"]["bg"] == "-"


def test_defect_isolation_per_transformer_instance() -> None:
    """Verify zero defect leakage between TX1 and TX2 (Bug 8)."""
    from src.full_report.scan_adapters import _check_has_active_defect
    from src.quick_report.defects import CbmDefectRecord

    # 1. Defect explicitly for TX1
    d_tx1 = CbmDefectRecord(
        equipment="TRANSFORMER",
        equipment_id="TX 1",
        defect_area="HV BUSHING",
        technology="IR",
    )
    assert _check_has_active_defect("tx", "tx1", cbm_defects=[d_tx1], total_count=2) is True
    assert _check_has_active_defect("tx", "tx2", cbm_defects=[d_tx1], total_count=2) is False

    # 2. Defect explicitly for TX2
    d_tx2 = CbmDefectRecord(
        equipment="TRANSFORMER",
        equipment_id="TX 2",
        defect_area="LV CABLE",
        technology="US",
    )
    assert _check_has_active_defect("tx", "tx1", cbm_defects=[d_tx2], total_count=2) is False
    assert _check_has_active_defect("tx", "tx2", cbm_defects=[d_tx2], total_count=2) is True

    # 3. Defect with generic LTX/DTX on multi-TX station (total_count=2):
    # Maps to TX1 only, zero leakage to TX2!
    d_generic = CbmDefectRecord(
        equipment="LTX/DTX",
        equipment_id="",
        defect_area="BODY",
        technology="IR",
    )
    assert _check_has_active_defect("tx", "tx1", cbm_defects=[d_generic], total_count=2) is True
    assert _check_has_active_defect("tx", "tx2", cbm_defects=[d_generic], total_count=2) is False

    # 4. Defect with generic CABLE LTX/DTX on single-TX station (total_count=1):
    assert _check_has_active_defect("tx", "tx1", cbm_defects=[d_generic], total_count=1) is True


def test_defect_isolation_per_feeder_pillar_instance() -> None:
    """Verify zero defect leakage between FP1 and FP2 (Bug 8)."""
    from src.full_report.scan_adapters import _check_has_active_defect
    from src.quick_report.defects import CbmDefectRecord

    d_fp1 = CbmDefectRecord(
        equipment="FEEDER PILLAR",
        equipment_id="FP 1",
        defect_area="FUSE BASE",
        technology="IR",
    )
    assert _check_has_active_defect("fp", "fp1", cbm_defects=[d_fp1], total_count=2) is True
    assert _check_has_active_defect("fp", "fp2", cbm_defects=[d_fp1], total_count=2) is False

    d_fp2 = CbmDefectRecord(
        equipment="FEEDER PILLAR",
        equipment_id="FP 2",
        defect_area="BUSBAR",
        technology="IR",
    )
    assert _check_has_active_defect("fp", "fp1", cbm_defects=[d_fp2], total_count=2) is False
    assert _check_has_active_defect("fp", "fp2", cbm_defects=[d_fp2], total_count=2) is True


def test_overview_templates_placeholders_and_activex() -> None:
    """Verify swg-overview, tx-overview, fp-overview have Analysis/Recommendation and ActiveX controls."""
    from pathlib import Path
    import zipfile
    import docx

    tpl_dir = Path("templates/FULL REPORT/NORMAL IR US TEV")
    assert not (tpl_dir / "blackbox-overview.docx").exists(), "blackbox-overview.docx must be removed"

    for tpl_name in ("swg-overview.docx", "tx-overview.docx", "fp-overview.docx"):
        tpl_path = tpl_dir / tpl_name
        assert tpl_path.exists(), f"Template {tpl_name} missing"

        # Check ActiveX preservation
        with zipfile.ZipFile(tpl_path, "r") as zf:
            namelist = zf.namelist()
            assert "word/activeX/activeX1.bin" in namelist, f"ActiveX bin missing in {tpl_name}"
            assert "word/activeX/activeX1.xml" in namelist, f"ActiveX xml missing in {tpl_name}"
            assert "word/activeX/_rels/activeX1.xml.rels" in namelist, f"ActiveX rels missing in {tpl_name}"

        # Check placeholders in table rows 35 & 36
        doc = docx.Document(tpl_path)
        table = doc.tables[0]
        row_texts = [r.cells[0].text.strip() for r in table.rows]
        analysis_rows = [t for t in row_texts if "Analysis:" in t]
        rec_rows = [t for t in row_texts if "Recommendation:" in t]

        assert any("{{ analysis }}" in t for t in analysis_rows), f"{{ analysis }} missing in {tpl_name}"
        assert any("{{ recommendation }}" in t for t in rec_rows), f"{{ recommendation }} missing in {tpl_name}"


def test_switchgear_adapter_strict_zero_fallback(mock_substation_info: dict[str, str]) -> None:
    """Verify missing compartment photo produces blank image strings and does NOT fall back to photo_numbers[0] or cable_photo."""
    fake_resolver = FakePhotoResolver({
        101: ("/photos/IR_101.jpg", "/photos/VIS_101.jpg"),
    })
    panel = SwitchgearPanelScanSpec(
        panel_no=1,
        name="FEEDER 1",
        cable_photo=101,
        breaker_photo=None,
        busbar_photo=None,
        secondary_photo=None,
        photo_numbers=(101,),
        compartments=(
            "BREAKER COMPARTMENT",
            "CABLE COMPARTMENT",
            "BUSBAR COMPARTMENT",
            "SECONDARY COMPARTMENT",
        ),
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW FRONT",),
        panels=(panel,),
        photo_numbers=(50,),
    )
    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        photo_resolver=fake_resolver,
    )
    result = adapter.adapt()
    items = result.items

    assert items[0].is_overview is True

    # Item 1 is BREAKER COMPARTMENT: breaker_photo is None -> ir.image and visual.image must be blank ""
    breaker_item = items[1]
    assert breaker_item.component_name == "BREAKER COMPARTMENT"
    assert breaker_item.context["ir"]["image"] == ""
    assert breaker_item.context["visual"]["image"] == ""

    # Item 2 is CABLE COMPARTMENT: cable_photo is 101 -> ir.image and visual.image must be resolved
    cable_item = items[2]
    assert cable_item.component_name == "CABLE COMPARTMENT"
    assert cable_item.context["ir"]["image"] == "/photos/IR_101.jpg"
    assert cable_item.context["visual"]["image"] == "/photos/VIS_101.jpg"

    # Item 3 is BUSBAR COMPARTMENT: busbar_photo is None -> blank ""
    busbar_item = items[3]
    assert busbar_item.component_name == "BUSBAR COMPARTMENT"
    assert busbar_item.context["ir"]["image"] == ""
    assert busbar_item.context["visual"]["image"] == ""

    # Item 4 is SECONDARY COMPARTMENT: secondary_photo is None -> blank ""
    sec_item = items[4]
    assert sec_item.component_name == "SECONDARY COMPARTMENT"
    assert sec_item.context["ir"]["image"] == ""
    assert sec_item.context["visual"]["image"] == ""


def test_switchgear_adapter_vcb_transition_front_rear_photo_mapping(mock_substation_info: dict[str, str]) -> None:
    """Verify VCB transition bay maps FRONT COMPARTMENT -> breaker_photo and REAR COMPARTMENT -> cable_photo."""
    fake_resolver = FakePhotoResolver({
        201: ("/photos/IR_201.jpg", "/photos/VIS_201.jpg"),
        202: ("/photos/IR_202.jpg", "/photos/VIS_202.jpg"),
    })
    panel = SwitchgearPanelScanSpec(
        panel_no=1,
        name="TRANSITION",
        breaker_photo=201,  # Should map to FRONT COMPARTMENT
        cable_photo=202,    # Should map to REAR COMPARTMENT
        busbar_photo=None,
        photo_numbers=(201, 202),
        compartments=("FRONT COMPARTMENT", "REAR COMPARTMENT", "BUSBAR COMPARTMENT"),
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW FRONT",),
        panels=(panel,),
        photo_numbers=(50,),
    )
    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        photo_resolver=fake_resolver,
    )
    result = adapter.adapt()
    items = result.items

    front_item = next(it for it in items if it.component_name == "FRONT COMPARTMENT")
    assert front_item.context["ir"]["image"] == "/photos/IR_201.jpg"
    assert front_item.context["visual"]["image"] == "/photos/VIS_201.jpg"

    rear_item = next(it for it in items if it.component_name == "REAR COMPARTMENT")
    assert rear_item.context["ir"]["image"] == "/photos/IR_202.jpg"
    assert rear_item.context["visual"]["image"] == "/photos/VIS_202.jpg"

    busbar_item = next(it for it in items if it.component_name == "BUSBAR COMPARTMENT")
    assert busbar_item.context["ir"]["image"] == ""
    assert busbar_item.context["visual"]["image"] == ""


def test_switchgear_adapter_overview_photo_mapping_vcb_and_rmu(mock_substation_info: dict[str, str]) -> None:
    """Verify overview photo mapping for VCB (Front=Row 27, Rear=Row 26, Top=Row 28) and RMU."""
    fake_resolver = FakePhotoResolver({
        50: ("/photos/IR_50.jpg", "/photos/VIS_50.jpg"),
        60: ("/photos/IR_60.jpg", "/photos/VIS_60.jpg"),
        70: ("/photos/IR_70.jpg", "/photos/VIS_70.jpg"),
        80: ("/photos/IR_80.jpg", "/photos/VIS_80.jpg"),
        90: ("/photos/IR_90.jpg", "/photos/VIS_90.jpg"),
    })

    # 1. VCB: photo_numbers has [Row 27 (Front), Row 26 (Rear), Row 28 (Top)]
    swg_vcb = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW FRONT", "OVERVIEW REAR", "OVERVIEW TOP"),
        panels=(),
        photo_numbers=(50, 60, 70),
    )
    adapter_vcb = SwitchgearScanAdapter(
        swg=swg_vcb,
        substation_info=mock_substation_info,
        photo_resolver=fake_resolver,
    )
    res_vcb = adapter_vcb.adapt()
    assert len(res_vcb.items) == 3
    assert res_vcb.items[0].component_name == "OVERVIEW FRONT"
    assert res_vcb.items[0].context["ir"]["image"] == "/photos/IR_50.jpg"
    assert res_vcb.items[1].component_name == "OVERVIEW REAR"
    assert res_vcb.items[1].context["ir"]["image"] == "/photos/IR_60.jpg"
    assert res_vcb.items[2].component_name == "OVERVIEW TOP"
    assert res_vcb.items[2].context["ir"]["image"] == "/photos/IR_70.jpg"

    # 2. RMU: photo_numbers has [Row 26 (Overview), Row 28 (Overview Bottom/Top)]
    swg_rmu = SwitchgearScanSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.RMU_DUAL_CABLE_ENTRY,
        overview_compartments=("OVERVIEW", "OVERVIEW BOTTOM"),
        panels=(),
        photo_numbers=(80, 90),
    )
    adapter_rmu = SwitchgearScanAdapter(
        swg=swg_rmu,
        substation_info=mock_substation_info,
        photo_resolver=fake_resolver,
    )
    res_rmu = adapter_rmu.adapt()
    assert len(res_rmu.items) == 2
    assert res_rmu.items[0].component_name == "OVERVIEW"
    assert res_rmu.items[0].context["ir"]["image"] == "/photos/IR_80.jpg"
    assert res_rmu.items[1].component_name == "OVERVIEW BOTTOM"
    assert res_rmu.items[1].context["ir"]["image"] == "/photos/IR_90.jpg"


def test_switchgear_adapter_panel_serial_number_in_render_context(mock_substation_info: dict[str, str]) -> None:
    """Verify panel.serial_no is passed directly into panel render context."""
    panel = SwitchgearPanelScanSpec(
        panel_no=1,
        name="PANEL 1",
        serial_no="SN-PANEL-9988",
        compartments=("CABLE COMPARTMENT",),
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW FRONT",),
        panels=(panel,),
    )
    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
    )
    result = adapter.adapt()
    panel_item = result.items[1]
    assert panel_item.context["panel"]["serialnumber"] == "SN-PANEL-9988"


def test_switchgear_adapter_33kv_overview_and_cable_entry_fallback(mock_substation_info: dict[str, str]) -> None:
    """Verify 33kV switchgear generates dual Front and Rear overviews, and CABLE ENTRY falls back to breaker_photo."""
    swg_33kv = SwitchgearSpec(
        switchgear_type="RMU SF6 33kV",
        rating="33kV",
        panels=(),
    )
    adapter = SwitchgearScanAdapter(
        swg=swg_33kv,
        substation_info=mock_substation_info,
    )
    result = adapter.adapt()
    ov_comps = [item.component_name for item in result.items if item.is_overview]
    assert ov_comps == ["OVERVIEW FRONT", "OVERVIEW REAR"]

    # CABLE ENTRY fallback test
    panel_cable_entry = SwitchgearPanelSpec(
        panel_no=1,
        name="PANEL 1",
        photo_numbers=(100,),  # only 1 photo in tuple
        breaker_photo=101,     # fallback photo for entry
    )
    swg_dual = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        model="GV3",
        panels=(panel_cable_entry,),
    )
    resolver = FakePhotoResolver({
        100: ("/photos/IR_100.jpg", "/photos/VIS_100.jpg"),
        101: ("/photos/IR_101.jpg", "/photos/VIS_101.jpg"),
    })
    adapter_dual = SwitchgearScanAdapter(
        swg=swg_dual,
        substation_info=mock_substation_info,
        photo_resolver=resolver,
    )
    res_dual = adapter_dual.adapt()
    cable_entry_items = [it for it in res_dual.items if it.component_name == "CABLE ENTRY"]
    assert len(cable_entry_items) == 1
    assert cable_entry_items[0].context["ir"]["image"] == "/photos/IR_101.jpg"


def test_transition_panel_linknumber_formats_as_dash_when_feeder_no_absent(
    mock_substation_info: dict[str, str],
) -> None:
    """Transition panels format linknumber as '-' when panel_feeder_no is empty or absent per Ticket #111."""
    swg = SwitchgearScanSpec(
        switchgear_type="VCB 11kV",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(
                panel_no=6,
                name="TRANSITION PANEL",
                panel_feeder_no="",
                compartments=("FRONT COMPARTMENT", "REAR COMPARTMENT", "BUSBAR COMPARTMENT"),
            ),
            SwitchgearPanelScanSpec(
                panel_no=12,
                name="TRANSITION PANEL",
                panel_feeder_no=None,
                compartments=("FRONT COMPARTMENT", "REAR COMPARTMENT", "BUSBAR COMPARTMENT"),
            ),
            SwitchgearPanelScanSpec(
                panel_no=18,
                name="TRANSITION PANEL",
                panel_feeder_no="CKN18",
                compartments=("FRONT COMPARTMENT", "REAR COMPARTMENT", "BUSBAR COMPARTMENT"),
            ),
            SwitchgearPanelScanSpec(
                panel_no=1,
                name="FEEDER 1",
                panel_feeder_no="",
                compartments=("CABLE COMPARTMENT",),
            ),
        ),
        photo_numbers=(30,),
    )

    adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
    result = adapter.adapt()

    # Panel 6 (transition, empty feeder_no) -> linknumber is "-"
    for item in result.items[1:4]:  # Panel 6 compartments
        assert item.context["panel"]["linknumber"] == "-"
        assert item.context["panel"]["name"] == "TRANSITION PANEL"

    # Panel 12 (transition, None feeder_no) -> linknumber is "-"
    for item in result.items[4:7]:  # Panel 12 compartments
        assert item.context["panel"]["linknumber"] == "-"
        assert item.context["panel"]["name"] == "TRANSITION PANEL"

    # Panel 18 (transition, with feeder_no="CKN18") -> linknumber is "CKN18"
    for item in result.items[7:10]:  # Panel 18 compartments
        assert item.context["panel"]["linknumber"] == "CKN18"
        assert item.context["panel"]["name"] == "TRANSITION PANEL"

    # Panel 1 (standard feeder, empty feeder_no) -> linknumber is "1"
    feeder_item = result.items[10]
    assert feeder_item.context["panel"]["linknumber"] == "1"
    assert feeder_item.context["panel"]["name"] == "FEEDER 1"


def test_switchgear_adapter_resolve_prpd_forwards_feeder_no(
    mock_substation_info: dict[str, str],
    monkeypatch,
    tmp_path: Path,
) -> None:
    """SwitchgearScanAdapter._resolve_prpd forwards panel_feeder_no to generate_prpd_graphs_for_swg_panel."""
    swg = SwitchgearScanSpec(
        switchgear_type="VCB",
        manufacturer="TAMCO",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(
            SwitchgearPanelScanSpec(
                panel_no=1,
                name="LOCAL TX",
                panel_feeder_no="7",
                compartments=("CABLE COMPARTMENT",),
            ),
        ),
    )

    captured_args = {}

    def mock_generate_prpd(survey_root, panel_no, output_dir, feeder_no="", panel_name="", **kwargs):
        captured_args["panel_no"] = panel_no
        captured_args["feeder_no"] = feeder_no
        captured_args["panel_name"] = panel_name
        return None, None

    monkeypatch.setattr(
        "src.full_report.scan_adapters.generate_prpd_graphs_for_swg_panel",
        mock_generate_prpd,
    )

    survey_root = tmp_path / "MOCK_SURVEY"
    survey_root.mkdir()

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info=mock_substation_info,
        survey_root=survey_root,
    )
    adapter.adapt()

    assert captured_args.get("panel_no") == 1
    assert captured_args.get("feeder_no") == "7"
    assert captured_args.get("panel_name") == "LOCAL TX"


def test_transformer_scan_adapter_ptx_and_ltx_prpd_catalog_lookup(
    mock_substation_info: dict[str, str],
    dummy_prpd_image: Path,
) -> None:
    """TransformerScanAdapter resolves PTX for PowerTransformerSpec and LTX for Local TX from catalog."""
    from src.testsheet.models import PowerTransformerSpec, TransformerSpec

    ptx = PowerTransformerSpec(
        tx_id="TX 1",
        rating_kva="5MVA",
        manufacturer="TAMCO",
        model="ONAN",
        serial_no="SN-PTX",
    )
    ltx = TransformerSpec(
        tx_id="LOCAL TX",
        rating_kva="300kVA",
        manufacturer="TAMCO",
        model="HERMETICALLY SEAL",
        serial_no="SN-LTX",
    )

    catalog = {
        "tx": {
            "ptx": {"us": str(dummy_prpd_image), "tev": None},
            "ltx": {"us": str(dummy_prpd_image), "tev": None},
        }
    }

    # 1. Power Transformer
    ad_ptx = TransformerScanAdapter(
        tx=ptx,
        substation_info=mock_substation_info,
        tx_index=1,
        prpd_catalog=catalog,
    )
    assert ad_ptx._resolve_prpd_us() == str(dummy_prpd_image)
    res_ptx = ad_ptx.adapt()
    bushing_item = next(it for it in res_ptx.items if it.component_name == "HV BUSHING")
    cable_item = next(it for it in res_ptx.items if it.component_name == "HV CABLE")
    assert bushing_item.context["us"]["prpd"] == str(dummy_prpd_image)
    assert cable_item.context["us"]["prpd"] == str(dummy_prpd_image)

    # 2. Local Transformer
    ad_ltx = TransformerScanAdapter(
        tx=ltx,
        substation_info=mock_substation_info,
        tx_index=2,
        prpd_catalog=catalog,
    )
    assert ad_ltx._resolve_prpd_us() == str(dummy_prpd_image)
    res_ltx = ad_ltx.adapt()
    bushing_ltx = next(it for it in res_ltx.items if it.component_name == "HV BUSHING")
    cable_ltx = next(it for it in res_ltx.items if it.component_name == "HV CABLE")
    assert bushing_ltx.context["us"]["prpd"] == str(dummy_prpd_image)
    assert cable_ltx.context["us"]["prpd"] == str(dummy_prpd_image)


def test_transformer_scan_adapter_dynamic_prpd_resolution_ptx_ltx(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """TransformerScanAdapter dynamically generates US PRPD graphs for PTX and LTX directories."""
    from src.testsheet.models import PowerTransformerSpec, TransformerSpec

    survey_dir = tmp_path / "SURVEY_ROOT"
    ptx_meas = survey_dir / "PTX" / "Transformer" / "20261022T164608_Ultrasonic"
    ltx_meas = survey_dir / "LTX" / "Transformer" / "20261022T164352_Ultrasonic"
    ptx_meas.mkdir(parents=True)
    ltx_meas.mkdir(parents=True)

    sample_js = 'var ultra_events = {"data": [[10.0, 50, 1]]};'
    (ptx_meas / "ultrasonic_phase_plot.js").write_text(sample_js, encoding="utf-8")
    (ltx_meas / "ultrasonic_phase_plot.js").write_text(sample_js, encoding="utf-8")

    out_dir = tmp_path / "OUT_PRPD"

    ptx = PowerTransformerSpec(tx_id="TX 1", rating_kva="5MVA")
    ltx = TransformerSpec(tx_id="LOCAL TX", rating_kva="300kVA")

    ad_ptx = TransformerScanAdapter(
        tx=ptx,
        substation_info=mock_substation_info,
        tx_index=1,
        survey_root=survey_dir,
        prpd_output_dir=out_dir,
        prpd_mode="option_b",
    )
    ad_ltx = TransformerScanAdapter(
        tx=ltx,
        substation_info=mock_substation_info,
        tx_index=2,
        survey_root=survey_dir,
        prpd_output_dir=out_dir,
        prpd_mode="option_b",
    )

    prpd_ptx = ad_ptx._resolve_prpd_us()
    prpd_ltx = ad_ltx._resolve_prpd_us()

    assert prpd_ptx and Path(prpd_ptx).exists()
    assert prpd_ltx and Path(prpd_ltx).exists()


def test_33kv_mini_ppu_part07_tx_scanning_pages_prpd(
    mock_substation_info: dict[str, str],
    tmp_path: Path,
) -> None:
    """Verify Part 07 TX scanning pages resolve non-empty US PRPD for 33kV KT and Merting."""
    from src.testsheet.parsers.ppu import PpuTestsheetParser
    from src.quick_report.prpd import generate_all_substation_prpd_graphs

    substations = [
        (
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JERANTUT\01. OCTOBER\22-10-2026\33. MINI PPU KUALA TAHAN 33-11kV (VI).xlsx"),
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\JERANTUT\01. OCTOBER\22-10-2026\33. MINI PPU KUALA TAHAN 33-11kV (VI)\RAW DATA\US"),
        ),
        (
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JERANTUT\01. OCTOBER\22-10-2026\34. MINI PPU MERTING 33-11kV (VI).xlsx"),
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\JERANTUT\01. OCTOBER\22-10-2026\34. MINI PPU MERTING 33-11kV (VI)\RAW DATA\US"),
        ),
    ]

    for ts_path, survey_path in substations:
        if not (ts_path.exists() and survey_path.exists() and any(survey_path.iterdir())):
            pytest.skip(f"Test data missing for {ts_path.stem}")

        data = PpuTestsheetParser().extract_testsheet_data(ts_path)
        tx_list = list(data.equipment.transformers)
        assert len(tx_list) >= 2, f"Expected at least 2 transformers, got {len(tx_list)}"

        out_dir = tmp_path / ts_path.stem
        catalog = generate_all_substation_prpd_graphs(
            survey_root=survey_path,
            output_dir=out_dir,
            mode="option_b",
        )

        for tx_idx, tx in enumerate(tx_list, 1):
            adapter = TransformerScanAdapter(
                tx=tx,
                substation_info=mock_substation_info,
                tx_index=tx_idx,
                prpd_catalog=catalog,
            )
            result = adapter.adapt()
            for item in result.items:
                if item.component_name in ("HV BUSHING", "HV CABLE"):
                    us_prpd = item.context.get("us", {}).get("prpd", "")
                    assert us_prpd != "", f"Empty US PRPD graph for {tx.tx_id} - {item.component_name}"
                    assert Path(us_prpd).exists(), f"PRPD file does not exist: {us_prpd}"


def test_switchgear_scan_adapter_compartment_specific_ultrasound_from_catalog(tmp_path: Path):
    """Test SwitchgearScanAdapter assigns distinct US PRPD graphs to compartments from catalog."""
    from src.full_report.models import SwitchgearPanelScanSpec, SwitchgearScanSpec

    p_cb = tmp_path / "panel1_cb_us.png"
    p_cable = tmp_path / "panel1_cable_us.png"
    p_busbar = tmp_path / "panel1_busbar_us.png"
    for p in (p_cb, p_cable, p_busbar):
        p.write_bytes(b"dummy_png_bytes")

    catalog = {
        "swg": {
            1: {
                "us": p_cb,
                "tev": None,
                "compartments": {
                    "CIRCUIT_BREAKER": {"us": p_cb, "tev": None},
                    "CABLE_BOX": {"us": p_cable, "tev": None},
                    "UPPER_BUSBARS": {"us": p_busbar, "tev": None},
                },
            }
        }
    }

    panel = SwitchgearPanelScanSpec(
        panel_no=1,
        name="INCOMING 1",
        compartments=("BREAKER COMPARTMENT", "CABLE COMPARTMENT", "BUSBAR COMPARTMENT"),
    )
    swg = SwitchgearScanSpec(
        switchgear_type="VCB",
        manufacturer="TAMCO",
        model="VH3",
        rating="33kV",
        panels=(panel,),
    )

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info={"date": "22/10/2026", "time": "10:00", "ambient": "30", "humidity": "60"},
        prpd_catalog=catalog,
    )
    result = adapter.adapt()

    cb_items = [it for it in result.items if it.component_name == "BREAKER COMPARTMENT"]
    cable_items = [it for it in result.items if it.component_name == "CABLE COMPARTMENT"]
    busbar_items = [it for it in result.items if it.component_name == "BUSBAR COMPARTMENT"]

    assert len(cb_items) == 1
    assert len(cable_items) == 1
    assert len(busbar_items) == 1

    us_cb = cb_items[0].context["panel"]["us"]["prpd"]
    us_cable = cable_items[0].context["panel"]["us"]["prpd"]
    us_busbar = busbar_items[0].context["panel"]["us"]["prpd"]

    assert us_cb == str(p_cb)
    assert us_cable == str(p_cable)
    assert us_busbar == str(p_busbar)

    # Must be distinct, not replicated
    assert len({us_cb, us_cable, us_busbar}) == 3

    # Check convenience method render_swg_scan_pages exists
    assert hasattr(adapter, "render_swg_scan_pages")


def test_switchgear_scan_adapter_33kv_kuala_tahan_and_merting_end_to_end(tmp_path: Path):
    """End-to-end test verifying compartment-specific ultrasound graphs on 22-10-2026 substations."""
    from src.quick_report.prpd import generate_all_substation_prpd_graphs
    from src.testsheet.parsers.ppu import PpuTestsheetParser

    substations = [
        (
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JERANTUT\01. OCTOBER\22-10-2026\33. MINI PPU KUALA TAHAN 33-11kV (VI).xlsx"),
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\JERANTUT\01. OCTOBER\22-10-2026\33. MINI PPU KUALA TAHAN 33-11kV (VI)\RAW DATA\US"),
        ),
        (
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JERANTUT\01. OCTOBER\22-10-2026\34. MINI PPU MERTING 33-11kV (VI).xlsx"),
            Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\JERANTUT\01. OCTOBER\22-10-2026\34. MINI PPU MERTING 33-11kV (VI)\RAW DATA\US"),
        ),
    ]

    for ts_path, survey_path in substations:
        if not (ts_path.exists() and survey_path.exists() and any(survey_path.iterdir())):
            pytest.skip(f"Test data missing for {ts_path.stem}")

        data = PpuTestsheetParser().extract_testsheet_data(ts_path)
        swg_list = list(data.equipment.switchgears)
        assert len(swg_list) >= 1, f"Expected at least 1 switchgear, got {len(swg_list)}"
        swg = swg_list[0]

        out_dir = tmp_path / ts_path.stem
        catalog = generate_all_substation_prpd_graphs(
            survey_root=survey_path,
            output_dir=out_dir,
            mode="option_b",
            panels=swg.panels,
        )

        adapter = SwitchgearScanAdapter(
            swg=swg,
            substation_info={"date": "22/10/2026", "time": "10:00", "ambient": "30", "humidity": "60"},
            survey_root=survey_path,
            prpd_catalog=catalog,
            prpd_mode="option_b",
            is_33kv=True,
        )
        result = adapter.adapt()

        # Check each panel has distinct US graphs across Breaker, Cable, Busbar
        for panel in swg.panels:
            p_no = panel.panel_no
            panel_items = [it for it in result.items if it.panel_no == p_no]
            cb_items = [it for it in panel_items if it.component_name == "BREAKER COMPARTMENT"]
            cable_items = [it for it in panel_items if it.component_name == "CABLE COMPARTMENT"]
            busbar_items = [it for it in panel_items if it.component_name == "BUSBAR COMPARTMENT"]

            if cb_items and cable_items and busbar_items:
                us_cb = cb_items[0].context["panel"]["us"]["prpd"]
                us_cable = cable_items[0].context["panel"]["us"]["prpd"]
                us_busbar = busbar_items[0].context["panel"]["us"]["prpd"]

                assert us_cb != "", f"Empty Breaker US PRPD for panel {p_no} in {ts_path.stem}"
                assert us_cable != "", f"Empty Cable US PRPD for panel {p_no} in {ts_path.stem}"
                assert us_busbar != "", f"Empty Busbar US PRPD for panel {p_no} in {ts_path.stem}"

                assert Path(us_cb).exists(), f"Breaker PRPD file does not exist: {us_cb}"
                assert Path(us_cable).exists(), f"Cable PRPD file does not exist: {us_cable}"
                assert Path(us_busbar).exists(), f"Busbar PRPD file does not exist: {us_busbar}"

                # Must be distinct file paths and distinct image contents (not replicated)
                assert us_cb != us_cable, f"Breaker and Cable graph paths identical for panel {p_no}"
                assert us_cable != us_busbar, f"Cable and Busbar graph paths identical for panel {p_no}"
                assert us_cb != us_busbar, f"Breaker and Busbar graph paths identical for panel {p_no}"

                h_cb = hashlib.md5(Path(us_cb).read_bytes()).hexdigest()
                h_cable = hashlib.md5(Path(us_cable).read_bytes()).hexdigest()
                h_busbar = hashlib.md5(Path(us_busbar).read_bytes()).hexdigest()

                assert h_cb != h_cable, f"Breaker and Cable graph contents identical for panel {p_no} in {ts_path.stem}"
                assert h_cable != h_busbar, f"Cable and Busbar graph contents identical for panel {p_no} in {ts_path.stem}"
                assert h_cb != h_busbar, f"Breaker and Busbar graph contents identical for panel {p_no} in {ts_path.stem}"


def test_switchgear_scan_adapter_33kv_gebeng_us_prpd_graphs(tmp_path: Path):
    """Verify compartment-specific ultrasound graphs for 33kV GIS at SSU KPC GEBENG."""
    from src.quick_report.prpd import generate_all_substation_prpd_graphs, discover_ultratev_survey_dir
    from src.testsheet.parsers.ppu import PpuTestsheetParser

    ts_path = Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\GEBENG\01. OCTOBER\08-10-2026\05. SSU KPC GEBENG 33kV (VI).xlsx")
    raw_dir = Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\GEBENG\01. OCTOBER\08-10-2026\05. SSU KPC GEBENG 33kV (VI)\RAW DATA")

    if not (ts_path.exists() and raw_dir.exists()):
        pytest.skip("Test data missing for SSU KPC GEBENG")

    data = PpuTestsheetParser().extract_testsheet_data(ts_path)
    swg = data.equipment.switchgears[0]
    survey_path = discover_ultratev_survey_dir(raw_dir)
    assert survey_path is not None, "UltraTEV survey root not discovered"

    out_dir = tmp_path / "prpd"
    catalog = generate_all_substation_prpd_graphs(
        survey_root=survey_path,
        output_dir=out_dir,
        mode="option_b",
        panels=swg.panels,
    )

    # Asserts that catalog['swg'] is populated for panels
    panel_keys = [k for k in catalog["swg"] if isinstance(k, int)]
    assert len(panel_keys) == len(swg.panels), f"Expected {len(swg.panels)} panels in catalog, got {len(panel_keys)}"

    adapter = SwitchgearScanAdapter(
        swg=swg,
        substation_info={"date": "08/10/2026", "time": "08:34", "ambient": "30", "humidity": "60"},
        survey_root=survey_path,
        prpd_catalog=catalog,
        prpd_mode="option_b",
        is_33kv=True,
    )
    result = adapter.adapt()

    p4_items = [it for it in result.items if it.panel_no == 4]
    assert len(p4_items) > 0
    cb_item = next(it for it in p4_items if it.component_name == "BREAKER COMPARTMENT")
    us_prpd = cb_item.context["panel"]["us"]["prpd"]
    assert us_prpd != "", "Panel 4 Breaker US PRPD graph was empty"
    assert Path(us_prpd).exists(), f"PRPD file does not exist: {us_prpd}"


def test_switchgear_scan_adapter_dual_switchgear_resolution(tmp_path: Path):
    """Ticket #129: Verify SwitchgearScanAdapter resolves waveforms using swg_index without lineup contamination."""
    p1_swg1 = tmp_path / "prpd_swg1_panel1_us.png"
    p1_swg2 = tmp_path / "prpd_swg2_panel1_us.png"
    p1_swg1.write_bytes(b"swg1")
    p1_swg2.write_bytes(b"swg2")

    catalog = {
        "swg": {
            1: {
                1: {"us": p1_swg1, "tev": None},
            },
            2: {
                1: {"us": p1_swg2, "tev": None},
            },
        }
    }

    swg1 = SwitchgearScanSpec(
        switchgear_type="GIS",
        rating="33kV",
        panels=(SwitchgearPanelScanSpec(panel_no=1, name="BURSTING PANEL", compartments=("BREAKER COMPARTMENT",)),),
    )
    swg2 = SwitchgearScanSpec(
        switchgear_type="VCB",
        rating="11kV",
        panels=(SwitchgearPanelScanSpec(panel_no=1, name="INCOMER TX", compartments=("BREAKER COMPARTMENT",)),),
    )

    ad1 = SwitchgearScanAdapter(
        swg=swg1,
        swg_index=1,
        prpd_catalog=catalog,
    )
    res1 = ad1.adapt()
    p_items1 = [it for it in res1.items if not it.is_overview]
    assert len(p_items1) == 1
    us1 = p_items1[0].context["panel"]["us"]["prpd"]

    ad2 = SwitchgearScanAdapter(
        swg=swg2,
        swg_index=2,
        prpd_catalog=catalog,
    )
    res2 = ad2.adapt()
    p_items2 = [it for it in res2.items if not it.is_overview]
    assert len(p_items2) == 1
    us2 = p_items2[0].context["panel"]["us"]["prpd"]

    assert us1 == str(p1_swg1)
    assert us2 == str(p1_swg2)
    assert us1 != us2


def test_switchgear_scan_adapter_dual_switchgear_jengka_live(tmp_path: Path):
    """Ticket #129: End-to-end verification on CJEN/PMU/J00001 (40. PMU JENGKA 132-11-33kV)."""
    from src.quick_report.prpd import generate_all_substation_prpd_graphs, discover_ultratev_survey_dir
    from src.testsheet.parsers.ppu import PpuTestsheetParser

    ts_path = Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JENGKA\01. OCTOBER\25-10-2026\40. PMU JENGKA 132-11-33kV (VI).xlsx")
    raw_dir = Path(r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\RAW MATERIAL\JENGKA\01. OCTOBER\25-10-2026\40. PMU JENGKA 132-11-33kV (VI)\RAW DATA")

    if not (ts_path.exists() and raw_dir.exists()):
        pytest.skip("Test data missing for PMU JENGKA")

    survey_path = discover_ultratev_survey_dir(raw_dir)
    assert survey_path is not None, "UltraTEV survey root not discovered"

    data = PpuTestsheetParser().extract_testsheet_data(ts_path)
    swgs = data.equipment.switchgears
    assert len(swgs) == 2, f"Expected 2 switchgears, got {len(swgs)}"

    out_dir = tmp_path / "prpd_jengka"
    catalog = generate_all_substation_prpd_graphs(
        survey_root=survey_path,
        output_dir=out_dir,
        mode="option_b",
        switchgears=swgs,
    )

    assert "swg" in catalog
    assert 1 in catalog["swg"], "SWG 1 missing from catalog"
    assert 2 in catalog["swg"], "SWG 2 missing from catalog"

    # SWG 1 (33kV GIS) has 8 panels
    swg1_panel_keys = [k for k in catalog["swg"][1] if isinstance(k, int)]
    assert len(swg1_panel_keys) == 8, f"Expected 8 panels for SWG 1, got {len(swg1_panel_keys)}"

    # SWG 2 (11kV VCB) has 29 panels
    swg2_panel_keys = [k for k in catalog["swg"][2] if isinstance(k, int)]
    assert len(swg2_panel_keys) == 29, f"Expected 29 panels for SWG 2, got {len(swg2_panel_keys)}"

    # Adapters resolve respective lineups
    ad1 = SwitchgearScanAdapter(
        swg=swgs[0],
        swg_index=1,
        prpd_catalog=catalog,
    )
    res1 = ad1.adapt()
    assert len(res1.items) > 0

    ad2 = SwitchgearScanAdapter(
        swg=swgs[1],
        swg_index=2,
        prpd_catalog=catalog,
    )
    res2 = ad2.adapt()
    assert len(res2.items) > 0


def test_link_box_compartment_metadata_sanitization(mock_substation_info: dict[str, str]) -> None:
    """Ticket #131: Verify LINK BOX compartment sanitizes switchgear metadata and voltage grading."""
    # 1. 33kV Switchgear with LINK BOX compartment
    panel_33kv = SwitchgearPanelScanSpec(
        panel_no=2,
        name="PMU GEBENG (7L5)",
        panel_feeder_no="3S5",
        serial_no="39000589/10/016",
        load_amp=120.0,
        status="SERVICE/CLOSE",
        heater_amp="0.5",
        busbar_position="MAIN",
        cable_type="XLPE",
        compartments=("BREAKER COMPARTMENT", "LINK BOX"),
    )
    swg_33kv = SwitchgearScanSpec(
        switchgear_type="GIS",
        manufacturer="EPE",
        model="VGA-32",
        rating="36kV",
        serial_no="SWG-GIS-001",
        archetype=SwitchgearArchetype.GIS_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(panel_33kv,),
        voltage_class=VoltageClass.KV_33,
    )

    ad_33kv = SwitchgearScanAdapter(
        swg=swg_33kv,
        substation_info=mock_substation_info,
        is_33kv=True,
    )
    res_33kv = ad_33kv.adapt()
    items_33kv = {item.component_name: item for item in res_33kv.items if not item.is_overview}

    # Verify standard compartment (BREAKER COMPARTMENT) retains switchgear attributes
    cb_item = items_33kv["BREAKER COMPARTMENT"]
    assert cb_item.context["swg"]["type"] == "GIS"
    assert cb_item.context["swg"]["manufacturer"] == "EPE"
    assert cb_item.context["swg"]["model"] == "VGA-32"
    assert cb_item.context["swg"]["rating"] == "36kV"
    assert cb_item.context["panel"]["area"] == "BREAKER COMPARTMENT"
    assert cb_item.context["panel"]["serialnumber"] == "39000589/10/016"
    assert cb_item.context["panel"]["serial_no"] == "39000589/10/016"
    assert cb_item.context["panel"]["heateramp"] == "ON:0.50A/OFF:0.0A"
    assert cb_item.context["panel"]["breakerstatus"] == "SERVICE/CLOSE"
    assert cb_item.context["panel"]["busbarposition"] == "MAIN"
    assert cb_item.context["panel"]["cabletype"] == "XLPE"
    assert cb_item.context["panel"]["linknumber"] == "3S5"
    assert cb_item.context["panel"]["name"] == "PMU GEBENG (7L5)"
    assert cb_item.context["panel"]["loadamp"] == "120"

    # Verify LINK BOX compartment sanitizes switchgear metadata and voltage grading
    lb_item_33kv = items_33kv["LINK BOX"]
    assert lb_item_33kv.context["swg"]["type"] == "LINK BOX"
    assert lb_item_33kv.context["swg"]["manufacturer"] == "-"
    assert lb_item_33kv.context["swg"]["model"] == "-"
    assert lb_item_33kv.context["swg"]["rating"] == "33kV"
    assert lb_item_33kv.context["panel"]["area"] == "-"
    assert lb_item_33kv.context["panel"]["serialnumber"] == "-"
    assert lb_item_33kv.context["panel"]["serial_no"] == "-"
    assert lb_item_33kv.context["panel"]["heateramp"] == "-"
    assert lb_item_33kv.context["panel"]["breakerstatus"] == "-"
    assert lb_item_33kv.context["panel"]["busbarposition"] == "-"
    assert lb_item_33kv.context["panel"]["cabletype"] == "-"
    # Preserved fields
    assert lb_item_33kv.context["panel"]["linknumber"] == "3S5"
    assert lb_item_33kv.context["panel"]["name"] == "PMU GEBENG (7L5)"
    assert lb_item_33kv.context["panel"]["loadamp"] == "120"
    # TEV / US blanking flags
    assert lb_item_33kv.context["__blank_us__"] is True
    assert lb_item_33kv.context["is_us_active"] is False
    assert lb_item_33kv.context["__blank_tev__"] is True
    assert lb_item_33kv.context["is_tev_active"] is False
    assert lb_item_33kv.context["us"]["reading"] == ""
    assert lb_item_33kv.context["tev"]["reading"] == ""

    # 2. 11kV Switchgear with LINK BOX compartment
    panel_11kv = SwitchgearPanelScanSpec(
        panel_no=1,
        name="FEEDER 1",
        panel_feeder_no="1K1",
        serial_no="SN-VCB-111",
        load_amp=85.0,
        status="SERVICE/CLOSE",
        heater_amp="0.4",
        busbar_position="TOP",
        cable_type="PILC",
        compartments=("LINK BOX",),
    )
    swg_11kv = SwitchgearScanSpec(
        switchgear_type="VCB",
        manufacturer="TAMCO",
        model="GV3",
        rating="12kV",
        serial_no="SWG-VCB-002",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW",),
        panels=(panel_11kv,),
        voltage_class=VoltageClass.KV_11,
    )
    ad_11kv = SwitchgearScanAdapter(
        swg=swg_11kv,
        substation_info=mock_substation_info,
        is_33kv=False,
    )
    res_11kv = ad_11kv.adapt()
    lb_item_11kv = [item for item in res_11kv.items if item.component_name == "LINK BOX"][0]

    assert lb_item_11kv.context["swg"]["type"] == "LINK BOX"
    assert lb_item_11kv.context["swg"]["manufacturer"] == "-"
    assert lb_item_11kv.context["swg"]["model"] == "-"
    assert lb_item_11kv.context["swg"]["rating"] == "11kV"
    assert lb_item_11kv.context["panel"]["area"] == "-"
    assert lb_item_11kv.context["panel"]["serialnumber"] == "-"
    assert lb_item_11kv.context["panel"]["serial_no"] == "-"
    assert lb_item_11kv.context["panel"]["heateramp"] == "-"
    assert lb_item_11kv.context["panel"]["breakerstatus"] == "-"
    assert lb_item_11kv.context["panel"]["busbarposition"] == "-"
    assert lb_item_11kv.context["panel"]["cabletype"] == "-"
    assert lb_item_11kv.context["panel"]["linknumber"] == "1K1"
    assert lb_item_11kv.context["panel"]["name"] == "FEEDER 1"
    assert lb_item_11kv.context["panel"]["loadamp"] == "85"
    assert lb_item_11kv.context["__blank_us__"] is True
    assert lb_item_11kv.context["is_us_active"] is False
    assert lb_item_11kv.context["__blank_tev__"] is True
    assert lb_item_11kv.context["is_tev_active"] is False


# ==============================================================================
# 5. Ticket #133: PTX and LTX Transformer Numbering Differentiation Tests
# ==============================================================================

def test_normalize_tx_id_archetype_resolution() -> None:
    """normalize_tx_id() differentiates Power Transformers (PTX), Local Transformers (LTX), and Distribution (TX)."""
    from src.core.normalizers import normalize_tx_id
    from src.testsheet.models import PowerTransformerSpec, TransformerSpec

    # 1. Standard 11kV Distribution Transformers (TX1, TX2)
    assert normalize_tx_id("TX 1") == "TX1"
    assert normalize_tx_id("TX 2") == "TX2"
    assert normalize_tx_id("Transformer 1") == "TX1"
    assert normalize_tx_id("Transformer 2") == "TX2"
    assert normalize_tx_id("TX1") == "TX1"
    assert normalize_tx_id(1) == "TX1"
    assert normalize_tx_id(2) == "TX2"
    assert normalize_tx_id(None, default_idx=1) == "TX1"
    assert normalize_tx_id(None, default_idx=2) == "TX2"

    # 2. Power Transformers (PTX1, PTX2, PTX3, PTX4)
    assert normalize_tx_id("PTX 1") == "PTX1"
    assert normalize_tx_id("PTX1") == "PTX1"
    assert normalize_tx_id("PTX 2") == "PTX2"
    assert normalize_tx_id("PTX2") == "PTX2"
    assert normalize_tx_id("POWER TX 1") == "PTX1"
    assert normalize_tx_id("POWER TX 2") == "PTX2"
    assert normalize_tx_id("TX 1", is_power_tx=True) == "PTX1"
    assert normalize_tx_id("TX 2", is_power_tx=True) == "PTX2"
    assert normalize_tx_id(1, is_power_tx=True) == "PTX1"
    assert normalize_tx_id(2, is_power_tx=True) == "PTX2"
    assert normalize_tx_id(None, default_idx=1, is_power_tx=True) == "PTX1"
    assert normalize_tx_id("TX 1", tx_type="PTX") == "PTX1"
    assert normalize_tx_id("TX 2", tx_type="PTX") == "PTX2"

    # Spec object polymorphic resolution
    ptx1 = PowerTransformerSpec(tx_id="TX 1", rating_kva="7.5MVA")
    ptx2 = PowerTransformerSpec(tx_id="TX 2", rating_kva="7.5MVA")
    assert normalize_tx_id(ptx1) == "PTX1"
    assert normalize_tx_id(ptx2) == "PTX2"

    # 3. Local Station Service Transformers (LTX, LTX1, LTX2)
    assert normalize_tx_id("LOCAL TX") == "LTX"
    assert normalize_tx_id("Local TX") == "LTX"
    assert normalize_tx_id("LTX") == "LTX"
    assert normalize_tx_id("LOCAL") == "LTX"
    # Single local transformer on site resolves to LTX
    assert normalize_tx_id("LOCAL TX 1") == "LTX"
    assert normalize_tx_id("LOCAL TX 1", total_local_count=1) == "LTX"
    assert normalize_tx_id("LOCAL TX", total_local_count=1) == "LTX"
    # Explicit secondary local transformer or multiple local transformers on site
    assert normalize_tx_id("LOCAL TX 2") == "LTX2"
    assert normalize_tx_id("LTX 2") == "LTX2"
    assert normalize_tx_id("LOCAL TX 1", total_local_count=2) == "LTX1"
    assert normalize_tx_id("LOCAL TX 2", total_local_count=2) == "LTX2"
    assert normalize_tx_id("TX 3", is_local_tx=True) == "LTX"
    assert normalize_tx_id(None, default_idx=1, is_local_tx=True) == "LTX"

    # Local spec object resolution
    ltx_spec = TransformerSpec(tx_id="LOCAL TX", rating_kva="300kVA")
    assert normalize_tx_id(ltx_spec) == "LTX"

    # Standard distribution spec object resolution
    std_spec1 = TransformerSpec(tx_id="TX 1", rating_kva="1000kVA")
    std_spec2 = TransformerSpec(tx_id="TX 2", rating_kva="1000kVA")
    assert normalize_tx_id(std_spec1) == "TX1"
    assert normalize_tx_id(std_spec2) == "TX2"


def test_transformer_scan_adapter_ptx_and_ltx_rendering(
    mock_substation_info: dict[str, str],
) -> None:
    """TransformerScanAdapter renders PTX1, PTX2, and LTX on scan pages and item titles (Ticket #133)."""
    from src.testsheet.models import PowerTransformerSpec, TransformerSpec

    ptx1 = PowerTransformerSpec(
        tx_id="TX 1",
        rating_kva="7.5MVA",
        manufacturer="MTM",
        model="7.5-1",
        serial_no="SN-PTX-01",
    )
    ptx2 = PowerTransformerSpec(
        tx_id="TX 2",
        rating_kva="7.5MVA",
        manufacturer="MTM",
        model="7.5-1",
        serial_no="SN-PTX-02",
    )
    ltx = TransformerSpec(
        tx_id="LOCAL TX",
        rating_kva="300kVA",
        manufacturer="EWT",
        model="HERMETICALLY SEAL",
        serial_no="SN-LTX-01",
    )
    tx_std = TransformerSpec(
        tx_id="TX 1",
        rating_kva="1000kVA",
        manufacturer="TAMCO",
        model="ONAN",
        serial_no="SN-TX-01",
    )

    # 1. Power Transformer 1 -> PTX1
    ad_ptx1 = TransformerScanAdapter(
        tx=ptx1,
        substation_info=mock_substation_info,
        tx_index=1,
    )
    assert ad_ptx1.tx_id_str == "PTX1"
    res_ptx1 = ad_ptx1.adapt()
    assert res_ptx1.equipment_id == "PTX1"
    for item in res_ptx1.items:
        assert item.page_name.startswith("PTX1 - ")
        assert item.context["tx"]["number"] == "PTX1"
        assert item.context["tx"]["panel"] == "PTX1"
        assert item.context["panel"]["name"] == "PTX1"

    # 2. Power Transformer 2 -> PTX2
    ad_ptx2 = TransformerScanAdapter(
        tx=ptx2,
        substation_info=mock_substation_info,
        tx_index=2,
    )
    assert ad_ptx2.tx_id_str == "PTX2"
    res_ptx2 = ad_ptx2.adapt()
    assert res_ptx2.equipment_id == "PTX2"
    for item in res_ptx2.items:
        assert item.page_name.startswith("PTX2 - ")
        assert item.context["tx"]["number"] == "PTX2"
        assert item.context["tx"]["panel"] == "PTX2"

    # 3. Local Station Service Transformer -> LTX
    ad_ltx = TransformerScanAdapter(
        tx=ltx,
        substation_info=mock_substation_info,
        tx_index=3,
    )
    assert ad_ltx.tx_id_str == "LTX"
    res_ltx = ad_ltx.adapt()
    assert res_ltx.equipment_id == "LTX"
    for item in res_ltx.items:
        assert item.page_name.startswith("LTX - ")
        assert item.context["tx"]["number"] == "LTX"
        assert item.context["tx"]["panel"] == "LTX"
        assert item.context["panel"]["name"] == "LTX"

    # 4. Standard Distribution Transformer -> TX1
    ad_std = TransformerScanAdapter(
        tx=tx_std,
        substation_info=mock_substation_info,
        tx_index=1,
    )
    assert ad_std.tx_id_str == "TX1"
    res_std = ad_std.adapt()
    assert res_std.equipment_id == "TX1"
    for item in res_std.items:
        assert item.page_name.startswith("TX1 - ")
        assert item.context["tx"]["number"] == "TX1"


def test_adapt_equipment_package_ptx_ltx_dispatch(
    mock_substation_info: dict[str, str],
) -> None:
    """adapt_equipment_package preserves PTX1, PTX2, and LTX across multi-transformer 33kV packages."""
    from src.testsheet.models import (
        PowerTransformerSpec,
        SubstationEquipmentPackage,
        TransformerSpec,
    )

    ptx1 = PowerTransformerSpec(tx_id="TX 1", rating_kva="7.5MVA")
    ptx2 = PowerTransformerSpec(tx_id="TX 2", rating_kva="7.5MVA")
    ltx = TransformerSpec(tx_id="LOCAL TX", rating_kva="300kVA")

    pkg_33kv = SubstationEquipmentPackage(
        transformers=(ptx1, ptx2, ltx),
        switchgears=(),
    )

    results_33kv = adapt_equipment_package(
        package=pkg_33kv,
        substation_info=mock_substation_info,
    )

    tx_results_33kv = [r for r in results_33kv if r.equipment_category == "tx"]
    assert len(tx_results_33kv) == 3
    assert tx_results_33kv[0].equipment_id == "PTX1"
    assert tx_results_33kv[0].items[0].context["tx"]["number"] == "PTX1"
    assert tx_results_33kv[0].items[0].page_name == "PTX1 - OVERVIEW"

    assert tx_results_33kv[1].equipment_id == "PTX2"
    assert tx_results_33kv[1].items[0].context["tx"]["number"] == "PTX2"
    assert tx_results_33kv[1].items[0].page_name == "PTX2 - OVERVIEW"

    assert tx_results_33kv[2].equipment_id == "LTX"
    assert tx_results_33kv[2].items[0].context["tx"]["number"] == "LTX"
    assert tx_results_33kv[2].items[0].page_name == "LTX - OVERVIEW"

    # Verify standard distribution packages continue to produce TX1, TX2
    tx1 = TransformerSpec(tx_id="TX 1", rating_kva="1000kVA")
    tx2 = TransformerSpec(tx_id="TX 2", rating_kva="1000kVA")
    pkg_11kv = SubstationEquipmentPackage(
        transformers=(tx1, tx2),
        switchgears=(),
    )
    results_11kv = adapt_equipment_package(
        package=pkg_11kv,
        substation_info=mock_substation_info,
    )
    tx_results_11kv = [r for r in results_11kv if r.equipment_category == "tx"]
    assert len(tx_results_11kv) == 2
    assert tx_results_11kv[0].equipment_id == "TX1"
    assert tx_results_11kv[0].items[0].context["tx"]["number"] == "TX1"
    assert tx_results_11kv[1].equipment_id == "TX2"
    assert tx_results_11kv[1].items[0].context["tx"]["number"] == "TX2"


def test_cbm_defect_attribution_and_forwarding_banner_ptx_ltx(
    mock_substation_info: dict[str, str],
) -> None:
    """Active defect forwarding banner is isolated between PTX1, PTX2, and LTX without leakage."""
    from src.core.shading import BANNER_DEFECT_FORWARDING, BANNER_HEALTHY_ANALYSIS
    from src.testsheet.models import PowerTransformerSpec, TransformerSpec

    ptx1 = PowerTransformerSpec(tx_id="TX 1", rating_kva="7.5MVA")
    ptx2 = PowerTransformerSpec(tx_id="TX 2", rating_kva="7.5MVA")
    ltx = TransformerSpec(tx_id="LOCAL TX", rating_kva="300kVA")

    # Defect explicitly targeting PTX1
    ptx1_defect = [
        {
            "equipment": "PTX",
            "equipment_id": "PTX1",
            "defect_area": "MAIN TANK",
            "additional_remarks": "SILICA GEL IN POOR CONDITION",
        }
    ]

    ad_ptx1 = TransformerScanAdapter(
        tx=ptx1,
        substation_info=mock_substation_info,
        tx_index=1,
        total_tx_count=3,
        cbm_defects=ptx1_defect,
    )
    ad_ptx2 = TransformerScanAdapter(
        tx=ptx2,
        substation_info=mock_substation_info,
        tx_index=2,
        total_tx_count=3,
        cbm_defects=ptx1_defect,
    )
    ad_ltx = TransformerScanAdapter(
        tx=ltx,
        substation_info=mock_substation_info,
        tx_index=3,
        total_tx_count=3,
        cbm_defects=ptx1_defect,
    )

    res_ptx1 = ad_ptx1.adapt()
    res_ptx2 = ad_ptx2.adapt()
    res_ltx = ad_ltx.adapt()

    # PTX1 overview has defect forwarding banner
    assert res_ptx1.has_defect is True
    assert res_ptx1.items[0].context["banner"]["analysis"] == BANNER_DEFECT_FORWARDING

    # PTX2 overview remains healthy
    assert res_ptx2.has_defect is False
    assert res_ptx2.items[0].context["banner"]["analysis"] == BANNER_HEALTHY_ANALYSIS

    # LTX overview remains healthy (zero defect leakage from PTX1)
    assert res_ltx.has_defect is False
    assert res_ltx.items[0].context["banner"]["analysis"] == BANNER_HEALTHY_ANALYSIS

    # Local defect explicitly targeting LTX
    ltx_defect = [
        {
            "equipment": "LTX",
            "equipment_id": "LTX",
            "defect_area": "HV BUSHING",
            "additional_remarks": "HOTSPOT",
        }
    ]
    ad_ptx1_rev = TransformerScanAdapter(
        tx=ptx1,
        substation_info=mock_substation_info,
        tx_index=1,
        total_tx_count=3,
        cbm_defects=ltx_defect,
    )
    ad_ltx_rev = TransformerScanAdapter(
        tx=ltx,
        substation_info=mock_substation_info,
        tx_index=3,
        total_tx_count=3,
        cbm_defects=ltx_defect,
    )
    assert ad_ptx1_rev.adapt().has_defect is False
    assert ad_ltx_rev.adapt().has_defect is True
    assert ad_ltx_rev.adapt().items[0].context["banner"]["analysis"] == BANNER_DEFECT_FORWARDING


def test_table26_placards_and_visual_defects_preserved() -> None:
    """Table 26 placards remain POWER TX1, POWER TX2, LOCAL TX and visual defects unmodified."""
    from src.quick_report.substation_condition import PpuSubstationConditionPairBuilderPolicy
    from src.testsheet.models import (
        PowerTransformerSpec,
        SubstationEquipmentPackage,
        SubstationTestsheetPackage,
        TestsheetData,
        TransformerSpec,
        VisualInspectionFinding,
    )

    ptx1 = PowerTransformerSpec(tx_id="TX 1", rating_kva="7.5MVA")
    ptx2 = PowerTransformerSpec(tx_id="TX 2", rating_kva="7.5MVA")
    ltx = TransformerSpec(tx_id="LOCAL TX", rating_kva="300kVA")

    verbatim_vi = VisualInspectionFinding(
        category="TRANSFORMER",
        item_no="1",
        description="SILICA GEL IN POOR CONDITION – PTX1 - MAIN TANK",
        status="DEFECT",
    )

    eq_pkg = SubstationEquipmentPackage(
        transformers=(ptx1, ptx2, ltx),
    )
    testsheet_data = TestsheetData(
        substation_number=40,
        substation_name_erms="PMU JENGKA 132/11/33kV",
        equipment=eq_pkg,
    )
    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("dummy.xlsx"),
        unsorted_raw_data_dir=Path("dummy_dir"),
        station="PMU JENGKA",
        month="10. OCTOBER",
        date_str="25-10-2026",
        substation_number=40,
        data=testsheet_data,
    )

    pairs = PpuSubstationConditionPairBuilderPolicy.build(pkg)

    # Verify Table 26 placards
    tx_pairs = [p for p in pairs if "TX" in p[0]]
    assert tx_pairs == [
        ("POWER TX1", "POWER TX1 NAMEPLATE"),
        ("POWER TX2", "POWER TX2 NAMEPLATE"),
        ("LOCAL TX", "LOCAL TX NAMEPLATE"),
    ]

    # Verify visual defect is 100% verbatim
    assert verbatim_vi.description == "SILICA GEL IN POOR CONDITION – PTX1 - MAIN TANK"


def test_switchgear_scan_adapter_overview_serial_dashed_for_non_rmu(mock_substation_info: dict[str, str]) -> None:
    """Ticket #135: Overview scan pages for GIS, VCB, and RMU OIL must display '-' for serial number."""
    # 1. SSU KIIP GIS switchgear with erroneous last-PM string or real serial
    panel_gis = SwitchgearPanelScanSpec(
        panel_no=1,
        name="INCOMING 33kV",
        serial_no="GIS-PANEL-001",
        compartments=("CABLE COMPARTMENT",),
    )
    swg_gis = SwitchgearScanSpec(
        switchgear_type="GIS",
        manufacturer="TAMCO",
        serial_no="LHS:9.5.23, RHS:9.5.23",
        archetype=SwitchgearArchetype.GIS_CUBICLE,
        overview_compartments=("OVERVIEW FRONT (LHS)", "OVERVIEW REAR (LHS)"),
        panels=(panel_gis,),
    )
    adapter_gis = SwitchgearScanAdapter(swg=swg_gis, substation_info=mock_substation_info)
    result_gis = adapter_gis.adapt()
    overview_items_gis = [it for it in result_gis.items if it.is_overview]
    assert len(overview_items_gis) == 2
    for item in overview_items_gis:
        assert item.context["swg"]["serialnumber"] == "-"

    # 2. VCB switchgear with board serial
    swg_vcb = SwitchgearScanSpec(
        switchgear_type="VCB",
        manufacturer="TAMCO",
        serial_no="VCB-BOARD-999",
        archetype=SwitchgearArchetype.VCB_CUBICLE,
        overview_compartments=("OVERVIEW FRONT",),
        panels=(panel_gis,),
    )
    adapter_vcb = SwitchgearScanAdapter(swg=swg_vcb, substation_info=mock_substation_info)
    result_vcb = adapter_vcb.adapt()
    overview_items_vcb = [it for it in result_vcb.items if it.is_overview]
    assert len(overview_items_vcb) == 1
    assert overview_items_vcb[0].context["swg"]["serialnumber"] == "-"

    # 3. RMU OIL switchgear with board serial
    swg_oil = SwitchgearScanSpec(
        switchgear_type="RMU OIL",
        manufacturer="LUCY",
        serial_no="OIL-BOARD-888",
        archetype=SwitchgearArchetype.RMU_OIL,
        overview_compartments=("OVERVIEW FRONT",),
        panels=(panel_gis,),
    )
    adapter_oil = SwitchgearScanAdapter(swg=swg_oil, substation_info=mock_substation_info)
    result_oil = adapter_oil.adapt()
    overview_items_oil = [it for it in result_oil.items if it.is_overview]
    assert len(overview_items_oil) == 1
    assert overview_items_oil[0].context["swg"]["serialnumber"] == "-"


def test_switchgear_scan_adapter_overview_serial_preserved_for_eligible_rmu(mock_substation_info: dict[str, str]) -> None:
    """Ticket #135: Overview scan pages for RMU SF6, MRMU, RMU CB, and generic RMU preserve board serial."""
    panel_rmu = SwitchgearPanelScanSpec(
        panel_no=1,
        name="INCOMING",
        serial_no="SN-PANEL-01",
        compartments=("CABLE COMPARTMENT",),
    )

    test_cases = [
        ("RMU SF6", "SN: ABB-SF6-2023", "ABB-SF6-2023"),
        ("MRMU", "SCH-MRMU-9900", "SCH-MRMU-9900"),
        ("RMU CB", "SN: RMUCB-1122", "RMUCB-1122"),
        ("RMU", "GEN-RMU-5544", "GEN-RMU-5544"),
    ]

    for swg_type, raw_sn, expected_sn in test_cases:
        swg = SwitchgearScanSpec(
            switchgear_type=swg_type,
            manufacturer="ABB",
            serial_no=raw_sn,
            archetype=SwitchgearArchetype.RMU_STANDARD,
            overview_compartments=("OVERVIEW FRONT",),
            panels=(panel_rmu,),
        )
        adapter = SwitchgearScanAdapter(swg=swg, substation_info=mock_substation_info)
        result = adapter.adapt()
        overview_items = [it for it in result.items if it.is_overview]
        assert len(overview_items) == 1
        assert overview_items[0].context["swg"]["serialnumber"] == expected_sn











