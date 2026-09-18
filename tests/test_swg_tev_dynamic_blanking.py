"""Comprehensive tests for Switchgear Panel TEV Dynamic Blanking & PRPD Optimization (Issue #41).

Tests:
1. Two-tier evaluation (contract awarded technologies & switchgear compartment eligibility).
2. Normalization of switchgear compartment names (including Breaker / VCB / Spout / Chamber, etc.).
3. OpenXML dynamic blanking on swg-panel.docx (rows 21-32, cols 11-22):
   - Cell text cleared
   - Shading (<w:shd>) removed
   - Borders set to <w:val="nil"/>
   - Surrounding US cells and IR cells untouched
4. PRPD generation optimization:
   - Non-TEV compartment or include_tev=False skips Chromium / TEV decoding and returns None.
5. Full Report Scan Adapter integration:
   - TAMCO RMU: Cable Compartment has TEV active; Cable Entry has TEV blanked.
   - Non-TEV contract: all compartments have TEV blanked.
6. Quick Report CBM render integration:
   - Cable Entry defect has TEV blanked and empty string context values.
   - Non-TEV contract has TEV blanked.
   - Breaker / Spout defect under 3-tech contract preserves TEV.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import docx
from docx.oxml.ns import qn
import pytest

from src.core.shading import (
    blank_swg_tev_cells,
    get_cell_shading,
    is_swg_compartment_tev_eligible,
    is_swg_tev_active,
    is_tev_contract_awarded,
    normalize_swg_compartment,
)
from src.full_report.models import (
    SwitchgearCategory,
    SwitchgearPanelScanSpec,
    SwitchgearScanSpec,
)
from src.full_report.scan_adapters import SwitchgearScanAdapter
from src.full_report.scan_render import FullReportScanPageRendererCore
from src.quick_report.cbm_render import (
    _build_swg_render_context,
    _render_docx_template,
)
from src.quick_report.defects import CbmDefectRecord
from src.quick_report.prpd import generate_prpd_graphs_for_swg_panel


SWG_PANEL_TEMPLATE = Path("templates/FULL REPORT/NORMAL IR US TEV/swg-panel.docx")
QR_SWG_PANEL_TEMPLATE = Path("templates/QUICK REPORT/DEFECT IR US TEV/swg-panel.docx")


# ==============================================================================
# 1. Two-Tier Model & Compartment Normalization Tests
# ==============================================================================

class TestSwitchgearCompartmentEligibility:
    """Test compartment normalization and TEV eligibility rules."""

    @pytest.mark.parametrize(
        "raw_comp,expected_norm",
        [
            ("Breaker Compartment", "BREAKER COMPARTMENT"),
            ("VCB", "BREAKER COMPARTMENT"),
            ("VCB PANEL 1", "BREAKER COMPARTMENT"),
            ("CB", "BREAKER COMPARTMENT"),
            ("CB Compartment", "BREAKER COMPARTMENT"),
            ("Spout", "BREAKER COMPARTMENT"),
            ("CONTACT SPOUT", "BREAKER COMPARTMENT"),
            ("Arc Chamber", "BREAKER COMPARTMENT"),
            ("Cable Compartment", "CABLE COMPARTMENT"),
            ("Cable Box", "CABLE COMPARTMENT"),
            ("Cable Termination", "CABLE COMPARTMENT"),
            ("Cable Lug", "CABLE COMPARTMENT"),
            ("Cable Entry", "CABLE ENTRY"),
            ("Cable Entry Bottom", "CABLE ENTRY"),
            ("Entry Cable", "CABLE ENTRY"),
            ("Cable Inlet", "CABLE ENTRY"),
            ("Busbar", "BUSBAR COMPARTMENT"),
            ("Busbar Compartment", "BUSBAR COMPARTMENT"),
            ("PT Compartment", "PT COMPARTMENT"),
            ("Voltage Transformer", "PT COMPARTMENT"),
            ("Potential Transformer", "PT COMPARTMENT"),
            ("VT", "PT COMPARTMENT"),
            ("VT Compartment", "PT COMPARTMENT"),
            ("Fuse Compartment", "FUSE COMPARTMENT"),
            ("Outgoing Fuse", "FUSE COMPARTMENT"),
            ("Secondary Compartment", "SECONDARY COMPARTMENT"),
            ("Back Compartment", "REAR COMPARTMENT"),
            ("BACK COMPARTMENT", "REAR COMPARTMENT"),
            ("Back", "REAR COMPARTMENT"),
            ("BACK", "REAR COMPARTMENT"),
            ("Rear Compartment", "REAR COMPARTMENT"),
            ("REAR COMPARTMENT", "REAR COMPARTMENT"),
            ("Rear", "REAR COMPARTMENT"),
            ("REAR", "REAR COMPARTMENT"),
            ("Front Compartment", "FRONT COMPARTMENT"),
            ("FRONT COMPARTMENT", "FRONT COMPARTMENT"),
            ("Front", "FRONT COMPARTMENT"),
            ("FRONT", "FRONT COMPARTMENT"),
        ],
    )
    def test_normalize_swg_compartment(self, raw_comp: str, expected_norm: str):
        assert normalize_swg_compartment(raw_comp) == expected_norm

    @pytest.mark.parametrize(
        "comp,expected_eligible",
        [
            ("BREAKER COMPARTMENT", True),
            ("Breaker", True),
            ("VCB", True),
            ("CB", True),
            ("CB Compartment", True),
            ("Spout", True),
            ("Chamber", True),
            ("CABLE COMPARTMENT", True),
            ("Cable Box", True),
            ("PT COMPARTMENT", True),
            ("Voltage Transformer", True),
            ("Potential Transformer", True),
            ("VT", True),
            ("VT Compartment", True),
            ("FUSE COMPARTMENT", True),
            ("Fuse", True),
            # Non-TEV / Blanked compartments
            ("CABLE ENTRY", False),
            ("Cable Entry", False),
            ("Entry Cable", False),
            ("Cable Inlet", False),
            ("BUSBAR COMPARTMENT", False),
            ("Busbar", False),
            ("SECONDARY COMPARTMENT", False),
            ("BACK COMPARTMENT", False),
            ("Back Compartment", False),
            ("REAR COMPARTMENT", False),
            ("Rear Compartment", False),
            ("FRONT COMPARTMENT", False),
            ("Front Compartment", False),
            ("", False),
            (None, False),
        ],
    )
    def test_is_swg_compartment_tev_eligible(self, comp: str | None, expected_eligible: bool):
        assert is_swg_compartment_tev_eligible(comp) is expected_eligible


class TestContractAwardedTechnologies:
    """Test Tier 1: contract technology evaluation."""

    @pytest.mark.parametrize(
        "techs,expected",
        [
            (None, True),  # Default 3-tech contract
            ([], True),
            (["IR", "US", "TEV"], True),
            ({"IR", "US", "TEV"}, True),
            (["TEV"], True),
            ("IR+US+TEV", True),
            (["IR", "US"], False),
            (["IR"], False),
            (["US"], False),
            ("IR, US", False),
            ("INFRARED+ULTRASOUND", False),
        ],
    )
    def test_is_tev_contract_awarded(self, techs, expected):
        assert is_tev_contract_awarded(techs) is expected

    def test_is_swg_tev_active_combined(self):
        # Tier 1 True + Tier 2 True -> True
        assert is_swg_tev_active(["IR", "US", "TEV"], "Cable Compartment") is True
        assert is_swg_tev_active(None, "Breaker Compartment") is True
        assert is_swg_tev_active(["TEV"], "Spout") is True
        assert is_swg_tev_active(["TEV"], "Fuse Compartment") is True

        # Tier 1 True + Tier 2 False -> False (compartment not eligible)
        assert is_swg_tev_active(["IR", "US", "TEV"], "Cable Entry") is False
        assert is_swg_tev_active(["IR", "US", "TEV"], "Busbar Compartment") is False
        assert is_swg_tev_active(None, "Secondary Compartment") is False

        # Tier 1 False + Tier 2 True -> False (contract lacks TEV)
        assert is_swg_tev_active(["IR", "US"], "Cable Compartment") is False
        assert is_swg_tev_active(["IR", "US"], "Breaker Compartment") is False
        assert is_swg_tev_active(["IR"], "Spout") is False

        # Tier 1 False + Tier 2 False -> False
        assert is_swg_tev_active(["IR", "US"], "Cable Entry") is False


# ==============================================================================
# 2. OpenXML Blanking Mechanism Tests
# ==============================================================================

class TestOpenXmlBlanking:
    """Validate OpenXML blanking modifies only TEV cells in swg-panel.docx."""

    @pytest.fixture
    def swg_panel_doc(self) -> docx.Document:
        if not SWG_PANEL_TEMPLATE.exists():
            pytest.skip(f"Template not found: {SWG_PANEL_TEMPLATE}")
        return docx.Document(SWG_PANEL_TEMPLATE)

    def test_blank_swg_tev_cells_modifies_expected_region(self, swg_panel_doc: docx.Document):
        table = swg_panel_doc.tables[0]
        assert len(table.rows) == 37
        assert len(table.columns) == 24

        # Apply blanking
        blank_swg_tev_cells(swg_panel_doc)

        # 1. Verify all TEV cells (rows 21..32, cols 11..22) are blanked via python-docx cell interface
        for r_idx in range(21, 33):
            for c_idx in range(11, 23):
                cell = table.rows[r_idx].cells[c_idx]
                # Text is cleared
                assert cell.text.strip() == "", f"Cell ({r_idx}, {c_idx}) text not cleared: {cell.text}"
                # Shading removed
                tcPr = cell._tc.get_or_add_tcPr()
                shd = tcPr.find(qn("w:shd"))
                assert shd is None or shd.get(qn("w:fill")) in (None, "FFFFFF", "clear"), (
                    f"Cell ({r_idx}, {c_idx}) has unexpected fill: {shd.get(qn('w:fill')) if shd is not None else None}"
                )
                # Borders set to nil
                tcBorders = tcPr.find(qn("w:tcBorders"))
                assert tcBorders is not None, f"Cell ({r_idx}, {c_idx}) missing tcBorders"
                for b_name in ("top", "left", "bottom", "right"):
                    b = tcBorders.find(qn(f"w:{b_name}"))
                    assert b is not None, f"Cell ({r_idx}, {c_idx}) missing border {b_name}"
                    assert b.get(qn("w:val")) == "nil", f"Cell ({r_idx}, {c_idx}) border {b_name} not nil"

        # 2. Verify raw XML <w:tc> elements in rows 21..32 (including vMerge=continue rows 24..26 for tev.prpd)
        for r_idx in range(21, 33):
            row = table.rows[r_idx]
            col_idx = 0
            for tc in row._tr.findall(qn("w:tc")):
                tcPr = tc.find(qn("w:tcPr"))
                gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                span = int(gridSpan_elem.attrib.get(qn("w:val"), 1)) if gridSpan_elem is not None else 1
                col_start = col_idx
                col_end = col_idx + span - 1
                col_idx += span

                # TEV cells (cols 11..22)
                if col_start >= 11 and col_end <= 22:
                    txt = "".join(tc.itertext()).strip()
                    assert txt == "", f"Row {r_idx} cols {col_start}..{col_end} XML text not empty: {txt}"
                    shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
                    assert shd is None or shd.get(qn("w:fill")) in (None, "FFFFFF", "clear")
                    tcBorders = tcPr.find(qn("w:tcBorders")) if tcPr is not None else None
                    assert tcBorders is not None, f"Row {r_idx} cols {col_start}..{col_end} missing tcBorders"
                    for b_name in ("top", "left", "bottom", "right"):
                        b = tcBorders.find(qn(f"w:{b_name}"))
                        assert b is not None, f"Row {r_idx} cols {col_start}..{col_end} missing border {b_name}"
                        assert b.get(qn("w:val")) == "nil", f"Row {r_idx} cols {col_start}..{col_end} border {b_name} not nil"

                # Adjacent left spacer (ending at col 10)
                elif col_end == 10:
                    tcBorders = tcPr.find(qn("w:tcBorders")) if tcPr is not None else None
                    assert tcBorders is not None, f"Row {r_idx} col 10 spacer missing tcBorders"
                    b_right = tcBorders.find(qn("w:right"))
                    assert b_right is not None, f"Row {r_idx} col 10 spacer missing right border"
                    assert b_right.get(qn("w:val")) == "nil", f"Row {r_idx} col 10 spacer right border not nil"

                # Adjacent right spacer (starting at col 23)
                elif col_start == 23:
                    tcBorders = tcPr.find(qn("w:tcBorders")) if tcPr is not None else None
                    assert tcBorders is not None, f"Row {r_idx} col 23 spacer missing tcBorders"
                    b_left = tcBorders.find(qn("w:left"))
                    assert b_left is not None, f"Row {r_idx} col 23 spacer missing left border"
                    assert b_left.get(qn("w:val")) == "nil", f"Row {r_idx} col 23 spacer left border not nil"

        # 3. Verify Row 20 top spacer bottom borders facing TEV block
        row_20 = table.rows[20]
        col_idx = 0
        for tc in row_20._tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1)) if gridSpan_elem is not None else 1
            col_start = col_idx
            col_end = col_idx + span - 1
            col_idx += span
            if max(col_start, 11) <= min(col_end, 22):
                tcBorders = tcPr.find(qn("w:tcBorders")) if tcPr is not None else None
                assert tcBorders is not None
                b_bot = tcBorders.find(qn("w:bottom"))
                assert b_bot is not None
                assert b_bot.get(qn("w:val")) == "nil", "Row 20 spacer bottom border facing TEV not nil"

        # 4. Verify Row 33 bottom spacer top borders facing TEV block
        row_33 = table.rows[33]
        col_idx = 0
        for tc in row_33._tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
            span = int(gridSpan_elem.attrib.get(qn("w:val"), 1)) if gridSpan_elem is not None else 1
            col_start = col_idx
            col_end = col_idx + span - 1
            col_idx += span
            if max(col_start, 11) <= min(col_end, 22):
                tcBorders = tcPr.find(qn("w:tcBorders")) if tcPr is not None else None
                assert tcBorders is not None
                b_top = tcBorders.find(qn("w:top"))
                assert b_top is not None
                assert b_top.get(qn("w:val")) == "nil", "Row 33 spacer top border facing TEV not nil"

        # 5. Verify US cells (rows 21..32, cols 1..8) are NOT blanked / corrupted
        cell_us_head = table.rows[21].cells[1]
        assert "ULTRASOUND" in cell_us_head.text.upper()
        tc_us_prpd = table.rows[23].cells[1]._tc
        tcBorders = tc_us_prpd.find(qn("w:tcPr")).find(qn("w:tcBorders"))
        assert tcBorders.find(qn("w:top")).get(qn("w:val")) == "single"
        assert tcBorders.find(qn("w:right")).get(qn("w:val")) == "single"

    def test_blank_swg_tev_cells_idempotent(self, swg_panel_doc: docx.Document):
        """Blanking twice should not raise errors."""
        blank_swg_tev_cells(swg_panel_doc)
        blank_swg_tev_cells(swg_panel_doc)
        table = swg_panel_doc.tables[0]
        cell = table.rows[25].cells[15]
        assert cell.text.strip() == ""

    def test_blank_swg_tev_cells_supports_iterable_of_documents(self):
        """blank_swg_tev_cells must support being passed a list/tuple of Document objects."""
        doc1 = docx.Document(SWG_PANEL_TEMPLATE)
        doc2 = docx.Document(SWG_PANEL_TEMPLATE)
        blank_swg_tev_cells([doc1, doc2])
        assert doc1.tables[0].rows[23].cells[11].text.strip() == ""
        assert doc2.tables[0].rows[23].cells[11].text.strip() == ""

    def test_clear_cell_text_removes_drawings_from_raw_tc(self):
        """clear_cell_text must strip drawing and pict elements from raw <w:tc>."""
        from docx.oxml import OxmlElement
        from src.core.shading import clear_cell_text
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=1)
        tc = table.cell(0, 0)._tc
        p = tc.find(qn("w:p"))
        drawing = OxmlElement("w:drawing")
        p.append(drawing)
        assert len(tc.findall(".//" + qn("w:drawing"))) == 1
        clear_cell_text(tc)
        assert len(tc.findall(".//" + qn("w:drawing"))) == 0

    def test_resolve_tc_strict_matching(self):
        """_resolve_tc extracts valid <w:tc> and rejects non-tc nodes or strings."""
        from docx.oxml import OxmlElement
        from src.core.shading import _resolve_tc
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        tc = cell._tc

        assert _resolve_tc(cell) is tc
        assert _resolve_tc(tc) is tc
        assert _resolve_tc("some_string") is None
        assert _resolve_tc(None) is None

        # Ruby text container (<w:rtc>) must not resolve as <w:tc>
        rtc = OxmlElement("w:rtc")
        assert _resolve_tc(rtc) is None

    def test_blank_swg_tev_cells_ignores_non_24_column_tables(self):
        """Tables not matching exactly 24 columns (e.g. 23 or 25 cols) must not be modified."""
        # 23-column table
        tx_doc = docx.Document("templates/FULL REPORT/NORMAL IR US TEV/tx-hv-sides.docx")
        t23 = tx_doc.tables[0]
        assert len(t23.columns) == 23
        us_text_before = t23.rows[21].cells[1].text
        blank_swg_tev_cells(tx_doc)
        assert t23.rows[21].cells[1].text == us_text_before

        # 25-column synthetic table
        synthetic_doc = docx.Document()
        t25 = synthetic_doc.add_table(rows=35, cols=25)
        t25.rows[21].cells[11].text = "DO_NOT_CLEAR"
        blank_swg_tev_cells(synthetic_doc)
        assert t25.rows[21].cells[11].text == "DO_NOT_CLEAR"

    def test_blank_swg_tev_cells_quick_report_template(self):
        """Verify Quick Report template swg-panel.docx TEV PRPD quadrant borders are all nil."""
        if not QR_SWG_PANEL_TEMPLATE.exists():
            pytest.skip("QR template swg-panel.docx not found.")
        doc = docx.Document(QR_SWG_PANEL_TEMPLATE)
        blank_swg_tev_cells(doc)
        table = doc.tables[0]
        # Inspect rows 21..32 raw XML <w:tc> elements for cols 11..22 (including {{ tev.prpd }})
        for r_idx in range(21, 33):
            row = table.rows[r_idx]
            col_idx = 0
            for tc in row._tr.findall(qn("w:tc")):
                tcPr = tc.find(qn("w:tcPr"))
                gridSpan_elem = tcPr.find(qn("w:gridSpan")) if tcPr is not None else None
                span = int(gridSpan_elem.attrib.get(qn("w:val"), 1)) if gridSpan_elem is not None else 1
                col_start = col_idx
                col_end = col_idx + span - 1
                col_idx += span

                if col_start >= 11 and col_end <= 22:
                    assert "".join(tc.itertext()).strip() == ""
                    tcBorders = tcPr.find(qn("w:tcBorders"))
                    assert tcBorders is not None
                    for b_name in ("top", "left", "bottom", "right"):
                        b = tcBorders.find(qn(f"w:{b_name}"))
                        assert b is not None and b.get(qn("w:val")) == "nil"
                elif col_end == 10:
                    tcBorders = tcPr.find(qn("w:tcBorders"))
                    assert tcBorders is not None
                    assert tcBorders.find(qn("w:right")).get(qn("w:val")) == "nil"
                elif col_start == 23:
                    tcBorders = tcPr.find(qn("w:tcBorders"))
                    assert tcBorders is not None
                    assert tcBorders.find(qn("w:left")).get(qn("w:val")) == "nil"


# ==============================================================================
# 3. PRPD Optimization Tests
# ==============================================================================

class TestPrpdOptimization:
    """Verify generate_prpd_graphs_for_swg_panel skips TEV for non-TEV compartments."""

    def test_include_tev_false_skips_tev(self, tmp_path: Path):
        mock_survey = tmp_path / "mock_survey"
        mock_feeder = mock_survey / "P01_INCOMING"
        mock_feeder.mkdir(parents=True)

        with patch("src.quick_report.prpd.find_swg_feeder_survey_dir", return_value=mock_feeder), \
             patch("src.quick_report.prpd.decode_tev_event_data") as mock_decode_tev, \
             patch("src.quick_report.prpd.render_prpd_option_c_image") as mock_render_c:

            us_png, tev_png = generate_prpd_graphs_for_swg_panel(
                survey_root=mock_survey,
                panel_no=1,
                output_dir=tmp_path,
                mode="option_b",
                compartment="CABLE ENTRY",
                include_tev=False,
            )

            # TEV generation must be skipped completely
            assert tev_png is None
            mock_decode_tev.assert_not_called()
            mock_render_c.assert_not_called()

    def test_compartment_cable_entry_skips_tev_automatically(self, tmp_path: Path):
        mock_survey = tmp_path / "mock_survey"
        mock_feeder = mock_survey / "P01_INCOMING"
        mock_feeder.mkdir(parents=True)

        with patch("src.quick_report.prpd.find_swg_feeder_survey_dir", return_value=mock_feeder), \
             patch("src.quick_report.prpd.decode_tev_event_data") as mock_decode_tev, \
             patch("src.quick_report.prpd.render_prpd_option_c_image") as mock_render_c:

            us_png, tev_png = generate_prpd_graphs_for_swg_panel(
                survey_root=mock_survey,
                panel_no=1,
                output_dir=tmp_path,
                mode="option_c",
                compartment="CABLE ENTRY",
                include_tev=True,  # Even if True, compartment logic suppresses it
            )

            assert tev_png is None
            mock_decode_tev.assert_not_called()
            # If render_prpd_option_c_image is called, it must not be called with TEV.html
            for call in mock_render_c.call_args_list:
                html_arg = str(call.kwargs.get("html_file") or call.args[0] if call.args else "")
                assert "TEV.html" not in html_arg


# ==============================================================================
# 4. Full Report Scan Adapter Integration Tests
# ==============================================================================

class TestFullReportScanAdapterTevBlanking:
    """Validate SwitchgearScanAdapter per-compartment TEV evaluation and rendering."""

    @pytest.fixture
    def tamco_scan_spec(self) -> SwitchgearScanSpec:
        panel = SwitchgearPanelScanSpec(
            panel_no=1,
            name="INCOMING 1",
            panel_type="TAMCO RMU",
            status="CLOSE",
            load_amp="120",
            heater_amp="0.8",
            cable_type="3Cx300mm2 XLPE",
            serial_no="SN-TAMCO-01",
            tev_reading="18",
            compartments=("CABLE COMPARTMENT", "CABLE ENTRY"),
        )
        return SwitchgearScanSpec(
            switchgear_type="RMU SF6",
            manufacturer="TAMCO",
            model="GV3",
            rating="11kV 630A",
            serial_no="SN-BOARD-01",
            category=SwitchgearCategory.TAMCO_LUCY,
            panels=[panel],
        )

    def test_tamco_cable_compartment_vs_cable_entry(self, tamco_scan_spec: SwitchgearScanSpec, tmp_path: Path):
        adapter = SwitchgearScanAdapter(
            swg=tamco_scan_spec,
            substation_info={"name_erms": "PE TEST MOCK"},
            project_technologies=["IR", "US", "TEV"],
        )
        items = adapter.adapt()
        panel_items = [it for it in items if not it.is_overview]
        assert len(panel_items) == 2

        item_cable_comp = panel_items[0]
        item_cable_entry = panel_items[1]

        # Cable Compartment: TEV is active
        assert item_cable_comp.component_name == "CABLE COMPARTMENT"
        assert item_cable_comp.context.get("__blank_tev__") is False
        assert item_cable_comp.context.get("is_tev_active") is True
        assert item_cable_comp.context["panel"]["tev"]["reading"] != ""

        # Cable Entry: TEV is blanked
        assert item_cable_entry.component_name == "CABLE ENTRY"
        assert item_cable_entry.context.get("__blank_tev__") is True
        assert item_cable_entry.context.get("is_tev_active") is False
        assert item_cable_entry.context["panel"]["tev"]["reading"] == ""
        assert item_cable_entry.context["panel"]["tev"]["char"] == ""
        assert item_cable_entry.context["panel"]["tev"]["ppc"] == ""
        assert item_cable_entry.context["panel"]["tev"]["bg"] == ""
        assert item_cable_entry.context["panel"]["tev"]["severity"] == ""

        # Render both pages and check OpenXML borders
        renderer = FullReportScanPageRendererCore()
        out_cable_comp = tmp_path / "cable_comp.docx"
        out_cable_entry = tmp_path / "cable_entry.docx"

        item_cable_comp.render(out_cable_comp, renderer=renderer)
        item_cable_entry.render(out_cable_entry, renderer=renderer)

        doc_entry = docx.Document(out_cable_entry)
        t_entry = doc_entry.tables[0]
        # Cable Entry TEV cells must have nil borders and empty text
        for r_idx in range(21, 33):
            for c_idx in range(11, 23):
                cell = t_entry.rows[r_idx].cells[c_idx]
                assert cell.text.strip() == ""
                tcBorders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
                assert tcBorders is not None
                assert tcBorders.find(qn("w:top")).get(qn("w:val")) == "nil"

    def test_contract_without_tev_blanks_all_compartments(self, tamco_scan_spec: SwitchgearScanSpec, tmp_path: Path):
        """If project awarded technologies is only ['IR', 'US'], even Cable Compartment is blanked."""
        adapter = SwitchgearScanAdapter(
            swg=tamco_scan_spec,
            substation_info={"name_erms": "PE TEST MOCK"},
            project_technologies=["IR", "US"],
        )
        items = adapter.adapt()
        panel_items = [it for it in items if not it.is_overview]

        for it in panel_items:
            assert it.context.get("__blank_tev__") is True
            assert it.context.get("is_tev_active") is False
            assert it.context["panel"]["tev"]["reading"] == ""

        # Render Cable Compartment docx and verify nil borders
        renderer = FullReportScanPageRendererCore()
        out_file = tmp_path / "no_tev_cable_comp.docx"
        panel_items[0].render(out_file, renderer=renderer)

        doc = docx.Document(out_file)
        t = doc.tables[0]
        cell_tev = t.rows[25].cells[15]
        assert cell_tev.text.strip() == ""
        tcBorders = cell_tev._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
        assert tcBorders is not None
        assert tcBorders.find(qn("w:top")).get(qn("w:val")) == "nil"

    def test_vcb_5_standard_compartments_tev_matrix(self):
        """Verify VCB standard 5 compartments: Breaker, Cable, PT active; Busbar, Secondary blanked."""
        from src.full_report.models import VCB_STANDARD_COMPARTMENTS
        panel = SwitchgearPanelScanSpec(
            panel_no=1,
            name="VCB 1",
            panel_type="VCB",
            status="CLOSE",
            load_amp="120",
            heater_amp="0.8",
            serial_no="SN-01",
            tev_reading="18",
            compartments=VCB_STANDARD_COMPARTMENTS,
        )
        swg = SwitchgearScanSpec(
            switchgear_type="VCB",
            manufacturer="TAMCO",
            model="GV3",
            rating="11kV 630A",
            serial_no="SN-BOARD-01",
            category=SwitchgearCategory.VCB,
            panels=[panel],
        )
        adapter = SwitchgearScanAdapter(
            swg=swg,
            substation_info={"name_erms": "PE TEST MOCK"},
            project_technologies=["IR", "US", "TEV"],
        )
        res = adapter.adapt()
        panel_items = {it.component_name: it for it in res.items if not it.is_overview}
        assert len(panel_items) == 5

        # Active TEV compartments
        for comp in ("BREAKER COMPARTMENT", "CABLE COMPARTMENT", "PT COMPARTMENT"):
            assert panel_items[comp].context.get("__blank_tev__") is False
            assert panel_items[comp].context.get("is_tev_active") is True
            assert panel_items[comp].context["panel"]["tev"]["reading"] != ""

        # Blanked compartments
        for comp in ("BUSBAR COMPARTMENT", "SECONDARY COMPARTMENT"):
            assert panel_items[comp].context.get("__blank_tev__") is True
            assert panel_items[comp].context.get("is_tev_active") is False
            assert panel_items[comp].context["panel"]["tev"]["reading"] == ""

    def test_vcb_5_transition_compartments_tev_matrix(self):
        """Verify VCB transition 5 compartments: PT active; Front, Rear, Busbar, Secondary blanked."""
        from src.full_report.models import VCB_TRANSITION_COMPARTMENTS
        panel = SwitchgearPanelScanSpec(
            panel_no=1,
            name="TRANSITION PANEL",
            panel_type="VCB",
            status="CLOSE",
            load_amp="0",
            heater_amp="0.8",
            serial_no="SN-TRANS-01",
            tev_reading="18",
            compartments=VCB_TRANSITION_COMPARTMENTS,
        )
        swg = SwitchgearScanSpec(
            switchgear_type="VCB",
            manufacturer="TAMCO",
            model="GV3",
            rating="11kV 630A",
            serial_no="SN-BOARD-01",
            category=SwitchgearCategory.VCB,
            panels=[panel],
        )
        adapter = SwitchgearScanAdapter(
            swg=swg,
            substation_info={"name_erms": "PE TEST MOCK"},
            project_technologies=["IR", "US", "TEV"],
        )
        res = adapter.adapt()
        panel_items = {it.component_name: it for it in res.items if not it.is_overview}
        assert len(panel_items) == 5

        # Active TEV compartment (only PT)
        assert panel_items["PT COMPARTMENT"].context.get("__blank_tev__") is False
        assert panel_items["PT COMPARTMENT"].context.get("is_tev_active") is True
        assert panel_items["PT COMPARTMENT"].context["panel"]["tev"]["reading"] != ""

        # Blanked compartments (Front, Rear, Busbar, Secondary)
        for comp in ("FRONT COMPARTMENT", "REAR COMPARTMENT", "BUSBAR COMPARTMENT", "SECONDARY COMPARTMENT"):
            assert panel_items[comp].context.get("__blank_tev__") is True
            assert panel_items[comp].context.get("is_tev_active") is False
            assert panel_items[comp].context["panel"]["tev"]["reading"] == ""


# ==============================================================================
# 5. Quick Report CBM Render Integration Tests
# ==============================================================================

class TestQuickReportCbmRenderTevBlanking:
    """Validate Quick Report _build_swg_render_context and _render_docx_template blanking."""

    def test_qr_cable_entry_defect_blanks_tev(self, tmp_path: Path):
        if not QR_SWG_PANEL_TEMPLATE.exists():
            pytest.skip("Template swg-panel.docx not found.")

        rec = CbmDefectRecord(
            equipment="VCB PANEL 1",
            technology="US",
            defect_area="Cable Entry",
            us_reading="18.5",
            us_char="TRACKING",
            tev_reading="25.0",
            tev_char="CONTINUOUS",
        )
        context = _build_swg_render_context(rec, overview=False)
        assert context.get("__blank_tev__") is True
        assert context.get("is_tev_active") is False
        assert context["panel"]["tev"]["reading"] == ""
        assert context["panel"]["tev"]["char"] == ""

        out_path = tmp_path / "qr_cable_entry_blanked.docx"
        _render_docx_template(QR_SWG_PANEL_TEMPLATE, out_path, context)

        doc = docx.Document(out_path)
        t = doc.tables[0]
        # Row 30, Col 18 is TEV severity cell; should be cleared and borders nil
        cell_tev = t.rows[30].cells[18]
        assert cell_tev.text.strip() == ""
        assert get_cell_shading(cell_tev) is None
        tcBorders = cell_tev._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
        assert tcBorders is not None
        assert tcBorders.find(qn("w:top")).get(qn("w:val")) == "nil"

    def test_qr_contract_without_tev_blanks_spout(self, tmp_path: Path):
        if not QR_SWG_PANEL_TEMPLATE.exists():
            pytest.skip("Template swg-panel.docx not found.")

        rec = CbmDefectRecord(
            equipment="VCB PANEL 1",
            technology="US",
            defect_area="Spout",
            us_reading="18.5",
            us_char="TRACKING",
            tev_reading="25.0",
        )
        pe_info = {"project_technologies": ["IR", "US"]}
        context = _build_swg_render_context(rec, overview=False, pe_info=pe_info)
        assert context.get("__blank_tev__") is True
        assert context.get("is_tev_active") is False
        assert context["panel"]["tev"]["reading"] == ""

        out_path = tmp_path / "qr_spout_contract_no_tev.docx"
        _render_docx_template(QR_SWG_PANEL_TEMPLATE, out_path, context)

        doc = docx.Document(out_path)
        t = doc.tables[0]
        cell_tev = t.rows[30].cells[18]
        assert cell_tev.text.strip() == ""
        assert get_cell_shading(cell_tev) is None
        tcBorders = cell_tev._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
        assert tcBorders is not None
        assert tcBorders.find(qn("w:top")).get(qn("w:val")) == "nil"

    def test_qr_spout_under_3_tech_contract_preserves_tev(self, tmp_path: Path):
        if not QR_SWG_PANEL_TEMPLATE.exists():
            pytest.skip("Template swg-panel.docx not found.")

        rec = CbmDefectRecord(
            equipment="VCB PANEL 1",
            technology="US",
            defect_area="Spout",
            us_reading="18.5",
            us_char="TRACKING",
            tev_reading="15.0",
            tev_char="NORMAL",
        )
        pe_info = {"project_technologies": ["IR", "US", "TEV"]}
        context = _build_swg_render_context(rec, overview=False, pe_info=pe_info)
        assert context.get("__blank_tev__") is False
        assert context.get("is_tev_active") is True
        assert context["panel"]["tev"]["reading"] != ""

        out_path = tmp_path / "qr_spout_3tech_preserved.docx"
        _render_docx_template(QR_SWG_PANEL_TEMPLATE, out_path, context)

        doc = docx.Document(out_path)
        t = doc.tables[0]
        # Row 30, Col 18 is TEV severity cell; healthy so shaded 00B050
        cell_tev = t.rows[30].cells[18]
        assert cell_tev.text.strip() == ""
        assert get_cell_shading(cell_tev) == "00B050"

    def test_qr_extract_project_technologies_from_metadata_objects(self):
        """Ensure ContractScope extracts awarded technologies from metadata or project_metadata objects."""
        from src.core.contract import ContractScope

        class MockMeta:
            technologies = ("IR", "US")

        pe_info_meta = {"metadata": MockMeta()}
        assert ContractScope.from_source(pe_info_meta).awarded_technologies == frozenset({"IR", "US"})

        pe_info_proj_meta = {"project_metadata": MockMeta()}
        assert ContractScope.from_source(pe_info_proj_meta).awarded_technologies == frozenset({"IR", "US"})

