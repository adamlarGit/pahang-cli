"""Unit test suite for Core Scan Page Renderer & Dynamic Shading Engine (Ticket #30 / T4.2a).

Verifies:
- FullReportScanPageRendererCore wrapping docxtpl.DocxTemplate.
- Dynamic shading of technology severity cells (IR, US, TEV) Green 00B050 (healthy) or Red EE0000 (defective), clearing text.
- Dynamic shading of Analysis & Recommendation banner Green 00B050 ('No Anomaly.') or Red EE0000 (defect forwarding).
- XML cell shading attributes (w:fill="00B050" and w:fill="EE0000") on mock and generated documents.
"""

from __future__ import annotations

from pathlib import Path
import tempfile
from typing import Any
import zipfile

import docx
from docx.oxml.ns import qn
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
from PIL import Image
import pytest

from src.full_report.scan_render import (
    BANNER_HEALTHY_ANALYSIS,
    COLOR_DEFECT,
    COLOR_HEALTHY,
    FullReportScanPageRendererCore,
    SEVERITY_MARKER_IR,
    SEVERITY_MARKER_TEV,
    SEVERITY_MARKER_US,
    _normalize_technologies,
    apply_banner_shading,
    apply_technology_severity_shading,
    detect_cell_technology,
    get_cell_shading,
    is_defect_forwarding_text,
    is_healthy_banner_text,
    render_scan_page,
)

TEMPLATES_DIR = Path("templates/FULL REPORT/NORMAL IR US TEV")


def _get_cell_fill(cell: Any) -> str | None:
    """Read w:fill attribute from a table cell's tcPr/w:shd XML element."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    return shd.attrib.get(qn("w:fill"))


def _read_document_xml(docx_path: Path) -> str:
    """Read word/document.xml content directly from .docx zip archive."""
    with zipfile.ZipFile(docx_path, "r") as z:
        return z.read("word/document.xml").decode("utf-8")


@pytest.fixture(scope="module")
def dummy_image_file() -> Path:
    """Create a temporary test image file with non-uniform pixels for image binding."""
    tmp = Path(tempfile.gettempdir()) / "test_scan_render_dummy.png"
    img = Image.new("RGB", (60, 45), color=(0, 176, 80))
    for x in range(10):
        img.putpixel((x, 0), (255, 0, 0))
        img.putpixel((x, 1), (0, 0, 255))
    img.save(tmp)
    return tmp


# ==============================================================================
# 1. Technology Detection & Text Matching Unit Tests
# ==============================================================================

def test_detect_cell_technology_markers():
    """Detect technology severity cells from Jinja sentinels and placeholders."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=6)

    t.cell(0, 0).text = SEVERITY_MARKER_IR
    t.cell(0, 1).text = "{{ ir.severity }}"
    t.cell(0, 2).text = SEVERITY_MARKER_US
    t.cell(0, 3).text = "{{ us.severity }}"
    t.cell(0, 4).text = SEVERITY_MARKER_TEV
    t.cell(0, 5).text = "{{ tev.severity }}"

    assert detect_cell_technology(t.cell(0, 0)) == "IR"
    assert detect_cell_technology(t.cell(0, 1)) == "IR"
    assert detect_cell_technology(t.cell(0, 2)) == "US"
    assert detect_cell_technology(t.cell(0, 3)) == "US"
    assert detect_cell_technology(t.cell(0, 4)) == "TEV"
    assert detect_cell_technology(t.cell(0, 5)) == "TEV"


def test_detect_cell_technology_tokens_and_none():
    """Detect technology from standard token names and return None for normal cells."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=5)

    t.cell(0, 0).text = "IR_SEVERITY"
    t.cell(0, 1).text = "U/S"
    t.cell(0, 2).text = "TEV"
    t.cell(0, 3).text = "PE SUBSTATION"
    t.cell(0, 4).text = ""

    assert detect_cell_technology(t.cell(0, 0)) == "IR"
    assert detect_cell_technology(t.cell(0, 1)) == "US"
    assert detect_cell_technology(t.cell(0, 2)) == "TEV"
    assert detect_cell_technology(t.cell(0, 3)) is None
    assert detect_cell_technology(t.cell(0, 4)) is None


def test_banner_text_helpers():
    """Verify healthy and defect forwarding text detection helpers."""
    assert is_healthy_banner_text("Analysis: No Anomaly.")
    assert is_healthy_banner_text("No Anomaly.")
    assert is_healthy_banner_text("No defect.")
    assert not is_healthy_banner_text("-")
    assert not is_healthy_banner_text("NORMAL")
    assert not is_healthy_banner_text("None")
    assert not is_healthy_banner_text("TIADA ANOMALI")
    assert not is_healthy_banner_text("Analysis: Please refer to the following page for details defect.")

    assert is_defect_forwarding_text("Analysis: Please refer to the following page for details defect.")
    assert is_defect_forwarding_text("Recommendation: Please refer to the following page for details defect.")
    assert is_defect_forwarding_text("Please refer to the following page for detail defect.")
    assert not is_defect_forwarding_text("No Anomaly.")


# ==============================================================================
# 2. DOM Post-Processor Unit Tests (Mock Documents & Tables)
# ==============================================================================

def test_apply_technology_severity_shading_single_cells():
    """Verify shading and text clearing on single cells for IR, US, and TEV."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)

    cell_ir = t.cell(0, 0)
    cell_ir.text = "{{ ir.severity }}"
    cell_us = t.cell(0, 1)
    cell_us.text = "{{ us.severity }}"
    cell_tev = t.cell(0, 2)
    cell_tev.text = "{{ tev.severity }}"

    # Apply healthy shading
    apply_technology_severity_shading(cell_ir, is_defective=False)
    apply_technology_severity_shading(cell_us, is_defective=False)
    apply_technology_severity_shading(cell_tev, is_defective=False)

    assert cell_ir.text == ""
    assert _get_cell_fill(cell_ir) == COLOR_HEALTHY
    assert "00B050" in cell_ir._tc.xml

    assert cell_us.text == ""
    assert _get_cell_fill(cell_us) == COLOR_HEALTHY
    assert "00B050" in cell_us._tc.xml

    assert cell_tev.text == ""
    assert _get_cell_fill(cell_tev) == COLOR_HEALTHY
    assert "00B050" in cell_tev._tc.xml

    # Apply defective shading
    cell_ir_def = t.cell(1, 0)
    cell_ir_def.text = SEVERITY_MARKER_IR
    cell_us_def = t.cell(1, 1)
    cell_us_def.text = SEVERITY_MARKER_US
    cell_tev_def = t.cell(1, 2)
    cell_tev_def.text = SEVERITY_MARKER_TEV

    apply_technology_severity_shading(cell_ir_def, is_defective=True)
    apply_technology_severity_shading(cell_us_def, is_defective=True)
    apply_technology_severity_shading(cell_tev_def, is_defective=True)

    assert cell_ir_def.text == ""
    assert _get_cell_fill(cell_ir_def) == COLOR_DEFECT
    assert "EE0000" in cell_ir_def._tc.xml

    assert cell_us_def.text == ""
    assert _get_cell_fill(cell_us_def) == COLOR_DEFECT
    assert "EE0000" in cell_us_def._tc.xml

    assert cell_tev_def.text == ""
    assert _get_cell_fill(cell_tev_def) == COLOR_DEFECT
    assert "EE0000" in cell_tev_def._tc.xml


def test_apply_technology_severity_shading_mock_table_healthy():
    """Verify table-level shading when all technologies are healthy."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "{{ ir.severity }}"
    t.cell(0, 1).text = "{{ us.severity }}"
    t.cell(0, 2).text = "{{ tev.severity }}"

    apply_technology_severity_shading(t, defective_technologies=None)

    for c_idx in range(3):
        cell = t.cell(0, c_idx)
        assert cell.text == ""
        assert _get_cell_fill(cell) == "00B050"
        assert 'w:fill="00B050"' in cell._tc.xml


def test_apply_technology_severity_shading_mock_table_selective_defects():
    """Verify table-level selective shading when specific technologies are defective."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = SEVERITY_MARKER_IR
    t.cell(0, 1).text = SEVERITY_MARKER_US
    t.cell(0, 2).text = SEVERITY_MARKER_TEV

    # IR and TEV defective, US healthy
    apply_technology_severity_shading(t, defective_technologies={"IR", "TEV"})

    assert t.cell(0, 0).text == ""
    assert _get_cell_fill(t.cell(0, 0)) == "EE0000"

    assert t.cell(0, 1).text == ""
    assert _get_cell_fill(t.cell(0, 1)) == "00B050"

    assert t.cell(0, 2).text == ""
    assert _get_cell_fill(t.cell(0, 2)) == "EE0000"


def test_apply_technology_severity_shading_string_syntax():
    """Verify defective_technologies supports string format 'IR, US' and 'IR+TEV'."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "{{ ir.severity }}"
    t.cell(0, 1).text = "{{ us.severity }}"
    t.cell(0, 2).text = "{{ tev.severity }}"

    apply_technology_severity_shading(t, defective_technologies="US+TEV")

    assert _get_cell_fill(t.cell(0, 0)) == "00B050"
    assert _get_cell_fill(t.cell(0, 1)) == "EE0000"
    assert _get_cell_fill(t.cell(0, 2)) == "EE0000"


def test_apply_banner_shading_single_cells():
    """Verify banner shading on individual cells for healthy and defective prose."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=2)

    healthy_cell = t.cell(0, 0)
    healthy_cell.text = "Analysis: No Anomaly."
    defective_cell = t.cell(0, 1)
    defective_cell.text = "Analysis: Please refer to the following page for details defect."

    apply_banner_shading(healthy_cell)
    apply_banner_shading(defective_cell)

    assert _get_cell_fill(healthy_cell) == "00B050"
    assert 'w:fill="00B050"' in healthy_cell._tc.xml

    assert _get_cell_fill(defective_cell) == "EE0000"
    assert 'w:fill="EE0000"' in defective_cell._tc.xml


def test_apply_banner_shading_mock_table_healthy_and_defective():
    """Verify table-level banner shading matching standard scan page banner rows."""
    # Healthy Table
    doc_healthy = docx.Document()
    t_h = doc_healthy.add_table(rows=3, cols=1)
    t_h.cell(0, 0).text = "Analysis & Recommendations:"
    t_h.cell(1, 0).text = "Analysis:  No Anomaly."
    t_h.cell(2, 0).text = "Recommendation:  -"

    apply_banner_shading(t_h)

    assert _get_cell_fill(t_h.cell(0, 0)) is None  # Heading unaffected
    assert _get_cell_fill(t_h.cell(1, 0)) == "00B050"  # Analysis Green
    assert 'w:fill="00B050"' in t_h.cell(1, 0)._tc.xml
    assert _get_cell_fill(t_h.cell(2, 0)) == "00B050"  # Recommendation Green
    assert 'w:fill="00B050"' in t_h.cell(2, 0)._tc.xml

    # Defective Table
    doc_def = docx.Document()
    t_d = doc_def.add_table(rows=3, cols=1)
    t_d.cell(0, 0).text = "Analysis & Recommendations:"
    t_d.cell(1, 0).text = "Analysis:  Please refer to the following page for details defect."
    t_d.cell(2, 0).text = "Recommendation:  Please refer to the following page for details defect."

    apply_banner_shading(t_d)

    assert _get_cell_fill(t_d.cell(0, 0)) is None  # Heading unaffected
    assert _get_cell_fill(t_d.cell(1, 0)) == "EE0000"  # Analysis Red
    assert _get_cell_fill(t_d.cell(2, 0)) == "EE0000"  # Recommendation Red
    assert 'w:fill="EE0000"' in t_d.cell(1, 0)._tc.xml
    assert 'w:fill="EE0000"' in t_d.cell(2, 0)._tc.xml


def test_apply_banner_shading_explicit_is_defective_flag():
    """Verify apply_banner_shading with explicit is_defective flag."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "Analysis: Custom defect analysis text."
    t.cell(1, 0).text = "Recommendation: Custom action."

    apply_banner_shading(t, is_defective=True)
    assert _get_cell_fill(t.cell(0, 0)) == "EE0000"
    assert _get_cell_fill(t.cell(1, 0)) == "EE0000"


# ==============================================================================
# 3. Core Scan Page Renderer Tests (FullReportScanPageRendererCore)
# ==============================================================================

def _build_test_context(dt: DocxTemplate, image_file: Path) -> dict[str, Any]:
    """Build mock context satisfying scanning template variables."""
    img = InlineImage(dt, str(image_file), width=Mm(40))
    return {
        "substation": {
            "name_erms": "PE TEST 11KV",
            "date": "10-08-2026",
            "time": "11:00 AM",
            "ambient": "31.0 °C",
            "humidity": "65%",
        },
        "swg": {
            "type": "RMU SF6",
            "manufacturer": "TAMCO",
            "model": "AIR",
            "rating": "12kV",
            "area": "OVERVIEW",
            "serialnumber": "SN-SWG-001",
        },
        "panel": {
            "name": "FEEDER 1",
            "linknumber": "1",
            "area": "CABLE COMPARTMENT",
            "serialnumber": "SN-PANEL-01",
            "heateramp": "0.5A",
            "breakerstatus": "CLOSE",
            "busbarposition": "TOP",
            "cabletype": "3C 300mm2 XLPE",
            "loadamp": "200A",
        },
        "tx": {
            "manufacturer": "SGB",
            "model": "HERMETICALLY SEAL",
            "rating": "1000kVA",
            "number": "TX1",
            "location": "HV SIDE",
            "area": "HV BUSHING",
            "serialnumber": "TX-100",
            "cabletype": "3C 300mm2 XLPE",
        },
        "fp": {
            "labelsource": "LVDB TX1",
            "manufacturer": "ALGEBRA",
            "model": "J-SLOTTED",
            "rating": "1600A",
            "area": "OVERVIEW",
            "serialnumber": "FP-01",
            "cabletype": "4C 300mm2 PVC",
            "feederno": "F01",
            "loadamp": "120A",
        },
        "battery": {
            "manufacturer": "SUNPOWER",
            "model": "SP-100",
            "number": "1",
            "serialnumber": "BATT-01",
            "rating": "240V/50Hz",
            "area": "OVERVIEW",
        },
        "batt": {
            "manufacturer": "SUNPOWER",
            "model": "SP-100",
            "number": "1",
            "serialnumber": "BATT-01",
            "rating": "240V/50Hz",
            "area": "OVERVIEW",
        },
        "ir": {
            "image": img,
            "sp1": "32.0 °C",
            "sp2": "31.5 °C",
            "ar1": "31.8 °C",
            "reading": "32.0 °C",
            "delta_t": "0.5 °C",
        },
        "us": {
            "reading": "10",
            "char": "NORMAL",
            "prpd": img,
        },
        "tev": {
            "bg": "3",
            "reading": "5",
            "ppc": "0.0",
            "prpd": img,
        },
        "visual": {
            "image": img,
        },
    }


def test_core_renderer_swg_panel_healthy(tmp_path: Path, dummy_image_file: Path):
    """Render swg-panel.docx when healthy: IR/US/TEV are 00B050 and banner is 00B050."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "rendered_swg_panel_healthy.docx"

    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    renderer = FullReportScanPageRendererCore()
    out = renderer.render(
        template_path,
        output_path,
        ctx,
        defective_technologies=None,
    )

    assert out.is_file()
    xml_str = _read_document_xml(out)

    # Document XML must contain healthy green fill
    assert 'w:fill="00B050"' in xml_str
    # Healthy scan page must NOT contain defect red fill
    assert 'w:fill="EE0000"' not in xml_str

    # Inspect generated python-docx elements
    doc = docx.Document(out)
    table = doc.tables[0]

    # Row 18: IR Severity (Cell 3)
    ir_cell = table.rows[18].cells[3]
    assert ir_cell.text.strip() == ""
    assert _get_cell_fill(ir_cell) == "00B050"

    # Row 30: US Severity (Cell 4) & TEV Severity (Cell 18)
    us_cell = table.rows[30].cells[4]
    assert us_cell.text.strip() == ""
    assert _get_cell_fill(us_cell) == "00B050"

    tev_cell = table.rows[30].cells[18]
    assert tev_cell.text.strip() == ""
    assert _get_cell_fill(tev_cell) == "00B050"

    # Banner analysis row
    analysis_cell = table.rows[35].cells[0]
    assert "No Anomaly." in analysis_cell.text
    assert _get_cell_fill(analysis_cell) == "00B050"

    # Banner recommendation row
    rec_cell = table.rows[36].cells[0]
    assert "Recommendation:" in rec_cell.text
    assert _get_cell_fill(rec_cell) == "00B050"


def test_core_renderer_swg_panel_ir_defect(tmp_path: Path, dummy_image_file: Path):
    """Render swg-panel.docx with IR defect: IR is EE0000, US/TEV are 00B050, banner is EE0000."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "rendered_swg_panel_ir_defect.docx"

    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    renderer = FullReportScanPageRendererCore(template_path)
    out = renderer.render(
        output_path,
        ctx,
        defective_technologies={"IR"},
    )

    assert out.is_file()
    xml_str = _read_document_xml(out)

    assert 'w:fill="EE0000"' in xml_str
    assert 'w:fill="00B050"' in xml_str

    doc = docx.Document(out)
    table = doc.tables[0]

    # Row 18: IR Severity is Red EE0000 with text cleared
    ir_cell = table.rows[18].cells[3]
    assert ir_cell.text.strip() == ""
    assert _get_cell_fill(ir_cell) == "EE0000"

    # Row 30: US & TEV Severity are Green 00B050 with text cleared
    us_cell = table.rows[30].cells[4]
    assert us_cell.text.strip() == ""
    assert _get_cell_fill(us_cell) == "00B050"

    tev_cell = table.rows[30].cells[18]
    assert tev_cell.text.strip() == ""
    assert _get_cell_fill(tev_cell) == "00B050"

    # Banner analysis & recommendation rows are Red EE0000
    analysis_cell = table.rows[35].cells[0]
    assert "Please refer to the following page" in analysis_cell.text
    assert _get_cell_fill(analysis_cell) == "EE0000"

    rec_cell = table.rows[36].cells[0]
    assert "Please refer to the following page" in rec_cell.text
    assert _get_cell_fill(rec_cell) == "EE0000"


def test_core_renderer_swg_panel_tev_defect(tmp_path: Path, dummy_image_file: Path):
    """Render swg-panel.docx with TEV defect: TEV is EE0000, IR/US are 00B050, banner is EE0000."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "rendered_swg_panel_tev_defect.docx"

    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    out = FullReportScanPageRendererCore.render_page(
        template_path,
        output_path,
        ctx,
        defective_technologies="TEV",
    )

    assert out.is_file()
    doc = docx.Document(out)
    table = doc.tables[0]

    # IR is healthy Green
    ir_cell = table.rows[18].cells[3]
    assert ir_cell.text.strip() == ""
    assert _get_cell_fill(ir_cell) == "00B050"

    # US is healthy Green
    us_cell = table.rows[30].cells[4]
    assert us_cell.text.strip() == ""
    assert _get_cell_fill(us_cell) == "00B050"

    # TEV is defective Red
    tev_cell = table.rows[30].cells[18]
    assert tev_cell.text.strip() == ""
    assert _get_cell_fill(tev_cell) == "EE0000"

    # Banner is Red
    assert _get_cell_fill(table.rows[35].cells[0]) == "EE0000"


def test_core_renderer_all_templates_smoke(tmp_path: Path, dummy_image_file: Path):
    """Verify FullReportScanPageRendererCore renders all 7 scanning templates cleanly."""
    templates = [
        "swg-overview.docx",
        "swg-panel.docx",
        "tx-overview.docx",
        "tx-hv-sides.docx",
        "tx-lv-sides.docx",
        "fp-overview.docx",
        "battery-overview.docx",
    ]

    for tpl_name in templates:
        tpl_file = TEMPLATES_DIR / tpl_name
        assert tpl_file.is_file()

        out_file = tmp_path / f"rendered_{tpl_name}"
        dt = DocxTemplate(str(tpl_file))
        ctx = _build_test_context(dt, dummy_image_file)

        render_scan_page(tpl_file, out_file, ctx)
        assert out_file.is_file()

        xml = _read_document_xml(out_file)
        if tpl_name not in ("swg-overview.docx", "tx-overview.docx", "fp-overview.docx"):
            assert 'w:fill="00B050"' in xml
        assert "{{" not in xml
        assert "}}" not in xml


def test_core_renderer_error_handling(tmp_path: Path):
    """Verify error handling on non-existent template or invalid paths."""
    renderer = FullReportScanPageRendererCore()
    with pytest.raises(FileNotFoundError):
        renderer.render("non_existent_template.docx", tmp_path / "out.docx", {})

    renderer_with_bad_tpl = FullReportScanPageRendererCore("missing.docx")
    with pytest.raises(FileNotFoundError):
        renderer_with_bad_tpl.render(tmp_path / "out.docx", {})


def test_core_renderer_missing_images_graceful_fallback(tmp_path: Path):
    """Verify missing images in context safely render as '' without crashing."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "missing_images.docx"

    ctx = {
        "substation": {"name_erms": "PE EMPTY IMAGES", "date": "10-08-2026"},
        "swg": {"area": "OVERVIEW"},
        "panel": {"name": "BAY 1", "area": "CABLE COMPARTMENT"},
        "ir": {"image": "non_existent_path.jpg", "reading": "30.0"},
        "visual": {"image": ""},
        "us": {"reading": "10", "prpd": "-"},
        "tev": {"reading": "5", "prpd": None},
    }

    renderer = FullReportScanPageRendererCore()
    out = renderer.render(template_path, output_path, ctx)

    assert out.is_file()
    xml = _read_document_xml(out)
    assert 'w:fill="00B050"' in xml
    assert "{{" not in xml


def test_core_renderer_all_technologies_defective(tmp_path: Path, dummy_image_file: Path):
    """Verify all 3 technology severity cells shaded Red EE0000 when all are defective."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "all_defective.docx"

    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    renderer = FullReportScanPageRendererCore()
    out = renderer.render(
        template_path,
        output_path,
        ctx,
        defective_technologies={"IR", "US", "TEV"},
    )

    doc = docx.Document(out)
    t = doc.tables[0]

    assert _get_cell_fill(t.rows[18].cells[3]) == "EE0000"
    assert _get_cell_fill(t.rows[30].cells[4]) == "EE0000"
    assert _get_cell_fill(t.rows[30].cells[18]) == "EE0000"
    assert _get_cell_fill(t.rows[35].cells[0]) == "EE0000"


def test_core_renderer_raw_filepath_images(tmp_path: Path, dummy_image_file: Path):
    """Verify passing string file paths for images works seamlessly without prior InlineImage wrapping."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "raw_path_images.docx"

    ctx = {
        "substation": {"name_erms": "PE PATH IMAGES", "date": "10-08-2026"},
        "swg": {"area": "OVERVIEW"},
        "panel": {"name": "BAY 1", "area": "CABLE COMPARTMENT"},
        "ir": {"image": str(dummy_image_file), "reading": "30.0"},
        "visual": {"image": dummy_image_file},
        "us": {"reading": "10", "prpd": str(dummy_image_file)},
        "tev": {"reading": "5", "prpd": dummy_image_file},
    }

    out = render_scan_page(template_path, output_path, ctx)
    assert out.is_file()

    doc = docx.Document(out)
    drawings = doc._body._element.xpath(".//w:drawing")
    assert len(drawings) >= 4


def test_apply_shading_document_level():
    """Verify apply_technology_severity_shading and apply_banner_shading accept docx.Document."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "{{ ir.severity }}"
    t.cell(0, 1).text = "Analysis: No Anomaly."

    apply_technology_severity_shading(doc, defective_technologies=None)
    apply_banner_shading(doc)

    assert _get_cell_fill(t.cell(0, 0)) == "00B050"
    assert t.cell(0, 0).text == ""
    assert _get_cell_fill(t.cell(0, 1)) == "00B050"


def test_core_renderer_keyword_argument_dispatch(tmp_path: Path, dummy_image_file: Path):
    """Verify render() works with keyword context= and output_path= patterns."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    renderer = FullReportScanPageRendererCore(template_path)

    # 1. Positional output, keyword context
    out1 = tmp_path / "kw_dispatch_1.docx"
    res1 = renderer.render(out1, context=ctx)
    assert res1.is_file()

    # 2. Both keyword arguments
    out2 = tmp_path / "kw_dispatch_2.docx"
    res2 = renderer.render(output_path=out2, context=ctx)
    assert res2.is_file()

    # 3. Explicit template_path keyword override
    out3 = tmp_path / "kw_dispatch_3.docx"
    res3 = renderer.render(template_path=template_path, output_path=out3, context=ctx)
    assert res3.is_file()


def test_core_renderer_corrupted_image_handling(tmp_path: Path):
    """Verify corrupted / 0-byte image files gracefully fall back to blank without crashing docxtpl."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "corrupt_img_test.docx"

    # Create corrupt non-image file and 0-byte file
    corrupt_img = tmp_path / "corrupt.png"
    corrupt_img.write_text("corrupted non-image content")

    zero_byte_img = tmp_path / "zero.jpg"
    zero_byte_img.write_bytes(b"")

    ctx = {
        "substation": {"name_erms": "PE CORRUPT IMG", "date": "10-08-2026"},
        "swg": {"area": "OVERVIEW"},
        "panel": {"name": "BAY 1", "area": "CABLE COMPARTMENT"},
        "ir": {"image": corrupt_img, "reading": "30.0"},
        "visual": {"image": zero_byte_img},
        "us": {"reading": "10", "prpd": corrupt_img},
        "tev": {"reading": "5", "prpd": zero_byte_img},
    }

    renderer = FullReportScanPageRendererCore(template_path)
    out = renderer.render(output_path, context=ctx)
    assert out.is_file()

    xml = _read_document_xml(out)
    assert 'w:fill="00B050"' in xml
    assert "{{" not in xml


def test_apply_technology_severity_shading_us_slash_normalization():
    """Verify defective_technologies={'U/S'} or ['U/S'] correctly shades US cell Red."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "{{ ir.severity }}"
    t.cell(0, 1).text = "{{ us.severity }}"
    t.cell(0, 2).text = "{{ tev.severity }}"

    apply_technology_severity_shading(t, defective_technologies={"U/S"})

    assert _get_cell_fill(t.cell(0, 0)) == "00B050"
    assert _get_cell_fill(t.cell(0, 1)) == "EE0000"
    assert _get_cell_fill(t.cell(0, 2)) == "00B050"


def test_apply_technology_severity_shading_non_severity_cell_unmodified():
    """Verify non-severity cell passed to apply_technology_severity_shading is left untouched."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    cell.text = "Substation PE TEST 11KV"

    apply_technology_severity_shading(cell)

    assert cell.text == "Substation PE TEST 11KV"
    assert _get_cell_fill(cell) is None


def test_apply_banner_shading_healthy_and_defect_recommendations():
    """Verify banner shading on both Analysis and Recommendation rows."""
    doc = docx.Document()
    t = doc.add_table(rows=3, cols=1)
    t.cell(0, 0).text = "Analysis & Recommendations:"
    t.cell(1, 0).text = "Analysis:  No Anomaly."
    t.cell(2, 0).text = "Recommendation:  -"

    apply_banner_shading(t, is_defective=False)
    assert _get_cell_fill(t.cell(0, 0)) is None
    assert _get_cell_fill(t.cell(1, 0)) == "00B050"
    assert _get_cell_fill(t.cell(2, 0)) == "00B050"

    doc_def = docx.Document()
    t_def = doc_def.add_table(rows=3, cols=1)
    t_def.cell(0, 0).text = "Analysis & Recommendations:"
    t_def.cell(1, 0).text = "Analysis:  Please refer to the following page for details defect."
    t_def.cell(2, 0).text = "Recommendation:  Please refer to the following page for details defect."

    apply_banner_shading(t_def, is_defective=True)
    assert _get_cell_fill(t_def.cell(0, 0)) is None
    assert _get_cell_fill(t_def.cell(1, 0)) == "EE0000"
    assert _get_cell_fill(t_def.cell(2, 0)) == "EE0000"


def test_apply_technology_severity_shading_us_string_and_mixed_delimiters():
    """Verify string representations of U/S ('U/S', 'u/s', 'IR, U/S', 'IR/U/S') normalize properly."""
    assert _normalize_technologies("U/S") == {"US"}
    assert _normalize_technologies("u/s") == {"US"}
    assert _normalize_technologies(["U/S"]) == {"US"}
    assert _normalize_technologies("IR, U/S") == {"IR", "US"}
    assert _normalize_technologies("IR/TEV") == {"IR", "TEV"}
    assert _normalize_technologies("IR/U/S") == {"IR", "US"}
    assert _normalize_technologies("IR+US+TEV") == {"IR", "US", "TEV"}

    doc = docx.Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "{{ ir.severity }}"
    t.cell(0, 1).text = "{{ us.severity }}"
    t.cell(0, 2).text = "{{ tev.severity }}"

    apply_technology_severity_shading(t, defective_technologies="U/S")
    assert _get_cell_fill(t.cell(0, 0)) == "00B050"
    assert _get_cell_fill(t.cell(0, 1)) == "EE0000"
    assert _get_cell_fill(t.cell(0, 2)) == "00B050"


def test_core_renderer_overview_downstream_defect_forwarding_d30(tmp_path: Path, dummy_image_file: Path):
    """Verify D30: Overview page with downstream defect shades banner Red EE0000 while Overview IR remains Green 00B050."""
    template_path = TEMPLATES_DIR / "swg-overview.docx"
    output_path = tmp_path / "overview_forwarding.docx"

    dt = DocxTemplate(str(template_path))
    ctx = _build_test_context(dt, dummy_image_file)

    renderer = FullReportScanPageRendererCore(template_path)
    out = renderer.render(
        output_path,
        ctx,
        is_defective=True,
        overview=True,
    )
    assert out.is_file()

    doc = docx.Document(out)
    table = doc.tables[0]

    # Row 18: IR Severity in swg-overview is '-' (overview photo has no technology severity)
    ir_cell = table.rows[18].cells[3]
    assert ir_cell.text.strip() == "-"
    assert _get_cell_fill(ir_cell) is None

    # Row 35: Analysis banner has forwarding text and is Red EE0000
    analysis_cell = table.rows[35].cells[0]
    assert is_defect_forwarding_text(analysis_cell.text)
    assert _get_cell_fill(analysis_cell) == "EE0000"

    # Row 36: Recommendation banner has forwarding text and is Red EE0000
    rec_cell = table.rows[36].cells[0]
    assert is_defect_forwarding_text(rec_cell.text)
    assert _get_cell_fill(rec_cell) == "EE0000"


def test_core_renderer_safe_path_resolution_and_template_override(tmp_path: Path, dummy_image_file: Path):
    """Verify positional template path override works safely without overwriting the template file."""
    default_tpl = TEMPLATES_DIR / "swg-overview.docx"
    override_tpl = TEMPLATES_DIR / "swg-panel.docx"
    out_file = tmp_path / "override_render.docx"

    dt = DocxTemplate(str(override_tpl))
    ctx = _build_test_context(dt, dummy_image_file)

    # Initial mtime of template
    mtime_before = override_tpl.stat().st_mtime_ns

    renderer = FullReportScanPageRendererCore(default_tpl)
    res = renderer.render(override_tpl, out_file, ctx)
    assert res.is_file()
    assert res.resolve() == out_file.resolve()

    # Template must NOT have been modified
    mtime_after = override_tpl.stat().st_mtime_ns
    assert mtime_before == mtime_after

    # Render attempting to overwrite template directly must raise ValueError
    with pytest.raises(ValueError, match="Output path cannot overwrite template path"):
        renderer.render(override_tpl, override_tpl, ctx)


def test_core_renderer_context_immutability(tmp_path: Path, dummy_image_file: Path):
    """Verify renderer does not mutate the caller's context dictionary or its nested structures."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "immut_test.docx"

    ctx = {
        "substation": {"name_erms": "PE IMMUT TEST", "date": "10-08-2026"},
        "swg": {"area": "OVERVIEW"},
        "panel": {"name": "BAY 1", "area": "CABLE COMPARTMENT"},
        "ir": {"image": str(dummy_image_file), "reading": "30.0", "severity": "NORMAL"},
        "visual": {"image": str(dummy_image_file)},
        "us": {"reading": "10", "prpd": str(dummy_image_file)},
        "tev": {"reading": "5", "prpd": str(dummy_image_file)},
    }

    # Deep snapshot before render
    ir_sev_before = ctx["ir"]["severity"]
    ir_img_before = ctx["ir"]["image"]
    keys_before = set(ctx.keys())

    renderer = FullReportScanPageRendererCore(template_path)
    renderer.render(output_path, ctx)

    assert ctx["ir"]["severity"] == ir_sev_before
    assert ctx["ir"]["image"] == ir_img_before
    assert set(ctx.keys()) == keys_before
    assert "banner" not in ctx


def test_get_cell_shading_helper():
    """Verify get_cell_shading utility reads shading correctly."""
    doc = docx.Document()
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"

    assert get_cell_shading(t.cell(0, 0)) is None
    assert get_cell_shading("not a cell") is None

    apply_technology_severity_shading(t.cell(0, 0), is_defective=False)
    assert get_cell_shading(t.cell(0, 0)) == "00B050"


def test_normalize_technologies_negative_sentinels(tmp_path: Path):
    """Verify negative sentinels ('-', 'NONE', 'NORMAL', 'N/A') are filtered out and do not flag defects."""
    assert _normalize_technologies("-") == set()
    assert _normalize_technologies("NONE") == set()
    assert _normalize_technologies("NORMAL") == set()
    assert _normalize_technologies("N/A") == set()
    assert _normalize_technologies(["-", "NONE", "NORMAL"]) == set()
    assert _normalize_technologies("IR, -") == {"IR"}
    assert _normalize_technologies("NONE+US") == {"US"}

    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "dash_healthy.docx"
    renderer = FullReportScanPageRendererCore(template_path)

    # Passing defective_technologies="-" must NOT turn page into a defect page
    renderer.render(output_path, {}, defective_technologies="-")
    assert output_path.is_file()

    doc = docx.Document(output_path)
    analysis_cell = doc.tables[0].rows[35].cells[0]
    assert BANNER_HEALTHY_ANALYSIS in analysis_cell.text
    assert get_cell_shading(analysis_cell) == "00B050"


def test_render_argument_dispatch_positional_output_with_keyword_template(tmp_path: Path):
    """Verify passing output positionally while specifying template as keyword works without FileNotFoundError."""
    default_tpl = TEMPLATES_DIR / "swg-overview.docx"
    override_tpl = TEMPLATES_DIR / "swg-panel.docx"
    out1 = tmp_path / "kw_tpl1.docx"
    out2 = tmp_path / "kw_tpl2.docx"

    # 1. Empty instance renderer + keyword template_path
    r1 = FullReportScanPageRendererCore()
    res1 = r1.render(out1, context={}, template_path=override_tpl)
    assert res1.is_file()

    # 2. Configured instance renderer + keyword template_path override
    r2 = FullReportScanPageRendererCore(default_tpl)
    res2 = r2.render(out2, template_path=override_tpl, context={})
    assert res2.is_file()

    # 3. Output path pointing to existing directory must raise ValueError
    with pytest.raises(ValueError, match="Output path cannot be an existing directory"):
        r1.render(tmp_path, context={}, template_path=override_tpl)


def test_render_none_context_structures_graceful_handling(tmp_path: Path):
    """Verify context with None values for banner, ir, us, tev does not crash with AttributeError/TypeError."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "none_context.docx"

    renderer = FullReportScanPageRendererCore(template_path)
    out = renderer.render(
        output_path,
        {
            "banner": None,
            "ir": None,
            "us": None,
            "tev": None,
        },
    )
    assert out.is_file()

    doc = docx.Document(out)
    table = doc.tables[0]
    analysis_cell = table.rows[35].cells[0]
    assert BANNER_HEALTHY_ANALYSIS in analysis_cell.text
    assert get_cell_shading(analysis_cell) == "00B050"


def test_is_healthy_banner_text_standalone_markers_and_shading():
    """Verify standalone healthy markers without banner prefix return False, while banner-prefixed text returns True."""
    assert is_healthy_banner_text("-") is False
    assert is_healthy_banner_text(" - ") is False
    assert is_healthy_banner_text("None") is False
    assert is_healthy_banner_text("NORMAL") is False
    assert is_healthy_banner_text("N/A") is False
    assert is_healthy_banner_text("Nil") is False
    assert is_healthy_banner_text("Tiada") is False
    assert is_healthy_banner_text("Cadangan: Tiada") is False

    # Valid banner-prefixed healthy texts
    assert is_healthy_banner_text("Analysis: -") is True
    assert is_healthy_banner_text("Recommendation: -") is True
    assert is_healthy_banner_text("Recommendation:  -") is True
    assert is_healthy_banner_text("Analysis: None") is True
    assert is_healthy_banner_text("Analysis: Normal") is True
    assert is_healthy_banner_text("Analysis: No Anomaly.") is True
    assert is_healthy_banner_text("Recommendation: N/A") is True

    doc = docx.Document()
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "-"
    t.cell(0, 1).text = "Analysis: -"

    # In multi-cell table mode, '-' is unshaded, while 'Analysis: -' is shaded Green
    apply_banner_shading(t)

    assert get_cell_shading(t.cell(0, 0)) is None
    assert get_cell_shading(t.cell(0, 1)) == "00B050"


def test_apply_shading_measurement_cells_remain_unshaded():
    """Verify tables with '-', 'NORMAL', and parameter cells remain unshaded while only severity and banner cells are shaded."""
    doc = docx.Document()
    table = doc.add_table(rows=6, cols=2)

    # Technology severity cell
    table.rows[0].cells[0].text = "IR Severity"
    table.rows[0].cells[1].text = SEVERITY_MARKER_IR

    # US severity cell
    table.rows[1].cells[0].text = "US Severity"
    table.rows[1].cells[1].text = SEVERITY_MARKER_US

    # Generic measurement cells that should NEVER be shaded
    table.rows[2].cells[0].text = "Breaker Status"
    table.rows[2].cells[1].text = "NORMAL"

    table.rows[3].cells[0].text = "Load Current"
    table.rows[3].cells[1].text = "-"

    # Banner cells
    table.rows[4].cells[0].text = "Analysis & Recommendations:"
    table.rows[4].cells[1].text = "Analysis: No Anomaly."

    table.rows[5].cells[0].text = "Recommendation:"
    table.rows[5].cells[1].text = "Recommendation: -"

    # Apply both shading passes as FullReportScanPageRendererCore does
    apply_technology_severity_shading(table, defective_technologies=None)
    apply_banner_shading(table, is_defective=False)

    # 1. Technology severity cells are shaded Green and text cleared
    assert get_cell_shading(table.rows[0].cells[1]) == COLOR_HEALTHY
    assert table.rows[0].cells[1].text == ""
    assert get_cell_shading(table.rows[1].cells[1]) == COLOR_HEALTHY
    assert table.rows[1].cells[1].text == ""

    # 2. Measurement / parameter cells MUST REMAIN UNSHADED (fill=None)
    assert get_cell_shading(table.rows[2].cells[0]) is None
    assert get_cell_shading(table.rows[2].cells[1]) is None
    assert table.rows[2].cells[1].text == "NORMAL"

    assert get_cell_shading(table.rows[3].cells[0]) is None
    assert get_cell_shading(table.rows[3].cells[1]) is None
    assert table.rows[3].cells[1].text == "-"

    # 3. Banner heading unaffected, banner content cells shaded Green
    assert get_cell_shading(table.rows[4].cells[0]) is None
    assert get_cell_shading(table.rows[4].cells[1]) == COLOR_HEALTHY
    assert get_cell_shading(table.rows[5].cells[1]) == COLOR_HEALTHY


def test_apply_shading_table_row_and_mixed_containers():
    """Verify apply_banner_shading and apply_technology_severity_shading process _Row objects."""
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)

    # Row 0: Banner cells
    row0 = t.rows[0]
    row0.cells[0].text = "Analysis: No Anomaly."
    row0.cells[1].text = "Recommendation: -"
    row0.cells[2].text = "Analysis: Please refer to the following page for details defect."

    apply_banner_shading(row0)
    assert get_cell_shading(row0.cells[0]) == "00B050"
    assert get_cell_shading(row0.cells[1]) == "00B050"
    assert get_cell_shading(row0.cells[2]) == "EE0000"

    # Row 1: Technology severity cells
    row1 = t.rows[1]
    row1.cells[0].text = "{{ ir.severity }}"
    row1.cells[1].text = "{{ us.severity }}"
    row1.cells[2].text = "{{ tev.severity }}"

    apply_technology_severity_shading(row1, defective_technologies="US")
    assert get_cell_shading(row1.cells[0]) == "00B050"
    assert get_cell_shading(row1.cells[1]) == "EE0000"
    assert get_cell_shading(row1.cells[2]) == "00B050"


def test_render_explicit_is_defective_false_override(tmp_path: Path):
    """Verify is_defective=False overrides defect indicators in context to render 100% healthy."""
    template_path = TEMPLATES_DIR / "swg-panel.docx"
    output_path = tmp_path / "forced_healthy.docx"

    renderer = FullReportScanPageRendererCore(template_path)
    # Context indicates IR defect, but caller explicitly passes is_defective=False
    renderer.render(
        output_path,
        {
            "ir": {"severity": "DEFECT"},
            "defective_technologies": "IR",
        },
        is_defective=False,
    )
    assert output_path.is_file()

    doc = docx.Document(output_path)
    table = doc.tables[0]

    # IR severity cell should be healthy Green 00B050
    ir_cell = table.rows[18].cells[3]
    assert ir_cell.text == ""
    assert get_cell_shading(ir_cell) == "00B050"

    # Banner analysis should be healthy Green 00B050
    analysis_cell = table.rows[35].cells[0]
    assert BANNER_HEALTHY_ANALYSIS in analysis_cell.text
    assert get_cell_shading(analysis_cell) == "00B050"


def test_cleanup_dash_measurement_units_and_tev_background(tmp_path: Path):
    """Verify cleanup_dash_measurement_units cleans -dB to - and renders cleanly in swg-panel.docx."""
    from src.full_report.scan_render import cleanup_dash_measurement_units

    # 1. Test unit cleanup on synthetic table
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].paragraphs[0].add_run("-dB")
    t.rows[0].cells[1].paragraphs[0].add_run("1dB")
    t.rows[0].cells[2].paragraphs[0].add_run("- dB")
    t.rows[1].cells[0].paragraphs[0].add_run("-°C")
    t.rows[1].cells[1].paragraphs[0].add_run("33.5 °C")
    t.rows[1].cells[2].paragraphs[0].add_run("-%")

    cleanup_dash_measurement_units(doc)

    assert t.rows[0].cells[0].text.strip() == "-"
    assert t.rows[0].cells[1].text.strip() == "1dB"
    assert t.rows[0].cells[2].text.strip() == "-"
    assert t.rows[1].cells[0].text.strip() == "-"
    assert t.rows[1].cells[1].text.strip() == "33.5 °C"
    assert t.rows[1].cells[2].text.strip() == "-"

    # 2. Test full render of swg-panel with missing TEV background (bg = "-")
    tpl = TEMPLATES_DIR / "swg-panel.docx"
    out = tmp_path / "panel_dash_bg.docx"
    renderer = FullReportScanPageRendererCore(tpl)
    renderer.render(
        out,
        {
            "substation": {"name_erms": "TEST PE", "date": "01-01-2026", "time": "-", "ambient": "-", "humidity": "-"},
            "swg": {"type": "RMU SF6", "manufacturer": "TAMCO", "model": "-", "rating": "-", "serialnumber": "-"},
            "panel": {
                "name": "PANEL 1",
                "linknumber": "1",
                "feeder_no": "P1",
                "area": "CABLE COMPARTMENT",
                "serialnumber": "-",
                "heateramp": "-",
                "breakerstatus": "CLOSE",
                "busbarposition": "-",
                "cabletype": "-",
                "loadamp": "-",
                "analysis": "No Anomaly.",
                "recommendation": "-",
                "ir": {"reading": "-", "severity": "NORMAL"},
                "us": {"reading": "-", "char": "NORMAL", "severity": "NORMAL"},
                "tev": {"bg": "-", "reading": "-", "ppc": "-", "char": "-", "severity": "NORMAL"},
            },
            "tev": {"bg": "-", "reading": "-", "severity": "NORMAL"},
            "us": {"reading": "-", "severity": "NORMAL"},
            "ir": {"reading": "-", "severity": "NORMAL"},
            "analysis": "No Anomaly.",
            "recommendation": "-",
        },
        is_defective=False,
    )
    assert out.is_file()
    rendered_doc = docx.Document(out)
    table = rendered_doc.tables[0]
    tev_bg_cell = table.rows[28].cells[17]
    assert tev_bg_cell.text.strip() == "-", f"Expected '-' for TEV background but found {tev_bg_cell.text!r}"


def test_bind_inline_images_rejects_blank_and_invalid_images(tmp_path: Path):
    """Verify _bind_inline_images replaces blank/white/corrupt image paths with ''."""
    from src.full_report.scan_render import _bind_inline_images

    # 1. Solid white image
    white_img = tmp_path / "solid_white.png"
    Image.new("RGB", (100, 100), color="white").save(white_img)

    # 2. Solid color image (zero variance)
    green_img = tmp_path / "solid_green.png"
    Image.new("RGB", (100, 100), color=(0, 176, 80)).save(green_img)

    # 3. Two-color image (near zero variance)
    two_col = tmp_path / "two_colors.png"
    im2 = Image.new("RGB", (50, 50), color="white")
    im2.putpixel((0, 0), (0, 0, 0))
    im2.save(two_col)

    # 4. Corrupt/empty file
    empty_file = tmp_path / "empty.png"
    empty_file.write_bytes(b"")

    # 5. Non-existent file
    missing_file = tmp_path / "missing.png"

    # 6. Valid multi-color image
    valid_img = tmp_path / "valid.png"
    v_im = Image.new("RGB", (50, 50), color="white")
    v_im.putpixel((0, 0), (255, 0, 0))
    v_im.putpixel((0, 1), (0, 255, 0))
    v_im.putpixel((0, 2), (0, 0, 255))
    v_im.save(valid_img)

    doc = DocxTemplate(TEMPLATES_DIR / "swg-panel.docx")
    context = {
        "ir": {"image": str(white_img)},
        "visual": {"image": green_img},
        "us": {"prpd": str(two_col)},
        "tev": {"prpd": empty_file},
        "missing": {"prpd": str(missing_file)},
        "good": {"prpd": str(valid_img)},
        "prebound_blank": {"image": InlineImage(doc, str(white_img))},
    }

    _bind_inline_images(doc, context)

    assert context["ir"]["image"] == ""
    assert context["visual"]["image"] == ""
    assert context["us"]["prpd"] == ""
    assert context["tev"]["prpd"] == ""
    assert context["missing"]["prpd"] == ""
    assert isinstance(context["good"]["prpd"], InlineImage)
    assert context["prebound_blank"]["image"] == ""





