"""Unit and integration tests for 33kV WhatsApp Daily Report generation (Ticket #099)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import docx
import openpyxl
import pandas as pd
import pytest

from src.project.environment import ProjectEnvironment
from src.project.models import ProjectMetadata
from src.project.storage import LocalWorkspaceStorage
from src.whatsapp.models import WhatsAppReportResources
from src.workflows.models import WhatsAppReportRequest
from src.workflows.whatsapp import (
    QUALIFYING_DOCX_PATTERN,
    WHATSAPP_ACTIVITY_11KV,
    WHATSAPP_ACTIVITY_33KV,
    WHATSAPP_LINKS_11KV,
    WHATSAPP_LINKS_33KV,
    WHATSAPP_PENYELIA_11KV,
    WHATSAPP_PENYELIA_33KV,
    WhatsAppReportAuditor,
    WhatsAppReportExtractor,
    WhatsAppReportFilter,
    WhatsAppReportLoader,
    WhatsAppReportPlan,
    WhatsAppReportPreflightGuard,
    WhatsAppReportTransformer,
    WhatsAppReportWorkflow,
    WhatsAppRawTuple,
    WhatsAppTargetEntity,
    run_generate_whatsapp_report,
)


def _create_33kv_total_pe(total_pe_path: Path) -> None:
    """Create synthetic 33kV TOTAL PE workbook with DataCycle1 matching production schema."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DataCycle1"
    ws.append(["PE NO", "FL NUMBER", "SUBSTATION NAME", "DATE", "TYPE", "WO", "SCOPE"])
    ws.append([1, "CGBG/PPU/G00001", "PPU APMM 33/11kV", datetime(2026, 10, 7), "PPU", "W1616898", "FULL"])
    ws.append([2, "CKTN/SSU/G00006", "SSU KIIP (SSU CAMEL POWER) 33kV", datetime(2026, 10, 7), "SSU", "W1616956", "FULL"])
    ws.append([3, "CGBG/SSU/G00003", "SSU LYNAS 33kV", datetime(2026, 10, 7), "SSU", "W1616922", "FULL"])
    ws.append([4, "CGBG/SSU/G00005", "SSU FERRO MANGANESE 33kV", datetime(2026, 10, 7), "SSU", "W1616924", "FULL"])
    # Row with float PE number and missing WO
    ws.append([5.0, "CGBG/SSU/G00006", "SSU KPC GEBENG 33kV", datetime(2026, 10, 8), "SSU", None, "FULL"])
    wb.save(total_pe_path)
    wb.close()


def test_qualifying_docx_pattern_internal_parentheses() -> None:
    """Verify regex preserves internal parentheses and anchors defects to trailing parentheses."""
    # Internal parenthesis with defect suffix
    m1 = QUALIFYING_DOCX_PATTERN.match("02. SSU KIIP (SSU CAMEL POWER) 33kV (VI).docx")
    assert m1 is not None
    assert m1.group(1) == "02"
    assert m1.group(2) == "SSU KIIP (SSU CAMEL POWER) 33kV"
    assert m1.group(3) == "VI"

    # Internal parenthesis without defect suffix
    m2 = QUALIFYING_DOCX_PATTERN.match("02. SSU KIIP (SSU CAMEL POWER) 33kV.docx")
    assert m2 is not None
    assert m2.group(1) == "02"
    assert m2.group(2) == "SSU KIIP (SSU CAMEL POWER) 33kV"
    assert m2.group(3) is None

    # Multiple parenthesized tokens in name with defect suffix
    m_nested = QUALIFYING_DOCX_PATTERN.match("02. SSU KIIP (SSU CAMEL POWER) (PHASE 1) (VI).docx")
    assert m_nested is not None
    assert m_nested.group(1) == "02"
    assert m_nested.group(2) == "SSU KIIP (SSU CAMEL POWER) (PHASE 1)"
    assert m_nested.group(3) == "VI"

    # Multi-technology defect suffix
    m3 = QUALIFYING_DOCX_PATTERN.match("01. PPU APMM 33/11kV (US+VI).docx")
    assert m3 is not None
    assert m3.group(1) == "01"
    assert m3.group(2) == "PPU APMM 33/11kV"
    assert m3.group(3) == "US+VI"

    m4 = QUALIFYING_DOCX_PATTERN.match("03. SSU FERRO MANGANESE 33kV (IR+VI).docx")
    assert m4 is not None
    assert m4.group(1) == "03"
    assert m4.group(2) == "SSU FERRO MANGANESE 33kV"
    assert m4.group(3) == "IR+VI"

    m_all = QUALIFYING_DOCX_PATTERN.match("04. SSU FERRO MANGANESE 33kV (IR+US+TEV+VI).docx")
    assert m_all is not None
    assert m_all.group(3) == "IR+US+TEV+VI"

    # Non-docx files should not match
    assert QUALIFYING_DOCX_PATTERN.match("01. PPU APMM.pdf") is None
    assert QUALIFYING_DOCX_PATTERN.match("desktop.ini") is None


def test_cross_station_date_folder_discovery_and_sorting(tmp_path: Path) -> None:
    """Verify WhatsAppReportExtractor aggregates across station directories and sorts by PE int."""
    qr_dir = tmp_path / "QUICK REPORT"
    date_str = "07-10-2026"
    
    # Interleaved multi-station layout: GEBENG (PE 1, 3, 4) and KUANTAN (PE 2)
    gebeng_dir = qr_dir / "GEBENG" / "10. OCTOBER" / date_str
    gebeng_dir.mkdir(parents=True)
    (gebeng_dir / "01. PPU APMM 33-11kV (VI).docx").write_bytes(b"dummy")
    (gebeng_dir / "04. SSU FERRO MANGANESE 33kV (VI).docx").write_bytes(b"dummy")
    (gebeng_dir / "03. SSU LYNAS 33kV (VI).docx").write_bytes(b"dummy")

    # Station 2: KUANTAN (PE 2)
    kuantan_dir = qr_dir / "KUANTAN" / "10. OCTOBER" / date_str
    kuantan_dir.mkdir(parents=True)
    (kuantan_dir / "02. SSU KIIP (SSU CAMEL POWER) 33kV (VI).docx").write_bytes(b"dummy")

    # Other date: should be excluded
    other_date_dir = qr_dir / "GEBENG" / "10. OCTOBER" / "08-10-2026"
    other_date_dir.mkdir(parents=True)
    (other_date_dir / "05. SSU KPC GEBENG 33kV.docx").write_bytes(b"dummy")

    total_pe_path = tmp_path / "TOTAL PE.xlsx"
    _create_33kv_total_pe(total_pe_path)

    metadata = ProjectMetadata(
        key="pahang_33kv_2026",
        name="Pahang 33kV 2026",
        po_number="PO42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="Cycle 1",
        technologies=("IR", "US", "TEV"),
        base_path=str(tmp_path),
    )
    storage = LocalWorkspaceStorage(tmp_path)
    env = ProjectEnvironment(metadata=metadata, storage=storage)

    resources = WhatsAppReportResources(
        quick_report_dir=qr_dir,
        save_dir=tmp_path / "PYTHON" / "WHATSAPP",
        template_path=tmp_path / "TEMPLATE.docx",
        total_pe_path=total_pe_path,
        station_mapping={"CGBG": "GEBENG", "CKTN": "KUANTAN"},
        voltage_type="33kV",
    )
    env.get_whatsapp_report_resources = lambda: resources

    extractor = WhatsAppReportExtractor()
    # Invoke by passing gebeng_dir as report_dir (date 07-10-2026)
    request = WhatsAppReportRequest(report_dir=gebeng_dir)
    raw = extractor.extract(env, request)

    # 4 files discovered across GEBENG and KUANTAN, 08-10-2026 excluded, sorted by PE integer
    assert len(raw.substation_docx_files) == 4
    file_names = [p.name for p in raw.substation_docx_files]
    assert file_names == [
        "01. PPU APMM 33-11kV (VI).docx",
        "02. SSU KIIP (SSU CAMEL POWER) 33kV (VI).docx",
        "03. SSU LYNAS 33kV (VI).docx",
        "04. SSU FERRO MANGANESE 33kV (VI).docx",
    ]
    assert raw.station_hints[raw.substation_docx_files[0]] == "GEBENG"
    assert raw.station_hints[raw.substation_docx_files[1]] == "KUANTAN"
    assert raw.station_hints[raw.substation_docx_files[2]] == "GEBENG"
    assert raw.station_hints[raw.substation_docx_files[3]] == "GEBENG"


def test_whatsapp_filter_33kv() -> None:
    """Verify WhatsAppReportFilter extracts PE integer, WONUM, defect suffix, handles non-defect parens and date fallback."""
    df = pd.DataFrame([
        {"PE NO": 1, "SUBSTATION NAME": "PPU APMM", "WO": "W1616898", "DATE": datetime(2026, 10, 7), "FL NUMBER": "CGBG/PPU/G00001"},
        {"PE NO": "2.0", "SUBSTATION NAME": "SSU LYNAS", "WO": 40012345.0, "DATE": datetime(2026, 10, 7), "FL NUMBER": "CGBG/SSU/G00002"},
        {"PE NO": 3, "SUBSTATION NAME": "SSU FERRO", "WO": None, "DATE": datetime(2026, 10, 7), "FL NUMBER": "CGBG/SSU/G00003"},
        # PE 4 has unpopulated DATE in TOTAL PE, should fallback to folder_date
        {"PE NO": 4, "SUBSTATION NAME": None, "WO": "W1616956", "DATE": None, "FL NUMBER": "CKTN/SSU/G00004"},
    ])
    raw = WhatsAppRawTuple(
        substation_docx_files=[
            Path("01. PPU APMM (VI).docx"),
            Path("02. SSU LYNAS (US+VI).docx"),
            Path("03. SSU FERRO.docx"),
            Path("04. SSU KULIM (OLD).docx"),  # Non-defect trailing parens
        ],
        total_pe_df=df,
        resources=WhatsAppReportResources(
            quick_report_dir=Path("/qr"),
            save_dir=Path("/save"),
            template_path=Path("/template.docx"),
            total_pe_path=Path("/total_pe.xlsx"),
            station_mapping={"CGBG": "GEBENG", "CKTN": "KUANTAN"},
            voltage_type="33kV",
        ),
        station_hints={
            Path("01. PPU APMM (VI).docx"): "GEBENG",
            Path("02. SSU LYNAS (US+VI).docx"): "GEBENG",
            Path("03. SSU FERRO.docx"): "GEBENG",
            Path("04. SSU KULIM (OLD).docx"): "KUANTAN",
        },
        folder_date="07-10-2026",
    )

    stage = WhatsAppReportFilter()
    targets = stage.filter(raw)

    assert len(targets) == 4
    # PE 1
    assert targets[0].substation_number == 1
    assert targets[0].substation_name == "PPU APMM"
    assert targets[0].defect == "VI"
    assert targets[0].msms == "W1616898"
    assert targets[0].station == "GEBENG"

    # PE 2: WO float normalized to clean integer string
    assert targets[1].substation_number == 2
    assert targets[1].defect == "US+VI"
    assert targets[1].msms == "40012345"

    # PE 3: Unpopulated WO rendered as "-" and defect as "-"
    assert targets[2].substation_number == 3
    assert targets[2].defect == "-"
    assert targets[2].msms == "-"

    # PE 4: (OLD) is not a defect -> retained in substation name, defect="-", date falls back to folder_date
    assert targets[3].substation_number == 4
    assert targets[3].substation_name == "SSU KULIM (OLD)"
    assert targets[3].defect == "-"
    assert targets[3].raw_date == "07-10-2026"
    assert targets[3].station == "KUANTAN"


def test_whatsapp_transformer_multi_station_grouping_and_links() -> None:
    """Verify multi-station grouping header, subheadings, continuous indexing, and 33kV links even with interleaved targets."""
    # Interleaved input: PE 1 (GEBENG), PE 2 (KUANTAN), PE 3 (GEBENG), PE 4 (GEBENG)
    targets = [
        WhatsAppTargetEntity(
            substation_number=1,
            substation_name="PPU APMM 33/11kV",
            defect="VI",
            msms="W1616898",
            raw_date=datetime(2026, 10, 7),
            fl_number="CGBG/PPU/G00001",
            station="GEBENG",
        ),
        WhatsAppTargetEntity(
            substation_number=2,
            substation_name="SSU KIIP (SSU CAMEL POWER) 33kV",
            defect="VI",
            msms="W1616956",
            raw_date=datetime(2026, 10, 7),
            fl_number="CKTN/SSU/G00004",
            station="KUANTAN",
        ),
        WhatsAppTargetEntity(
            substation_number=3,
            substation_name="SSU LYNAS 33kV",
            defect="VI",
            msms="W1616922",
            raw_date=datetime(2026, 10, 7),
            fl_number="CGBG/SSU/G00002",
            station="GEBENG",
        ),
        WhatsAppTargetEntity(
            substation_number=4,
            substation_name="SSU FERRO MANGANESE 33kV",
            defect="VI",
            msms="W1616924",
            raw_date=datetime(2026, 10, 7),
            fl_number="CGBG/SSU/G00003",
            station="GEBENG",
        ),
    ]

    resources = WhatsAppReportResources(
        quick_report_dir=Path("/qr"),
        save_dir=Path("/save"),
        template_path=Path("/template.docx"),
        total_pe_path=Path("/total_pe.xlsx"),
        station_mapping={"CGBG": "GEBENG", "CKTN": "KUANTAN"},
        voltage_type="33kV",
    )

    transformer = WhatsAppReportTransformer()
    plan = transformer.transform(targets, resources)

    assert plan.station_name == "GEBENG/KUANTAN"
    assert plan.report_date == "07/10/2026"
    assert plan.voltage_type == "33kV"

    ctx = plan.context
    assert ctx["activity"] == WHATSAPP_ACTIVITY_33KV  # "CBM 33KV"
    assert ctx["penyelia"] == WHATSAPP_PENYELIA_33KV  # "ZULHUSNI"
    assert ctx["station"] == "GEBENG/KUANTAN"
    assert ctx["show_subheadings"] is True
    assert ctx["link_laporan"] == WHATSAPP_LINKS_33KV["laporan"]
    assert ctx["link_engr"] == WHATSAPP_LINKS_33KV["engr"]

    # Verify grouping: all GEBENG items grouped together first (items 1, 2, 3), then KUANTAN (item 4)
    items = ctx["items"]
    assert len(items) == 4
    # Item 1: GEBENG (PE 1)
    assert items[0]["index"] == 1
    assert items[0]["name"] == "PPU APMM 33/11kV"
    assert items[0]["station_header"] == "GEBENG"
    # Item 2: GEBENG (PE 3)
    assert items[1]["index"] == 2
    assert items[1]["name"] == "SSU LYNAS 33kV"
    assert items[1]["station_header"] is None
    # Item 3: GEBENG (PE 4)
    assert items[2]["index"] == 3
    assert items[2]["name"] == "SSU FERRO MANGANESE 33kV"
    assert items[2]["station_header"] is None
    # Item 4: KUANTAN (PE 2)
    assert items[3]["index"] == 4
    assert items[3]["name"] == "SSU KIIP (SSU CAMEL POWER) 33kV"
    assert items[3]["station_header"] == "KUANTAN"

    # Verify station_groups structure
    assert len(ctx["station_groups"]) == 2
    assert ctx["station_groups"][0]["station"] == "GEBENG"
    assert len(ctx["station_groups"][0]["items"]) == 3
    assert ctx["station_groups"][1]["station"] == "KUANTAN"
    assert len(ctx["station_groups"][1]["items"]) == 1


def test_whatsapp_transformer_single_station_no_subheadings() -> None:
    """Verify single-station on a date renders without subheadings."""
    targets = [
        WhatsAppTargetEntity(
            substation_number=1,
            substation_name="SSU KPC GEBENG 33kV",
            defect="VI",
            msms="W1616933",
            raw_date=datetime(2026, 10, 8),
            fl_number="CGBG/SSU/G00005",
            station="GEBENG",
        ),
        WhatsAppTargetEntity(
            substation_number=2,
            substation_name="SSU FLEXSYS CHEMICAL 33kV",
            defect="VI",
            msms="W1616923",
            raw_date=datetime(2026, 10, 8),
            fl_number="CGBG/SSU/G00006",
            station="GEBENG",
        ),
    ]

    resources = WhatsAppReportResources(
        quick_report_dir=Path("/qr"),
        save_dir=Path("/save"),
        template_path=Path("/template.docx"),
        total_pe_path=Path("/total_pe.xlsx"),
        station_mapping={"CGBG": "GEBENG"},
        voltage_type="33kV",
    )

    transformer = WhatsAppReportTransformer()
    plan = transformer.transform(targets, resources)

    assert plan.station_name == "GEBENG"
    assert plan.context["station"] == "GEBENG"
    assert plan.context["show_subheadings"] is False
    assert plan.context["items"][0]["station_header"] is None
    assert plan.context["items"][1]["station_header"] is None


def test_whatsapp_loader_output_filename_33kv(tmp_path: Path) -> None:
    """Verify 33kV filename is formatted strictly as '{NN}. {clean_date}.docx'."""
    template_path = tmp_path / "template.docx"
    doc = docx.Document()
    doc.add_paragraph("Test WhatsApp Report")
    doc.save(template_path)

    save_dir = tmp_path / "PYTHON" / "WHATSAPP"
    loader = WhatsAppReportLoader()

    # 1. 33kV Plan: {NN}. {clean_date}.docx (e.g. 01. 07-10-2026.docx)
    plan_33kv = WhatsAppReportPlan(
        template_path=template_path,
        save_dir=save_dir,
        context={"date": "07/10/2026", "station": "GEBENG/KUANTAN", "items": []},
        substations_count=0,
        station_name="GEBENG/KUANTAN",
        report_date="07/10/2026",
        voltage_type="33kV",
    )
    out1 = loader.load(plan_33kv)
    assert out1.name == "01. 07-10-2026.docx"

    # Second file on next date increments to 02.
    plan_33kv_2 = WhatsAppReportPlan(
        template_path=template_path,
        save_dir=save_dir,
        context={"date": "08/10/2026", "station": "GEBENG", "items": []},
        substations_count=0,
        station_name="GEBENG",
        report_date="08/10/2026",
        voltage_type="33kV",
    )
    out2 = loader.load(plan_33kv_2)
    assert out2.name == "02. 08-10-2026.docx"

    # 2. 11kV Plan: preserves {NN}. {station} {clean_date}.docx
    plan_11kv = WhatsAppReportPlan(
        template_path=template_path,
        save_dir=save_dir,
        context={"date": "01/05/2026", "station": "CAMERON HIGHLAND", "items": []},
        substations_count=0,
        station_name="CAMERON HIGHLAND",
        report_date="01/05/2026",
        voltage_type="11kV",
    )
    out3 = loader.load(plan_11kv)
    assert out3.name == "03. CAMERON HIGHLAND 01-05-2026.docx"


def test_whatsapp_preflight_guard_33kv(tmp_path: Path) -> None:
    """Verify preflight guard cross-station date directory validation in 33kV."""
    qr_dir = tmp_path / "QUICK REPORT"
    target_date = "07-10-2026"
    gebeng_dir = qr_dir / "GEBENG" / "10. OCTOBER" / target_date
    gebeng_dir.mkdir(parents=True)
    (gebeng_dir / "01. PPU APMM 33-11kV (VI).docx").write_bytes(b"dummy")

    total_pe_path = tmp_path / "TOTAL PE.xlsx"
    _create_33kv_total_pe(total_pe_path)

    template_path = tmp_path / "TEMPLATE.docx"
    docx.Document().save(template_path)

    metadata = ProjectMetadata(
        key="pahang_33kv",
        name="Pahang 33kV",
        po_number="PO42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="Cycle 1",
        technologies=("IR", "VI"),
        base_path=str(tmp_path),
    )
    storage = LocalWorkspaceStorage(tmp_path)
    env = ProjectEnvironment(metadata=metadata, storage=storage)

    resources = WhatsAppReportResources(
        quick_report_dir=qr_dir,
        save_dir=tmp_path / "PYTHON" / "WHATSAPP",
        template_path=template_path,
        total_pe_path=total_pe_path,
        station_mapping={"CGBG": "GEBENG"},
        voltage_type="33kV",
    )
    env.get_whatsapp_report_resources = lambda: resources

    guard = WhatsAppReportPreflightGuard()
    # 1. Valid: gebeng_dir exists with qualifying files
    guard.validate(env, WhatsAppReportRequest(report_dir=gebeng_dir))

    # 2. Valid: invoked with target_date string as report_dir
    guard.validate(env, WhatsAppReportRequest(report_dir=Path(target_date)))

    # 3. Invalid: date with no docx files anywhere raises ValueError
    empty_date = qr_dir / "GEBENG" / "10. OCTOBER" / "09-10-2026"
    empty_date.mkdir(parents=True)
    with pytest.raises(ValueError, match="No qualifying PE DOCX files"):
        guard.validate(env, WhatsAppReportRequest(report_dir=empty_date))

    # 4. Invalid: date folder does not exist raises FileNotFoundError
    with pytest.raises(FileNotFoundError, match="Directory not found"):
        guard.validate(env, WhatsAppReportRequest(report_dir=Path("31-12-2099")))


def test_e2e_generate_whatsapp_report_33kv(tmp_path: Path) -> None:
    """Full end-to-end integration test rendering 33kV multi-station daily WhatsApp report."""
    root_path = tmp_path / "workspace"
    qr_dir = root_path / "QUICK REPORT"
    python_dir = root_path / "PYTHON"
    target_date = "07-10-2026"

    # Multi-station layout with interleaved PE numbers: GEBENG (PE 1, 3, 4) and KUANTAN (PE 2)
    gebeng_dir = qr_dir / "GEBENG" / "10. OCTOBER" / target_date
    kuantan_dir = qr_dir / "KUANTAN" / "10. OCTOBER" / target_date
    gebeng_dir.mkdir(parents=True)
    kuantan_dir.mkdir(parents=True)
    python_dir.mkdir(parents=True)

    (gebeng_dir / "01. PPU APMM 33-11kV (VI).docx").write_bytes(b"dummy")
    (gebeng_dir / "03. SSU LYNAS 33kV (VI).docx").write_bytes(b"dummy")
    (gebeng_dir / "04. SSU FERRO MANGANESE 33kV (VI).docx").write_bytes(b"dummy")
    (kuantan_dir / "02. SSU KIIP (SSU CAMEL POWER) 33kV (VI).docx").write_bytes(b"dummy")

    total_pe_path = python_dir / "TOTAL PE.xlsx"
    _create_33kv_total_pe(total_pe_path)

    # Use master template from repo
    real_template = Path(__file__).parent.parent / "templates" / "WHATSAPP" / "TEMPLATE WHATSAPP PYTHON.docx"
    template_path = tmp_path / "TEMPLATE WHATSAPP PYTHON.docx"
    template_path.write_bytes(real_template.read_bytes())

    metadata = ProjectMetadata(
        key="pahang_33kv_2026",
        name="Pahang 33kV 2026",
        po_number="PO42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="Cycle 1",
        technologies=("IR", "US", "TEV"),
        base_path=str(root_path),
    )
    storage = LocalWorkspaceStorage(root_path, templates_dir=template_path.parent)
    env = ProjectEnvironment(metadata=metadata, storage=storage)

    resources = WhatsAppReportResources(
        quick_report_dir=qr_dir,
        save_dir=python_dir / "WHATSAPP",
        template_path=template_path,
        total_pe_path=total_pe_path,
        station_mapping={"CGBG": "GEBENG", "CKTN": "KUANTAN"},
        voltage_type="33kV",
    )
    env.get_whatsapp_report_resources = lambda: resources

    result = run_generate_whatsapp_report(env, gebeng_dir)

    assert result.substations_count == 4
    assert result.output_path is not None
    assert result.output_path.exists()
    assert result.output_path.name == "01. 07-10-2026.docx"

    # Verify rendered document contents
    output_doc = docx.Document(result.output_path)
    doc_text = "\n".join(p.text for p in output_doc.paragraphs if p.text.strip())

    assert "AKTIVITI : *CBM 33KV*" in doc_text
    assert "TARIKH : *07/10/2026*" in doc_text
    assert "STATION : *GEBENG/KUANTAN*" in doc_text
    assert "PENYELIA : *ZULHUSNI*" in doc_text

    # Grouping subheadings and continuous sequential numbering
    assert doc_text.count("*GEBENG*") == 1
    assert doc_text.count("*KUANTAN*") == 1
    assert "*GEBENG*" in doc_text
    assert "1. PPU APMM 33/11kV" in doc_text
    assert "2. SSU LYNAS 33kV" in doc_text
    assert "3. SSU FERRO MANGANESE 33kV" in doc_text
    assert "*KUANTAN*" in doc_text
    assert "4. SSU KIIP (SSU CAMEL POWER) 33kV" in doc_text

    # Defect and MSMS
    assert "DEFECT : VI" in doc_text
    assert "MSMS : W1616898" in doc_text
    assert "MSMS : W1616956" in doc_text

    # Links
    assert "Link Laporan:" in doc_text
    assert WHATSAPP_LINKS_33KV["laporan"] in doc_text
    assert "Link ENGR:" in doc_text
    assert WHATSAPP_LINKS_33KV["engr"] in doc_text


def test_whatsapp_date_normalization() -> None:
    """Verify _normalize_date_str correctly normalizes date strings with slashes, dashes, backslashes, and directory paths."""
    assert WhatsAppReportPreflightGuard._normalize_date_str("07/10/2026") == "07-10-2026"
    assert WhatsAppReportPreflightGuard._normalize_date_str(r"07\10\2026") == "07-10-2026"
    assert WhatsAppReportPreflightGuard._normalize_date_str(Path("07/10/2026")) == "07-10-2026"
    assert WhatsAppReportPreflightGuard._normalize_date_str("07-10-2026") == "07-10-2026"
    assert WhatsAppReportPreflightGuard._normalize_date_str("QUICK REPORT/GEBENG/10. OCTOBER/07-10-2026") == "07-10-2026"
    assert WhatsAppReportPreflightGuard._normalize_date_str("C:/path/to/07-10-2026") == "07-10-2026"


def test_whatsapp_url_xml_escaping_idempotent(tmp_path: Path) -> None:
    """Verify WhatsAppReportLoader handles XML entity escaping idempotently without double-escaping."""
    template_path = tmp_path / "template.docx"
    doc = docx.Document()
    doc.add_paragraph("Laporan: {{link_laporan}}")
    doc.add_paragraph("ENGR: {{link_engr}}")
    doc.save(template_path)

    save_dir = tmp_path / "save"
    loader = WhatsAppReportLoader()

    # Pass pre-escaped and unescaped URLs
    plan = WhatsAppReportPlan(
        template_path=template_path,
        save_dir=save_dir,
        context={
            "link_laporan": "https://test.com/view?a=1&amp;b=2",
            "link_engr": "https://test.com/view?c=3&d=4",
        },
        substations_count=0,
        station_name="TEST",
        report_date="07/10/2026",
        voltage_type="33kV",
    )
    out_path = loader.load(plan)
    rendered_doc = docx.Document(out_path)
    full_text = "\n".join(p.text for p in rendered_doc.paragraphs)

    assert "https://test.com/view?a=1&b=2" in full_text
    assert "&amp;" not in full_text
    assert "https://test.com/view?c=3&d=4" in full_text


def test_whatsapp_station_resolution_all_districts(tmp_path: Path) -> None:
    """Verify discovery extracts station names for all Pahang districts including Cameron Highland and Lipis alias."""
    qr_dir = tmp_path / "QUICK REPORT"
    date_str = "16-10-2026"

    chl_dir = qr_dir / "CAMERON HIGHLAND" / "10. OCTOBER" / date_str
    chl_dir.mkdir(parents=True)
    (chl_dir / "26. MINI PPU TELANOK 33-11kV (US+VI).docx").write_bytes(b"dummy")

    lipis_dir = qr_dir / "LIPIS" / "10. OCTOBER" / date_str
    lipis_dir.mkdir(parents=True)
    (lipis_dir / "27. PMU BENTA 132-11kV (VI).docx").write_bytes(b"dummy")

    extractor = WhatsAppReportExtractor()
    files, hints = extractor._scan_qualifying_docx_files_33kv(qr_dir, date_str)

    assert len(files) == 2
    assert hints[files[0]] == "CAMERON HIGHLAND"
    assert hints[files[1]] == "KUALA LIPIS"

