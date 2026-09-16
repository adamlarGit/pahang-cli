"""Unit tests for Quick Report front page generator and post_process toggle."""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from src.core.shading import get_cell_shading
from src.quick_report.cbm_render import (
    _build_swg_render_context,
    _render_docx_template,
)
from src.quick_report.defects import CbmDefectRecord
from src.quick_report.front_page import generate_front_page


def test_generate_front_page_preserves_tev_cells(tmp_path: Path):
    """Verify generate_front_page preserves TEV text and transparent shading on Row 14 and Row 19."""
    template_path = Path("templates/QUICK REPORT/1. FRONT PAGE TEMPLATE IR US TEV BOX Jinja2 updated.docx")
    if not template_path.exists():
        pytest.skip(f"Template not found: {template_path}")

    pe_info = {
        "substation_name": "PE TEST 132KV",
        "substation_id": "123456",
    }
    out_path = generate_front_page(pe_info, template_path, tmp_path, substation_number=1)

    assert out_path.exists()
    assert out_path.name == "001_01_front_page.docx"

    doc = docx.Document(out_path)
    # Table 1 (index 1) contains the front page form fields
    assert len(doc.tables) >= 2
    table = doc.tables[1]

    # Verify Row 14: SCANNED BY -> TEV entry
    row_14_cell_0 = table.rows[14].cells[0]
    assert row_14_cell_0.text.strip() == "TEV"
    assert get_cell_shading(row_14_cell_0) is None

    # Verify Row 19: EQUIPMENT/TOOLS -> TEV entry
    row_19_cell_0 = table.rows[19].cells[0]
    assert row_19_cell_0.text.strip() == "TEV"
    assert get_cell_shading(row_19_cell_0) is None


def test_render_docx_template_post_process_suppression(tmp_path: Path):
    """Verify _render_docx_template with post_process=False suppresses severity and banner shading."""
    template_path = Path("templates/QUICK REPORT/DEFECT IR US TEV/swg-panel.docx")
    if not template_path.exists():
        pytest.skip(f"Template not found: {template_path}")

    rec = CbmDefectRecord(
        equipment="VCB PANEL 1",
        technology="IR",
        defect_area="Cable Box",
        ir_reading="55.4",
    )
    context = _build_swg_render_context(rec, overview=False)

    # 1. With post_process=True (default), severity cell is shaded EE0000 and text cleared
    out_true = tmp_path / "test_swg_post_process_true.docx"
    _render_docx_template(template_path, out_true, context, post_process=True)
    doc_true = docx.Document(out_true)
    cell_ir_true = doc_true.tables[0].rows[18].cells[3]
    assert get_cell_shading(cell_ir_true) == "EE0000"
    assert cell_ir_true.text.strip() == ""
    # Banner cells (Row 35 and 36) are shaded red EE0000
    assert get_cell_shading(doc_true.tables[0].rows[35].cells[0]) == "EE0000"
    assert get_cell_shading(doc_true.tables[0].rows[36].cells[0]) == "EE0000"

    # 2. With post_process=False, severity cell shading and banner shading are suppressed
    out_false = tmp_path / "test_swg_post_process_false.docx"
    _render_docx_template(template_path, out_false, context, post_process=False)
    doc_false = docx.Document(out_false)
    cell_ir_false = doc_false.tables[0].rows[18].cells[3]
    assert get_cell_shading(cell_ir_false) is None
    assert cell_ir_false.text.strip() == "__SEVERITY_IR__"
    # Banner cells are not shaded
    assert get_cell_shading(doc_false.tables[0].rows[35].cells[0]) is None
    assert get_cell_shading(doc_false.tables[0].rows[36].cells[0]) is None
