"""Unit tests for Full Report substation attribution and structural guards (Ticket #44 / Seam 1)."""

from __future__ import annotations

from pathlib import Path
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import pytest

from src.full_report.attribution import (
    is_substation_attribution_match,
    normalize_substation_tokens,
    verify_cbm_defect_attribution,
    verify_condition_pages_structure,
    verify_front_page_attribution,
    verify_sticker_page_structure,
    verify_vi_defect_pages_structure,
    verify_vi_summary_structure,
)
from src.full_report.defect_parser import CbmDefectSliceMetadata


def test_normalize_substation_tokens_strips_prefixes_suffixes_and_pe_numbers():
    """Verify normalize_substation_tokens strips PE prefixes, electrical tags, and parenthesized suffixes."""
    assert normalize_substation_tokens("005. TALAPIA") == {"TALAPIA"}
    assert normalize_substation_tokens("PE 144 TELEKOM TANAH PUTIH") == {"TELEKOM", "TANAH", "PUTIH"}
    assert normalize_substation_tokens("P/E 179 CENDERAWASIH NO.1") == {"CENDERAWASIH", "NO", "1"} or normalize_substation_tokens("P/E 179 CENDERAWASIH NO.1") == {"CENDERAWASIH", "1"}
    assert normalize_substation_tokens("SSU KOTA PERDANA") == {"KOTA", "PERDANA"}
    assert normalize_substation_tokens("PPU BENTONG") == {"BENTONG"}
    assert normalize_substation_tokens("PMU MARAN") == {"MARAN"}
    assert normalize_substation_tokens("S/S JAYA") == {"JAYA"}
    assert normalize_substation_tokens("TALAPIA (VCB)") == {"TALAPIA"}
    assert normalize_substation_tokens("TALAPIA (IR+VI)") == {"TALAPIA"}
    assert normalize_substation_tokens("PERPUSTAKAAN AWAM (IR+US+TEV)") == {"PERPUSTAKAAN", "AWAM"}
    assert normalize_substation_tokens("TAMAN TAS (F/P)") == {"TAMAN", "TAS"}


def test_is_substation_attribution_match():
    """Verify is_substation_attribution_match matches targets and rejects foreign substations."""
    # Target subset or exact match
    assert is_substation_attribution_match("TALAPIA", "005. PE TALAPIA (IR+VI)") is True
    assert is_substation_attribution_match("TALAPIA", "TALAPIA") is True
    assert is_substation_attribution_match("TALAPIA", "PE 5 TALAPIA RAUB") is True
    assert is_substation_attribution_match("PE 5 TALAPIA", "TALAPIA") is True
    assert is_substation_attribution_match("CENDERAWASIH NO.1", "PE 179 CENDERAWASIH NO. 1") is True

    # Benign blank/empty/dash values allowed
    assert is_substation_attribution_match("PERPUSTAKAAN AWAM", "") is True
    assert is_substation_attribution_match("PERPUSTAKAAN AWAM", "-") is True
    assert is_substation_attribution_match("PERPUSTAKAAN AWAM", None) is True

    # Explicit foreign conflict rejected
    assert is_substation_attribution_match("PERPUSTAKAAN AWAM", "TELEKOM TANAH PUTIH") is False
    assert is_substation_attribution_match("PERPUSTAKAAN AWAM", "PE 5 TALAPIA") is False
    assert is_substation_attribution_match("TALAPIA", "CENDERAWASIH NO.1") is False


def test_verify_cbm_defect_attribution():
    """Verify CBM defect slice metadata is checked against target substation."""
    meta_matching = CbmDefectSliceMetadata(
        equipment_category="swg",
        equipment_instance="swg1",
        sequence="p01",
        equipment_id="VCB1",
        defect_area="CABLE_BOX",
        substation="PE 5 TALAPIA",
    )
    assert verify_cbm_defect_attribution(meta_matching, "TALAPIA") is True

    meta_blank = CbmDefectSliceMetadata(
        equipment_category="swg",
        equipment_instance="swg1",
        sequence="p01",
        equipment_id="VCB1",
        defect_area="CABLE_BOX",
        substation="",
    )
    # Blank allowed with warning
    assert verify_cbm_defect_attribution(meta_blank, "TALAPIA") is True

    meta_foreign = CbmDefectSliceMetadata(
        equipment_category="swg",
        equipment_instance="swg1",
        sequence="p01",
        equipment_id="VCB1",
        defect_area="CABLE_BOX",
        substation="TELEKOM TANAH PUTIH",
    )
    assert verify_cbm_defect_attribution(meta_foreign, "TALAPIA") is False


def test_verify_front_page_attribution(tmp_path: Path):
    """Verify Table 1 in front_page.docx is checked for both ERMS and SITE names."""
    doc = docx.Document()
    doc.add_table(rows=1, cols=3)  # Table 0
    t1 = doc.add_table(rows=5, cols=3)  # Table 1
    t1.rows[0].cells[0].text = "FUNCTIONAL LOCATION (ERMS)"
    t1.rows[1].cells[0].text = "FUNCTIONAL LOCATION (SITE)"
    t1.rows[2].cells[0].text = "SUBSTATION NAME (ERMS)"
    t1.rows[2].cells[2].text = "PE 12 PERPUSTAKAAN AWAM"
    t1.rows[3].cells[0].text = "SUBSTATION NAME (SITE)"
    t1.rows[3].cells[2].text = "PERPUSTAKAAN AWAM"

    front_valid = tmp_path / "front_valid.docx"
    doc.save(front_valid)

    valid, reason = verify_front_page_attribution(front_valid, "PERPUSTAKAAN AWAM")
    assert valid is True
    assert reason == ""

    # Mutate to foreign substation
    t1.rows[2].cells[2].text = "PE 144 TELEKOM TANAH PUTIH"
    t1.rows[3].cells[2].text = "TELEKOM TANAH PUTIH"
    front_foreign = tmp_path / "front_foreign.docx"
    doc.save(front_foreign)

    valid_f, reason_f = verify_front_page_attribution(front_foreign, "PERPUSTAKAAN AWAM")
    assert valid_f is False
    assert "TELEKOM TANAH PUTIH" in reason_f


def test_verify_structural_guards(tmp_path: Path):
    """Verify structural guards for condition_pages, sticker_page, vi_summary, and vi_defect_pages."""
    # 1. Condition pages: requires >= 1 table and >= 1 image
    doc_cond = docx.Document()
    doc_cond.add_table(rows=2, cols=2)
    # Add a mock drawing element to satisfy image check
    p = doc_cond.add_paragraph()
    run = p.add_run()
    drawing_elm = parse_xml(r'<w:drawing %s><w:inline><a:graphic %s><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic %s><pic:blipFill><a:blip r:embed="rId1"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></w:inline></w:drawing>' % (nsdecls('w'), nsdecls('a'), nsdecls('pic', 'r')))
    run._r.append(drawing_elm)
    cond_path = tmp_path / "condition_pages.docx"
    doc_cond.save(cond_path)

    valid, reason = verify_condition_pages_structure(cond_path)
    assert valid is True

    # 2. Sticker page: requires NORMAL STICKER or DEFECT STICKER
    doc_stk = docx.Document()
    doc_stk.add_paragraph("SUBSTATION NORMAL STICKER ATTACHED")
    stk_path = tmp_path / "sticker_page.docx"
    doc_stk.save(stk_path)

    valid_stk, _ = verify_sticker_page_structure(stk_path)
    assert valid_stk is True

    doc_bad_stk = docx.Document()
    doc_bad_stk.add_paragraph("JUST SOME RANDOM TEXT")
    bad_stk_path = tmp_path / "bad_sticker.docx"
    doc_bad_stk.save(bad_stk_path)
    valid_bad_stk, reason = verify_sticker_page_structure(bad_stk_path)
    assert valid_bad_stk is False

    # 3. VI summary: table rows match expected count
    doc_visum = docx.Document()
    t_visum = doc_visum.add_table(rows=3, cols=4)  # 1 header + 2 data rows
    t_visum.rows[0].cells[0].text = "NO."
    t_visum.rows[1].cells[0].text = "Defect 1"
    t_visum.rows[2].cells[0].text = "Defect 2"
    visum_path = tmp_path / "vi_summary.docx"
    doc_visum.save(visum_path)

    valid_vi, _ = verify_vi_summary_structure(visum_path, expected_count=2)
    assert valid_vi is True

    valid_vi_bad, reason = verify_vi_summary_structure(visum_path, expected_count=5)
    assert valid_vi_bad is False

    # 4. VI defect pages: requires images
    doc_videf = docx.Document()
    p_vi = doc_videf.add_paragraph()
    r_vi = p_vi.add_run()
    r_vi._r.append(parse_xml(r'<w:drawing %s><w:inline><a:graphic %s><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic %s><pic:blipFill><a:blip r:embed="rId2"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></w:inline></w:drawing>' % (nsdecls('w'), nsdecls('a'), nsdecls('pic', 'r'))))
    videf_path = tmp_path / "vi_defect_pages.docx"
    doc_videf.save(videf_path)

    valid_videf, _ = verify_vi_defect_pages_structure(videf_path)
    assert valid_videf is True


def test_inspect_deliverable_attribution_and_quarantine(tmp_path: Path):
    """Verify post-compilation headless OpenXML inspection detects foreign headers and quarantines file."""
    from src.full_report.attribution import inspect_deliverable_attribution, quarantine_deliverable

    # 1. Clean deliverable matching target
    doc_clean = docx.Document()
    t_clean = doc_clean.add_table(rows=2, cols=3)
    t_clean.rows[0].cells[0].text = "SUBSTATION NAME"
    t_clean.rows[0].cells[2].text = "PE 12 PERPUSTAKAAN AWAM"
    t_clean.rows[1].cells[0].text = "FEEDER NAME"
    t_clean.rows[1].cells[2].text = "FEEDER TO PE TALAPIA"  # Feeder line ignored!

    clean_path = tmp_path / "012. PERPUSTAKAAN AWAM.docx"
    doc_clean.save(clean_path)

    valid, reason = inspect_deliverable_attribution(clean_path, "PERPUSTAKAAN AWAM")
    assert valid is True
    assert reason == ""

    # 2. Compromised deliverable with foreign header
    doc_foreign = docx.Document()
    t_front = doc_foreign.add_table(rows=1, cols=3)
    t_front.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_front.rows[0].cells[2].text = "PERPUSTAKAAN AWAM"
    # Leaked foreign defect page table!
    t_leak = doc_foreign.add_table(rows=2, cols=2)
    t_leak.rows[0].cells[0].text = "Substation"
    t_leak.rows[0].cells[1].text = "TELEKOM TANAH PUTIH"
    t_leak.rows[1].cells[0].text = "Equipment"
    t_leak.rows[1].cells[1].text = "VCB"

    compromised_path = tmp_path / "compromised_report.docx"
    doc_foreign.save(compromised_path)

    valid_f, reason_f = inspect_deliverable_attribution(compromised_path, "PERPUSTAKAAN AWAM")
    assert valid_f is False
    assert "TELEKOM TANAH PUTIH" in reason_f

    # 3. Quarantine
    quarantine_dir = tmp_path / ".quarantine"
    quarantined_path = quarantine_deliverable(compromised_path, quarantine_dir)

    assert not compromised_path.exists()
    assert quarantined_path.exists()
    assert quarantined_path.parent == quarantine_dir
    assert quarantined_path.name == "compromised_report.docx"


def _add_mock_image(cell: Any) -> None:
    """Helper to attach a mock drawing XML element to a table cell."""
    p = cell.add_paragraph()
    r = p.add_run()
    drawing_xml = parse_xml(
        r'<w:drawing %s><w:inline><a:graphic %s><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        r'<pic:pic %s><pic:blipFill><a:blip r:embed="rIdMock"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></w:inline></w:drawing>'
        % (nsdecls("w"), nsdecls("a"), nsdecls("pic", "r"))
    )
    r._r.append(drawing_xml)


def test_inspect_deliverable_attribution_ignores_defect_summaries_and_photo_grids(tmp_path: Path):
    """Regression test (Issue #45): Ensure defect tables with 'SUBSTATION' equipment & photo grids don't false alarm."""
    from src.full_report.attribution import inspect_deliverable_attribution

    doc = docx.Document()

    # 1. Front page header
    t_front = doc.add_table(rows=1, cols=3)
    t_front.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_front.rows[0].cells[2].text = "PERPUSTAKAAN AWAM(VCB)"

    # 2. Visual Inspection Defect Summary Table (with 'SUBSTATION' equipment)
    t_defect = doc.add_table(rows=3, cols=4)
    t_defect.rows[0].cells[0].text = "NO."
    t_defect.rows[0].cells[1].text = "EQUIPMENT"
    t_defect.rows[0].cells[2].text = "DEFECT DESCRIPTION"
    t_defect.rows[0].cells[3].text = "ADDITIONAL REMARKS"

    t_defect.rows[1].cells[0].text = "1"
    t_defect.rows[1].cells[1].text = "SUBSTATION"
    t_defect.rows[1].cells[2].text = "CPR POSTER OLD VERSION"
    t_defect.rows[1].cells[3].text = "SWG ROOM"

    t_defect.rows[2].cells[0].text = "2"
    t_defect.rows[2].cells[1].text = "SIGNBOARD"
    t_defect.rows[2].cells[2].text = "NO FUNCTIONAL LOCATION"
    t_defect.rows[2].cells[3].text = "-"

    # 3. Defect Photo Grid Table (with 'SUBSTATION' & 'SIGNBOARD' headers and images)
    t_grid = doc.add_table(rows=2, cols=2)
    t_grid.rows[0].cells[0].text = "SUBSTATION"
    t_grid.rows[0].cells[1].text = "SIGNBOARD"
    t_grid.rows[1].cells[0].text = "CPR POSTER OLD VERSION - SWG ROOM"
    t_grid.rows[1].cells[1].text = "NO FUNCTIONAL LOCATION"
    _add_mock_image(t_grid.rows[1].cells[0])
    _add_mock_image(t_grid.rows[1].cells[1])

    # 4. Photo Grid with Arbitrary/Non-standard category (e.g. 'EARTHING & LIGHTNING ARRESTOR')
    t_grid2 = doc.add_table(rows=2, cols=2)
    t_grid2.rows[0].cells[0].text = "SUBSTATION"
    t_grid2.rows[0].cells[1].text = "EARTHING & LIGHTNING ARRESTOR"
    t_grid2.rows[1].cells[0].text = "Photo A"
    t_grid2.rows[1].cells[1].text = "Photo B"
    _add_mock_image(t_grid2.rows[1].cells[0])
    _add_mock_image(t_grid2.rows[1].cells[1])

    # 5. Valid CBM Test Sheet Table
    t_cbm = doc.add_table(rows=2, cols=4)
    t_cbm.rows[0].cells[0].text = "Substation"
    t_cbm.rows[0].cells[1].text = "PERPUSTAKAAN AWAM(VCB)"
    t_cbm.rows[0].cells[2].text = "Date: 25/08/2026"
    t_cbm.rows[1].cells[0].text = "Equipment"
    t_cbm.rows[1].cells[1].text = "VCB"

    report_path = tmp_path / "157. PERPUSTAKAAN AWAM(VCB) (VI).docx"
    doc.save(report_path)

    is_valid, reason = inspect_deliverable_attribution(report_path, "PERPUSTAKAAN AWAM(VCB)")
    assert is_valid is True
    assert reason == ""


def test_verify_front_page_attribution_relaxed_dual_name(tmp_path: Path):
    """Regression test (Issue #46): Front page dual-name attribution allows matching at least one non-empty field."""
    # 1. ERMS='TALAPIA' and SITE='KPRC S/B' -> passes for target='TALAPIA'
    doc1 = docx.Document()
    t1 = doc1.add_table(rows=2, cols=3)
    t1.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t1.rows[0].cells[2].text = "TALAPIA"
    t1.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t1.rows[1].cells[2].text = "KPRC S/B"
    doc1_path = tmp_path / "front_dual_match.docx"
    doc1.save(doc1_path)

    valid1, reason1 = verify_front_page_attribution(doc1_path, "TALAPIA")
    assert valid1 is True
    assert reason1 == ""

    # 2. ERMS='KPRC S/B' and SITE='TALAPIA' -> passes for target='TALAPIA'
    doc2 = docx.Document()
    t2 = doc2.add_table(rows=2, cols=3)
    t2.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t2.rows[0].cells[2].text = "KPRC S/B"
    t2.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t2.rows[1].cells[2].text = "TALAPIA"
    doc2_path = tmp_path / "front_dual_match_site.docx"
    doc2.save(doc2_path)

    valid2, reason2 = verify_front_page_attribution(doc2_path, "TALAPIA")
    assert valid2 is True
    assert reason2 == ""

    # 3. ERMS='TELEKOM TANAH PUTIH' and SITE='KPRC S/B' -> fails for target='TALAPIA'
    doc3 = docx.Document()
    t3 = doc3.add_table(rows=2, cols=3)
    t3.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t3.rows[0].cells[2].text = "TELEKOM TANAH PUTIH"
    t3.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t3.rows[1].cells[2].text = "KPRC S/B"
    doc3_path = tmp_path / "front_dual_mismatch.docx"
    doc3.save(doc3_path)

    valid3, reason3 = verify_front_page_attribution(doc3_path, "TALAPIA")
    assert valid3 is False
    assert "ERMS='TELEKOM TANAH PUTIH', SITE='KPRC S/B'" in reason3
    assert "(target: 'TALAPIA')" in reason3

    # 4. Only ERMS present and non-matching -> fails citing ERMS field
    doc4 = docx.Document()
    t4 = doc4.add_table(rows=1, cols=3)
    t4.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t4.rows[0].cells[2].text = "TELEKOM TANAH PUTIH"
    doc4_path = tmp_path / "front_single_erms_mismatch.docx"
    doc4.save(doc4_path)

    valid4, reason4 = verify_front_page_attribution(doc4_path, "TALAPIA")
    assert valid4 is False
    assert "front page ERMS field: 'TELEKOM TANAH PUTIH'" in reason4

    # 6. Foreign ERMS with SITE='-' placeholder -> fails citing ERMS field
    doc6 = docx.Document()
    t6 = doc6.add_table(rows=2, cols=3)
    t6.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t6.rows[0].cells[2].text = "TELEKOM TANAH PUTIH"
    t6.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t6.rows[1].cells[2].text = "-"
    doc6_path = tmp_path / "front_foreign_erms_site_dash.docx"
    doc6.save(doc6_path)

    valid6, reason6 = verify_front_page_attribution(doc6_path, "TALAPIA")
    assert valid6 is False
    assert "front page ERMS field: 'TELEKOM TANAH PUTIH'" in reason6


def test_verify_vi_summary_structure_trailing_blank_rows(tmp_path: Path):
    """Regression test (Issue #46): VI summary table with trailing blank rows passes when non-empty rows match."""
    doc = docx.Document()
    # 1 header row + 2 data rows + 10 trailing blank rows = 13 rows total
    table = doc.add_table(rows=13, cols=4)
    # Header row
    table.rows[0].cells[0].text = "NO."
    table.rows[0].cells[1].text = "EQUIPMENT"
    table.rows[0].cells[2].text = "DEFECT DESCRIPTION"
    table.rows[0].cells[3].text = "REMARKS"

    # Defect row 1
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "FEEDER PILLAR"
    table.rows[1].cells[2].text = "DOOR HINGE BROKEN"
    table.rows[1].cells[3].text = "REPAIR"

    # Defect row 2
    table.rows[2].cells[0].text = "2"
    table.rows[2].cells[1].text = "SWITCHGEAR"
    table.rows[2].cells[2].text = "HIGH TEV DISCHARGE"
    table.rows[2].cells[3].text = "URGENT"

    # Rows 3 to 12 are left blank (all empty strings)
    for r in table.rows[3:]:
        for c in r.cells:
            c.text = "   "  # whitespace-only cells should also be considered empty

    vi_path = tmp_path / "vi_summary_with_blanks.docx"
    doc.save(vi_path)

    # Passes for expected_count=2
    valid, reason = verify_vi_summary_structure(vi_path, expected_count=2)
    assert valid is True
    assert reason == ""

    # Fails for expected_count=5
    valid_bad, reason_bad = verify_vi_summary_structure(vi_path, expected_count=5)
    assert valid_bad is False
    assert "2 data rows" in reason_bad
    assert "expected defect count 5" in reason_bad

    # Header-only table (0 data rows) fails when expected_count=1
    doc_empty = docx.Document()
    t_empty = doc_empty.add_table(rows=1, cols=4)
    t_empty.rows[0].cells[0].text = "NO."
    t_empty.rows[0].cells[1].text = "EQUIPMENT"
    t_empty.rows[0].cells[2].text = "DEFECT DESCRIPTION"
    t_empty.rows[0].cells[3].text = "REMARKS"
    empty_path = tmp_path / "vi_summary_header_only.docx"
    doc_empty.save(empty_path)

    valid_zero, reason_zero = verify_vi_summary_structure(empty_path, expected_count=1)
    assert valid_zero is False
    assert "0 data rows" in reason_zero


def test_inspect_deliverable_attribution_real_world_talapia_dual_name(tmp_path: Path):
    """Regression test (Issue #46): Deliverable with front page ERMS='TALAPIA' and SITE='KPRC S/B' passes inspection for 'TALAPIA'."""
    from src.full_report.attribution import inspect_deliverable_attribution

    doc = docx.Document()
    # Front page table
    t_fp = doc.add_table(rows=2, cols=3)
    t_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_fp.rows[0].cells[2].text = "TALAPIA"
    t_fp.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t_fp.rows[1].cells[2].text = "KPRC S/B"

    # CBM test sheet table
    t_cbm = doc.add_table(rows=2, cols=4)
    t_cbm.rows[0].cells[0].text = "Substation:"
    t_cbm.rows[0].cells[1].text = "PE TALAPIA"
    t_cbm.rows[1].cells[0].text = "Equipment:"
    t_cbm.rows[1].cells[1].text = "VCB"

    report_path = tmp_path / "005. TALAPIA.docx"
    doc.save(report_path)

    is_valid, reason = inspect_deliverable_attribution(report_path, "TALAPIA")
    assert is_valid is True
    assert reason == ""


def test_inspect_deliverable_attribution_real_world_telekom_photo_grid(tmp_path: Path):
    """Regression test (Issue #46): Deliverable with photo grid headers ['SUBSTATION', '', 'FP/LVDB'] passes inspection."""
    from src.full_report.attribution import inspect_deliverable_attribution

    doc = docx.Document()
    # Front page table
    t_fp = doc.add_table(rows=2, cols=3)
    t_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_fp.rows[0].cells[2].text = "PE 144 TELEKOM TANAH PUTIH"
    t_fp.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t_fp.rows[1].cells[2].text = "TELEKOM TANAH PUTIH"

    # Photo grid table with ['SUBSTATION', '', 'FP/LVDB'] and image shapes
    t_photo = doc.add_table(rows=2, cols=3)
    t_photo.rows[0].cells[0].text = "SUBSTATION"
    t_photo.rows[0].cells[1].text = ""
    t_photo.rows[0].cells[2].text = "FP/LVDB"
    _add_mock_image(t_photo.rows[1].cells[0])
    _add_mock_image(t_photo.rows[1].cells[2])

    # CBM test sheet
    t_cbm = doc.add_table(rows=2, cols=2)
    t_cbm.rows[0].cells[0].text = "Substation"
    t_cbm.rows[0].cells[1].text = "PE 144 TELEKOM TANAH PUTIH"

    report_path = tmp_path / "144. TELEKOM TANAH PUTIH.docx"
    doc.save(report_path)

    is_valid, reason = inspect_deliverable_attribution(report_path, "TELEKOM TANAH PUTIH")
    assert is_valid is True
    assert reason == ""


def test_inspect_deliverable_attribution_arbitrary_photo_grid_captions(tmp_path: Path):
    """Ensure arbitrary custom photo grid captions are skipped by image-based photo grid classification."""
    from src.full_report.attribution import inspect_deliverable_attribution

    doc = docx.Document()
    t_fp = doc.add_table(rows=1, cols=3)
    t_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_fp.rows[0].cells[2].text = "PE TALAPIA"

    # Photo grid with arbitrary captions and image shapes
    t_grid = doc.add_table(rows=3, cols=2)
    t_grid.rows[0].cells[0].text = "SUBSTATION"
    t_grid.rows[0].cells[1].text = "EARTHING & LIGHTNING ARRESTOR"
    _add_mock_image(t_grid.rows[1].cells[0])
    _add_mock_image(t_grid.rows[1].cells[1])
    t_grid.rows[2].cells[0].text = "CORROSION DETECTED"
    t_grid.rows[2].cells[1].text = "GOOD CONDITION"

    # Another photo grid with custom component name
    t_grid2 = doc.add_table(rows=2, cols=2)
    t_grid2.rows[0].cells[0].text = "SUBSTATION"
    t_grid2.rows[0].cells[1].text = "CUSTOM AIR DUCT SYSTEM"
    _add_mock_image(t_grid2.rows[1].cells[0])
    _add_mock_image(t_grid2.rows[1].cells[1])

    report_path = tmp_path / "custom_captions_report.docx"
    doc.save(report_path)

    is_valid, reason = inspect_deliverable_attribution(report_path, "TALAPIA")
    assert is_valid is True
    assert reason == ""


def test_inspect_deliverable_attribution_detects_true_foreign_front_page_leak(tmp_path: Path):
    """True foreign front page leak is detected when neither ERMS nor SITE matches target."""
    from src.full_report.attribution import inspect_deliverable_attribution

    doc = docx.Document()
    t_fp = doc.add_table(rows=2, cols=3)
    t_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_fp.rows[0].cells[2].text = "PE 144 TELEKOM TANAH PUTIH"
    t_fp.rows[1].cells[0].text = "SUBSTATION NAME (SITE)"
    t_fp.rows[1].cells[2].text = "KPRC S/B"

    report_path = tmp_path / "foreign_front_page_leak.docx"
    doc.save(report_path)

    is_valid, reason = inspect_deliverable_attribution(report_path, "TALAPIA")
    assert is_valid is False
    assert "ERMS='PE 144 TELEKOM TANAH PUTIH', SITE='KPRC S/B'" in reason
    assert "(target: 'TALAPIA')" in reason


def test_inspect_deliverable_attribution_detects_true_foreign_cbm_test_sheet_leak(tmp_path: Path):
    """True foreign CBM test sheet leak is detected across text and image-bearing measurement tables."""
    from src.full_report.attribution import inspect_deliverable_attribution

    # Case 1: Standard CBM test sheet header table leak
    doc1 = docx.Document()
    t_fp = doc1.add_table(rows=1, cols=3)
    t_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t_fp.rows[0].cells[2].text = "PE TALAPIA"

    t_cbm = doc1.add_table(rows=2, cols=2)
    t_cbm.rows[0].cells[0].text = "Substation:"
    t_cbm.rows[0].cells[1].text = "TELEKOM TANAH PUTIH"  # Foreign!
    t_cbm.rows[1].cells[0].text = "Equipment:"
    t_cbm.rows[1].cells[1].text = "VCB"

    report1_path = tmp_path / "foreign_cbm_leak.docx"
    doc1.save(report1_path)

    is_valid1, reason1 = inspect_deliverable_attribution(report1_path, "TALAPIA")
    assert is_valid1 is False
    assert "TELEKOM TANAH PUTIH" in reason1

    # Case 2: CBM measurement table containing images (thermal/spot temp) with foreign substation
    doc2 = docx.Document()
    t2_fp = doc2.add_table(rows=1, cols=3)
    t2_fp.rows[0].cells[0].text = "SUBSTATION NAME (ERMS)"
    t2_fp.rows[0].cells[2].text = "PE TALAPIA"

    t2_cbm = doc2.add_table(rows=3, cols=4)
    t2_cbm.rows[0].cells[0].text = "Station Name:"
    t2_cbm.rows[0].cells[1].text = "CENDERAWASIH NO.1"  # Foreign!
    t2_cbm.rows[0].cells[2].text = "Date:"
    t2_cbm.rows[0].cells[3].text = "25/08/2026"
    t2_cbm.rows[1].cells[0].text = "Spot Temp: 55.4 C"
    t2_cbm.rows[1].cells[1].text = "Ambient Temp: 30.1 C"
    t2_cbm.rows[1].cells[2].text = "Delta T: 25.3 C"
    t2_cbm.rows[1].cells[3].text = "Load: 120A"
    _add_mock_image(t2_cbm.rows[2].cells[0])
    _add_mock_image(t2_cbm.rows[2].cells[1])

    report2_path = tmp_path / "foreign_cbm_thermal_image_leak.docx"
    doc2.save(report2_path)

    is_valid2, reason2 = inspect_deliverable_attribution(report2_path, "TALAPIA")
    assert is_valid2 is False
    assert "CENDERAWASIH NO.1" in reason2
