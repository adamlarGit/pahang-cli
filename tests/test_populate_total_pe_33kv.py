"""Unit and integration tests for 33kV Populate TOTAL PE workflow (Ticket #098).

Verifies:
1. Canonical metadata extraction via PpuTestsheetParser / PopulateTotalPeExtractor:
   - Substation number resolution from stem / filename prefix / V2
   - FL ERMS from cell D7 on 33kV (Main)
   - Date from cell D8 on 33kV (Main)
   - Fallback substation name (cell D6) and substation type (stem / FL token)
2. QR02 CBA master metadata enrichment:
   - Resolves statewide ENGR-750-39 via storage.get_engr_33kv_cba_path()
   - Indexes by Column G (FUNCTIONAL LOCATION)
   - Enriches Column H (SUBSTATION NAME) and Column I (TYPE)
   - Graceful fallback to cell D6 / stem token with logged warning when FL not in CBA or CBA missing
3. Row lookup dictionary in LocalExcelTotalPeRepository.upsert_packages:
   - 2-digit zero-padded keys (f"{pe_no:02d}", dt_str)
   - Functional Location keys (fl_num, dt_str) and normalized variants
4. Safe writes into TOTAL PE.xlsx DataCycle1:
   - Columns A-F updated directly (PE NO, FL NUMBER, SUBSTATION NAME, DATE, TYPE, WO)
   - Date formatted as DD-MMM-YYYY
   - Columns beyond Col F preserved untouched
   - Formulas preserved during _sort_datacycle_sheet
   - Ghost formatting sanitized
5. PopulateTotalPeAction runs transparently in 33kV ProjectEnvironment
6. Full end-to-end workflow execution (Auto, All, Specific Folders).
"""

from __future__ import annotations

from datetime import date, datetime
import logging
from pathlib import Path
from typing import Any
import openpyxl
import pytest

from src.core.normalizers import normalize_date_str
from src.master.total_pe import LocalExcelTotalPeRepository
from src.project.environment import ProjectEnvironment
from src.project.models import ProjectMetadata
from src.project.storage import LocalWorkspaceStorage
from src.project_workflow_actions import PopulateTotalPeAction
from src.testsheet.models import SubstationTestsheetPackage, TestsheetData
from src.testsheet.parsers.ppu import PpuTestsheetParser
from src.testsheet.repository import SubstationTestsheetRepository
from src.workflows.models import PopulateMode, PopulateTotalPeRequest
from src.workflows.populate_total_pe import (
    PopulateTotalPeExtractor,
    PopulateTotalPeFilter,
    PopulateTotalPeWorkflow,
    load_qr02_cba_33kv_metadata,
)
from tests.fixtures.ppu_workbook_factory import create_ppu_testsheet_workbook


def _create_sample_cba_33kv(path: Path, rows: list[list[Any]]) -> None:
    """Helper to create a synthetic 33kV QR02 CBA master workbook."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QR02 CBA"
    # Row 1: Headers
    # Col A(1)..Col F(6): Arbitrary
    # Col G(7): FUNCTIONAL LOCATION
    # Col H(8): SUBSTATION NAME
    # Col I(9): TYPE
    # Col J(10): SYSTEM
    ws.append([
        "NO", "DISTRICT", "ZONE", "STATION CODE", "CYCLE", "VOLTAGE",
        "FUNCTIONAL LOCATION", "SUBSTATION NAME", "TYPE", "SYSTEM",
    ])
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _create_sample_total_pe(path: Path, rows: list[list[Any]]) -> None:
    """Helper to create a synthetic TOTAL PE.xlsx workbook with DataCycle1."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DataCycle1"
    ws.append([
        "PE NO", "FL NUMBER", "SUBSTATION NAME", "DATE", "TYPE", "WO", "SCOPE", "FORMULA_COL",
    ])
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


@pytest.fixture
def mock_env_33kv(tmp_path: Path) -> ProjectEnvironment:
    """Provide an isolated 33kV ProjectEnvironment backed by tmp_path."""
    meta = ProjectMetadata(
        key="test_33kv_thunder",
        name="Contract PO 42394550 - THUNDER",
        po_number="PO 42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="Cycle 4",
        technologies=("IR", "DG", "US", "TEV", "VI"),
        base_path=str(tmp_path),
    )
    storage = LocalWorkspaceStorage(tmp_path, voltage_type="33kV")
    return ProjectEnvironment(metadata=meta, storage=storage)


# =============================================================================
# Seam 1: PpuTestsheetParser Metadata Extraction
# =============================================================================

def test_ppu_metadata_extraction_canonical_fields(tmp_path: Path) -> None:
    """Verify PpuTestsheetParser extracts substation_number, fl_erms, date, fallback name and type."""
    ts_file = tmp_path / "01. PPU APMM 33-11kV (VI).xlsx"
    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=1,
        substation_name="PPU APMM 33/11kV",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
        v2_substation_num_33kv=1,
    )
    wb.save(ts_file)
    wb.close()

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_metadata(ts_file)

    assert data.substation_number == 1
    assert data.fl_erms == "CGBG/PPU/G00001"
    assert data.date_str == "07-10-2026"
    assert data.substation_name_erms == "PPU APMM 33/11kV"
    assert data.substation_type == "PPU"


def test_ppu_metadata_extraction_fallback_type_from_fl_token(tmp_path: Path) -> None:
    """Verify substation_type resolves from FL token when name/filename lack explicit type."""
    ts_file = tmp_path / "02. APMM SPECIAL.xlsx"
    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=2,
        substation_name="APMM SPECIAL SITE",
        fl_number="CGBG/SSU/G00002",
        date_val="2026-10-08",
        v2_substation_num_33kv=2,
    )
    wb.save(ts_file)
    wb.close()

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_metadata(ts_file)

    assert data.substation_number == 2
    assert data.substation_type == "SSU"


def test_ppu_metadata_extraction_number_from_parent_folder(tmp_path: Path) -> None:
    """Verify substation_number resolves from parent folder when V2 is empty and file name lacks prefix."""
    folder = tmp_path / "05. PPU BENTONG"
    folder.mkdir(parents=True)
    ts_file = folder / "Testsheet.xlsx"

    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=None,
        substation_name="PPU BENTONG",
        fl_number="CBTG/PPU/G00005",
        date_val="2026-10-09",
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=None,
    )
    wb.save(ts_file)
    wb.close()

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_metadata(ts_file)

    assert data.substation_number == 5


# =============================================================================
# Seam 2: QR02 CBA Master Metadata Enrichment
# =============================================================================

def test_qr02_cba_metadata_loader(tmp_path: Path) -> None:
    """Verify load_qr02_cba_33kv_metadata reads and indexes Col G, H, I."""
    cba_path = tmp_path / "ENGR-750-39 -CBA-PAHANG - 2026.xlsx"
    _create_sample_cba_33kv(cba_path, [
        [1, "KUANTAN", "CGBG", "APMM", 4, "33kV", "CGBG/PPU/G00001", "PPU APMM GEBENG", "PPU", "33/11"],
        [2, "KUANTAN", "CGBG", "BASF", 4, "33kV", "CGBG/SSU/G00002", "SSU BASF PETRONAS", "SSU", "33/11"],
    ])

    meta = load_qr02_cba_33kv_metadata(cba_path)
    assert "CGBG/PPU/G00001" in meta
    assert meta["CGBG/PPU/G00001"]["substation_name"] == "PPU APMM GEBENG"
    assert meta["CGBG/PPU/G00001"]["substation_type"] == "PPU"

    assert "CGBG/SSU/G00002" in meta
    assert meta["CGBG/SSU/G00002"]["substation_name"] == "SSU BASF PETRONAS"
    assert meta["CGBG/SSU/G00002"]["substation_type"] == "SSU"


def test_enrich_packages_with_cba_success(mock_env_33kv: ProjectEnvironment) -> None:
    """Verify PopulateTotalPeExtractor enriches package data from QR02 CBA."""
    cba_path = mock_env_33kv.storage.get_engr_33kv_cba_path("2026")
    _create_sample_cba_33kv(cba_path, [
        [1, "KUANTAN", "CGBG", "APMM", 4, "33kV", "CGBG/PPU/G00001", "PPU APMM OFFICIAL", "MPPU", "33/11"],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("dummy.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM 33/11kV",
            substation_type="PPU",
            date_str="07-10-2026",
        ),
    )

    extractor = PopulateTotalPeExtractor()
    enriched = extractor.enrich_packages_with_cba([pkg], environment=mock_env_33kv)

    assert len(enriched) == 1
    assert enriched[0].data is not None
    assert enriched[0].data.substation_name_erms == "PPU APMM OFFICIAL"
    assert enriched[0].data.substation_type == "MPPU"


def test_enrich_packages_with_cba_fallback_on_unmatched_fl(
    mock_env_33kv: ProjectEnvironment, caplog: pytest.LogCaptureFixture
) -> None:
    """Verify graceful fallback to cell D6 name and stem type with warning when FL not in CBA."""
    cba_path = mock_env_33kv.storage.get_engr_33kv_cba_path("2026")
    _create_sample_cba_33kv(cba_path, [
        [1, "KUANTAN", "CGBG", "OTHER", 4, "33kV", "CGBG/PPU/G99999", "OTHER PPU", "PPU", "33/11"],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("dummy.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM 33/11kV",
            substation_type="PPU",
            date_str="07-10-2026",
        ),
    )

    extractor = PopulateTotalPeExtractor()
    with caplog.at_level(logging.WARNING):
        enriched = extractor.enrich_packages_with_cba([pkg], environment=mock_env_33kv)

    assert len(enriched) == 1
    assert enriched[0].data is not None
    # Maintained fallbacks
    assert enriched[0].data.substation_name_erms == "PPU APMM 33/11kV"
    assert enriched[0].data.substation_type == "PPU"
    assert "not found in QR02 CBA" in caplog.text


def test_enrich_packages_with_cba_missing_cba_file_graceful(
    mock_env_33kv: ProjectEnvironment, caplog: pytest.LogCaptureFixture
) -> None:
    """Verify graceful fallback without exception when CBA workbook file does not exist."""
    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("dummy.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM 33/11kV",
            substation_type="PPU",
            date_str="07-10-2026",
        ),
    )

    extractor = PopulateTotalPeExtractor()
    with caplog.at_level(logging.WARNING):
        enriched = extractor.enrich_packages_with_cba([pkg], environment=mock_env_33kv)

    assert len(enriched) == 1
    assert enriched[0].data.substation_name_erms == "PPU APMM 33/11kV"
    assert "not found" in caplog.text


# =============================================================================
# Seam 3: Row Lookup & Safe Upsert in LocalExcelTotalPeRepository
# =============================================================================

def test_upsert_packages_2digit_zero_padded_key_lookup(tmp_path: Path) -> None:
    """Verify upsert_packages updates row matched via 2-digit zero-padded key ('01', '07-10-2026')."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"

    # Pre-populate row with PE 1 formatted as '01'
    _create_sample_total_pe(total_pe_file, [
        ["01", "CGBG/PPU/G00001", "OLD NAME", "07-10-2026", "OLD TYPE", None, "FULL", "=A2+1"],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("01. PPU APMM.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM UPDATED",
            substation_type="PPU",
            date_str="07-10-2026",
            wo_number="45001234",
        ),
    )

    new_cnt, upd_cnt = repo.upsert_packages(total_pe_file, [pkg])
    assert new_cnt == 0
    assert upd_cnt == 1

    wb = openpyxl.load_workbook(total_pe_file)
    ws = wb["DataCycle1"]
    assert ws.cell(2, 1).value == 1
    assert ws.cell(2, 2).value == "CGBG/PPU/G00001"
    assert ws.cell(2, 3).value == "PPU APMM UPDATED"
    assert ws.cell(2, 4).number_format == "DD-MMM-YYYY"
    assert ws.cell(2, 5).value == "PPU"
    assert str(ws.cell(2, 6).value) == "45001234"
    assert ws.cell(2, 7).value == "FULL"  # Col G preserved
    assert ws.cell(2, 8).value == "=A2+1"  # Formula preserved
    wb.close()


def test_upsert_packages_fl_key_lookup(tmp_path: Path) -> None:
    """Verify upsert_packages updates row matched via Functional Location key."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"

    # Pre-populate row with empty PE NO but matching FL
    _create_sample_total_pe(total_pe_file, [
        [None, "CGBG/SSU/G00002", "OLD SSU", "08-10-2026", "SSU", None, "FULL", None],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("02. SSU BASF.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="08-10-2026",
        substation_number=2,
        data=TestsheetData(
            substation_number=2,
            fl_erms="CGBG/SSU/G00002",
            substation_name_erms="SSU BASF PETRONAS",
            substation_type="SSU",
            date_str="08-10-2026",
            wo_number="45005678",
        ),
    )

    new_cnt, upd_cnt = repo.upsert_packages(total_pe_file, [pkg])
    assert new_cnt == 0
    assert upd_cnt == 1

    wb = openpyxl.load_workbook(total_pe_file)
    ws = wb["DataCycle1"]
    assert ws.cell(2, 1).value == 2
    assert ws.cell(2, 2).value == "CGBG/SSU/G00002"
    assert ws.cell(2, 3).value == "SSU BASF PETRONAS"
    assert str(ws.cell(2, 6).value) == "45005678"
    wb.close()


def test_upsert_packages_with_enriched_metadata(tmp_path: Path) -> None:
    """Verify upsert_packages stores metadata enriched from CBA into TOTAL PE."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"
    cba_file = tmp_path / "ENGR-750-39 -CBA-PAHANG - 2026.xlsx"

    _create_sample_cba_33kv(cba_file, [
        [1, "KUANTAN", "CGBG", "APMM", 4, "33kV", "CGBG/PPU/G00001", "PPU APMM CBA NAME", "MPPU", "33/11"],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("01. PPU APMM.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM CELL D6",
            substation_type="PPU",
            date_str="07-10-2026",
        ),
    )

    extractor = PopulateTotalPeExtractor()
    enriched = extractor.enrich_packages_with_cba([pkg], cba_path=cba_file)
    new_cnt, upd_cnt = repo.upsert_packages(total_pe_file, enriched)
    assert new_cnt == 1

    wb = openpyxl.load_workbook(total_pe_file)
    ws = wb["DataCycle1"]
    assert ws.cell(2, 1).value == 1
    assert ws.cell(2, 3).value == "PPU APMM CBA NAME"
    assert ws.cell(2, 5).value == "MPPU"
    wb.close()


def test_get_existing_auto_keys_indexes_2digit_and_fl(tmp_path: Path) -> None:
    """Verify get_existing_auto_keys includes 2-digit zero-padded and FL keys."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"

    _create_sample_total_pe(total_pe_file, [
        [6, "CGBG/PPU/G00006", "PPU SIX", "07-10-2026", "PPU", None, None, None],
    ])

    keys = repo.get_existing_auto_keys(total_pe_file)
    # Check 2-digit zero-padded key
    assert ("06", "07-10-2026") in keys
    assert ("6", "07-10-2026") in keys
    assert ("006", "07-10-2026") in keys
    # Check FL key
    assert ("CGBG/PPU/G00006", "07-10-2026") in keys
    assert ("CGBGPPUG00006", "07-10-2026") in keys


# =============================================================================
# Seam 4: End-to-End Workflow in 33kV ProjectEnvironment
# =============================================================================

def test_populate_total_pe_workflow_33kv_e2e(mock_env_33kv: ProjectEnvironment) -> None:
    """Verify full workflow on 33kV workspace with QR02 CBA enrichment and AUTO mode deduplication."""
    # 1. Setup CBA master file
    cba_path = mock_env_33kv.storage.get_engr_33kv_cba_path("2026")
    _create_sample_cba_33kv(cba_path, [
        [1, "KUANTAN", "CGBG", "APMM", 4, "33kV", "CGBG/PPU/G00001", "PPU APMM 33/11kV OFFICIAL", "PPU", "33/11"],
    ])

    # 2. Setup TESTSHEET folder structure
    ts_dir = mock_env_33kv.storage.get_testsheet_dir() / "KUANTAN" / "10. OCTOBER" / "07-10-2026"
    ts_dir.mkdir(parents=True)
    (ts_dir / "UNSORTED RAW DATA").mkdir()

    ts_file = ts_dir / "01. PPU APMM 33-11kV (VI).xlsx"
    wb_ppu = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=1,
        substation_name="PPU APMM SITE D6",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
        v2_substation_num_33kv=1,
    )
    wb_ppu.save(ts_file)
    wb_ppu.close()

    # 3. Setup TOTAL PE.xlsx
    total_pe_path = mock_env_33kv.storage.get_total_pe_path()
    _create_sample_total_pe(total_pe_path, [])

    # 4. Execute workflow in AUTO mode
    workflow = PopulateTotalPeWorkflow()
    req = PopulateTotalPeRequest(mode=PopulateMode.AUTO)
    res = workflow.execute(mock_env_33kv, req)

    assert res.new_rows_added == 1

    # Verify content in DataCycle1
    wb_pe = openpyxl.load_workbook(total_pe_path)
    ws_pe = wb_pe["DataCycle1"]
    assert ws_pe.cell(2, 1).value == 1
    assert ws_pe.cell(2, 2).value == "CGBG/PPU/G00001"
    assert ws_pe.cell(2, 3).value == "PPU APMM 33/11kV OFFICIAL"  # Enriched from CBA
    assert ws_pe.cell(2, 4).number_format == "DD-MMM-YYYY"
    assert ws_pe.cell(2, 5).value == "PPU"
    wb_pe.close()

    # 5. Re-run AUTO mode -> should be skipped because ("01", "07-10-2026") exists
    res_second = workflow.execute(mock_env_33kv, req)
    assert res_second.new_rows_added == 0


def test_populate_total_pe_action_registration_and_execution(
    mock_env_33kv: ProjectEnvironment, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Verify PopulateTotalPeAction runs smoothly in 33kV active project environment."""
    # Setup minimal TOTAL PE and TESTSHEET
    total_pe_path = mock_env_33kv.storage.get_total_pe_path()
    _create_sample_total_pe(total_pe_path, [])

    ts_dir = mock_env_33kv.storage.get_testsheet_dir() / "KUANTAN" / "10. OCTOBER" / "07-10-2026"
    ts_dir.mkdir(parents=True)
    (ts_dir / "UNSORTED RAW DATA").mkdir()

    ts_file = ts_dir / "01. PPU APMM 33-11kV (VI).xlsx"
    wb_ppu = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=1,
        substation_name="PPU APMM",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
    )
    wb_ppu.save(ts_file)
    wb_ppu.close()

    # Mock interactive menu selection to choose 'auto'
    import src.cli_selectors as cli_selectors
    monkeypatch.setattr(cli_selectors, "select_one", lambda title, options: "auto")

    action = PopulateTotalPeAction("Populate TOTAL PE (from testsheets)")
    result = action.execute(mock_env_33kv)

    assert result is not None
    assert getattr(result, "new_rows_added", None) == 1


def test_ppu_metadata_extraction_date_parent_folder_not_matched_as_substation(tmp_path: Path) -> None:
    """Verify parent folder formatted as a date (DD-MM-YYYY) is NOT treated as a numeric substation prefix."""
    date_folder = tmp_path / "07-10-2026"
    date_folder.mkdir(parents=True)
    ts_file = date_folder / "Testsheet.xlsx"

    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=None,
        substation_name="PPU UNPREFIXED",
        fl_number="CGBG/PPU/G00007",
        date_val="2026-10-07",
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=None,
    )
    wb.save(ts_file)
    wb.close()

    parser = PpuTestsheetParser()
    # Should raise ValueError because neither cell V2, nor filename, nor parent (date folder) has a valid substation prefix
    with pytest.raises(ValueError, match="Unable to determine substation number"):
        parser.extract_testsheet_metadata(ts_file)


def test_ppu_metadata_extraction_iso_date_parent_folder_not_matched_as_substation(tmp_path: Path) -> None:
    """Verify parent folder formatted as ISO date (YYYY-MM-DD) is NOT treated as a numeric substation prefix (e.g. 2026)."""
    date_folder = tmp_path / "2026-10-07"
    date_folder.mkdir(parents=True)
    ts_file = date_folder / "Testsheet.xlsx"

    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=None,
        substation_name="PPU UNPREFIXED",
        fl_number="CGBG/PPU/G00007",
        date_val="2026-10-07",
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=None,
    )
    wb.save(ts_file)
    wb.close()

    parser = PpuTestsheetParser()
    with pytest.raises(ValueError, match="Unable to determine substation number"):
        parser.extract_testsheet_metadata(ts_file)


def test_repository_discovery_iso_date_parent_folder_default_substation_number(tmp_path: Path) -> None:
    """Verify SubstationTestsheetRepository does not treat ISO date parent folder (2026-10-07) as substation number 2026."""
    date_folder = tmp_path / "KUANTAN" / "10. OCTOBER" / "2026-10-07"
    date_folder.mkdir(parents=True)
    (date_folder / "UNSORTED RAW DATA").mkdir()
    ts_file = date_folder / "Testsheet.xlsx"

    wb = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=None,
        substation_name="PPU UNPREFIXED",
        fl_number="CGBG/PPU/G00007",
        date_val="2026-10-07",
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=None,
    )
    wb.save(ts_file)
    wb.close()

    repo = SubstationTestsheetRepository()
    pkgs = repo.discover_packages(tmp_path / "KUANTAN")
    assert len(pkgs) == 1
    assert pkgs[0].substation_number == 1  # Default fallback 1, NOT 2026!


def test_ppu_metadata_extraction_date_formats_and_malaysian_fallback(tmp_path: Path) -> None:
    """Verify various date formats in D8 parse cleanly or fall back without crashing."""
    ts_file = tmp_path / "01. PPU DATES.xlsx"

    # Case 1: Slash ISO date "2026/10/07"
    wb1 = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=1,
        substation_name="PPU DATES",
        fl_number="CGBG/PPU/G00001",
        date_val="2026/10/07",
        v2_substation_num_33kv=1,
    )
    wb1.save(ts_file)
    wb1.close()

    parser = PpuTestsheetParser()
    data1 = parser.extract_testsheet_metadata(ts_file)
    assert data1.date_str == "07-10-2026"

    # Case 2: Malaysian irregular text date "7hb Oktober 2026"
    wb2 = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=1,
        substation_name="PPU DATES",
        fl_number="CGBG/PPU/G00001",
        date_val="7hb Oktober 2026",
        v2_substation_num_33kv=1,
    )
    wb2.save(ts_file)
    wb2.close()

    data2 = parser.extract_testsheet_metadata(ts_file)
    assert data2.date_str == "7hb Oktober 2026"


def test_upsert_packages_multirow_sort_and_formula_preservation(tmp_path: Path) -> None:
    """Verify _sort_datacycle_sheet properly sorts multiple rows without losing static text or corrupting formulas."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"

    # Row 2: PE 2, formula in Col H, static value in Col G
    # Row 3: PE 1, formula in Col H, static value in Col G
    _create_sample_total_pe(total_pe_file, [
        [2, "CGBG/PPU/G00002", "PPU TWO", "07-10-2026", "PPU", "45000002", "SCOPE TWO", "=A2*10"],
        [1, "CGBG/PPU/G00001", "PPU ONE", "07-10-2026", "PPU", "45000001", "SCOPE ONE", "=A3*10"],
    ])

    # Run sort by performing a no-op upsert of an existing row
    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("01. PPU ONE.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU ONE UPDATED",
            substation_type="PPU",
            date_str="07-10-2026",
            wo_number="45000001",
        ),
    )

    new_cnt, upd_cnt = repo.upsert_packages(total_pe_file, [pkg])
    assert new_cnt == 0
    assert upd_cnt == 1

    wb = openpyxl.load_workbook(total_pe_file)
    ws = wb["DataCycle1"]
    # Row 2 must now be PE 1 (sorted)
    assert ws.cell(2, 1).value == 1
    assert ws.cell(2, 2).value == "CGBG/PPU/G00001"
    assert ws.cell(2, 3).value == "PPU ONE UPDATED"
    assert ws.cell(2, 7).value == "SCOPE ONE"  # Preserved from PE 1 row!
    assert ws.cell(2, 8).value == "=A2*10"    # Formula preserved on row 2

    # Row 3 must now be PE 2 (sorted)
    assert ws.cell(3, 1).value == 2
    assert ws.cell(3, 2).value == "CGBG/PPU/G00002"
    assert ws.cell(3, 3).value == "PPU TWO"
    assert ws.cell(3, 7).value == "SCOPE TWO"  # Preserved from PE 2 row!
    assert ws.cell(3, 8).value == "=A3*10"    # Formula preserved on row 3
    wb.close()


def test_upsert_packages_preserves_existing_work_order(tmp_path: Path) -> None:
    """Verify upsert_packages preserves existing WO number in Col F if incoming package has empty WO."""
    repo = LocalExcelTotalPeRepository()
    total_pe_file = tmp_path / "TOTAL PE.xlsx"

    _create_sample_total_pe(total_pe_file, [
        [1, "CGBG/PPU/G00001", "PPU OLD", "07-10-2026", "PPU", "45009999", "FULL", None],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("01. PPU APMM.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU NEW NAME",
            substation_type="PPU",
            date_str="07-10-2026",
            wo_number="",  # Empty WO in incoming package
        ),
    )

    new_cnt, upd_cnt = repo.upsert_packages(total_pe_file, [pkg])
    assert new_cnt == 0
    assert upd_cnt == 1

    wb = openpyxl.load_workbook(total_pe_file)
    ws = wb["DataCycle1"]
    assert ws.cell(2, 3).value == "PPU NEW NAME"
    assert ws.cell(2, 6).value == "45009999"  # Preserved!
    wb.close()


def test_qr02_cba_sentinel_strings_fall_back_gracefully(mock_env_33kv: ProjectEnvironment) -> None:
    """Verify CBA entries with sentinel values ('-', 'None', 'N/A') fall back to testsheet cell D6 and stem."""
    cba_path = mock_env_33kv.storage.get_engr_33kv_cba_path("2026")
    _create_sample_cba_33kv(cba_path, [
        [1, "KUANTAN", "CGBG", "APMM", 4, "33kV", "CGBG/PPU/G00001", "-", "None", "33/11"],
    ])

    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("dummy.xlsx"),
        unsorted_raw_data_dir=Path("raw"),
        station="KUANTAN",
        month="10. OCTOBER",
        date_str="07-10-2026",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            fl_erms="CGBG/PPU/G00001",
            substation_name_erms="PPU APMM D6 ORIGINAL",
            substation_type="PPU",
            date_str="07-10-2026",
        ),
    )

    extractor = PopulateTotalPeExtractor()
    enriched = extractor.enrich_packages_with_cba([pkg], environment=mock_env_33kv)

    assert len(enriched) == 1
    assert enriched[0].data is not None
    assert enriched[0].data.substation_name_erms == "PPU APMM D6 ORIGINAL"
    assert enriched[0].data.substation_type == "PPU"


def test_auto_mode_deduplication_with_v2_substation_number(mock_env_33kv: ProjectEnvironment) -> None:
    """Verify AUTO mode deduplication works when package initially lacks prefix but cell V2 specifies substation number."""
    total_pe_path = mock_env_33kv.storage.get_total_pe_path()
    # TOTAL PE already has PE 5 on 07-10-2026
    _create_sample_total_pe(total_pe_path, [
        [5, "CGBG/PPU/G00005", "PPU FIVE", "07-10-2026", "PPU", None, None, None],
    ])

    ts_dir = mock_env_33kv.storage.get_testsheet_dir() / "KUANTAN" / "10. OCTOBER" / "07-10-2026"
    ts_dir.mkdir(parents=True)
    (ts_dir / "UNSORTED RAW DATA").mkdir()

    # File without numeric prefix in filename, but cell V2 = 5
    ts_file = ts_dir / "PPU FIVE SPECIAL.xlsx"
    wb_ppu = create_ppu_testsheet_workbook(
        path=ts_file,
        substation_number=None,
        substation_name="PPU FIVE",
        fl_number="CGBG/PPU/G00005",
        date_val="2026-10-07",
        v2_substation_num_33kv=5,
    )
    wb_ppu.save(ts_file)
    wb_ppu.close()

    workflow = PopulateTotalPeWorkflow()
    req = PopulateTotalPeRequest(mode=PopulateMode.AUTO)
    res = workflow.execute(mock_env_33kv, req)

    # Must be deduplicated and not re-added
    assert res.new_rows_added == 0


def test_populate_total_pe_workflow_all_and_specific_folders_33kv(
    mock_env_33kv: ProjectEnvironment,
) -> None:
    """Verify workflow execution in ALL mode and SPECIFIC_FOLDERS mode on 33kV workspace."""
    total_pe_path = mock_env_33kv.storage.get_total_pe_path()
    _create_sample_total_pe(total_pe_path, [])

    # Setup two date folders
    ts_dir_1 = mock_env_33kv.storage.get_testsheet_dir() / "KUANTAN" / "10. OCTOBER" / "07-10-2026"
    ts_dir_1.mkdir(parents=True)
    (ts_dir_1 / "UNSORTED RAW DATA").mkdir()

    ts_file_1 = ts_dir_1 / "01. PPU APMM 33-11kV (VI).xlsx"
    wb_1 = create_ppu_testsheet_workbook(
        path=ts_file_1,
        substation_number=1,
        substation_name="PPU APMM",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
    )
    wb_1.save(ts_file_1)
    wb_1.close()

    ts_dir_2 = mock_env_33kv.storage.get_testsheet_dir() / "KUANTAN" / "10. OCTOBER" / "08-10-2026"
    ts_dir_2.mkdir(parents=True)
    (ts_dir_2 / "UNSORTED RAW DATA").mkdir()

    ts_file_2 = ts_dir_2 / "02. SSU BASF (VI).xlsx"
    wb_2 = create_ppu_testsheet_workbook(
        path=ts_file_2,
        substation_number=2,
        substation_name="SSU BASF",
        fl_number="CGBG/SSU/G00002",
        date_val="2026-10-08",
    )
    wb_2.save(ts_file_2)
    wb_2.close()

    workflow = PopulateTotalPeWorkflow()

    # 1. SPECIFIC_FOLDERS mode matching only 07-10-2026
    req_spec = PopulateTotalPeRequest(
        mode=PopulateMode.SPECIFIC_FOLDERS,
        target_folder_names=["07-10-2026"],
    )
    res_spec = workflow.execute(mock_env_33kv, req_spec)
    assert res_spec.new_rows_added == 1

    wb = openpyxl.load_workbook(total_pe_path)
    ws = wb["DataCycle1"]
    assert ws.max_row == 2
    assert ws.cell(2, 1).value == 1
    wb.close()

    # 2. ALL mode re-evaluates all packages and updates/adds remaining
    req_all = PopulateTotalPeRequest(mode=PopulateMode.ALL)
    res_all = workflow.execute(mock_env_33kv, req_all)
    assert res_all.new_rows_added == 1  # SSU BASF added, PPU APMM updated

    wb = openpyxl.load_workbook(total_pe_path)
    ws = wb["DataCycle1"]
    assert ws.max_row == 3
    assert ws.cell(2, 1).value == 1
    assert ws.cell(3, 1).value == 2
    wb.close()

