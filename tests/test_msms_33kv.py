"""Unit and integration tests for 33kV MSMS pipeline and TestsheetReadingMapper for T1SUB33.

Ticket #100:
- Support Maximo job plan T1SUB33 across MSMS CSV ingestion and validation.
- Polymorphic ReadingMapperStrategy seam in TestsheetReadingMapper (PceReadingMapper & PpuReadingMapper).
- Correlate bay-level TNBLOCATION hierarchy paths to parent substation package.
- Reading update logic in PopulateDataMsmsWorkflow (GAUGE measurements and CHARACTERISTIC defect parity).
- Move completed CSVs to PYTHON/MSMS/COMPLETED/.
"""

from __future__ import annotations

import csv
from pathlib import Path
import openpyxl
import pytest

from src.core.normalizers import (
    extract_gas_pressure,
    extract_heater_amp,
    format_load_amp,
)
from src.project.environment import ProjectEnvironment
from src.project.models import ProjectMetadata
from src.project.storage import LocalWorkspaceStorage
from src.quick_report.defects import ViDefectRecord
from src.testsheet.mapper import (
    PceReadingMapper,
    PpuReadingMapper,
    ReadingMapperStrategy,
    TestsheetReadingMapper,
    build_ppu_bay_map,
    extract_root_substation_token,
    parse_equipment_index,
)
from src.workflows.ingest_msms_csv import (
    IngestMsmsCsvExtractor,
    IngestMsmsCsvWorkflow,
    extract_date_from_filename,
)
from src.workflows.models import (
    IngestMsmsCsvRequest,
    PopulateDataMsmsRequest,
    PopulateMode,
)
from src.workflows.populate_data_msms import (
    PopulateDataMsmsWorkflow,
    match_vi_defect,
)
from tests.fixtures.ppu_workbook_factory import create_ppu_testsheet_workbook


@pytest.fixture
def mock_33kv_env(tmp_path: Path) -> ProjectEnvironment:
    meta = ProjectMetadata(
        key="thunder_2026",
        name="Pahang 33kV Thunder Test",
        po_number="PO42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="4",
        technologies=("IR", "DG", "US", "TEV", "VI"),
        base_path=str(tmp_path),
    )
    storage = LocalWorkspaceStorage(tmp_path)
    return ProjectEnvironment(metadata=meta, storage=storage)


# ==============================================================================
# 1. Filename & CSV Schema Validation (T1SUB33)
# ==============================================================================

def test_extract_date_from_33kv_filename_conventions() -> None:
    assert extract_date_from_filename("07-10-2026_001.csv") == "07-10-2026"
    assert extract_date_from_filename("08-10-2026_002.csv") == "08-10-2026"
    assert extract_date_from_filename("33KV_T1SUB33_14-10-2026_001.csv") == "14-10-2026"
    assert extract_date_from_filename("15.10.2026_001.csv") == "15-10-2026"


def test_validate_t1sub33_maximo_headers(tmp_path: Path) -> None:
    extractor = IngestMsmsCsvExtractor()
    csv_dir = tmp_path / "RAW"
    csv_dir.mkdir()

    # Valid T1SUB33 CSV containing all 7 required Maximo headers
    t1sub33_valid = csv_dir / "07-10-2026_001.csv"
    valid_header = "ACTFINISH,ACTSTART,JPNUM,WONUM,TNBLOCATION,METERNAME,TNBNEWREADING,DESCRIPTION\n"
    valid_row = ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,VI_NER_EARTH,NO,Defect?\n"
    t1sub33_valid.write_text(valid_header + valid_row, encoding="utf-8")

    files = extractor.extract_files(csv_dir)
    assert len(files) == 1
    assert files[0] == t1sub33_valid

    # Missing JPNUM in a T1SUB33-context file (or missing required header)
    t1sub33_invalid = csv_dir / "invalid_33kv.csv"
    # Missing ACTFINISH
    invalid_header = "ACTSTART,JPNUM,WONUM,TNBLOCATION,METERNAME,TNBNEWREADING\n"
    invalid_row = "2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,VI_NER_EARTH,NO\n"
    t1sub33_invalid.write_text(invalid_header + invalid_row, encoding="utf-8")

    with pytest.raises(ValueError, match="Missing required CSV headers"):
        extractor.extract_files(csv_dir)


def test_generic_jpnum_header_not_treated_as_t1sub33(tmp_path: Path) -> None:
    extractor = IngestMsmsCsvExtractor()
    csv_dir = tmp_path / "RAW_11KV"
    csv_dir.mkdir()

    # 11kV export containing generic JPNUM column with non-T1SUB33 job plan
    export_11kv = csv_dir / "07-10-2026_11kv.csv"
    header_11kv = "WONUM,TNBLOCATION,METERNAME,JPNUM\n"
    row_11kv = "W123456,CCHL/PCEJ00001,TH_S11_BR_AVG_PE13R,T1SUB11\n"
    export_11kv.write_text(header_11kv + row_11kv, encoding="utf-8")

    files = extractor.extract_files(csv_dir)
    assert len(files) == 1
    assert files[0] == export_11kv


# ==============================================================================
# 2. Correlate Bay-level TNBLOCATION Hierarchy Paths to Parent Substation Package
# ==============================================================================

def test_extract_root_substation_token() -> None:
    # 33kV bay segments
    assert extract_root_substation_token("CGBG/PPUAPMM/33KV/3P5") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CGBG/PPUAPMM/33KV/1HO") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CGBG/SSULNAS/33KV/5S5") == "CGBG/SSULNAS"
    assert extract_root_substation_token("CKTN/SSUKIIP/33KV/6S5") == "CKTN/SSUKIIP"
    assert extract_root_substation_token("CGBG/SSUFERRO/33KV/BT1") == "CGBG/SSUFERRO"

    # 11kV bay segments
    assert extract_root_substation_token("CGBG/PPUAPMM/11KV/5K5") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CGBG/PPUAPMM/11KV/31") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CGBG/PPUAPMM/11KV/TRANS") == "CGBG/PPUAPMM"

    # Transformer segments
    assert extract_root_substation_token("CGBG/PPUAPMM/TX/PTX1") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CGBG/PPUAPMM/TX/PTX2") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CCHL/PCEJ00024/TX/DTX1") == "CCHL/PCEJ00024"

    # Root substation without bay segments
    assert extract_root_substation_token("CGBG/PPUAPMM") == "CGBG/PPUAPMM"
    assert extract_root_substation_token("CKTN/SSUKIIP") == "CKTN/SSUKIIP"
    assert extract_root_substation_token("") == ""


def test_parse_equipment_index_33kv() -> None:
    # 33kV switchgear
    cat, idx = parse_equipment_index("CGBG/PPUAPMM/33KV/1")
    assert cat == "33KV" and idx == 1

    cat, idx = parse_equipment_index("CGBG/PPUAPMM/33KV/3P5")
    assert cat == "33KV" and idx == 3

    # Transformer
    cat, idx = parse_equipment_index("CGBG/PPUAPMM/TX/PTX1")
    assert cat == "TX" and idx == 1

    cat, idx = parse_equipment_index("CGBG/PPUAPMM/TX/PTX2")
    assert cat == "TX" and idx == 2

    # 11kV switchgear
    cat, idx = parse_equipment_index("CGBG/PPUAPMM/11KV/5K5")
    assert cat == "11KV" and idx == 5


# ==============================================================================
# 3. Polymorphic ReadingMapperStrategy Seam & PpuReadingMapper Resolution
# ==============================================================================

def test_reading_mapper_strategy_polymorphism() -> None:
    pce_mapper = PceReadingMapper()
    ppu_mapper = PpuReadingMapper()

    assert isinstance(pce_mapper, ReadingMapperStrategy)
    assert isinstance(ppu_mapper, ReadingMapperStrategy)

    # TestsheetReadingMapper adapts strategy based on voltage_type or explicit param
    facade_11kv = TestsheetReadingMapper(voltage_type="11kV")
    assert isinstance(facade_11kv.strategy, PceReadingMapper)

    facade_33kv = TestsheetReadingMapper(voltage_type="33kV")
    assert isinstance(facade_33kv.strategy, PpuReadingMapper)

    facade_explicit = TestsheetReadingMapper(strategy=ppu_mapper)
    assert facade_explicit.strategy is ppu_mapper


def test_pce_mapper_preserves_11kv_parity() -> None:
    mapper = PceReadingMapper()

    # Background metadata on PCE Testsheet
    assert mapper.get_target("BG_ROOM_TV") == ("PCE Testsheet", "P6")
    assert mapper.get_target("BG_ROOM_HUM") == ("PCE Testsheet", "S6")
    assert mapper.get_target("BG_ROOM_TEM") == ("PCE Testsheet", "W6")

    # RMU Switchgear slot 1
    assert mapper.get_target("TH_S11_RMUCBL1_AVG_PE13R", "CCHL/PCEJ00002/11KV/1") == ("PCE Testsheet", "N10")
    assert mapper.get_target("TH_S11_RMUCBL1_MAX_PE13R", "CCHL/PCEJ00002/11KV/1") == ("PCE Testsheet", "L10")
    assert mapper.get_target("TV_S11_CBL_PE13R", "CCHL/PCEJ00002/11KV/1") == ("PCE Testsheet", "T10")


def test_ppu_mapper_coordinate_resolution() -> None:
    mapper = PpuReadingMapper()

    # 1. 33kV switchgear meters -> 33kV (Main)
    # Panel 1 (row 15..20):
    # Offset 0: Breaker
    assert mapper.get_target("TH_S33_BR_AVG_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "O15")
    assert mapper.get_target("TH_S33_BR_MAX_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "M15")
    assert mapper.get_target("TH_S33_BR_REF_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "L15")
    assert mapper.get_target("TH_S33_BR_DEL_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "N15")
    assert mapper.get_target("US_S33_FP_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "T15")
    assert mapper.get_target("TEV_S33_FP_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "R15")
    assert mapper.get_target("TEV_S33_FP_PUL_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "S15")

    # Offset 1: Cable Compartment / Back Panel
    assert mapper.get_target("TH_S33_CBL_AVG_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "O16")
    assert mapper.get_target("TH_S33_CBL_MAX_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "M16")
    assert mapper.get_target("US_S33_CBL_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "T16")
    assert mapper.get_target("TEV_S33_BP_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "R16")
    assert mapper.get_target("TEV_S33_BP_PUL_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "S16")

    # Offset 2: Busbar Top
    assert mapper.get_target("TH_S33_BB_TREF_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "L17")
    assert mapper.get_target("TH_S33_BB_TMAX_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "M17")
    assert mapper.get_target("TH_S33_BB_TDEL_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "N17")
    assert mapper.get_target("TH_S33_BB_TAVG_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "O17")
    assert mapper.get_target("US_S33_TOP_P31VV", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "T17")

    # Gas Pressure
    assert mapper.get_target("SG_33_GIS_PRES_CB", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "V16")
    assert mapper.get_target("SG_33_GIS_PRES_MBB", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "V17")

    # Offset 3: PT Compartment
    assert mapper.get_target("TH_S33_PT_AVG", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "O18")
    assert mapper.get_target("US_S33_PT", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "T18")
    assert mapper.get_target("TEV_S33_PT", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "R18")

    # Offset 5: Link Box
    assert mapper.get_target("TH_S33_LB_AVG", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "O20")

    # Loading & Heater
    assert mapper.get_target("SG_33_LOADING", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "I15")
    assert mapper.get_target("HEATER_33KV_AMP", "CGBG/PPUAPMM/33KV/1") == ("33kV (Main)", "J15")

    # Panel 2 (row 21..26):
    assert mapper.get_target("TH_S33_BR_AVG_P31VV", "CGBG/PPUAPMM/33KV/2") == ("33kV (Main)", "O21")
    assert mapper.get_target("SG_33_LOADING", "CGBG/PPUAPMM/33KV/2") == ("33kV (Main)", "I21")

    # 2. 11kV switchgear meters -> 11kV (Main)
    assert mapper.get_target("SG_11_LOADING", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "I15")
    assert mapper.get_target("TH_S11_BR_AVG_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "O15")
    assert mapper.get_target("US_S11_BR_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "T15")
    assert mapper.get_target("TH_S11_CBL_AVG_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "O16")
    assert mapper.get_target("US_S11_CBL_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "T16")
    assert mapper.get_target("TEV_S11_BP_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "R16")
    assert mapper.get_target("TEV_S11_BP_PUL_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "S16")
    assert mapper.get_target("TH_S11_TBB_AVG_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "O17")
    assert mapper.get_target("US_S11_TOP_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "T17")
    assert mapper.get_target("TEV_S11_FP_P31VV", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "R15")
    assert mapper.get_target("HEATER_11KV_AMP", "CGBG/PPUAPMM/11KV/1") == ("11kV (Main)", "J15")
    assert mapper.get_target("BG_TEM_11_ROOM_P31VV", "CGBG/PPUAPMM") == ("11kV (Main)", "U6")
    assert mapper.get_target("BG_HUM_11_ROOM_P31VV", "CGBG/PPUAPMM") == ("11kV (Main)", "U7")
    assert mapper.get_target("BG_TEV_11_ROOM_P31VV", "CGBG/PPUAPMM") == ("11kV (Main)", "U8")

    # 3. Power Transformer meters -> TX
    # PTX1 (Part 1 row 14, Part 2 row 33):
    assert mapper.get_target("TH_PTXHV_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "U14")
    assert mapper.get_target("TH_PTXHV_MAX_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "S14")
    assert mapper.get_target("TH_PTXLV_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "Z14")
    assert mapper.get_target("TH_PTXLV_MAX_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "X14")
    assert mapper.get_target("US_PTX_HV_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "J14")
    assert mapper.get_target("US_PTX_LV_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "N14")

    assert mapper.get_target("TH_PTXMT_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "U33")
    assert mapper.get_target("TH_PTXMT_MAX_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "S33")
    assert mapper.get_target("TH_PTXFIN_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "AA33")
    assert mapper.get_target("TH_PTXFIN_MAX_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "Y33")
    assert mapper.get_target("BG_TEM_TX1_BAY_P31VV", "CGBG/PPUAPMM/TX/PTX1") == ("TX", "U6")

    # PTX2 (Part 1 row 17, Part 2 row 36):
    assert mapper.get_target("TH_PTXHV_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX2") == ("TX", "U17")
    assert mapper.get_target("TH_PTXMT_AVG_P31VV", "CGBG/PPUAPMM/TX/PTX2") == ("TX", "U36")

    # 4. Visual Inspection and checklist status meters -> None (handled semantically via visual findings, matching PCE)
    assert mapper.get_target("VI_NER_EARTH", "CGBG/PPUAPMM") is None
    assert mapper.get_target("VI_PTX_OILLEAK1", "CGBG/PPUAPMM/TX/PTX1") is None
    assert mapper.get_target("HEATER_33KV_STATUS", "CGBG/PPUAPMM/33KV/1") is None
    assert mapper.get_target("HEATER_11KV_STATUS", "CGBG/PPUAPMM/11KV/1") is None


def test_ppu_mapper_excludes_pce_distribution_meters_and_delegates() -> None:
    ppu_mapper = PpuReadingMapper()
    # 1. PpuReadingMapper directly returns None for PCE distribution meters
    assert ppu_mapper.get_target("TH_S11_RMUCBL1_AVG_PE13R", "CCHL/PCEJ00002/11KV/1") is None
    assert ppu_mapper.get_target("TH_S11_BR_AVG_PE13V", "CCHL/PCEJ00002/11KV/1") is None
    assert ppu_mapper.get_target("TV_S11_CBL_PE13R", "CCHL/PCEJ00002/11KV/1") is None
    assert ppu_mapper.get_target("TH_S11_BR_MAX_PE13O", "CCHL/PCEJ00002/11KV/1") is None

    # 2. TestsheetReadingMapper in 33kV mode delegates them to PceReadingMapper via cross-voltage fallback
    facade_33kv = TestsheetReadingMapper(voltage_type="33kV")
    assert facade_33kv.get_target("TH_S11_RMUCBL1_AVG_PE13R", "CCHL/PCEJ00002/11KV/1") == ("PCE Testsheet", "N10")
    assert facade_33kv.get_target("TV_S11_CBL_PE13R", "CCHL/PCEJ00002/11KV/1") == ("PCE Testsheet", "T10")


def test_ppu_bay_map_construction() -> None:
    wb = openpyxl.Workbook()
    ws_33 = wb.active
    ws_33.title = "33kV (Main)"

    # Panel 1: 3P5 SPARE
    ws_33["B15"] = 1
    ws_33["C15"] = "3P5"
    ws_33["D15"] = "SPARE"

    # Panel 2: 1P5 PMU GBID (1L5)
    ws_33["B21"] = 2
    ws_33["C21"] = "1P5"
    ws_33["D21"] = "PMU GBID (1L5)"

    # Panel 3: 1HO TRANS 1
    ws_33["B27"] = 3
    ws_33["C27"] = "1HO"
    ws_33["D27"] = "TRANSFORMER FEEDER"

    # 11kV sheet
    ws_11 = wb.create_sheet(title="11kV (Main)")
    # Panel 1: 1K5 FEEDER 1
    ws_11["B15"] = 1
    ws_11["C15"] = "1K5"
    ws_11["D15"] = "FEEDER 1"

    # Panel 2: 1WO BUS COUPLER
    ws_11["B21"] = 2
    ws_11["C21"] = "1WO"
    ws_11["D21"] = "BUS COUPLER"

    # Panel 3: 16K5 LOCAL TX
    ws_11["B27"] = 3
    ws_11["C27"] = "16K5"
    ws_11["D27"] = "LOCAL TRANSFORMER"

    # Panel 4 on 33kV: Feeder code in Col B (p_no) due to column alignment
    ws_33["B33"] = "4P5"
    ws_33["C33"] = None
    ws_33["D33"] = "FEEDER 4"

    # Panel 5 on 11kV: 5K5 (must NOT be hijacked by Panel 1's 1K5 trailing digit!)
    ws_11["B39"] = 5
    ws_11["C39"] = "5K5"
    ws_11["D39"] = "FEEDER 5"

    # Panel 10 on 11kV: Pure numeric code 10 should NOT claim token 1
    ws_11["B45"] = 10
    ws_11["C45"] = "10"
    ws_11["D45"] = "SPARE 10"

    bay_map = build_ppu_bay_map(wb)

    assert bay_map[("33KV", "3P5")] == ("33kV (Main)", 15)
    assert bay_map[("33KV", "3")] == ("33kV (Main)", 15)  # extracted from 3P5 via regex
    assert bay_map[("33KV", "1P5")] == ("33kV (Main)", 21)
    assert bay_map[("33KV", "1L5")] == ("33kV (Main)", 21)
    assert bay_map[("33KV", "1HO")] == ("33kV (Main)", 27)
    assert bay_map[("33KV", "4P5")] == ("33kV (Main)", 33)  # from Col B (p_no)
    assert bay_map[("33KV", "4")] == ("33kV (Main)", 33)

    assert bay_map[("11KV", "1K5")] == ("11kV (Main)", 15)
    assert bay_map[("11KV", "1")] == ("11kV (Main)", 15)  # extracted from 1K5 via regex
    assert bay_map[("11KV", "1WO")] == ("11kV (Main)", 21)
    assert bay_map[("11KV", "16K5")] == ("11kV (Main)", 27)
    assert bay_map[("11KV", "16")] == ("11kV (Main)", 27)  # extracted from 16K5 via regex
    assert bay_map[("11KV", "LTX")] == ("11kV (Main)", 27)
    assert bay_map[("11KV", "5")] == ("11kV (Main)", 39)
    assert bay_map[("11KV", "5K5")] == ("11kV (Main)", 39)
    assert bay_map[("11KV", "10")] == ("11kV (Main)", 45)


def test_normalizers_33kv() -> None:
    # 1. format_load_amp
    assert format_load_amp("119/600") == "119"
    assert format_load_amp("0/600") == "0"
    assert format_load_amp("150A") == "150"
    assert format_load_amp("-") == "-"

    # 2. extract_gas_pressure (resilient against digit collision in equipment/labels)
    assert extract_gas_pressure("B: 1.41 Bar") == 1.41
    assert extract_gas_pressure("S: 6.0 Bar") == 6.0
    assert extract_gas_pressure("SF6: 1.41 Bar") == 1.41
    assert extract_gas_pressure("CB1: 1.41 Bar") == 1.41
    assert extract_gas_pressure("33kV: 1.41 Bar") == 1.41
    assert extract_gas_pressure("11kV: 1.41 Bar") == 1.41
    assert extract_gas_pressure("GIS1: 1.41 Bar") == 1.41
    assert extract_gas_pressure("SF6: 1.41") == 1.41
    assert extract_gas_pressure("CB2: 1.41") == 1.41
    assert extract_gas_pressure("CB-1: 1.41") == 1.41
    assert extract_gas_pressure("SF-6: 1.41") == 1.41
    assert extract_gas_pressure("CB#1: 1.41") == 1.41
    assert extract_gas_pressure("BB1: 1.41") == 1.41
    assert extract_gas_pressure("BB-1: 1.41") == 1.41
    assert extract_gas_pressure("BUS 1: 1.41") == 1.41
    assert extract_gas_pressure("MBB1: 1.41") == 1.41
    assert extract_gas_pressure("MBB-1: 1.41") == 1.41
    assert extract_gas_pressure("RBB1: 1.41") == 1.41
    assert extract_gas_pressure("B: 1.41 kg/cm²") == 1.41
    assert extract_gas_pressure("0.5 MPa") == 0.5
    assert extract_gas_pressure("B: 1.41 psi") == 1.41
    assert extract_gas_pressure(1.45) == 1.45
    assert extract_gas_pressure("NONE") is None
    assert extract_gas_pressure(None) is None

    # 3. extract_heater_amp
    assert extract_heater_amp("ON") is None
    assert extract_heater_amp("OFF") is None
    assert extract_heater_amp("0.85A") == 0.85
    assert extract_heater_amp("0.85 A") == 0.85
    assert extract_heater_amp("ON: 0.85A") == 0.85
    assert extract_heater_amp("ON:0.85A") == 0.85
    assert extract_heater_amp(0.5) == 0.5
    assert extract_heater_amp("-") is None


def test_match_33kv_vi_defect_null_safety() -> None:
    from src.testsheet.models import VisualInspectionFinding, SubstationTestsheetPackage, TestsheetData
    wf = PopulateDataMsmsWorkflow()
    pkg = SubstationTestsheetPackage(
        testsheet_path=Path("test.xlsx"),
        unsorted_raw_data_dir=Path("."),
        station="TEST",
        month="OCTOBER",
        date_str="2026-10-07",
        substation_number=1,
        data=TestsheetData(
            substation_number=1,
            substation_name_erms="PPU TEST",
            visual_findings=(
                VisualInspectionFinding(category=None, item_no=None, description="Fence condition", status=None, remarks=None),
                VisualInspectionFinding(category="CIVIL", item_no="1.2", description="Gate", status="YES", remarks=None),
                VisualInspectionFinding(category="CIVIL", item_no="1.3", description="Wall", status=None, remarks="Crack on wall"),
            ),
        ),
    )
    is_defect, comment = wf.transformer._match_33kv_vi_defect("VI_CIVIL_FENCE", "Fence condition", "CGBG/TEST", pkg, {})
    assert is_defect is False

    is_defect, comment = wf.transformer._match_33kv_vi_defect("VI_CIVIL_GATE", "Gate", "CGBG/TEST", pkg, {})
    assert is_defect is False

    is_defect, comment = wf.transformer._match_33kv_vi_defect("VI_CIVIL_WALL", "Wall", "CGBG/TEST", pkg, {})
    assert is_defect is True
    assert "Crack on wall" in comment


# ==============================================================================
# 4. Reading Population End-to-End Workflow with 33kV Workbooks
# ==============================================================================

def test_populate_data_msms_33kv_end_to_end(mock_33kv_env: ProjectEnvironment, tmp_path: Path) -> None:
    # 1. Create synthetic 33kV testsheet workbook using ppu_workbook_factory
    date_str = "07-10-2026"
    testsheet_dir = tmp_path / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / date_str
    testsheet_dir.mkdir(parents=True)
    (testsheet_dir / "UNSORTED RAW DATA").mkdir()

    wb_path = testsheet_dir / "01. PPU APMM 33-11kV (VI).xlsx"
    create_ppu_testsheet_workbook(
        path=wb_path,
        substation_number=1,
        substation_name="PPU APMM 33-11kV",
        fl_number="CGBG/PPUAPMM",
        date_val="2026-10-07",
        ambient_temp="30.5",
        humidity="65%",
        background_tev="5",
        has_33kv=True,
        has_11kv=True,
        panels_33kv=2,
        panels_11kv=2,
        include_tx_sheet=True,
        transformer_ratings=["10MVA", "10MVA"],
        visual_checklist_items=[
            {"category": "4. NER", "item_no": "4.1", "description": "Improper NER earth installation", "status": "NO", "remarks": "Loose NER earth link"},
            {"category": "1. CIVIL", "item_no": "1.1", "description": "Fence & gate condition", "status": "YES", "remarks": ""},
            {"category": "2. EARTHING", "item_no": "2.1", "description": "Substation earthing bar", "status": "YES", "remarks": "Re-tightened earthing bar"},
            {"category": "3. SWITCHGEAR", "item_no": "3.5", "description": "Switchgear heater failure", "status": "NO", "remarks": "Heater broken at 11kV bay"},
        ],
    )

    # Inject specific measurement values into workbook for testing
    wb = openpyxl.load_workbook(wb_path)
    ws_33 = wb["33kV (Main)"]
    ws_33["I15"] = "119/600"  # SG_33_LOADING (ratio format)
    ws_33["J15"] = "ON"       # HEATER_33KV_AMP (status text, must NOT be written)
    ws_33["L17"] = 27.5      # TH_S33_BB_TREF_P31VV
    ws_33["O15"] = 28.5     # TH_S33_BR_AVG_P31VV
    ws_33["M15"] = 29.0     # TH_S33_BR_MAX_P31VV
    ws_33["T15"] = 0        # US_S33_FP_P31VV
    ws_33["R15"] = 5        # TEV_S33_FP_P31VV
    ws_33["S15"] = 0        # TEV_S33_FP_PUL_P31VV
    ws_33["R16"] = "7 dB"   # TEV_S33_BP_P31VV (Cable Compartment TEV)
    ws_33["S16"] = 0        # TEV_S33_BP_PUL_P31VV (Cable Compartment Pulse)
    ws_33["V16"] = "B: 1.41 Bar"  # Gas Pressure

    ws_11 = wb["11kV (Main)"]
    ws_11["I15"] = "220"    # SG_11_LOADING
    ws_11["J15"] = "0.85A"  # HEATER_11KV_AMP (numeric ampere)
    ws_11["O15"] = 31.2     # TH_S11_BR_AVG_P31VV
    ws_11["T15"] = 0        # US_S11_BR_P31VV
    ws_11["R16"] = "10dB"   # TEV_S11_BP_P31VV (Cable Compartment TEV)
    ws_11["S16"] = 15.0     # TEV_S11_BP_PUL_P31VV (Cable Compartment Pulse)

    ws_tx = wb["TX"]
    ws_tx["U14"] = 36.8     # TH_PTXHV_AVG_P31VV (TX 1)
    ws_tx["U33"] = 43.5     # TH_PTXMT_AVG_P31VV (TX 1)
    ws_tx["J14"] = 0        # US_PTX_HV_P31VV (TX 1)

    wb.save(wb_path)
    wb.close()

    # 2. Setup TOTAL PE.xlsx
    total_pe_path = mock_33kv_env.storage.get_total_pe_path()
    mock_33kv_env.storage.ensure_directory(total_pe_path.parent)
    wb_pe = openpyxl.Workbook()
    ws_pe = wb_pe.active
    ws_pe.title = "DataCycle1"
    ws_pe.append(["PE NO", "FL NUMBER", "SUBSTATION NAME", "DATE", "TYPE", "WO", "SCOPE"])
    ws_pe.append([1, "CGBG/PPUAPMM", "PPU APMM 33-11kV", date_str, "PPU", "W1616898", "CBM"])
    wb_pe.save(total_pe_path)
    wb_pe.close()

    # 3. Setup MSMS CSV in TO BE FILLED/
    filled_dir = mock_33kv_env.storage.get_msms_to_be_filled_dir()
    completed_dir = mock_33kv_env.storage.get_msms_completed_dir()
    mock_33kv_env.storage.ensure_directory(filled_dir)
    mock_33kv_env.storage.ensure_directory(completed_dir)

    csv_path = filled_dir / "07-10-2026_001.csv"
    csv_rows = [
        "ACTFINISH,ACTSTART,JPNUM,WONUM,TNBLOCATION,METERNAME,TNBNEWREADING,TNBNEWREADINGDATE,TNBCOMMENTS,METERTYPE\n",
        # Characteristic defect meter (Defect present in Visual Inspec. -> update to YES)
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,VI_NER_EARTH,NO,,,CHARACTERISTIC\n",
        # Characteristic sound meter (No defect -> remain NO, untouched)
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,VI_BLDG_ROOF,NO,,,CHARACTERISTIC\n",
        # Characteristic sound meter with informational remarks (Sound -> remain NO, untouched)
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,VI_SGERT_LINK,NO,,,CHARACTERISTIC\n",
        # Heater status meters (Sound -> remain NO untouched; Defect -> update to YES)
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,HEATER_33KV_STATUS,NO,,,CHARACTERISTIC\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,HEATER_11KV_STATUS,NO,,,CHARACTERISTIC\n",
        # Heater ampere meters (ON text -> unpopulated; 0.85A -> 0.85)
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,HEATER_33KV_AMP,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,HEATER_11KV_AMP,,,,GAUGE\n",
        # 33kV switchgear gauge meters
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,SG_33_LOADING,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,SG_33_GIS_PRES_CB,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TH_S33_BB_TREF_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TH_S33_BR_AVG_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TH_S33_BR_MAX_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,US_S33_FP_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_FP_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_BP_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/33KV/1,TEV_S33_BP_PUL_P31VV,,,,GAUGE\n",
        # 11kV switchgear gauge meters
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,SG_11_LOADING,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,TH_S11_BR_AVG_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,US_S11_BR_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,TEV_S11_BP_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/11KV/1,TEV_S11_BP_PUL_P31VV,,,,GAUGE\n",
        # Power Transformer gauge meters
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/TX/PTX1,TH_PTXHV_AVG_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/TX/PTX1,TH_PTXMT_AVG_P31VV,,,,GAUGE\n",
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM/TX/PTX1,US_PTX_HV_P31VV,,,,GAUGE\n",
        # Background temperature meter
        ",2026-10-07T12:00:00+08:00,T1SUB33,W1616898,CGBG/PPUAPMM,BG_TEM_11_ROOM_P31VV,,,,GAUGE\n",
    ]
    csv_path.write_text("".join(csv_rows), encoding="utf-8")

    # 4. Execute workflow with move_to_completed=True
    workflow = PopulateDataMsmsWorkflow()
    request = PopulateDataMsmsRequest(
        mode=PopulateMode.AUTO,
        overwrite=True,
        move_to_completed=True,
    )
    result = workflow.execute(mock_33kv_env, request)

    assert result.csv_files_processed == 1
    assert result.rows_populated > 0
    assert len(result.errors) == 0

    # Verify completed CSV moved to PYTHON/MSMS/COMPLETED/
    completed_csv = completed_dir / "07-10-2026_001.csv"
    assert completed_csv.exists()
    assert not csv_path.exists()

    # 5. Read back populated CSV and verify readings
    with open(completed_csv, mode="r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        records = {r["METERNAME"]: r for r in reader}

    # VI defect meter: updated to YES with comments
    assert records["VI_NER_EARTH"]["TNBNEWREADING"] == "YES"
    assert "Loose NER earth link" in records["VI_NER_EARTH"]["TNBCOMMENTS"]

    # Sound VI meter: left untouched as NO
    assert records["VI_BLDG_ROOF"]["TNBNEWREADING"] == "NO"

    # Sound VI meter with informational remarks ("Re-tightened earthing bar"): left untouched as NO
    assert records["VI_SGERT_LINK"]["TNBNEWREADING"] == "NO"

    # Heater status meters:
    # HEATER_33KV_STATUS (sound): left untouched as NO (protected from overwriting with D3 header!)
    assert records["HEATER_33KV_STATUS"]["TNBNEWREADING"] == "NO"
    # HEATER_11KV_STATUS (defect): updated to YES with comments
    assert records["HEATER_11KV_STATUS"]["TNBNEWREADING"] == "YES"
    assert "Heater broken" in records["HEATER_11KV_STATUS"]["TNBCOMMENTS"]

    # Heater ampere meters:
    # HEATER_33KV_AMP (status text "ON"): left blank / unpopulated
    assert records["HEATER_33KV_AMP"]["TNBNEWREADING"] == ""
    # HEATER_11KV_AMP (numeric "0.85A"): preserved 2-decimal precision as "0.85"
    assert records["HEATER_11KV_AMP"]["TNBNEWREADING"] == "0.85"

    # Gas pressure meter: preserved 2-decimal precision as "1.41"
    assert records["SG_33_GIS_PRES_CB"]["TNBNEWREADING"] == "1.41"

    # Busbar thermal meter (_TREF):
    assert records["TH_S33_BB_TREF_P31VV"]["TNBNEWREADING"] == "27.5"

    # 33kV switchgear meters populated (ratio 119/600 -> 119)
    assert records["SG_33_LOADING"]["TNBNEWREADING"] == "119"
    assert records["TH_S33_BR_AVG_P31VV"]["TNBNEWREADING"] == "28.5"
    assert records["TH_S33_BR_MAX_P31VV"]["TNBNEWREADING"] in ("29", "29.0")
    assert records["US_S33_FP_P31VV"]["TNBNEWREADING"] == "0"
    assert records["TEV_S33_FP_P31VV"]["TNBNEWREADING"] == "5"
    assert records["TEV_S33_BP_P31VV"]["TNBNEWREADING"] == "7"
    assert records["TEV_S33_BP_PUL_P31VV"]["TNBNEWREADING"] == "0"

    # 11kV switchgear meters populated
    assert records["SG_11_LOADING"]["TNBNEWREADING"] == "220"
    assert records["TH_S11_BR_AVG_P31VV"]["TNBNEWREADING"] == "31.2"
    assert records["US_S11_BR_P31VV"]["TNBNEWREADING"] == "0"
    assert records["TEV_S11_BP_P31VV"]["TNBNEWREADING"] == "10"
    assert records["TEV_S11_BP_PUL_P31VV"]["TNBNEWREADING"] == "15"

    # Power Transformer meters populated
    assert records["TH_PTXHV_AVG_P31VV"]["TNBNEWREADING"] == "36.8"
    assert records["TH_PTXMT_AVG_P31VV"]["TNBNEWREADING"] == "43.5"
    assert records["US_PTX_HV_P31VV"]["TNBNEWREADING"] == "0"

    # Background ambient temperature populated
    assert records["BG_TEM_11_ROOM_P31VV"]["TNBNEWREADING"] == "30.5"
