"""End-to-end integration test suite and verification harness for 33kV MSMS workflows 8 through 12.

Ticket #108:
- Exercises complete 33kV pipeline:
    1. Workflow 8 (ConsolidateMsmsAction / ConsolidateMsmsWorkflow): Maximo HTML .xls -> DATA MSMS.xlsx
    2. Workflow 9 (EnrichMsmsAction / EnrichMsmsWorkflow): TOTAL PE.xlsx metadata -> DATA MSMS.xlsx Cols D-G
    3. Workflow 10 (PropagateWoAction / PropagateWoWorkflow): DATA MSMS.xlsx WOs -> TOTAL PE.xlsx Col F
    4. Workflow 11 (IngestMsmsCsvAction / IngestMsmsCsvWorkflow): Raw T1SUB33 Maximo CSVs -> TO BE FILLED/
    5. Workflow 12 (PopulateDataMsmsAction / PopulateDataMsmsWorkflow): Testsheet packages -> Populated CSVs in COMPLETED/
- Validates that no work order mappings are lost, pe_records are fully populated in TOTAL PE,
  and readings/defects are correctly mapped from testsheet packages into ingested CSVs.
- Verifies clean execution against the Thunder Cycle 4 benchmark dataset.
"""

from __future__ import annotations

import csv
from pathlib import Path
import shutil
from typing import Any
from unittest.mock import patch
import openpyxl
import pytest

from src.project.environment import ProjectEnvironment
from src.project.models import ProjectMetadata
from src.project.repository import JsonFileProjectRepository
from src.project.storage import LocalWorkspaceStorage
from src.project_workflow_actions import (
    ConsolidateMsmsAction,
    EnrichMsmsAction,
    IngestMsmsCsvAction,
    PopulateDataMsmsAction,
    PropagateWoAction,
)
from src.workflows.consolidate_msms import ConsolidateMsmsWorkflow
from src.workflows.enrich_msms import EnrichMsmsWorkflow
from src.workflows.ingest_msms_csv import IngestMsmsCsvWorkflow
from src.workflows.models import (
    ConsolidateMsmsRequest,
    EnrichMsmsRequest,
    IngestMsmsCsvRequest,
    PopulateDataMsmsRequest,
    PopulateMode,
    PropagateWoRequest,
)
from src.workflows.populate_data_msms import PopulateDataMsmsWorkflow
from src.workflows.propagate_wo import PropagateWoWorkflow
from src.workflows.service import WorkflowService
from tests.fixtures.ppu_workbook_factory import create_ppu_testsheet_workbook


# ==============================================================================
# Helper Factories for 33kV MSMS Test Data
# ==============================================================================

def _create_html_xls_file(path: Path, rows: list[tuple[str, str, str, str]]) -> None:
    """Create a simulated Maximo HTML-formatted .xls file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    html_lines = [
        "<html><body><table border='1'>",
        "<tr><th>Wonum</th><th>Status</th><th>Location</th><th>Description</th></tr>",
    ]
    for wonum, status, location, desc in rows:
        html_lines.append(f"<tr><td>{wonum}</td><td>{status}</td><td>{location}</td><td>{desc}</td></tr>")
    html_lines.append("</table></body></html>")
    path.write_text("\n".join(html_lines), encoding="utf-8")


def _init_data_msms_file(path: Path) -> None:
    """Initialize an empty DATA MSMS.xlsx with standard headers."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DATA MSMS"
    ws.append([
        "Work Order",
        "Location",
        "Description",
        "Substation Name ERMS",
        "FL ERMS",
        "Cycle Date",
        "Substation Number",
    ])
    wb.save(path)
    wb.close()


def _init_total_pe_file(path: Path, rows: list[list[Any]]) -> None:
    """Initialize a TOTAL PE.xlsx with DataCycle1 sheet."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DataCycle1"
    ws.append(["PE NO", "FL NUMBER", "SUBSTATION NAME", "DATE", "TYPE", "WO", "SCOPE"])
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _create_raw_t1sub33_csv(path: Path, rows: list[dict[str, str]]) -> None:
    """Create a raw Maximo CSV file with T1SUB33 headers."""
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = [
        "ACTFINISH",
        "ACTSTART",
        "JPNUM",
        "WONUM",
        "TNBLOCATION",
        "METERNAME",
        "TNBNEWREADING",
        "TNBNEWREADINGDATE",
        "TNBCOMMENTS",
        "DESCRIPTION",
        "METERTYPE",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for r in rows:
            writer.writerow({
                "ACTFINISH": r.get("ACTFINISH", ""),
                "ACTSTART": r.get("ACTSTART", "2026-10-07T12:00:00+08:00"),
                "JPNUM": r.get("JPNUM", "T1SUB33"),
                "WONUM": r.get("WONUM", ""),
                "TNBLOCATION": r.get("TNBLOCATION", ""),
                "METERNAME": r.get("METERNAME", ""),
                "TNBNEWREADING": r.get("TNBNEWREADING", ""),
                "TNBNEWREADINGDATE": r.get("TNBNEWREADINGDATE", ""),
                "TNBCOMMENTS": r.get("TNBCOMMENTS", ""),
                "DESCRIPTION": r.get("DESCRIPTION", ""),
                "METERTYPE": r.get("METERTYPE", "GAUGE"),
            })


# ==============================================================================
# Fixtures
# ==============================================================================

@pytest.fixture
def synthetic_33kv_env(tmp_path: Path) -> ProjectEnvironment:
    """Fixture providing an isolated 33kV ProjectEnvironment with full folder tree."""
    meta = ProjectMetadata(
        key="thunder_cycle_4_synthetic",
        name="Synthetic 33kV Thunder Project",
        po_number="PO42394550",
        state="Pahang",
        voltage_type="33kV",
        year="2026",
        cycle="4",
        technologies=("IR", "DG", "US", "TEV", "VI"),
        base_path=str(tmp_path),
    )
    storage = LocalWorkspaceStorage(tmp_path)
    # Ensure standard directories exist
    storage.ensure_directory(storage.get_python_dir())
    storage.ensure_directory(storage.get_msms_dir())
    storage.ensure_directory(storage.get_msms_raw_data_dir())
    storage.ensure_directory(storage.get_msms_to_be_filled_dir())
    storage.ensure_directory(storage.get_msms_completed_dir())
    storage.ensure_directory(storage.get_testsheet_dir())
    return ProjectEnvironment(metadata=meta, storage=storage)


# ==============================================================================
# Test Cases
# ==============================================================================

def test_full_33kv_msms_pipeline_end_to_end_synthetic(synthetic_33kv_env: ProjectEnvironment) -> None:
    """Exercise complete 33kV MSMS pipeline sequentially from Workflow 8 through 12.

    Stages tested:
      1. Consolidate MSMS: Read HTML .xls files -> populate DATA MSMS.xlsx -> move to COMPLETED/.
      2. Enrich MSMS: Read TOTAL PE.xlsx -> synchronize SAP FLs in Col E & metadata in Cols D, F, G.
      3. Propagate WO: Read DATA MSMS.xlsx -> propagate WOs into TOTAL PE.xlsx Col F.
      4. Ingest MSMS CSV: Read raw CSV from RAW DATA/ -> validate T1SUB33 schema -> rename to DD-MM-YYYY_NNN.csv in TO BE FILLED/.
      5. Populate Data MSMS: Read testsheet packages & CSVs in TO BE FILLED/ -> populate TNBNEWREADING -> move to COMPLETED/.
    """
    env = synthetic_33kv_env
    storage = env.storage
    service = WorkflowService()
    date_str = "07-10-2026"

    # --------------------------------------------------------------------------
    # 0. Setup Ground-Truth Test Assets
    # --------------------------------------------------------------------------
    # 3 Substations:
    # 1. PPU APMM 33/11kV (WO: W1616898, SAP FL: CGBG/PPU/G00001, Mnemonic: CGBG/PPUAPMM)
    # 2. SSU KIIP (WO: W1616956, SAP FL: CKTN/SSU/G00006, Mnemonic: CKTN/SSUKIIP)
    # 3. SSU LYNAS (WO: W1616922, SAP FL: CGBG/SSU/G00003, Mnemonic: CGBG/SSULNAS)

    # A. Initialize DATA MSMS.xlsx with header only
    data_msms_path = storage.get_data_msms_path()
    _init_data_msms_file(data_msms_path)

    # B. Initialize TOTAL PE.xlsx (WO column is initially empty / None)
    total_pe_path = storage.get_total_pe_path()
    _init_total_pe_file(
        total_pe_path,
        [
            [1, "CGBG/PPU/G00001", "PPU APMM 33/11kV", date_str, "PPU", None, "CBM"],
            [2, "CKTN/SSU/G00006", "SSU KIIP (SSU CAMEL POWER) 33kV", date_str, "SSU", None, "CBM"],
            [3, "CGBG/SSU/G00003", "SSU LYNAS 33kV", date_str, "SSU", None, "CBM"],
        ],
    )

    # C. Setup 2 scattered Maximo HTML .xls files in PYTHON/MSMS/
    msms_dir = storage.get_msms_dir()
    xls1_path = msms_dir / "maximo_export_01.xls"
    xls2_path = msms_dir / "maximo_export_02.xls"

    _create_html_xls_file(
        xls1_path,
        [
            ("W1616898", "INPRG", "CGBG/PPUAPMM", "PPU APMM 33/11KV"),
            ("W1616956", "INPRG", "CKTN/SSUKIIP", "SSU KIIP 33kV"),
        ],
    )
    _create_html_xls_file(
        xls2_path,
        [
            ("W1616922", "INPRG", "CGBG/SSULNAS", "SSU LYNAS 33KV"),
            # Duplicate entry to test deduplication
            ("W1616898", "INPRG", "CGBG/PPUAPMM", "PPU APMM 33/11KV"),
        ],
    )

    # D. Setup Raw T1SUB33 Maximo CSV in PYTHON/MSMS/RAW DATA/
    raw_csv_path = storage.get_msms_raw_data_dir() / "33KV_T1SUB33_07-10-2026_RAW.csv"
    _create_raw_t1sub33_csv(
        raw_csv_path,
        [
            # PPU APMM - Characteristic defect meter (Loose earth -> YES + comments)
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM", "METERNAME": "VI_NER_EARTH", "TNBNEWREADING": "NO", "METERTYPE": "CHARACTERISTIC"},
            # PPU APMM - Characteristic sound meter (No defect -> remain NO)
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM", "METERNAME": "VI_BLDG_ROOF", "TNBNEWREADING": "NO", "METERTYPE": "CHARACTERISTIC"},
            # PPU APMM - 33kV Switchgear Panel 1 meters
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "SG_33_LOADING", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "SG_33_GIS_PRES_CB", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "TH_S33_BR_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "TH_S33_BB_TREF_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "US_S33_FP_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "TEV_S33_FP_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "HEATER_33KV_AMP", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "HEATER_33KV_STATUS", "TNBNEWREADING": "NO", "METERTYPE": "CHARACTERISTIC"},
            # PPU APMM - 11kV Switchgear Panel 1 meters
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/11KV/1", "METERNAME": "SG_11_LOADING", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/11KV/1", "METERNAME": "TH_S11_BR_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/11KV/1", "METERNAME": "HEATER_11KV_AMP", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/11KV/1", "METERNAME": "HEATER_11KV_STATUS", "TNBNEWREADING": "NO", "METERTYPE": "CHARACTERISTIC"},
            # PPU APMM - Transformer 1 meters
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/TX/PTX1", "METERNAME": "TH_PTXHV_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/TX/PTX1", "METERNAME": "TH_PTXMT_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            # PPU APMM - Ambient temperature
            {"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM", "METERNAME": "BG_TEM_11_ROOM_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            # SSU KIIP - 33kV meters
            {"WONUM": "W1616956", "TNBLOCATION": "CKTN/SSUKIIP/33KV/1", "METERNAME": "SG_33_LOADING", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616956", "TNBLOCATION": "CKTN/SSUKIIP/33KV/1", "METERNAME": "TH_S33_BR_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            # SSU LYNAS - 33kV meters
            {"WONUM": "W1616922", "TNBLOCATION": "CGBG/SSULNAS/33KV/1", "METERNAME": "SG_33_LOADING", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
            {"WONUM": "W1616922", "TNBLOCATION": "CGBG/SSULNAS/33KV/1", "METERNAME": "TH_S33_BR_AVG_P31VV", "TNBNEWREADING": "", "METERTYPE": "GAUGE"},
        ],
    )

    # E. Setup Testsheet Packages in TESTSHEET/
    testsheet_date_dir = storage.get_testsheet_dir() / "GEBENG" / "01. OCTOBER" / date_str
    testsheet_date_dir.mkdir(parents=True)
    (testsheet_date_dir / "UNSORTED RAW DATA").mkdir()

    # Package 1: PPU APMM 33-11kV
    pkg1_path = testsheet_date_dir / "01. PPU APMM 33-11kV (VI).xlsx"
    create_ppu_testsheet_workbook(
        path=pkg1_path,
        substation_number=1,
        substation_name="PPU APMM 33-11kV",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
        ambient_temp="30.5",
        humidity="65%",
        background_tev="5",
        has_33kv=True,
        has_11kv=True,
        panels_33kv=2,
        panels_11kv=2,
        include_tx_sheet=True,
        transformer_ratings=["10MVA", "10MVA"],
        visual_checklist_items=[
            {"category": "4. NER", "item_no": "4.1", "description": "Improper NER earth installation", "status": "NO", "remarks": "Loose NER earth link"},
            {"category": "1. CIVIL", "item_no": "1.1", "description": "Fence & gate condition", "status": "YES", "remarks": ""},
            {"category": "3. SWITCHGEAR", "item_no": "3.5", "description": "Switchgear heater failure", "status": "NO", "remarks": "Heater broken at 11kV bay"},
        ],
    )
    # Populate specific values into Package 1
    wb1 = openpyxl.load_workbook(pkg1_path)
    ws_33 = wb1["33kV (Main)"]
    ws_33["I15"] = "119/600"       # SG_33_LOADING
    ws_33["J15"] = "ON"            # HEATER_33KV_AMP (status string -> leave blank)
    ws_33["L17"] = 27.5           # TH_S33_BB_TREF_P31VV
    ws_33["O15"] = 28.5           # TH_S33_BR_AVG_P31VV
    ws_33["T15"] = 0             # US_S33_FP_P31VV
    ws_33["R15"] = 5             # TEV_S33_FP_P31VV
    ws_33["V16"] = "B: 1.41 Bar"  # SG_33_GIS_PRES_CB

    ws_11 = wb1["11kV (Main)"]
    ws_11["I15"] = "220"          # SG_11_LOADING
    ws_11["J15"] = "0.85A"        # HEATER_11KV_AMP (numeric -> 0.85)
    ws_11["O15"] = 31.2           # TH_S11_BR_AVG_P31VV

    ws_tx = wb1["TX"]
    ws_tx["U14"] = 36.8           # TH_PTXHV_AVG_P31VV
    ws_tx["U33"] = 43.5           # TH_PTXMT_AVG_P31VV

    wb1.save(pkg1_path)
    wb1.close()

    # Package 2: SSU KIIP (SSU CAMEL POWER) 33kV
    pkg2_path = testsheet_date_dir / "02. SSU KIIP (SSU CAMEL POWER) 33kV (VI).xlsx"
    create_ppu_testsheet_workbook(
        path=pkg2_path,
        substation_number=2,
        substation_name="SSU KIIP (SSU CAMEL POWER) 33kV",
        fl_number="CKTN/SSU/G00006",
        date_val="2026-10-07",
        ambient_temp="31.0",
        has_33kv=True,
        has_11kv=False,
        panels_33kv=1,
        include_tx_sheet=False,
    )
    wb2 = openpyxl.load_workbook(pkg2_path)
    wb2["33kV (Main)"]["I15"] = "85"
    wb2["33kV (Main)"]["O15"] = 29.8
    wb2.save(pkg2_path)
    wb2.close()

    # Package 3: SSU LYNAS 33kV
    pkg3_path = testsheet_date_dir / "03. SSU LYNAS 33kV (VI).xlsx"
    create_ppu_testsheet_workbook(
        path=pkg3_path,
        substation_number=3,
        substation_name="SSU LYNAS 33kV",
        fl_number="CGBG/SSU/G00003",
        date_val="2026-10-07",
        ambient_temp="29.5",
        has_33kv=True,
        has_11kv=False,
        panels_33kv=1,
        include_tx_sheet=False,
    )
    wb3 = openpyxl.load_workbook(pkg3_path)
    wb3["33kV (Main)"]["I15"] = "140"
    wb3["33kV (Main)"]["O15"] = 32.1
    wb3.save(pkg3_path)
    wb3.close()

    # ==========================================================================
    # WORKFLOW 8: Consolidate MSMS
    # ==========================================================================
    wf8_result = service.run_consolidate_msms(env, ConsolidateMsmsRequest())
    assert wf8_result.files_processed == 2
    assert wf8_result.rows_appended == 3
    assert wf8_result.duplicates_skipped == 1
    assert len(wf8_result.errors) == 0

    # Verify .xls files moved to COMPLETED/
    completed_xls_dir = storage.get_msms_completed_dir()
    assert (completed_xls_dir / "maximo_export_01.xls").exists()
    assert (completed_xls_dir / "maximo_export_02.xls").exists()
    assert not xls1_path.exists()
    assert not xls2_path.exists()

    # Verify DATA MSMS.xlsx has 3 appended records
    wb_msms = openpyxl.load_workbook(data_msms_path)
    ws_msms = wb_msms.active
    assert ws_msms.max_row == 4
    wos_in_msms = [ws_msms.cell(r, 1).value for r in range(2, 5)]
    assert set(wos_in_msms) == {"W1616898", "W1616956", "W1616922"}
    wb_msms.close()

    # ==========================================================================
    # WORKFLOW 9: Enrich MSMS
    # ==========================================================================
    wf9_result = service.run_enrich_msms(env, EnrichMsmsRequest())
    assert wf9_result.matched_count == 3
    assert wf9_result.unmatched_count == 0
    assert wf9_result.updated_cells_count == 12  # 3 rows * 4 metadata columns (D, E, F, G)

    # Verify DATA MSMS.xlsx has synchronized SAP FLs in Col E and metadata in D, F, G
    wb_msms = openpyxl.load_workbook(data_msms_path)
    ws_msms = wb_msms.active
    msms_data = {
        ws_msms.cell(r, 1).value: {
            "name": ws_msms.cell(r, 4).value,
            "fl": ws_msms.cell(r, 5).value,
            "date": ws_msms.cell(r, 6).value,
            "sub_num": ws_msms.cell(r, 7).value,
        }
        for r in range(2, 5)
    }
    wb_msms.close()

    assert msms_data["W1616898"]["fl"] == "CGBG/PPU/G00001"
    assert msms_data["W1616898"]["name"] == "PPU APMM 33/11kV"
    assert msms_data["W1616898"]["sub_num"] == 1

    assert msms_data["W1616956"]["fl"] == "CKTN/SSU/G00006"
    assert msms_data["W1616956"]["sub_num"] == 2

    assert msms_data["W1616922"]["fl"] == "CGBG/SSU/G00003"
    assert msms_data["W1616922"]["sub_num"] == 3

    # ==========================================================================
    # WORKFLOW 10: Propagate Work Orders
    # ==========================================================================
    wf10_result = service.run_propagate_wo(env, PropagateWoRequest())
    assert wf10_result.matched_count == 3
    assert wf10_result.updated_count == 3
    assert wf10_result.unmatched_count == 0

    # Verify TOTAL PE.xlsx Column F now contains valid Work Orders
    wb_pe = openpyxl.load_workbook(total_pe_path)
    ws_pe = wb_pe["DataCycle1"]
    pe_wos = {ws_pe.cell(r, 2).value: ws_pe.cell(r, 6).value for r in range(2, 5)}
    wb_pe.close()

    assert pe_wos["CGBG/PPU/G00001"] == "W1616898"
    assert pe_wos["CKTN/SSU/G00006"] == "W1616956"
    assert pe_wos["CGBG/SSU/G00003"] == "W1616922"

    # ==========================================================================
    # WORKFLOW 11: Ingest MSMS CSV
    # ==========================================================================
    wf11_result = service.run_ingest_msms_csv(env, IngestMsmsCsvRequest())
    assert wf11_result.files_ingested == 1
    assert wf11_result.duplicates_skipped == 0
    assert len(wf11_result.errors) == 0

    # Ingested CSV should exist in TO BE FILLED/ as canonical 07-10-2026_001.csv
    to_be_filled_csv = storage.get_msms_to_be_filled_dir() / "07-10-2026_001.csv"
    assert to_be_filled_csv.exists()
    assert not raw_csv_path.exists()

    # ==========================================================================
    # WORKFLOW 12: Populate Data MSMS
    # ==========================================================================
    wf12_result = service.run_populate_data_msms(
        env,
        PopulateDataMsmsRequest(
            mode=PopulateMode.AUTO,
            overwrite=True,
            move_to_completed=True,
        ),
    )
    assert wf12_result.csv_files_processed == 1
    assert wf12_result.rows_populated > 0
    assert len(wf12_result.errors) == 0

    # Verify CSV moved to PYTHON/MSMS/COMPLETED/
    completed_csv = storage.get_msms_completed_dir() / "07-10-2026_001.csv"
    assert completed_csv.exists()
    assert not to_be_filled_csv.exists()

    # ==========================================================================
    # Comprehensive Verification of Populated CSV Readings
    # ==========================================================================
    with open(completed_csv, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        records = {(r["WONUM"], r["METERNAME"]): r for r in reader}

    # 1. PPU APMM Visual Defect Findings
    assert records[("W1616898", "VI_NER_EARTH")]["TNBNEWREADING"] == "YES"
    assert "Loose NER earth link" in records[("W1616898", "VI_NER_EARTH")]["TNBCOMMENTS"]
    assert records[("W1616898", "VI_BLDG_ROOF")]["TNBNEWREADING"] == "NO"

    # 2. PPU APMM Heaters
    assert records[("W1616898", "HEATER_33KV_STATUS")]["TNBNEWREADING"] == "NO"
    assert records[("W1616898", "HEATER_11KV_STATUS")]["TNBNEWREADING"] == "YES"
    assert "Heater broken" in records[("W1616898", "HEATER_11KV_STATUS")]["TNBCOMMENTS"]

    # Heater ampere: status 'ON' left blank; numeric 0.85A formatted to 0.85
    assert records[("W1616898", "HEATER_33KV_AMP")]["TNBNEWREADING"] == ""
    assert records[("W1616898", "HEATER_11KV_AMP")]["TNBNEWREADING"] == "0.85"

    # 3. PPU APMM 33kV Switchgear Panel 1
    assert records[("W1616898", "SG_33_LOADING")]["TNBNEWREADING"] == "119"
    assert records[("W1616898", "SG_33_GIS_PRES_CB")]["TNBNEWREADING"] == "1.41"
    assert records[("W1616898", "TH_S33_BB_TREF_P31VV")]["TNBNEWREADING"] == "27.5"
    assert records[("W1616898", "TH_S33_BR_AVG_P31VV")]["TNBNEWREADING"] == "28.5"
    assert records[("W1616898", "US_S33_FP_P31VV")]["TNBNEWREADING"] == "0"
    assert records[("W1616898", "TEV_S33_FP_P31VV")]["TNBNEWREADING"] == "5"

    # 4. PPU APMM 11kV Switchgear Panel 1
    assert records[("W1616898", "SG_11_LOADING")]["TNBNEWREADING"] == "220"
    assert records[("W1616898", "TH_S11_BR_AVG_P31VV")]["TNBNEWREADING"] == "31.2"

    # 5. PPU APMM Power Transformer 1
    assert records[("W1616898", "TH_PTXHV_AVG_P31VV")]["TNBNEWREADING"] == "36.8"
    assert records[("W1616898", "TH_PTXMT_AVG_P31VV")]["TNBNEWREADING"] == "43.5"

    # 6. PPU APMM Ambient Temperature
    assert records[("W1616898", "BG_TEM_11_ROOM_P31VV")]["TNBNEWREADING"] == "30.5"

    # 7. SSU KIIP 33kV Readings
    assert records[("W1616956", "SG_33_LOADING")]["TNBNEWREADING"] == "85"
    assert records[("W1616956", "TH_S33_BR_AVG_P31VV")]["TNBNEWREADING"] == "29.8"

    # 8. SSU LYNAS 33kV Readings
    assert records[("W1616922", "SG_33_LOADING")]["TNBNEWREADING"] == "140"
    assert records[("W1616922", "TH_S33_BR_AVG_P31VV")]["TNBNEWREADING"] == "32.1"


def test_msms_33kv_actions_pipeline_presentation_adapters(synthetic_33kv_env: ProjectEnvironment, capsys: pytest.CaptureFixture[str]) -> None:
    """Verify CLI Presentation Adapters (ConsolidateMsmsAction, EnrichMsmsAction, PropagateWoAction, IngestMsmsCsvAction, PopulateDataMsmsAction) execute cleanly."""
    env = synthetic_33kv_env
    storage = env.storage
    date_str = "07-10-2026"

    # Setup seed data
    _init_data_msms_file(storage.get_data_msms_path())
    _init_total_pe_file(
        storage.get_total_pe_path(),
        [[1, "CGBG/PPU/G00001", "PPU APMM 33/11kV", date_str, "PPU", None, "CBM"]],
    )
    _create_html_xls_file(
        storage.get_msms_dir() / "sample_action.xls",
        [("W1616898", "INPRG", "CGBG/PPUAPMM", "PPU APMM 33/11KV")],
    )
    _create_raw_t1sub33_csv(
        storage.get_msms_raw_data_dir() / "raw_action_07-10-2026.csv",
        [{"WONUM": "W1616898", "TNBLOCATION": "CGBG/PPUAPMM/33KV/1", "METERNAME": "SG_33_LOADING", "TNBNEWREADING": "", "METERTYPE": "GAUGE"}],
    )
    pkg_dir = storage.get_testsheet_dir() / "GEBENG" / "01. OCTOBER" / date_str
    pkg_dir.mkdir(parents=True)
    pkg_path = pkg_dir / "01. PPU APMM 33-11kV (VI).xlsx"
    create_ppu_testsheet_workbook(
        path=pkg_path,
        substation_number=1,
        substation_name="PPU APMM 33-11kV",
        fl_number="CGBG/PPU/G00001",
        date_val="2026-10-07",
    )
    wb = openpyxl.load_workbook(pkg_path)
    wb["33kV (Main)"]["I15"] = "175"
    wb.save(pkg_path)
    wb.close()

    # 1. ConsolidateMsmsAction
    action_cons = ConsolidateMsmsAction("Consolidate MSMS")
    res_cons = action_cons.execute(env)
    assert res_cons.files_processed == 1
    assert res_cons.rows_appended == 1
    assert "Files processed: 1, Rows appended: 1, Duplicates skipped: 0" in capsys.readouterr().out

    # 2. EnrichMsmsAction
    action_enrich = EnrichMsmsAction("Enrich MSMS")
    res_enrich = action_enrich.execute(env)
    assert res_enrich.matched_count == 1
    assert res_enrich.updated_cells_count == 4
    assert "Cells updated: 4, Matched: 1, Unmatched: 0" in capsys.readouterr().out

    # 3. PropagateWoAction
    action_prop = PropagateWoAction("Propagate Work Orders")
    res_prop = action_prop.execute(env)
    assert res_prop.matched_count == 1
    assert res_prop.updated_count == 1
    assert "Work Orders propagated: 1 updated, 1 matched, 0 already populated." in capsys.readouterr().out

    # 4. IngestMsmsCsvAction
    action_ingest = IngestMsmsCsvAction("Ingest MSMS CSVs")
    res_ingest = action_ingest.execute(env)
    assert res_ingest.files_ingested == 1
    assert "Files ingested: 1, Duplicates skipped: 0" in capsys.readouterr().out

    # 5. PopulateDataMsmsAction (auto mode simulated)
    action_populate = PopulateDataMsmsAction("Populate Data MSMS")
    with patch("src.cli_selectors.select_one", return_value="auto"):
        res_populate = action_populate.execute(env)
        assert res_populate.csv_files_processed == 1
        assert res_populate.rows_populated == 1
        out = capsys.readouterr().out
        assert "CSV files processed: 1, Rows populated: 1" in out


def test_dry_run_thunder_cycle_4_benchmark_replica(tmp_path: Path) -> None:
    """Dry-run verification against an isolated sandbox replica of thunder_cycle_4_2026_4 benchmark dataset.

    Validates that:
      - ConsolidateMsms processes scattered .xls exports (e.g. 56212856.xls).
      - EnrichMsms matches all 25 substations in DATA MSMS.xlsx with TOTAL PE.xlsx.
      - PropagateWo propagates all 25 Work Orders into TOTAL PE.xlsx without data loss.
      - PopulateDataMsms correctly correlates testsheet packages by WO and populates 33kV meters.
      - No errors, warnings, or unhandled exceptions occur.
    """
    repo = JsonFileProjectRepository()
    try:
        live_meta = repo.get("thunder_cycle_4_2026_4")
    except KeyError:
        pytest.skip("thunder_cycle_4_2026_4 not configured in .cli_config.json")

    live_base = Path(live_meta.base_path)
    if not live_base.exists():
        pytest.skip(f"Live benchmark directory not found: {live_base}")

    # Create temporary isolated sandbox
    sandbox_path = tmp_path / "thunder_sandbox"
    sandbox_path.mkdir(parents=True)

    # Copy PYTHON directory containing DATA MSMS.xlsx, TOTAL PE.xlsx, and MSMS/
    shutil.copytree(live_base / "PYTHON", sandbox_path / "PYTHON")

    # Copy a representative date of TESTSHEET (07-10-2026 Gebeng with 4 substations)
    testsheet_src = live_base / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / "07-10-2026"
    testsheet_dst = sandbox_path / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / "07-10-2026"
    testsheet_dst.mkdir(parents=True)
    for item in testsheet_src.iterdir():
        if item.is_file() and item.suffix == ".xlsx":
            shutil.copy2(item, testsheet_dst)
        elif item.is_dir() and item.name == "UNSORTED RAW DATA":
            shutil.copytree(item, testsheet_dst / "UNSORTED RAW DATA")

    # Configure sandbox environment
    sandbox_meta = ProjectMetadata(
        key=live_meta.key,
        name=live_meta.name,
        po_number=live_meta.po_number,
        state=live_meta.state,
        voltage_type=live_meta.voltage_type,
        year=live_meta.year,
        cycle=live_meta.cycle,
        technologies=live_meta.technologies,
        base_path=str(sandbox_path),
    )
    sandbox_storage = LocalWorkspaceStorage(sandbox_path)
    sandbox_env = ProjectEnvironment(metadata=sandbox_meta, storage=sandbox_storage)
    service = WorkflowService()

    # Step 1: Consolidate MSMS
    res_cons = service.run_consolidate_msms(sandbox_env)
    assert res_cons.files_processed >= 1
    assert res_cons.rows_appended >= 1
    assert len(res_cons.errors) == 0

    # Step 2: Enrich MSMS
    res_enrich = service.run_enrich_msms(sandbox_env)
    assert res_enrich.matched_count == 25
    assert res_enrich.updated_cells_count == 100
    # W1616869 from 56212856.xls (PMU JERANTUT) is an off-cycle export not in Cycle 4 TOTAL PE
    assert res_enrich.unmatched_count in (0, 1)
    if res_enrich.unmatched_count == 1:
        assert res_enrich.unmatched_wos == ("W1616869",)

    # Verify DATA MSMS.xlsx has SAP FLs in Col E
    wb_msms = openpyxl.load_workbook(sandbox_storage.get_data_msms_path())
    ws_msms = wb_msms.active
    col_e_fls = [ws_msms.cell(r, 5).value for r in range(2, 27)]
    wb_msms.close()
    assert all(fl and "/" in str(fl) for fl in col_e_fls)
    assert "CGBG/PPU/G00001" in col_e_fls
    assert "CKTN/SSU/G00006" in col_e_fls

    # Step 3: Propagate Work Orders
    res_prop = service.run_propagate_wo(sandbox_env)
    assert res_prop.matched_count == 25
    assert res_prop.updated_count == 25

    # Verify TOTAL PE.xlsx Column F contains valid Work Orders for the 25 matched substations
    wb_pe = openpyxl.load_workbook(sandbox_storage.get_total_pe_path())
    ws_pe = wb_pe["DataCycle1"]
    pe_wos = [ws_pe.cell(r, 6).value for r in range(2, ws_pe.max_row + 1) if ws_pe.cell(r, 6).value is not None]
    wb_pe.close()
    assert len(pe_wos) == 25
    assert all(str(wo).startswith("W") for wo in pe_wos)

    # Step 4: Populate Data MSMS for 07-10-2026_001.csv
    # Retain only 07-10-2026_001.csv in TO BE FILLED for fast targeted run
    to_be_filled_dir = sandbox_storage.get_msms_to_be_filled_dir()
    for f in to_be_filled_dir.glob("*.csv"):
        if f.name != "07-10-2026_001.csv":
            f.unlink()

    res_pop = service.run_populate_data_msms(
        sandbox_env,
        PopulateDataMsmsRequest(
            mode=PopulateMode.AUTO,
            overwrite=True,
            move_to_completed=True,
        ),
    )
    assert res_pop.csv_files_processed == 1
    assert res_pop.rows_populated > 0
    assert len(res_pop.errors) == 0

    # Verify completed CSV moved to PYTHON/MSMS/COMPLETED/
    completed_csv = sandbox_storage.get_msms_completed_dir() / "07-10-2026_001.csv"
    assert completed_csv.exists()
    assert not (to_be_filled_dir / "07-10-2026_001.csv").exists()

    # Verify rows in completed CSV actually received readings
    with open(completed_csv, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        populated_count = sum(1 for r in reader if r.get("TNBNEWREADING"))
    assert populated_count >= 800


def test_msms_33kv_work_order_mapping_integrity(synthetic_33kv_env: ProjectEnvironment) -> None:
    """Validate that no Work Order mappings are lost, corrupted, or misaligned across complex 33kV domain edge cases."""
    env = synthetic_33kv_env
    storage = env.storage
    service = WorkflowService()
    date_str = "08-10-2026"

    # Edge cases:
    # 1. Gebeng Cabin vs GIS domain aliases
    # 2. Parenthesized commercial names: SSU PUTRA SQUARE 2 ( ZENITH )
    # 3. Voltage stripping variations: 275/132/33/11kV, 132/33kV
    # 4. Acronym and token subset matching
    data_msms_path = storage.get_data_msms_path()
    _init_data_msms_file(data_msms_path)

    total_pe_path = storage.get_total_pe_path()
    _init_total_pe_file(
        total_pe_path,
        [
            [7, "CGBG/PMU/G00002", "PMU GBID CABIN 33kV", date_str, "PMU", None, "CBM"],
            [8, "CGBG/PMU/G00004", "PMU GBID 275/132/33/11kV", date_str, "PMU", None, "CBM"],
            [16, "CKTN/SSU/G00004", "SSU PUTRA SQUARE 2 (SSU ZENITH) 33kV", date_str, "SSU", None, "CBM"],
            [28, "CCHL/MPU/G00001", "MINI PPU TERLA 33/11kV", date_str, "MPPU", None, "CBM"],
        ],
    )

    xls_path = storage.get_msms_dir() / "edge_cases.xls"
    _create_html_xls_file(
        xls_path,
        [
            ("W1616864", "INPRG", "CGBG/PMUGBID", "PMU GBID ( T ) 132/33KV"),
            ("W1616865", "INPRG", "CGBG/PMUGBIG", "PMU GBIG GIS 132/33/11kV"),
            ("W1616928", "INPRG", "CKTN/SSUPUTS", "SSU PUTRA SQUARE 2 ( ZENITH ) 33KV"),
            ("W1616896", "INPRG", "CCHL/MPUMPTERLA", "MINI PPU TERLA 33/11KV"),
        ],
    )

    # 1. Consolidate MSMS
    res_cons = service.run_consolidate_msms(env)
    assert res_cons.files_processed == 1
    assert res_cons.rows_appended == 4

    # 2. Enrich MSMS
    res_enrich = service.run_enrich_msms(env)
    assert res_enrich.matched_count == 4
    assert res_enrich.unmatched_count == 0
    assert res_enrich.updated_cells_count == 16

    # Verify SAP FLs assigned correctly
    wb_msms = openpyxl.load_workbook(data_msms_path)
    ws_msms = wb_msms.active
    msms_fls = {ws_msms.cell(r, 1).value: ws_msms.cell(r, 5).value for r in range(2, 6)}
    wb_msms.close()

    assert msms_fls["W1616864"] == "CGBG/PMU/G00002"  # Cabin
    assert msms_fls["W1616865"] == "CGBG/PMU/G00004"  # GIS
    assert msms_fls["W1616928"] == "CKTN/SSU/G00004"  # Zenith
    assert msms_fls["W1616896"] == "CCHL/MPU/G00001"  # Terla

    # 3. Propagate Work Orders
    res_prop = service.run_propagate_wo(env)
    assert res_prop.matched_count == 4
    assert res_prop.updated_count == 4

    # Verify TOTAL PE.xlsx has exact Work Orders
    wb_pe = openpyxl.load_workbook(total_pe_path)
    ws_pe = wb_pe["DataCycle1"]
    pe_wos = {ws_pe.cell(r, 2).value: ws_pe.cell(r, 6).value for r in range(2, 6)}
    wb_pe.close()

    assert pe_wos["CGBG/PMU/G00002"] == "W1616864"
    assert pe_wos["CGBG/PMU/G00004"] == "W1616865"
    assert pe_wos["CKTN/SSU/G00004"] == "W1616928"
    assert pe_wos["CCHL/MPU/G00001"] == "W1616896"


def test_msms_33kv_cli_execution_dry_run_active_project(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """Verify that CLI execution dry-run against the active project replica executes cleanly without errors."""
    repo = JsonFileProjectRepository()
    try:
        live_meta = repo.get("thunder_cycle_4_2026_4")
    except KeyError:
        pytest.skip("thunder_cycle_4_2026_4 not configured in .cli_config.json")

    live_base = Path(live_meta.base_path)
    if not live_base.exists():
        pytest.skip(f"Live benchmark directory not found: {live_base}")

    # Create temporary isolated sandbox
    sandbox_path = tmp_path / "thunder_cli_sandbox"
    sandbox_path.mkdir(parents=True)

    # Copy PYTHON directory
    shutil.copytree(live_base / "PYTHON", sandbox_path / "PYTHON")

    # Copy a representative date of TESTSHEET (07-10-2026 Gebeng)
    testsheet_src = live_base / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / "07-10-2026"
    testsheet_dst = sandbox_path / "TESTSHEET" / "GEBENG" / "01. OCTOBER" / "07-10-2026"
    testsheet_dst.mkdir(parents=True)
    for item in testsheet_src.iterdir():
        if item.is_file() and item.suffix == ".xlsx":
            shutil.copy2(item, testsheet_dst)
        elif item.is_dir() and item.name == "UNSORTED RAW DATA":
            shutil.copytree(item, testsheet_dst / "UNSORTED RAW DATA")

    # Configure sandbox environment
    sandbox_meta = ProjectMetadata(
        key=live_meta.key,
        name=live_meta.name,
        po_number=live_meta.po_number,
        state=live_meta.state,
        voltage_type=live_meta.voltage_type,
        year=live_meta.year,
        cycle=live_meta.cycle,
        technologies=live_meta.technologies,
        base_path=str(sandbox_path),
    )
    sandbox_storage = LocalWorkspaceStorage(sandbox_path)
    sandbox_env = ProjectEnvironment(metadata=sandbox_meta, storage=sandbox_storage)

    # Execute CLI presentation adapters in sequence
    action_cons = ConsolidateMsmsAction("Consolidate MSMS")
    res_cons = action_cons.execute(sandbox_env)
    assert res_cons is not None
    assert "Files processed:" in capsys.readouterr().out

    action_enrich = EnrichMsmsAction("Enrich MSMS")
    res_enrich = action_enrich.execute(sandbox_env)
    assert res_enrich is not None
    assert "Cells updated:" in capsys.readouterr().out

    action_prop = PropagateWoAction("Propagate Work Orders")
    res_prop = action_prop.execute(sandbox_env)
    assert res_prop is not None
    assert "Work Orders propagated:" in capsys.readouterr().out

    # Retain only 07-10-2026_001.csv for fast execution
    to_be_filled_dir = sandbox_storage.get_msms_to_be_filled_dir()
    for f in to_be_filled_dir.glob("*.csv"):
        if f.name != "07-10-2026_001.csv":
            f.unlink()

    action_pop = PopulateDataMsmsAction("Populate Data MSMS")
    with patch("src.cli_selectors.select_one", return_value="auto"):
        res_pop = action_pop.execute(sandbox_env)
        assert res_pop is not None
        out = capsys.readouterr().out
        assert "CSV files processed: 1" in out
        assert "Rows populated:" in out

