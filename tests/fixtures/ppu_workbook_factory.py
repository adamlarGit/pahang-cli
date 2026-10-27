"""Shared PPU 33kV testsheet workbook test fixture factory."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Mapping, Sequence
import openpyxl


_DEFAULT_SENTINEL = object()


def create_ppu_testsheet_workbook(
    path: Path | str | None = None,
    substation_number: int | None = 1,
    substation_name: str = "PPU APMM 33-11kV",
    fl_number: str = "CGBG-PPU-APMM",
    date_val: str | datetime | date = "2026-10-07",
    time_val: str = "10:30 AM",
    ambient_temp: str = "30.5",
    humidity: str = "65%",
    background_tev: str = "5",
    gps_coordinate: str = "3.8123, 103.3214",
    has_33kv: bool = True,
    has_11kv: bool = True,
    panels_33kv: int | Sequence[Mapping[str, Any]] = 3,
    panels_11kv: int | Sequence[Mapping[str, Any]] = 3,
    extra_panels_33kv: Sequence[Sequence[Mapping[str, Any]]] = (),
    extra_panels_11kv: Sequence[Sequence[Mapping[str, Any]]] = (),
    swg_33kv_details: Mapping[str, Any] | None = None,
    swg_11kv_details: Mapping[str, Any] | None = None,
    overview_rhs_photos_33kv: Sequence[int] = (),
    overview_lhs_photos_33kv: Sequence[int] = (),
    overview_rhs_photos_11kv: Sequence[int] = (),
    overview_lhs_photos_11kv: Sequence[int] = (),
    visual_checklist_items: Sequence[Mapping[str, Any]] | None = None,
    other_findings: Sequence[str] = (),
    raw_photo_bounds: Mapping[str, Any] | None = None,
    v2_substation_num_33kv: Any = _DEFAULT_SENTINEL,
    v2_substation_num_11kv: Any = _DEFAULT_SENTINEL,
    include_tx_sheet: bool = True,
    tx_sheet_name: str = "TX",
    vi_sheet_name: str = "Visual Inspec.",
    transformer_ratings: Sequence[str] = (),
    transformers: Sequence[Mapping[str, Any]] | None = None,
    local_tx: Mapping[str, Any] | None = None,
    lvdb_details: Mapping[str, Any] | None = None,
    include_signature_placeholders: bool = False,
) -> openpyxl.Workbook:
    """Build a synthetic PPU 33kV/11kV openpyxl workbook for testing."""
    wb = openpyxl.Workbook()
    # Remove default sheet once we add real sheets
    default_sheet = wb.active

    def _populate_switchgear_sheet(
        ws: openpyxl.worksheet.worksheet.Worksheet,
        voltage_label: str,
        panels_input: int | Sequence[Mapping[str, Any]],
        swg_details: Mapping[str, Any] | None,
        v2_val: Any,
        overview_rhs: Sequence[int],
        overview_lhs: Sequence[int],
    ) -> None:
        # Cell V2
        if v2_val is not None:
            ws["V2"] = v2_val

        # Metadata Header
        ws["D6"] = substation_name
        ws["D7"] = fl_number
        ws["D8"] = date_val
        ws["D9"] = time_val
        ws["U6"] = ambient_temp
        ws["U7"] = humidity
        ws["U8"] = background_tev

        # Switchgear Details (L6..L11)
        details = swg_details or {}
        ws["L6"] = details.get("manufacturer", "TAMCO")
        ws["L7"] = details.get("manufactured_year", "2020" if voltage_label == "33kV" else "2018")
        ws["L8"] = details.get("model", "GV3" if voltage_label == "33kV" else "VH3")
        ws["L9"] = details.get("serial_no", f"SN-{voltage_label}-001")
        ws["L10"] = details.get("switchgear_type", "GIS" if voltage_label == "33kV" else "VCB")
        ws["L11"] = details.get("rating", "1250A" if voltage_label == "33kV" else "630A")

        # Normalize panels list
        panel_specs: list[Mapping[str, Any]]
        if isinstance(panels_input, int):
            panel_specs = [
                {
                    "feeder_no": f"{'3P' if voltage_label == '33kV' else '1K'}{i + 1}",
                    "name": f"FEEDER {voltage_label} {i + 1}",
                    "status": "CLOSE",
                    "load_amp": f"{100 + i * 20}A",
                    "cable_type": "XLPE 1C 500mm2" if voltage_label == "33kV" else "XLPE 3C 240mm2",
                    "heater_amp": "0.5A",
                    "panel_type": "GIS" if voltage_label == "33kV" else "VCB",
                    "breaker_serial": f"CB-{voltage_label}-{i + 1}",
                    "breaker_photo": 100 + i * 10 + 1,
                    "cable_photo": 100 + i * 10 + 2,
                    "busbar_photo": 100 + i * 10 + 3,
                    "pt_photo": 100 + i * 10 + 4,
                    "cable_entry_photo": 100 + i * 10 + 5,
                    "link_box_photo": 100 + i * 10 + 6,
                    "tev_reading": "0",
                    "tev_ppc": "0",
                    "us_reading": "0",
                }
                for i in range(panels_input)
            ]
        else:
            panel_specs = list(panels_input)

        # Write panels (stride = 6 rows, starting row 15)
        for idx, p in enumerate(panel_specs):
            if idx >= 8:
                break
            r = 15 + idx * 6
            ws[f"B{r}"] = idx + 1
            ws[f"C{r}"] = p.get("feeder_no", "")
            ws[f"D{r}"] = p.get("name", "")
            ws[f"E{r}"] = p.get("cable_type", "")

            raw_status = p.get("status", "CLOSE")
            if "/" in raw_status:
                b_pos, b_stat = raw_status.split("/", 1)
                ws[f"F{r}"] = b_pos.strip()
                ws[f"G{r}"] = b_stat.strip()
            else:
                ws[f"F{r}"] = p.get("breaker_position", "")
                ws[f"G{r}"] = raw_status

            ws[f"H{r}"] = p.get("busbar_position", "MAIN")
            ws[f"I{r}"] = p.get("load_amp", "")
            ws[f"J{r}"] = p.get("heater_amp", "")
            ws[f"K{r}"] = "BREAKER"
            ws[f"V{r}"] = p.get("breaker_serial", "")

            # TEV and Ultrasound readings
            if "tev_reading" in p and p["tev_reading"] is not None:
                ws[f"R{r}"] = p["tev_reading"]
            if "tev_ppc" in p and p["tev_ppc"] is not None:
                ws[f"S{r}"] = p["tev_ppc"]
            if "us_reading" in p and p["us_reading"] is not None:
                ws[f"T{r}"] = p["us_reading"]
            if "tev_char" in p and p["tev_char"] is not None:
                ws[f"U{r}"] = p["tev_char"]
            elif "us_char" in p and p["us_char"] is not None:
                ws[f"U{r}"] = p["us_char"]

            if "cable_tev_reading" in p and p["cable_tev_reading"] is not None:
                ws[f"R{r + 1}"] = p["cable_tev_reading"]
            if "cable_tev_ppc" in p and p["cable_tev_ppc"] is not None:
                ws[f"S{r + 1}"] = p["cable_tev_ppc"]
            if "cable_us_reading" in p and p["cable_us_reading"] is not None:
                ws[f"T{r + 1}"] = p["cable_us_reading"]

            if "sub_rows" in p and isinstance(p["sub_rows"], dict):
                for offset, col_vals in p["sub_rows"].items():
                    sub_row = r + int(offset)
                    for col_letter, cell_val in col_vals.items():
                        ws[f"{col_letter}{sub_row}"] = cell_val

            # Photos in Col P (Row 1..6)
            if "breaker_photo" in p and p["breaker_photo"] is not None:
                ws[f"P{r}"] = p["breaker_photo"]
            if "cable_photo" in p and p["cable_photo"] is not None:
                ws[f"P{r + 1}"] = p["cable_photo"]
            if "busbar_photo" in p and p["busbar_photo"] is not None:
                ws[f"P{r + 2}"] = p["busbar_photo"]
            if "pt_photo" in p and p["pt_photo"] is not None:
                ws[f"P{r + 3}"] = p["pt_photo"]
            if "cable_entry_photo" in p and p["cable_entry_photo"] is not None:
                ws[f"P{r + 4}"] = p["cable_entry_photo"]
            if "link_box_photo" in p and p["link_box_photo"] is not None:
                ws[f"P{r + 5}"] = p["link_box_photo"]

            # Secondary panel photo in Col Q regex SP: <photo>
            if "secondary_photo" in p and p["secondary_photo"] is not None:
                ws[f"Q{r}"] = f"SP: {p['secondary_photo']}"
            elif "remarks" in p and p["remarks"]:
                ws[f"Q{r}"] = p["remarks"]

            # PT compartment label and measurement in Row 4
            if ws[f"K{r + 3}"].value is None:
                ws[f"K{r + 3}"] = "PT COMPARTMENT"
            if p.get("has_pt_measurement"):
                ws[f"L{r + 3}"] = "24.5"

        # Overview Bays (Rows 63-74)
        # RHS Overview (Rows 63-68)
        rhs_labels = ["FRONT", "BACK", "TOP", "CABLE CELLAR", "MBB", "RBB"]
        for o_idx, label in enumerate(rhs_labels):
            row_num = 63 + o_idx
            ws[f"C{row_num}"] = label
            if o_idx < len(overview_rhs):
                ws[f"P{row_num}"] = overview_rhs[o_idx]

        # LHS Overview (Rows 69-74)
        lhs_labels = ["FRONT", "BACK", "TOP", "CABLE CELLAR", "MBB", "RBB"]
        for o_idx, label in enumerate(lhs_labels):
            row_num = 69 + o_idx
            ws[f"C{row_num}"] = label
            if o_idx < len(overview_lhs):
                ws[f"P{row_num}"] = overview_lhs[o_idx]

        # Signature placeholders (Rows 74-80)
        if include_signature_placeholders:
            ws["E77"] = "{{signvendor}}"
            ws["L77"] = "{{signtnb}}"

    # Resolve v2 default values
    v2_33 = substation_number if v2_substation_num_33kv is _DEFAULT_SENTINEL else v2_substation_num_33kv
    v2_11 = substation_number if v2_substation_num_11kv is _DEFAULT_SENTINEL else v2_substation_num_11kv

    # 1. 33kV (Main)
    if has_33kv:
        ws_33 = wb.create_sheet(title="33kV (Main)")
        _populate_switchgear_sheet(
            ws_33,
            voltage_label="33kV",
            panels_input=panels_33kv,
            swg_details=swg_33kv_details,
            v2_val=v2_33,
            overview_rhs=overview_rhs_photos_33kv,
            overview_lhs=overview_lhs_photos_33kv,
        )

        # 33kV (Extra) sheets
        for e_idx, extra_panels in enumerate(extra_panels_33kv):
            sheet_title = "33kV (Extra)" if e_idx == 0 else f"33kV (Extra) ({e_idx + 1})"
            ws_extra = wb.create_sheet(title=sheet_title)
            _populate_switchgear_sheet(
                ws_extra,
                voltage_label="33kV",
                panels_input=extra_panels,
                swg_details=swg_33kv_details,
                v2_val=None,
                overview_rhs=(),
                overview_lhs=(),
            )

    # 2. 11kV (Main)
    if has_11kv:
        ws_11 = wb.create_sheet(title="11kV (Main)")
        _populate_switchgear_sheet(
            ws_11,
            voltage_label="11kV",
            panels_input=panels_11kv,
            swg_details=swg_11kv_details,
            v2_val=v2_11,
            overview_rhs=overview_rhs_photos_11kv,
            overview_lhs=overview_lhs_photos_11kv,
        )

        # 11kV (Extra) sheets
        for e_idx, extra_panels in enumerate(extra_panels_11kv):
            sheet_title = "11kV (Extra)" if e_idx == 0 else f"11kV (Extra) ({e_idx + 1})"
            ws_extra = wb.create_sheet(title=sheet_title)
            _populate_switchgear_sheet(
                ws_extra,
                voltage_label="11kV",
                panels_input=extra_panels,
                swg_details=swg_11kv_details,
                v2_val=None,
                overview_rhs=(),
                overview_lhs=(),
            )

    # 3. Visual Inspection sheet
    ws_vi = wb.create_sheet(title=vi_sheet_name)
    ws_vi["B3"] = "Substation Name"
    ws_vi["D3"] = substation_name
    ws_vi["L3"] = "GPS:"
    ws_vi["M3"] = gps_coordinate

    # Default checklist items if none provided
    if visual_checklist_items is None:
        items = [
            {"category": "1. CIVIL", "item_no": "1.1", "description": "Fence & gate condition", "status": "YES", "remarks": ""},
            {"category": "1. CIVIL", "item_no": "1.2", "description": "Trench cover missing/broken", "status": "NO", "remarks": "Broken trench slab"},
            {"category": "2. SWITCHGEAR", "item_no": "2.1", "description": "Earthing connection intact", "status": "YES", "remarks": "Re-tightened earthing bar"},
            {"category": "2. SWITCHGEAR", "item_no": "2.2", "description": "Gas pressure gauge normal", "status": "YES", "remarks": ""},
        ]
    else:
        items = list(visual_checklist_items)

    curr_row = 8
    last_cat = ""
    for item in items:
        cat = item.get("category", "")
        if cat and cat != last_cat:
            ws_vi[f"B{curr_row}"] = cat
            curr_row += 1
            last_cat = cat

        ws_vi[f"B{curr_row}"] = item.get("item_no", "")
        ws_vi[f"C{curr_row}"] = item.get("description", "")
        status = item.get("status", "YES")
        if status == "NO":
            ws_vi[f"D{curr_row}"] = ""
            ws_vi[f"E{curr_row}"] = "X"
        else:
            ws_vi[f"D{curr_row}"] = "X"
            ws_vi[f"E{curr_row}"] = ""
        ws_vi[f"F{curr_row}"] = item.get("remarks", "")
        curr_row += 1

    # Category 8 "Other Findings" (Rows 78-80)
    for o_idx, remark in enumerate(other_findings):
        r_num = 78 + o_idx
        if r_num <= 80:
            ws_vi[f"B{r_num}"] = f"8.{o_idx + 1}"
            ws_vi[f"C{r_num}"] = remark
            ws_vi[f"F{r_num}"] = remark

    # Signature placeholders (Rows 80-87)
    if include_signature_placeholders:
        ws_vi["D82"] = "{{signvendor}}"
        ws_vi["M82"] = "{{signtnb}}"

    # 4. RAW DATA sheet
    bounds = raw_photo_bounds if raw_photo_bounds is not None else {"ir_start": 1, "ir_end": 50, "dg_start": 51, "dg_end": 90}
    ws_raw = wb.create_sheet(title="RAW DATA")
    ws_raw.append(["", "START", "END"])
    ws_raw.append(["IR", bounds.get("ir_start", 1), bounds.get("ir_end", 50)])
    ws_raw.append(["DG", bounds.get("dg_start", 51), bounds.get("dg_end", 90)])

    # 5. TX sheet
    if include_tx_sheet:
        ws_tx = wb.create_sheet(title=tx_sheet_name)
        # Header labels
        ws_tx["B5"] = "1. Substation Details:"
        ws_tx["D6"] = substation_name
        ws_tx["G5"] = "2. Transformer Details:"

        # Column blocks for slots 1..5
        # Slot 1: TX 1
        ws_tx["G6"] = "TX 1"
        ws_tx["I6"] = "Model:"
        ws_tx["G7"] = "2.1 Manufacturer:"
        ws_tx["G8"] = "2.2 Rating (MVA):"
        ws_tx["G9"] = "2.3 Serial No. :"

        # Slot 2: TX 2
        ws_tx["K6"] = "TX 2"
        ws_tx["L6"] = "Model:"
        ws_tx["K7"] = "2.1 Manufacturer:"
        ws_tx["K8"] = "2.2 Rating (MVA):"
        ws_tx["K9"] = "2.3 Serial No. :"

        # Slot 3: TX 3
        ws_tx["N6"] = "TX 3"
        ws_tx["O6"] = "Model:"
        ws_tx["N7"] = "2.1 Manufacturer:"
        ws_tx["N8"] = "2.2 Rating (MVA):"
        ws_tx["N9"] = "2.3 Serial No. :"

        # Slot 4: TX 4
        ws_tx["Q6"] = "TX 4"
        ws_tx["R6"] = "Model:"
        ws_tx["Q7"] = "2.1 Manufacturer:"
        ws_tx["Q8"] = "2.2 Rating (KVA):"
        ws_tx["Q9"] = "2.3 Serial No. :"

        # Slot 5: LOCAL TX
        ws_tx["T6"] = "LOCAL TX"
        ws_tx["U6"] = "Model:"
        ws_tx["T7"] = "2.1 Manufacturer:"
        ws_tx["T8"] = "2.2 Rating (KVA):"
        ws_tx["T9"] = "2.3 Serial No. :"

        # Part 1 and Part 2 section headers
        ws_tx["B10"] = "TRANSFORMER"
        ws_tx["B11"] = "DETAILS"
        ws_tx["J11"] = "ULTRASOUND"
        ws_tx["R11"] = "THERMOGRAPHY"
        ws_tx["B12"] = "No."
        ws_tx["C12"] = "TX No."
        ws_tx["F12"] = "Load ( A) "
        ws_tx["H12"] = "Cable Type"
        ws_tx["J12"] = "HV Cable"
        ws_tx["N12"] = "LV Cable"
        ws_tx["R12"] = "HV Cable"
        ws_tx["W12"] = "LV Cable"

        ws_tx["B30"] = "DETAILS"
        ws_tx["F30"] = "THERMOGRAPHY"
        ws_tx["B31"] = "No"
        ws_tx["C31"] = "TX No"
        ws_tx["F31"] = "Marshalling Kiosk"
        ws_tx["L31"] = "Conservator Tank"
        ws_tx["R31"] = "Body"
        ws_tx["X31"] = "Cooling Fin"
        ws_tx["AC31"] = "Remarks"

        # Resolve transformers list
        tx_list: list[Mapping[str, Any]] = []
        if transformers is not None:
            tx_list = list(transformers)
        elif transformer_ratings:
            tx_list = [
                {
                    "slot_id": i + 1,
                    "tx_id": f"TX {i + 1}",
                    "rating": r,
                    "model": "15-201",
                    "manufacturer": "MTM",
                    "serial_no": f"SN-TX-{i + 1}",
                    "load_amp": f"{100 + i * 20}A",
                    "cable_type": "XLPE",
                    "hv_ultrasound": "0",
                    "lv_ultrasound": "0",
                    "hv_photo": 560 + i * 15,
                    "lv_photo": 554 + i * 15,
                    "mk_photo": 550 + i * 15,
                    "ct_photo": 547 + i * 15,
                    "body_photo": 555 + i * 15,
                    "cf_photo": f"{546 + i * 15}/{548 + i * 15}",
                    "hv_cable_thermal": {"tmin": "36.4", "tmax": "37.2", "delta_t": "0.8", "avg": "36.8"},
                    "lv_cable_thermal": {"tmin": "36.0", "tmax": "36.1", "delta_t": "0.1", "avg": "36.05"},
                    "marshalling_kiosk_temp": {"tmin": "35.8", "tmax": "36.0", "delta_t": "0.2", "avg": "35.9"},
                    "conservator_temp": {"tmin": "41.0", "tmax": "41.5", "delta_t": "0.5", "avg": "41.25"},
                    "body_temp": {"tmin": "43.0", "tmax": "43.9", "delta_t": "0.9", "avg": "43.45"},
                    "cooling_fin_temp": {"tmin": "43.3", "tmax": "43.7", "delta_t": "0.4", "avg": "43.5"},
                }
                for i, r in enumerate(transformer_ratings)
            ]

        slot_configs = [
            {"slot": 1, "default_id": "TX 1", "model_col": "J", "mfg_col": "I", "rating_header": "G8", "rating_col": "I8", "sn_col": "I", "p1_row": 14, "p2_row": 33},
            {"slot": 2, "default_id": "TX 2", "model_col": "M", "mfg_col": "L", "rating_header": "K8", "rating_col": "L8", "sn_col": "L", "p1_row": 17, "p2_row": 36},
            {"slot": 3, "default_id": "TX 3", "model_col": "P", "mfg_col": "O", "rating_header": "N8", "rating_col": "O8", "sn_col": "O", "p1_row": 20, "p2_row": 39},
            {"slot": 4, "default_id": "TX 4", "model_col": "S", "mfg_col": "R", "rating_header": "Q8", "rating_col": "R8", "sn_col": "R", "p1_row": 23, "p2_row": 42},
        ]

        for idx, t in enumerate(tx_list):
            if idx >= len(slot_configs):
                break
            cfg = slot_configs[idx]
            p1_r = cfg["p1_row"]
            p2_r = cfg["p2_row"]
            tx_name = t.get("tx_id", cfg["default_id"])

            # Header metadata
            if "model" in t:
                ws_tx[f"{cfg['model_col']}6"] = t["model"]
            if "manufacturer" in t:
                ws_tx[f"{cfg['mfg_col']}7"] = t["manufacturer"]
            if "rating" in t:
                ws_tx[cfg["rating_header"]] = t["rating"]
                ws_tx[cfg["rating_col"]] = t["rating"]
            if "serial_no" in t:
                ws_tx[f"{cfg['sn_col']}9"] = t["serial_no"]

            # Part 1: Rows 14..28
            ws_tx[f"B{p1_r}"] = idx + 1
            ws_tx[f"C{p1_r}"] = tx_name
            ws_tx[f"F{p1_r}"] = t.get("load_amp", "")
            ws_tx[f"H{p1_r}"] = t.get("cable_type", "")
            ws_tx[f"J{p1_r}"] = t.get("hv_ultrasound", "")
            ws_tx[f"N{p1_r}"] = t.get("lv_ultrasound", "")

            hv_th = t.get("hv_cable_thermal", {})
            ws_tx[f"R{p1_r}"] = hv_th.get("tmin", "")
            ws_tx[f"S{p1_r}"] = hv_th.get("tmax", "")
            ws_tx[f"T{p1_r}"] = hv_th.get("delta_t", "")
            ws_tx[f"U{p1_r}"] = hv_th.get("avg", "")
            if "hv_photo" in t and t["hv_photo"] is not None:
                ws_tx[f"V{p1_r}"] = t["hv_photo"]

            lv_th = t.get("lv_cable_thermal", {})
            ws_tx[f"W{p1_r}"] = lv_th.get("tmin", "")
            ws_tx[f"X{p1_r}"] = lv_th.get("tmax", "")
            ws_tx[f"Y{p1_r}"] = lv_th.get("delta_t", "")
            ws_tx[f"Z{p1_r}"] = lv_th.get("avg", "")
            if "lv_photo" in t and t["lv_photo"] is not None:
                ws_tx[f"AA{p1_r}"] = t["lv_photo"]

            # Part 2: Rows 33..48
            ws_tx[f"B{p2_r}"] = idx + 1
            ws_tx[f"C{p2_r}"] = tx_name

            mk_th = t.get("marshalling_kiosk_temp", {})
            ws_tx[f"F{p2_r}"] = mk_th.get("tmin", "")
            ws_tx[f"G{p2_r}"] = mk_th.get("tmax", "")
            ws_tx[f"H{p2_r}"] = mk_th.get("delta_t", "")
            ws_tx[f"I{p2_r}"] = mk_th.get("avg", "")
            if "mk_photo" in t and t["mk_photo"] is not None:
                ws_tx[f"J{p2_r}"] = t["mk_photo"]

            ct_th = t.get("conservator_temp", {})
            ws_tx[f"L{p2_r}"] = ct_th.get("tmin", "")
            ws_tx[f"M{p2_r}"] = ct_th.get("tmax", "")
            ws_tx[f"N{p2_r}"] = ct_th.get("delta_t", "")
            ws_tx[f"O{p2_r}"] = ct_th.get("avg", "")
            if "ct_photo" in t and t["ct_photo"] is not None:
                ws_tx[f"P{p2_r}"] = t["ct_photo"]

            b_th = t.get("body_temp", {})
            ws_tx[f"R{p2_r}"] = b_th.get("tmin", "")
            ws_tx[f"S{p2_r}"] = b_th.get("tmax", "")
            ws_tx[f"T{p2_r}"] = b_th.get("delta_t", "")
            ws_tx[f"U{p2_r}"] = b_th.get("avg", "")
            if "body_photo" in t and t["body_photo"] is not None:
                ws_tx[f"V{p2_r}"] = t["body_photo"]

            cf_th = t.get("cooling_fin_temp", {})
            ws_tx[f"X{p2_r}"] = cf_th.get("tmin", "")
            ws_tx[f"Y{p2_r}"] = cf_th.get("tmax", "")
            ws_tx[f"Z{p2_r}"] = cf_th.get("delta_t", "")
            ws_tx[f"AA{p2_r}"] = cf_th.get("avg", "")
            if "cf_photo" in t and t["cf_photo"] is not None:
                ws_tx[f"AB{p2_r}"] = t["cf_photo"]

            if "remarks" in t:
                ws_tx[f"AC{p2_r}"] = t["remarks"]

        # Local TX (Slot 5: Row 26 & Row 45)
        if local_tx is not None:
            ltx = local_tx
            if "model" in ltx:
                ws_tx["V6"] = ltx["model"]
            if "manufacturer" in ltx:
                ws_tx["U7"] = ltx["manufacturer"]
            if "rating" in ltx:
                ws_tx["T8"] = ltx["rating"]
                ws_tx["U8"] = ltx["rating"]
            if "serial_no" in ltx:
                ws_tx["U9"] = ltx["serial_no"]

            # Part 1 (Row 26)
            ws_tx["B26"] = 5
            ws_tx["C26"] = ltx.get("tx_id", "Local TX")
            ws_tx["F26"] = ltx.get("load_amp", "")
            ws_tx["H26"] = ltx.get("cable_type", "")
            ws_tx["J26"] = ltx.get("hv_ultrasound", "")
            hv_th = ltx.get("hv_cable_thermal", {})
            ws_tx["R26"] = hv_th.get("tmin", "")
            ws_tx["S26"] = hv_th.get("tmax", "")
            ws_tx["T26"] = hv_th.get("delta_t", "")
            ws_tx["U26"] = hv_th.get("avg", "")
            if "hv_photo" in ltx and ltx["hv_photo"] is not None:
                ws_tx["V26"] = ltx["hv_photo"]

            lv_th = ltx.get("lv_cable_thermal", {})
            ws_tx["W26"] = lv_th.get("tmin", "")
            ws_tx["X26"] = lv_th.get("tmax", "")
            ws_tx["Y26"] = lv_th.get("delta_t", "")
            ws_tx["Z26"] = lv_th.get("avg", "")
            if "lv_photo" in ltx and ltx["lv_photo"] is not None:
                ws_tx["AA26"] = ltx["lv_photo"]

            # Part 2 (Row 45)
            ws_tx["B45"] = 5
            ws_tx["C45"] = ltx.get("tx_id", "LOCAL TX")
            b_th = ltx.get("body_temp", {})
            ws_tx["R45"] = b_th.get("tmin", "")
            ws_tx["S45"] = b_th.get("tmax", "")
            ws_tx["T45"] = b_th.get("delta_t", "")
            ws_tx["U45"] = b_th.get("avg", "")
            if "body_photo" in ltx and ltx["body_photo"] is not None:
                ws_tx["V45"] = ltx["body_photo"]

        # Station Service LVDB (Rows 49..54)
        if lvdb_details is not None:
            ws_tx["B49"] = "LV DISTRIBUTION BOARD/FEEDER PILLAR/OTHERS"
            ws_tx["B51"] = "Rating (A):"
            ws_tx["D51"] = lvdb_details.get("rating", "")
            if "photo" in lvdb_details and lvdb_details["photo"] is not None:
                ws_tx["G51"] = f"IR: {lvdb_details['photo']}" if isinstance(lvdb_details['photo'], int) else str(lvdb_details['photo'])
            ws_tx["B52"] = "Serial No.: "
            ws_tx["D52"] = lvdb_details.get("serial_no", "")
            ws_tx["B53"] = "Type:"
            ws_tx["D53"] = lvdb_details.get("model", "")
            ws_tx["B54"] = "Manufacturer:"
            ws_tx["D54"] = lvdb_details.get("manufacturer", "")

        # Signature placeholders (Rows 70-78)
        if include_signature_placeholders:
            ws_tx["F73"] = "{{signvendor}}"
            ws_tx["R73"] = "{{signtnb}}"

    # Remove the initial blank sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    if path is not None:
        save_path = Path(path)
        save_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(save_path)

    return wb
