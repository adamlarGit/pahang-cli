"""Tests for testsheet parser strategy seam and factory."""

from pathlib import Path
import openpyxl
import pytest

from src.testsheet.models import TestsheetData, RawPhotoRanges
from src.testsheet.parsers.base import TestsheetParserStrategy
from src.testsheet.parsers.factory import get_testsheet_parser
from src.testsheet.parsers.pce import PceTestsheetParser
from src.testsheet.parsers.ppu import PpuTestsheetParser
from src.testsheet.extractor import TestsheetExtractor
from src.testsheet.repository import SubstationTestsheetRepository


def test_factory_resolves_by_voltage_type():
    parser_11kv = get_testsheet_parser(voltage_type="11kV")
    assert isinstance(parser_11kv, PceTestsheetParser)
    assert isinstance(parser_11kv, TestsheetParserStrategy)

    parser_33kv = get_testsheet_parser(voltage_type="33kV")
    assert isinstance(parser_33kv, PpuTestsheetParser)
    assert isinstance(parser_33kv, TestsheetParserStrategy)


def test_factory_raises_for_unsupported_voltage_type():
    with pytest.raises(ValueError, match="Unsupported voltage type: 132kV"):
        get_testsheet_parser(voltage_type="132kV")


def test_factory_raises_when_no_voltage_and_no_workbook():
    with pytest.raises(
        ValueError,
        match="Either voltage_type or workbook_path must be provided to resolve testsheet parser",
    ):
        get_testsheet_parser()


def test_factory_resolves_by_workbook_sheet_inspection(tmp_path: Path):
    # 1. 11kV workbook with "PCE Testsheet"
    wb_11 = openpyxl.Workbook()
    wb_11.active.title = "PCE Testsheet"
    path_11 = tmp_path / "test_11kv.xlsx"
    wb_11.save(path_11)
    wb_11.close()

    parser_11 = get_testsheet_parser(workbook_path=path_11)
    assert isinstance(parser_11, PceTestsheetParser)

    # 2. 33kV workbook with "33kV (Main)"
    wb_33_main = openpyxl.Workbook()
    wb_33_main.active.title = "33kV (Main)"
    path_33_main = tmp_path / "test_33kv_main.xlsx"
    wb_33_main.save(path_33_main)
    wb_33_main.close()

    parser_33_main = get_testsheet_parser(workbook_path=path_33_main)
    assert isinstance(parser_33_main, PpuTestsheetParser)

    # 3. 33kV workbook with "33kV (Extra)"
    wb_33_extra = openpyxl.Workbook()
    wb_33_extra.active.title = "33kV (Extra)"
    path_33_extra = tmp_path / "test_33kv_extra.xlsx"
    wb_33_extra.save(path_33_extra)
    wb_33_extra.close()

    parser_33_extra = get_testsheet_parser(workbook_path=path_33_extra)
    assert isinstance(parser_33_extra, PpuTestsheetParser)


def test_factory_raises_for_unrecognized_workbook_sheets(tmp_path: Path):
    wb_unrecognized = openpyxl.Workbook()
    wb_unrecognized.active.title = "RandomSheet"
    path_unrecognized = tmp_path / "unrecognized.xlsx"
    wb_unrecognized.save(path_unrecognized)
    wb_unrecognized.close()

    with pytest.raises(ValueError, match="Unrecognized testsheet format in .* missing required sheets"):
        get_testsheet_parser(workbook_path=path_unrecognized)


def test_ppu_parser_can_handle_sheets():
    parser = PpuTestsheetParser()
    assert parser.can_handle_sheets(["33kV (Main)", "Sheet1"]) is True
    assert parser.can_handle_sheets(["33kV (Extra)"]) is True
    assert parser.can_handle_sheets(["PCE Testsheet"]) is False
    assert parser.can_handle_sheets(["Other"]) is False


def test_ppu_parser_placeholder_raises_not_implemented(tmp_path: Path):
    parser = PpuTestsheetParser()
    dummy_path = tmp_path / "dummy.xlsx"

    with pytest.raises(NotImplementedError, match="PpuTestsheetParser will be implemented in Ticket #92"):
        parser.extract_testsheet_data(dummy_path)

    with pytest.raises(NotImplementedError, match="PpuTestsheetParser will be implemented in Ticket #92"):
        parser.extract_testsheet_metadata(dummy_path)

    with pytest.raises(NotImplementedError, match="PpuTestsheetParser will be implemented in Ticket #92"):
        parser.extract_photo_ranges(dummy_path)


def test_pce_parser_can_handle_sheets():
    parser = PceTestsheetParser()
    assert parser.can_handle_sheets(["PCE Testsheet", "PCE VI"]) is True
    assert parser.can_handle_sheets(["33kV (Main)"]) is False


def test_extractor_delegates_to_strategy(tmp_path: Path):
    # Dummy mock strategy
    class MockStrategy:
        def __init__(self):
            self.called_data = False
            self.called_meta = False
            self.called_ranges = False

        def extract_testsheet_data(self, workbook_path, station_hint="", date_hint=""):
            self.called_data = True
            return TestsheetData(substation_number=99, station_name="MOCK", substation_name_erms="MOCK")

        def extract_testsheet_metadata(self, workbook_path, station_hint="", date_hint=""):
            self.called_meta = True
            return TestsheetData(substation_number=99, station_name="MOCK_META", substation_name_erms="MOCK_META")

        def extract_photo_ranges(self, workbook_path):
            self.called_ranges = True
            return RawPhotoRanges()

        def can_handle_sheets(self, sheet_names):
            return True

    mock_strat = MockStrategy()
    extractor = TestsheetExtractor(strategy=mock_strat)

    data = extractor.extract_testsheet_data("dummy.xlsx")
    assert data.substation_number == 99
    assert mock_strat.called_data is True

    meta = extractor.extract_testsheet_metadata("dummy.xlsx")
    assert meta.station_name == "MOCK_META"
    assert mock_strat.called_meta is True

    ranges = extractor.extract_photo_ranges("dummy.xlsx")
    assert isinstance(ranges, RawPhotoRanges)
    assert mock_strat.called_ranges is True


def test_extractor_dynamic_resolution_via_workbook(tmp_path: Path):
    wb = openpyxl.Workbook()
    wb.active.title = "PCE Testsheet"
    path = tmp_path / "1. PE TEST.xlsx"
    wb.save(path)
    wb.close()

    extractor = TestsheetExtractor()
    data = extractor.extract_testsheet_data(path)
    assert isinstance(data, TestsheetData)
    assert data.substation_number == 1


def test_repository_accepts_voltage_type():
    repo = SubstationTestsheetRepository(voltage_type="33kV")
    assert repo.extractor.voltage_type == "33kV"
