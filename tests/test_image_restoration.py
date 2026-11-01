"""Behavior tests for Source Image Restoration (ticket #146, ADR 0009).

Seam: ``restore_source_images(docx_bytes, native_pdf, output_pdf) -> RestorationReport`` plus
``verify_invariance`` (the gate). Fixtures are synthetic (python-docx + PyMuPDF); no client data.
"""

from __future__ import annotations

import io
import random
import struct
from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from PIL import Image, ImageDraw

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz

from src.postprocessing import image_restoration as ir
from src.postprocessing.image_restoration import (
    RestorationError,
    restore_source_images,
    verify_invariance,
)

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


# ---------------------------------------------------------------- fixture helpers


def graph(seed: int, size=(600, 400)) -> Image.Image:
    rnd = random.Random(seed)
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    n = 20
    bw = size[0] // n
    for i in range(n):
        h = rnd.randint(40, int(size[1] * 0.7))
        col = (rnd.randint(0, 200), rnd.randint(0, 200), rnd.randint(0, 200))
        d.rectangle([i * bw + 2, size[1] - h, (i + 1) * bw - 2, size[1]], fill=col)
    return im


def png_bytes(im: Image.Image) -> bytes:
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def jpeg_bytes(im: Image.Image, q=60) -> bytes:
    b = io.BytesIO()
    im.convert("RGB").save(b, "JPEG", quality=q)
    return b.getvalue()


def docx_with_pictures(pictures: list[bytes]) -> bytes:
    doc = Document()
    for p in pictures:
        doc.add_picture(io.BytesIO(p))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_with_srccrop(picture: bytes, l=0, t=0, r=0, b=0) -> bytes:
    doc = Document()
    doc.add_picture(io.BytesIO(picture))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    blip_fill = doc.element.body.xpath(".//pic:blipFill")[0]
    blip = blip_fill.xpath("./a:blip")[0]
    src = parse_xml(
        f'<a:srcRect xmlns:a="{ns["a"]}" l="{l}" t="{t}" r="{r}" b="{b}"/>'
    )
    blip.addnext(src)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_with_vml_crop(picture: bytes, **crop: str) -> bytes:
    doc = Document()
    rid, _ = doc.part.get_or_add_image(io.BytesIO(picture))
    attrs = " ".join(f'{k}="{v}"' for k, v in crop.items())
    run = parse_xml(
        f'<w:r {NS}><w:pict><v:shape style="width:100pt;height:80pt">'
        f'<v:imagedata r:id="{rid}" {attrs}/></v:shape></w:pict></w:r>'
    )
    doc.add_paragraph()._p.append(run)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def wmf_with_dib(im: Image.Image) -> bytes:
    """Fake placeable-WMF blob carrying one 24-bit uncompressed DIB (bottom-up)."""
    w, h = im.size
    stride = ((w * 24 + 31) // 32) * 4
    rows = []
    rgb = im.convert("RGB")
    for y in range(h - 1, -1, -1):
        row = bytearray()
        for x in range(w):
            r, g, b = rgb.getpixel((x, y))
            row += bytes((b, g, r))
        row += b"\x00" * (stride - len(row))
        rows.append(bytes(row))
    hdr = struct.pack("<IiiHHIIiiII", 40, w, h, 1, 24, 0, stride * h, 0, 0, 0, 0)
    return b"\xd7\xcd\xc6\x9a" + b"\x00" * 22 + hdr + b"".join(rows)


def docx_with_wmf(wmf: bytes) -> bytes:
    doc = Document()
    part = Part(PackURI("/word/media/obj1.wmf"), "image/x-wmf", wmf, doc.part.package)
    rid = doc.part.relate_to(part, RT.IMAGE)
    run = parse_xml(
        f'<w:r {NS}><w:pict><v:shape style="width:100pt;height:80pt">'
        f'<v:imagedata r:id="{rid}"/></v:shape></w:pict></w:r>'
    )
    doc.add_paragraph()._p.append(run)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def native_pdf(path: Path, images: list[bytes], link: str | None = "https://example.com/x") -> Path:
    """A PDF as Word would have exported: pictures embedded at the given (already downsampled) bytes."""
    doc = fitz.open()
    for i, data in enumerate(images):
        page = doc.new_page(width=595, height=842)
        page.insert_image(fitz.Rect(50, 100, 350, 300), stream=data)
        page.insert_text((50, 60), f"Page {i + 1} text")
        if link and i == 0:
            page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(50, 40, 150, 70), "uri": link})
    doc.save(path)
    doc.close()
    return path


def pdf_images(path: Path) -> list[tuple[int, int, str]]:
    out = []
    with fitz.open(path) as doc:
        seen = set()
        smasks = {im[1] for p in doc for im in p.get_images(full=True) if im[1]}
        for p in doc:
            for im in p.get_images(full=True):
                if im[0] in seen or im[0] in smasks:
                    continue
                seen.add(im[0])
                out.append((im[2], im[3], doc.xref_get_key(im[0], "Filter")[1]))
    return out


def pdf_pixels(path: Path, index: int = 0) -> Image.Image:
    with fitz.open(path) as doc:
        xrefs = []
        for p in doc:
            for im in p.get_images(full=True):
                if im[0] not in xrefs:
                    xrefs.append(im[0])
        pix = fitz.Pixmap(doc, xrefs[index])
        if pix.alpha:
            pix = fitz.Pixmap(pix, 0)
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def mean_abs_diff(a: Image.Image, b: Image.Image) -> float:
    assert a.size == b.size
    pa, pb = a.convert("RGB").tobytes(), b.convert("RGB").tobytes()
    return sum(abs(x - y) for x, y in zip(pa, pb)) / len(pa)


@pytest.fixture
def paths(tmp_path: Path):
    return tmp_path / "native.pdf", tmp_path / "out.pdf"


def leftovers(dir_: Path) -> list[Path]:
    return [p for p in dir_.iterdir() if p.name.startswith(".tmp_restore_")]


# ---------------------------------------------------------------- restoration behavior


def test_downsampled_png_is_restored_to_source_pixels(paths):
    native, out = paths
    src = graph(1)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    report = restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert report.restored == 1 and report.images_total == 1
    assert pdf_images(out)[0][:2] == (600, 400)
    assert mean_abs_diff(pdf_pixels(out), src) == 0
    assert report.entries[0].action == "restored"


def test_jpeg_reencoded_lossless_source_is_restored_losslessly(paths):
    native, out = paths
    src = graph(2)
    native_pdf(native, [jpeg_bytes(src, q=40)])
    assert "DCT" in pdf_images(native)[0][2]
    report = restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert report.restored == 1
    assert "DCT" not in pdf_images(out)[0][2]
    assert mean_abs_diff(pdf_pixels(out), src) == 0


def test_drawingml_cropped_source_is_restored_to_cropped_region(paths):
    native, out = paths
    src = graph(3)
    cropped = src.crop((60, 40, 540, 360))  # l=10% t=10% r=10% b=10%
    native_pdf(native, [png_bytes(cropped.resize((240, 160), Image.LANCZOS))])
    docx = docx_with_srccrop(png_bytes(src), l=10000, t=10000, r=10000, b=10000)
    report = restore_source_images(docx, native, out)
    assert report.restored == 1
    assert pdf_images(out)[0][:2] == (480, 320)
    assert mean_abs_diff(pdf_pixels(out), cropped) == 0


def test_vml_cropped_source_is_restored_to_cropped_region(paths):
    native, out = paths
    src = graph(4)
    cropped = src.crop((60, 40, 540, 360))
    native_pdf(native, [png_bytes(cropped.resize((240, 160), Image.LANCZOS))])
    docx = docx_with_vml_crop(
        png_bytes(src), croptop="6554f", cropbottom="6554f", cropleft="6554f", cropright="6554f"
    )
    report = restore_source_images(docx, native, out)
    assert report.restored == 1
    w, h = pdf_images(out)[0][:2]
    assert abs(w - 480) <= 2 and abs(h - 320) <= 2


def test_dib_inside_wmf_is_restored(paths):
    native, out = paths
    src = graph(5, size=(320, 240))
    native_pdf(native, [jpeg_bytes(src.resize((160, 120), Image.LANCZOS), q=80)])
    report = restore_source_images(docx_with_wmf(wmf_with_dib(src)), native, out)
    assert report.restored == 1
    assert pdf_images(out)[0][:2] == (320, 240)
    assert mean_abs_diff(pdf_pixels(out), src) == 0


def test_look_alike_graphs_each_map_to_their_own_source(paths):
    native, out = paths
    srcs = [graph(s) for s in (11, 12, 13)]
    native_pdf(native, [png_bytes(s.resize((300, 200), Image.LANCZOS)) for s in srcs])
    # DOCX lists the sources in a different order than the PDF pages
    docx = docx_with_pictures([png_bytes(srcs[2]), png_bytes(srcs[0]), png_bytes(srcs[1])])
    report = restore_source_images(docx, native, out)
    assert report.restored == 3
    for i, s in enumerate(srcs):
        assert mean_abs_diff(pdf_pixels(out, i), s) == 0


def test_ambiguous_match_is_skipped(paths):
    native, out = paths
    a = graph(21)
    b = a.copy()
    ImageDraw.Draw(b).rectangle([10, 10, 69, 49], fill="black")  # ~1% strongly different
    native_pdf(native, [png_bytes(a.resize((300, 200), Image.LANCZOS))])
    report = restore_source_images(docx_with_pictures([png_bytes(a), png_bytes(b)]), native, out)
    assert report.ambiguous == 1 and report.restored == 0
    assert pdf_images(out)[0][:2] == (300, 200)


def test_unrelated_image_is_unmatched(paths):
    native, out = paths
    native_pdf(native, [png_bytes(graph(31).resize((300, 200), Image.LANCZOS))])
    report = restore_source_images(docx_with_pictures([png_bytes(graph(32))]), native, out)
    assert report.unmatched == 1 and report.restored == 0


def test_smaller_source_never_replaces_larger_pdf_image(paths):
    native, out = paths
    small = graph(41, size=(150, 100))
    big = small.resize((600, 400), Image.LANCZOS)
    native_pdf(native, [png_bytes(big)])
    report = restore_source_images(docx_with_pictures([png_bytes(small)]), native, out)
    assert report.restored == 0
    assert pdf_images(out)[0][:2] == (600, 400)


def test_alpha_is_preserved(paths):
    native, out = paths
    src = graph(51).convert("RGBA")
    ImageDraw.Draw(src).rectangle([0, 0, 9, 9], fill=(255, 255, 255, 0))
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    report = restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert report.restored == 1
    with fitz.open(out) as doc:
        im = doc[0].get_images(full=True)[0]
        assert im[1] != 0  # SMask present
        assert (im[2], im[3]) == (600, 400)
        mask = fitz.Pixmap(doc, im[1])
        assert (mask.width, mask.height) == (600, 400)
        assert mask.samples[0] == 0 and mask.samples[mask.width * 200 + 300] == 255


def test_self_check_rejects_payload_that_does_not_match(paths, monkeypatch):
    native, out = paths
    src = graph(61)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    # Private seam on purpose: a payload that disagrees with Word's pixels cannot be produced through the
    # public interface, because the encoder derives it from the very source that was matched.
    monkeypatch.setattr(ir, "_encode_payload", lambda source, pixels: png_bytes(graph(99)))
    report = restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert report.self_check_rejected == 1 and report.restored == 0
    assert pdf_images(out)[0][:2] == (300, 200)


def test_restoring_a_restored_pdf_is_a_no_op(paths, tmp_path):
    native, out = paths
    src = graph(71)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    docx = docx_with_pictures([png_bytes(src)])
    restore_source_images(docx, native, out)
    again = restore_source_images(docx, out, tmp_path / "again.pdf")
    assert again.restored == 0


def test_cropped_jpeg_is_reencoded_as_jpeg_not_lossless(paths):
    native, out = paths
    src = graph(81)
    jpg = jpeg_bytes(src, q=90)
    native_pdf(native, [jpeg_bytes(Image.open(io.BytesIO(jpg)).crop((60, 40, 540, 360)).resize((240, 160)))])
    docx = docx_with_srccrop(jpg, l=10000, t=10000, r=10000, b=10000)
    report = restore_source_images(docx, native, out)
    assert report.restored == 1
    assert "DCT" in pdf_images(out)[0][2]
    assert pdf_images(out)[0][:2] == (480, 320)


# ---------------------------------------------------------------- safety


def test_inputs_are_never_modified_and_no_temp_left(paths):
    native, out = paths
    src = graph(91)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    before = native.read_bytes()
    restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert native.read_bytes() == before
    assert leftovers(out.parent) == []


def test_corrupt_docx_snapshot_raises_and_leaves_no_files(paths):
    native, out = paths
    native_pdf(native, [png_bytes(graph(1).resize((300, 200)))])
    with pytest.raises(RestorationError):
        restore_source_images(b"not a zip", native, out)
    assert not out.exists()
    assert leftovers(out.parent) == []


def test_gate_failure_raises_and_leaves_no_files(paths, monkeypatch):
    native, out = paths
    src = graph(2)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])

    def boom(*_a, **_k):
        raise RestorationError("tampered")

    monkeypatch.setattr(ir, "verify_invariance", boom)
    with pytest.raises(RestorationError):
        restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert not out.exists()
    assert leftovers(out.parent) == []


def test_invariance_gate_accepts_identical_and_rejects_tampered(tmp_path):
    a = native_pdf(tmp_path / "a.pdf", [png_bytes(graph(1))])
    same = tmp_path / "same.pdf"
    same.write_bytes(a.read_bytes())
    verify_invariance(a, same)  # no raise

    def tampered(name: str, mutate) -> Path:
        p = tmp_path / name
        with fitz.open(a) as d:
            mutate(d)
            d.save(p)
        return p

    cases = {
        "text": lambda d: d[0].insert_text((50, 400), "extra"),
        "link": lambda d: d[0].insert_link(
            {"kind": fitz.LINK_URI, "from": fitz.Rect(1, 1, 20, 20), "uri": "https://evil"}
        ),
        "annot": lambda d: d[0].add_text_annot((300, 400), "note"),
        "page": lambda d: d.new_page(),
        "placement": lambda d: d[0].insert_image(fitz.Rect(400, 400, 500, 500), stream=png_bytes(graph(2))),
    }
    for name, mutate in cases.items():
        with pytest.raises(RestorationError):
            verify_invariance(a, tampered(f"{name}.pdf", mutate))


# ---------------------------------------------------------------- review follow-ups


def test_smaller_lossless_source_never_replaces_larger_jpeg_pdf_image(paths):
    native, out = paths
    small = graph(101, size=(150, 100))
    native_pdf(native, [jpeg_bytes(small.resize((600, 400), Image.LANCZOS), q=70)])
    report = restore_source_images(docx_with_pictures([png_bytes(small)]), native, out)
    assert report.restored == 0
    assert pdf_images(out)[0][:2] == (600, 400)
    assert "DCT" in pdf_images(out)[0][2]


def test_idempotent_for_lossless_over_jpeg_and_cropped_jpeg(paths, tmp_path):
    native, out = paths
    src = graph(111)
    native_pdf(native, [jpeg_bytes(src, q=40)])
    docx = docx_with_pictures([png_bytes(src)])
    assert restore_source_images(docx, native, out).restored == 1
    assert restore_source_images(docx, out, tmp_path / "again.pdf").restored == 0

    jpg = jpeg_bytes(graph(112), q=90)
    crop = Image.open(io.BytesIO(jpg)).crop((60, 40, 540, 360))
    native2, out2 = tmp_path / "n2.pdf", tmp_path / "o2.pdf"
    native_pdf(native2, [jpeg_bytes(crop.resize((240, 160)))])
    docx2 = docx_with_srccrop(jpg, l=10000, t=10000, r=10000, b=10000)
    assert restore_source_images(docx2, native2, out2).restored == 1
    assert restore_source_images(docx2, out2, tmp_path / "again2.pdf").restored == 0


def test_failure_while_rewriting_an_xobject_raises_instead_of_reporting_unmatched(paths, monkeypatch):
    native, out = paths
    src = graph(121)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    real = fitz.Document.xref_set_key
    calls = {"n": 0}

    def flaky(self, xref, key, value):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("disk on fire")
        return real(self, xref, key, value)

    monkeypatch.setattr(fitz.Document, "xref_set_key", flaky)
    with pytest.raises(RestorationError):
        restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert not out.exists()
    assert leftovers(out.parent) == []


def test_invariance_gate_rejects_changed_internal_link(tmp_path):
    def build(name: str, target: int) -> Path:
        d = fitz.open()
        d.new_page()
        d.new_page()
        d[0].insert_link({"kind": fitz.LINK_GOTO, "from": fitz.Rect(10, 10, 60, 40), "page": target})
        p = tmp_path / name
        d.save(p)
        d.close()
        return p

    with pytest.raises(RestorationError):
        verify_invariance(build("a.pdf", 1), build("b.pdf", 0))


def test_opaque_source_removes_old_smask(paths):
    native, out = paths
    src = graph(131)
    rgba = src.resize((300, 200), Image.LANCZOS).convert("RGBA")
    native_pdf(native, [png_bytes(rgba)])
    with fitz.open(native) as d:
        assert d[0].get_images(full=True)[0][1] != 0  # Word-style copy carries an SMask
    assert restore_source_images(docx_with_pictures([png_bytes(src)]), native, out).restored == 1
    with fitz.open(out) as d:
        assert d[0].get_images(full=True)[0][1] == 0


def test_same_xref_on_two_pages_is_rewritten_once_and_smasks_are_not_counted(paths):
    native, out = paths
    src = graph(141).convert("RGBA")
    data = png_bytes(src.resize((300, 200), Image.LANCZOS))
    doc = fitz.open()
    first = doc.new_page()
    xref = first.insert_image(fitz.Rect(50, 100, 350, 300), stream=data)
    second = doc.new_page()
    second.insert_image(fitz.Rect(50, 100, 350, 300), xref=xref)
    doc.save(native)
    doc.close()
    report = restore_source_images(docx_with_pictures([png_bytes(src)]), native, out)
    assert report.images_total == 1 and report.restored == 1  # SMask not counted, shared xref once


def test_negative_crop_and_external_link_are_skipped(paths):
    native, out = paths
    src = graph(151)
    native_pdf(native, [png_bytes(src.resize((300, 200), Image.LANCZOS))])
    docx = docx_with_srccrop(png_bytes(src), l=-5000)  # padding crop: not representable
    assert restore_source_images(docx, native, out).unmatched == 1
    assert ir.harvest_sources(docx) == []


def test_header_pictures_are_harvested():
    doc = Document()
    hdr = doc.sections[0].header
    hdr.paragraphs[0].add_run().add_picture(io.BytesIO(png_bytes(graph(161))))
    out = io.BytesIO()
    doc.save(out)
    assert len(ir.harvest_sources(out.getvalue())) == 1
