"""Tests for Full Report normal component scanning templates (Ticket #29 / T4.1).

Validates transparent Jinja2 templates under templates/FULL REPORT/NORMAL IR US TEV/
ensuring placeholder consistency, inline image binding, and decoupling from Quick Report.
"""

from __future__ import annotations

from pathlib import Path
import tempfile
from typing import Any
import zipfile

from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm
import docx
from PIL import Image
import pytest

from src.quick_report.cbm_render import _build_jinja_env

TEMPLATES_DIR = Path("templates/FULL REPORT/NORMAL IR US TEV")
QUICK_REPORT_DIR = Path("templates/QUICK REPORT")

EXPECTED_TEMPLATES: tuple[str, ...] = (
    "swg-overview.docx",
    "swg-panel.docx",
    "tx-overview.docx",
    "tx-hv-sides.docx",
    "tx-lv-sides.docx",
    "fp-overview.docx",
    "battery-overview.docx",
)

# Declarative namespace expectations per template matching canonical Quick Report templates
REQUIRED_NAMESPACES_MAP: dict[str, set[str]] = {
    "swg-overview.docx": {"substation", "swg", "visual"},
    "swg-panel.docx": {"substation", "swg", "panel", "visual", "ir", "us", "tev", "analysis", "recommendation"},
    "tx-overview.docx": {"substation", "tx", "visual"},
    "tx-hv-sides.docx": {"substation", "tx", "visual", "ir", "us", "analysis", "recommendation"},
    "tx-lv-sides.docx": {"substation", "tx", "visual", "ir", "analysis", "recommendation"},
    "fp-overview.docx": {"substation", "fp", "visual"},
    "battery-overview.docx": {"substation", "batt", "visual", "ir", "analysis", "recommendation"},
}

# Expected drawing count when fully bound with mock images (visual.image, us.prpd, tev.prpd).
# IR is an embedded ActiveX <w:object> control rather than an inline DrawingML drawing.
EXPECTED_DRAWINGS_MAP: dict[str, int] = {
    "swg-overview.docx": 1,
    "swg-panel.docx": 3,      # VI + US PRPD + TEV PRPD
    "tx-overview.docx": 1,
    "tx-hv-sides.docx": 2,    # VI + US PRPD (no TEV on TX)
    "tx-lv-sides.docx": 1,
    "fp-overview.docx": 1,
    "battery-overview.docx": 1,
}

# Declarative specifications for restored FLIR ActiveX CIRViewer OLE controls
ACTIVEX_SPEC_MAP: dict[str, dict[str, str]] = {
    "swg-overview.docx": {
        "control_name": "CIRViewer121111",
        "shape_id": "_x0000_i1027",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "swg-panel.docx": {
        "control_name": "CIRViewer1211111131211",
        "shape_id": "_x0000_i1029",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "tx-overview.docx": {
        "control_name": "CIRViewer1211111131",
        "shape_id": "_x0000_i1029",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "tx-hv-sides.docx": {
        "control_name": "CIRViewer1211111131",
        "shape_id": "_x0000_i1029",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "tx-lv-sides.docx": {
        "control_name": "CIRViewer1211111131",
        "shape_id": "_x0000_i1029",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "fp-overview.docx": {
        "control_name": "CIRViewer121111113121",
        "shape_id": "_x0000_i1027",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
    "battery-overview.docx": {
        "control_name": "CIRViewer121111113121111",
        "shape_id": "_x0000_i1030",
        "control_rel": "rId8",
        "image_rel": "rId7",
    },
}


@pytest.fixture(scope="module")
def dummy_image_path() -> Path:
    """Create a temporary test image file for docxtpl inline image binding."""
    tmp = Path(tempfile.gettempdir()) / "test_normal_template_dummy.png"
    img = Image.new("RGB", (120, 90), color=(0, 176, 80))
    img.save(tmp)
    return tmp


def _build_full_mock_context(dt: DocxTemplate, image_path: Path) -> dict[str, Any]:
    """Construct a full mock context satisfying all 7 template namespaces."""
    inline_img = InlineImage(dt, str(image_path), width=Mm(40))
    return {
        "substation": {
            "name_erms": "PE TEST STATION",
            "date": "01-01-2026",
            "time": "10:30 AM",
            "ambient": "30.5 °C",
            "humidity": "60%",
        },
        "swg": {
            "type": "RMU SF6",
            "manufacturer": "TAMCO",
            "model": "AIR",
            "rating": "12kV",
            "area": "OVERVIEW",
            "serialnumber": "SWG-12345",
        },
        "panel": {
            "name": "INCOMING 1",
            "linknumber": "1",
            "area": "CABLE COMPARTMENT",
            "serialnumber": "PANEL-001",
            "heateramp": "0.5A",
            "breakerstatus": "CLOSE",
            "busbarposition": "TOP",
            "cabletype": "3C 300mm2 XLPE",
            "loadamp": "250A",
        },
        "tx": {
            "manufacturer": "SGB",
            "model": "HERMETICALLY SEAL",
            "rating": "1000kVA",
            "number": "TX1",
            "location": "HV SIDE",
            "area": "HV BUSHING",
            "serialnumber": "TX-99999",
            "cabletype": "3C 300mm2 XLPE",
        },
        "fp": {
            "labelsource": "LVDB TX1",
            "manufacturer": "ALGEBRA",
            "model": "J-SLOTTED",
            "rating": "1600A",
            "area": "OVERVIEW",
            "serialnumber": "FP-55555",
            "cabletype": "4C 300mm2 PVC",
            "feederno": "F01",
            "loadamp": "150A",
        },
        "battery": {
            "manufacturer": "SUNPOWER",
            "model": "SP-100",
            "number": "1",
            "serialnumber": "BATT-777",
            "rating": "240V/50Hz",
            "area": "OVERVIEW",
        },
        "batt": {
            "manufacturer": "SUNPOWER",
            "model": "SP-100",
            "number": "1",
            "serialnumber": "BATT-777",
            "rating": "240V/50Hz",
            "area": "OVERVIEW",
        },
        "ir": {
            "image": inline_img,
            "sp1": "31.5 °C",
            "sp2": "31.0 °C",
            "ar1": "31.2 °C",
            "reading": "31.5 °C",
            "delta_t": "0.5 °C",
            "severity": "NORMAL",
        },
        "us": {
            "reading": "12",
            "char": "NORMAL",
            "severity": "NORMAL",
            "prpd": inline_img,
        },
        "tev": {
            "bg": "4",
            "reading": "6",
            "ppc": "0.1",
            "severity": "NORMAL",
            "prpd": inline_img,
        },
        "visual": {
            "image": inline_img,
        },
        "banner": {
            "analysis": "No Anomaly.",
            "recommendation": "-",
        },
        "analysis": "No Anomaly.",
        "recommendation": "-",
    }


def _read_document_xml(docx_path: Path) -> str:
    """Read word/document.xml content directly from .docx zip archive."""
    with zipfile.ZipFile(docx_path, "r") as z:
        return z.read("word/document.xml").decode("utf-8")


@pytest.mark.parametrize("filename", EXPECTED_TEMPLATES)
def test_all_seven_templates_exist_and_decoupled(filename: str):
    """Confirms all 7 templates exist in NORMAL IR US TEV/ and are decoupled from QUICK REPORT/."""
    template_file = TEMPLATES_DIR / filename
    assert template_file.is_file(), f"Missing template file: {filename}"

    # Verify physical file decoupling from QUICK REPORT
    quick_report_file = QUICK_REPORT_DIR / "DEFECT IR US TEV" / filename
    assert template_file.resolve() != quick_report_file.resolve(), (
        f"{filename} must not be a symlink or share the same path as QUICK REPORT"
    )


@pytest.mark.parametrize("filename", EXPECTED_TEMPLATES)
def test_template_placeholders_consistency(filename: str):
    """Verifies all 7 templates have consistent Jinja placeholders for metadata, photos, parameters, and US/TEV."""
    template_path = TEMPLATES_DIR / filename
    doc = DocxTemplate(template_path)
    undeclared = doc.get_undeclared_template_variables()

    # 1. Verify required namespaces match canonical Quick Report template definitions
    expected_namespaces = REQUIRED_NAMESPACES_MAP[filename]
    assert expected_namespaces.issubset(undeclared), (
        f"Template {filename} missing expected namespaces: {expected_namespaces - undeclared}. Found: {sorted(undeclared)}"
    )

    # 2. Leaf placeholder and ActiveX container verification via document XML
    xml_content = _read_document_xml(template_path)
    assert "visual.image" in xml_content, f"Template {filename} missing visual.image placeholder"
    assert "<w:control" in xml_content, f"Template {filename} missing <w:control> ActiveX element"
    assert "<w:object" in xml_content, f"Template {filename} missing <w:object> container"
    assert "ir.image" not in xml_content, f"Template {filename} should not contain raw ir.image placeholder"
    assert "substation.name_erms" in xml_content, f"Template {filename} missing substation.name_erms"
    assert "substation.date" in xml_content, f"Template {filename} missing substation.date"

    if filename in ("swg-panel.docx", "tx-hv-sides.docx", "tx-lv-sides.docx", "battery-overview.docx"):
        assert "ir.severity" in xml_content, f"Template {filename} missing ir.severity"
        assert "analysis" in xml_content, f"Template {filename} missing analysis placeholder"
        assert "recommendation" in xml_content, f"Template {filename} missing recommendation placeholder"

    # Exact ActiveX specification verification per template
    spec = ACTIVEX_SPEC_MAP[filename]
    assert f'w:name="{spec["control_name"]}"' in xml_content, (
        f"Template {filename} missing exact control name {spec['control_name']}"
    )
    assert f'w:shapeid="{spec["shape_id"]}"' in xml_content, (
        f"Template {filename} missing exact shapeid {spec['shape_id']}"
    )
    assert f'r:id="{spec["control_rel"]}"' in xml_content, (
        f"Template {filename} missing {spec['control_rel']} control relationship"
    )
    assert f'r:id="{spec["image_rel"]}"' in xml_content, (
        f"Template {filename} missing {spec['image_rel']} image relationship"
    )


@pytest.mark.parametrize("filename", EXPECTED_TEMPLATES)
def test_template_inline_image_binding(filename: str, tmp_path: Path, dummy_image_path: Path):
    """Verifies inline image placeholders (visual.image, us.prpd, tev.prpd) bind InlineImage cleanly and ActiveX survives."""
    template_path = TEMPLATES_DIR / filename
    dt = DocxTemplate(template_path)
    ctx = _build_full_mock_context(dt, dummy_image_path)

    dt.render(ctx, jinja_env=_build_jinja_env())
    output_file = tmp_path / f"test_img_{filename}"
    dt.save(output_file)

    rendered_doc = docx.Document(output_file)
    drawings = rendered_doc._body._element.xpath(".//w:drawing")
    expected_count = EXPECTED_DRAWINGS_MAP[filename]
    assert len(drawings) >= expected_count, (
        f"{filename} expected at least {expected_count} image drawings, found {len(drawings)}"
    )
    spec = ACTIVEX_SPEC_MAP[filename]
    controls = rendered_doc._body._element.xpath(".//w:control")
    assert len(controls) == 1, f"{filename} missing ActiveX <w:control> after render"
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['w']}}}name") == spec["control_name"]
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['w']}}}shapeid") == spec["shape_id"]
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['r']}}}id") == spec["control_rel"]

    objects = rendered_doc._body._element.xpath(".//w:object")
    assert len(objects) == 1, f"{filename} missing ActiveX <w:object> after render"


@pytest.mark.parametrize("filename", EXPECTED_TEMPLATES)
def test_render_smoke_test_all_seven_templates(filename: str, tmp_path: Path, dummy_image_path: Path):
    """Render smoke test validating that docxtpl compiles each template cleanly with mock context payloads."""
    template_path = TEMPLATES_DIR / filename
    dt = DocxTemplate(template_path)
    ctx = _build_full_mock_context(dt, dummy_image_path)

    dt.render(ctx, jinja_env=_build_jinja_env())
    output_file = tmp_path / f"smoke_{filename}"
    dt.save(output_file)

    assert output_file.is_file()
    rendered_doc = docx.Document(output_file)

    # Inspect all paragraphs and all tables in the document
    para_text = "\n".join(p.text for p in rendered_doc.paragraphs)
    table_text = "\n".join(c.text.strip() for t in rendered_doc.tables for r in t.rows for c in r.cells)
    full_text = f"{para_text}\n{table_text}"

    # No raw unrendered Jinja brackets
    assert "{{" not in full_text, f"Unrendered Jinja tag in {filename}: {full_text}"
    assert "}}" not in full_text, f"Unrendered Jinja tag in {filename}: {full_text}"

    # Key metadata and banner rendered cleanly
    assert "PE TEST STATION" in full_text
    assert "01-01-2026" in full_text
    if filename in ("swg-panel.docx", "tx-hv-sides.docx", "tx-lv-sides.docx", "battery-overview.docx"):
        assert "No Anomaly." in full_text
        assert "NORMAL" in full_text

    # Verify FLIR ActiveX CIRViewer container and control are fully preserved
    spec = ACTIVEX_SPEC_MAP[filename]
    controls = rendered_doc._body._element.xpath(".//w:control")
    assert len(controls) == 1, f"{filename} missing ActiveX <w:control> after smoke render"
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['w']}}}name") == spec["control_name"]
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['w']}}}shapeid") == spec["shape_id"]
    assert controls[0].attrib.get(f"{{{controls[0].nsmap['r']}}}id") == spec["control_rel"]

    objects = rendered_doc._body._element.xpath(".//w:object")
    assert len(objects) == 1, f"{filename} missing ActiveX <w:object> after smoke render"
