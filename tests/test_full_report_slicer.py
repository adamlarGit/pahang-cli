"""Tests for Full Report DocumentSlicer core, FakeDocumentSlicer, and slicing protocol (Ticket #22 / T1.3a)."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch
import docx
import pytest

from src.full_report.slicer import (
    CbmDefectSliceMetadata,
    DocumentSlicer,
    FakeDocumentSlicer,
    ParagraphBoundary,
    SlicedSections,
    SlicingError,
    WordComDocumentSlicer,
    _find_paragraph,
    _replace_front_page_title,
    _slice_range_to_doc,
    get_temp_parts_dir,
    temp_parts_workspace,
)


def test_sliced_sections_dataclass():
    """Verify SlicedSections captures required and optional section paths."""
    front = Path("/tmp/front_page.docx")
    cond = Path("/tmp/condition_pages.docx")
    sticker = Path("/tmp/sticker_page.docx")
    vi_sum = Path("/tmp/vi_summary.docx")
    vi_def = Path("/tmp/vi_defect_pages.docx")

    sections = SlicedSections(
        station="TALAPIA",
        front_page=front,
        condition_pages=cond,
        sticker_page=sticker,
        vi_summary=vi_sum,
        vi_defect_pages=vi_def,
    )

    assert sections.station == "TALAPIA"
    assert sections.front_page == front
    assert sections.condition_pages == cond
    assert sections.sticker_page == sticker
    assert sections.vi_summary == vi_sum
    assert sections.vi_defect_pages == vi_def
    assert sections.cbm_defect_pages == ()


def test_paragraph_boundary_dataclass():
    """Verify ParagraphBoundary stores immutable character offsets and metadata."""
    boundary = ParagraphBoundary(start=100, end=150, text="SUBSTATION CONDITION\r", page_number=3)
    assert boundary.start == 100
    assert boundary.end == 150
    assert boundary.text == "SUBSTATION CONDITION\r"
    assert boundary.page_number == 3


def test_fake_document_slicer_conforms_to_protocol():
    """Verify FakeDocumentSlicer conforms to DocumentSlicer protocol."""
    slicer = FakeDocumentSlicer()
    assert isinstance(slicer, DocumentSlicer)


def test_fake_document_slicer_slice_sections_creates_stub_files(tmp_path: Path):
    """Verify FakeDocumentSlicer creates valid stub docx parts in target_dir."""
    source_doc = tmp_path / "source_quick_report.docx"
    source_doc.touch()
    target_dir = tmp_path / "temp_parts" / "TALAPIA"

    slicer = FakeDocumentSlicer()
    result = slicer.slice_sections(
        source_path=source_doc,
        target_dir=target_dir,
        station="TALAPIA",
        has_vi_summary=True,
        has_vi_defects=True,
    )

    assert isinstance(result, SlicedSections)
    assert result.station == "TALAPIA"
    assert result.front_page.exists()
    assert result.condition_pages.exists()
    assert result.sticker_page.exists()
    assert result.vi_summary is not None and result.vi_summary.exists()
    assert result.vi_defect_pages is not None and result.vi_defect_pages.exists()

    # Verify stub payloads are written
    assert result.front_page.read_bytes().startswith(b"PK")
    assert result.condition_pages.read_bytes().startswith(b"PK")
    assert result.sticker_page.read_bytes().startswith(b"PK")
    assert result.cbm_defect_pages == ()
    assert slicer.slice_cbm_defects(source_doc, target_dir) == ()


def test_fake_document_slicer_omits_optional_vi_sections_when_absent(tmp_path: Path):
    """Verify FakeDocumentSlicer omits vi_summary and vi_defect_pages when disabled."""
    source_doc = tmp_path / "source_quick_report.docx"
    source_doc.touch()
    target_dir = tmp_path / "temp_parts" / "CENDERAWASIH"

    slicer = FakeDocumentSlicer()
    result = slicer.slice_sections(
        source_path=source_doc,
        target_dir=target_dir,
        station="CENDERAWASIH",
        has_vi_summary=False,
        has_vi_defects=False,
    )

    assert result.vi_summary is None
    assert result.vi_defect_pages is None
    assert result.front_page.exists()
    assert result.condition_pages.exists()
    assert result.sticker_page.exists()


def test_get_temp_parts_dir(tmp_path: Path):
    """Verify get_temp_parts_dir resolves to .temp/temp_parts/<STATION>/."""
    res = get_temp_parts_dir("TALAPIA", base_dir=tmp_path)
    assert res == tmp_path / ".temp" / "temp_parts" / "TALAPIA"


def test_temp_parts_workspace_cleanup_by_default(tmp_path: Path):
    """Verify temp_parts_workspace removes station temp directory upon context exit."""
    station = "TALAPIA"
    with temp_parts_workspace(station=station, base_dir=tmp_path) as temp_dir:
        assert temp_dir.exists()
        assert temp_dir == tmp_path / ".temp" / "temp_parts" / station
        (temp_dir / "sample.docx").touch()
        assert (temp_dir / "sample.docx").exists()

    # After exit, station temp dir must be deleted
    assert not temp_dir.exists()


def test_temp_parts_workspace_keep_temp_preserves_files(tmp_path: Path):
    """Verify temp_parts_workspace preserves files when keep_temp=True."""
    station = "TALAPIA"
    with temp_parts_workspace(station=station, base_dir=tmp_path, keep_temp=True) as temp_dir:
        assert temp_dir.exists()
        sample_file = temp_dir / "preserved.docx"
        sample_file.touch()

    # Must still exist because keep_temp=True
    assert temp_dir.exists()
    assert sample_file.exists()


def test_temp_parts_workspace_cleans_up_on_exception(tmp_path: Path):
    """Verify temp_parts_workspace cleans up temp directory even if an exception occurs."""
    station = "TALAPIA"
    temp_dir_ref: Path | None = None
    with pytest.raises(RuntimeError, match="Simulated error during slicing"):
        with temp_parts_workspace(station=station, base_dir=tmp_path, keep_temp=False) as temp_dir:
            temp_dir_ref = temp_dir
            (temp_dir / "sample.docx").touch()
            raise RuntimeError("Simulated error during slicing")

    assert temp_dir_ref is not None
    assert not temp_dir_ref.exists()


def test_word_com_slicer_raises_when_win32com_missing(tmp_path: Path):
    """Verify WordComDocumentSlicer raises RuntimeError when win32com is missing."""
    slicer = WordComDocumentSlicer()
    p1 = tmp_path / "qr.docx"
    p1.touch()
    out = tmp_path / "temp_parts" / "TALAPIA"

    with patch("src.full_report.slicer.win32com", None):
        with pytest.raises(RuntimeError, match="win32com is required"):
            slicer.slice_sections(p1, out)


def test_word_com_slicer_session_reuses_word_app(tmp_path: Path):
    """Verify WordComDocumentSlicer reuses existing word_app without quitting it."""
    mock_word = MagicMock()
    slicer = WordComDocumentSlicer(word_app=mock_word)

    with slicer.session() as s:
        assert s._word_app is mock_word

    mock_word.Quit.assert_not_called()


def test_word_com_slicer_conforms_to_protocol():
    """Verify WordComDocumentSlicer conforms to DocumentSlicer protocol."""
    slicer = WordComDocumentSlicer(word_app=MagicMock())
    assert isinstance(slicer, DocumentSlicer)


def test_word_com_slicer_missing_substation_condition_raises_slicing_error(tmp_path: Path):
    """Verify WordComDocumentSlicer raises SlicingError when SUBSTATION CONDITION is missing."""
    p1 = tmp_path / "qr.docx"
    p1.touch()
    out = tmp_path / "temp_parts" / "TALAPIA"

    mock_word = MagicMock()
    mock_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_doc
    mock_doc.Content.Start = 0
    mock_doc.Content.End = 100
    mock_doc.ComputeStatistics.return_value = 2
    mock_doc.GoTo.return_value = MagicMock(Start=50)

    # Empty paragraphs list -> SUBSTATION CONDITION not found
    mock_doc.Paragraphs = []
    mock_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)
    with pytest.raises(SlicingError, match="SUBSTATION CONDITION"):
        slicer.slice_sections(p1, out)


def test_word_com_slicer_missing_sticker_page_raises_slicing_error(tmp_path: Path):
    """Verify WordComDocumentSlicer raises SlicingError when NORMAL/DEFECT STICKER is missing."""
    p1 = tmp_path / "qr.docx"
    p1.touch()
    out = tmp_path / "temp_parts" / "TALAPIA"

    mock_word = MagicMock()
    mock_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_doc
    mock_doc.Content.Start = 0
    mock_doc.Content.End = 100
    mock_doc.ComputeStatistics.return_value = 2
    mock_doc.GoTo.return_value = MagicMock(Start=50)

    # Only SUBSTATION CONDITION is present, STICKER is missing
    p_cond = MagicMock()
    p_cond.Range.Start = 60
    p_cond.Range.End = 80
    p_cond.Range.Text = "SUBSTATION CONDITION\r"
    p_cond.Range.Information.return_value = 2

    mock_doc.Paragraphs = [p_cond]
    mock_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)
    with pytest.raises(SlicingError, match="NORMAL/DEFECT STICKER"):
        slicer.slice_sections(p1, out)


def test_word_com_slicer_mock_slices_all_five_sections(tmp_path: Path):
    """Verify WordComDocumentSlicer boundary detection and slicing logic with mocked COM objects."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "TALAPIA"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_new_doc = MagicMock()

    mock_word.Documents.Open.return_value = mock_source_doc
    mock_word.Documents.Add.return_value = mock_new_doc

    mock_source_doc.Content.Start = 0
    mock_source_doc.Content.End = 500
    mock_source_doc.ComputeStatistics.return_value = 5

    # Page 2 start at 50
    mock_p2 = MagicMock(Start=50)
    # Page 3 start at 100
    mock_p3 = MagicMock(Start=100)
    mock_source_doc.GoTo.side_effect = lambda what, which, count: mock_p2 if count == 2 else mock_p3

    # Paragraph boundaries:
    p_visum = MagicMock()
    p_visum.Range.Start = 50
    p_visum.Range.End = 75
    p_visum.Range.Text = "VISUAL DEFECT SUMMARY\r"
    p_visum.Range.Information.return_value = 2

    p_cond = MagicMock()
    p_cond.Range.Start = 100
    p_cond.Range.End = 130
    p_cond.Range.Text = "SUBSTATION CONDITION\r"
    p_cond.Range.Information.return_value = 3

    p_vi = MagicMock()
    p_vi.Range.Start = 200
    p_vi.Range.End = 220
    p_vi.Range.Text = "VISUAL DEFECT\r"
    p_vi.Range.Information.return_value = 4

    p_sticker = MagicMock()
    p_sticker.Range.Start = 300
    p_sticker.Range.End = 330
    p_sticker.Range.Text = "NORMAL/DEFECT STICKER\r"
    p_sticker.Range.Information.return_value = 5

    mock_source_doc.Paragraphs = [p_visum, p_cond, p_vi, p_sticker]
    mock_source_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)
    sections = slicer.slice_sections(
        source_path=source_doc_path,
        target_dir=target_dir,
        station="TALAPIA",
        has_vi_summary=True,
        has_vi_defects=True,
    )

    assert sections.station == "TALAPIA"
    assert sections.front_page == target_dir / "front_page.docx"
    assert sections.vi_summary == target_dir / "vi_summary.docx"
    assert sections.condition_pages == target_dir / "condition_pages.docx"
    assert sections.vi_defect_pages == target_dir / "vi_defect_pages.docx"
    assert sections.sticker_page == target_dir / "sticker_page.docx"

    # Verify 5 sliced parts were created via Documents.Add()
    assert mock_word.Documents.Add.call_count == 5
    # Verify SaveAs2 was called 5 times
    assert mock_new_doc.SaveAs2.call_count == 5


def test_word_com_slicer_mock_omits_vi_when_not_in_doc(tmp_path: Path):
    """Verify WordComDocumentSlicer omits vi_summary and vi_defect_pages when paragraphs absent."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "CENDERAWASIH"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_new_doc = MagicMock()

    mock_word.Documents.Open.return_value = mock_source_doc
    mock_word.Documents.Add.return_value = mock_new_doc

    mock_source_doc.Content.Start = 0
    mock_source_doc.Content.End = 300
    mock_source_doc.ComputeStatistics.return_value = 3
    mock_source_doc.GoTo.return_value = MagicMock(Start=50)

    p_cond = MagicMock()
    p_cond.Range.Start = 50
    p_cond.Range.End = 80
    p_cond.Range.Text = "SUBSTATION CONDITION\r"
    p_cond.Range.Information.return_value = 2

    p_sticker = MagicMock()
    p_sticker.Range.Start = 150
    p_sticker.Range.End = 180
    p_sticker.Range.Text = "NORMAL/DEFECT STICKER\r"
    p_sticker.Range.Information.return_value = 3

    # No visual defect summary or visual defect paragraphs
    mock_source_doc.Paragraphs = [p_cond, p_sticker]
    mock_source_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)
    sections = slicer.slice_sections(
        source_path=source_doc_path,
        target_dir=target_dir,
        station="CENDERAWASIH",
        has_vi_summary=True,
        has_vi_defects=True,
    )

    assert sections.station == "CENDERAWASIH"
    assert sections.front_page == target_dir / "front_page.docx"
    assert sections.vi_summary is None
    assert sections.condition_pages == target_dir / "condition_pages.docx"
    assert sections.vi_defect_pages is None
    assert sections.sticker_page == target_dir / "sticker_page.docx"

    # Only 3 sliced parts created (front, condition, sticker)
    assert mock_word.Documents.Add.call_count == 3


def test_replace_front_page_title_invokes_find_replace():
    """Verify _replace_front_page_title executes Word COM Find & Replace per D46."""
    mock_doc = MagicMock()
    mock_find = MagicMock()
    mock_doc.Content.Find = mock_find

    _replace_front_page_title(mock_doc)

    mock_find.ClearFormatting.assert_called_once()
    mock_find.Replacement.ClearFormatting.assert_called_once()
    mock_find.Execute.assert_called_once()
    # Check that replacement target is FULL SCANNING REPORT
    args, kwargs = mock_find.Execute.call_args
    assert "QUICK SCANNING REPORT" in args or kwargs.get("FindText") == "QUICK SCANNING REPORT"
    assert "FULL SCANNING REPORT" in args or kwargs.get("ReplaceWith") == "FULL SCANNING REPORT"


def test_replace_front_page_title_raises_slicing_error_on_failure():
    """Verify _replace_front_page_title raises SlicingError when Find.Execute fails."""
    mock_doc = MagicMock()
    mock_find = MagicMock()
    mock_doc.Content.Find = mock_find
    mock_find.Execute.side_effect = RuntimeError("COM Find failed")

    with pytest.raises(SlicingError, match="Word COM Find.Execute failed"):
        _replace_front_page_title(mock_doc)


def test_find_paragraph_respects_start_pos_when_end_pos_is_none():
    """Verify _find_paragraph respects start_pos even when end_pos is None."""
    mock_doc = MagicMock()
    mock_doc.Content.End = 500
    mock_rng = MagicMock()
    mock_doc.Range.return_value = mock_rng

    p = MagicMock()
    p.Range.Start = 200
    p.Range.End = 230
    p.Range.Text = "NORMAL/DEFECT STICKER\r"
    p.Range.Information.return_value = 5

    mock_rng.Paragraphs.return_value = p
    mock_rng.Paragraphs.Count = 1
    mock_rng.Find.Execute.return_value = True

    result = _find_paragraph(mock_doc, "NORMAL/DEFECT STICKER", start_pos=150, end_pos=None)

    assert result is not None
    assert result.start == 200
    # Range must have been constructed starting at 150 up to document end
    mock_doc.Range.assert_called_with(150, 500)


def test_slice_range_preserves_drawingml_and_shapes(tmp_path: Path):
    """Verify _slice_range_to_doc calls Range.Copy() and PasteAndFormat(16) preserving DrawingML."""
    mock_word = MagicMock()
    mock_source = MagicMock()
    mock_target = MagicMock()
    mock_rng = MagicMock()
    mock_dest_rng = MagicMock()

    mock_source.Range.return_value = mock_rng
    mock_target.Range.return_value = mock_dest_rng
    mock_word.Documents.Add.return_value = mock_target

    out_file = tmp_path / "condition_pages.docx"
    _slice_range_to_doc(mock_word, mock_source, 100, 200, out_file)

    mock_rng.Copy.assert_called_once()
    mock_dest_rng.PasteAndFormat.assert_called_once_with(16)  # WD_FORMAT_ORIGINAL = 16
    mock_target.SaveAs2.assert_called_once_with(str(out_file.resolve()))


def test_fake_document_slicer_with_mock_cbm_defects(tmp_path: Path):
    """Verify FakeDocumentSlicer populates and returns mock CBM defect paths."""
    source_doc = tmp_path / "source.docx"
    source_doc.touch()
    target_dir = tmp_path / "temp_parts" / "CENDERAWASIH"

    mock_defect_1 = tmp_path / "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx"
    mock_defect_1.write_bytes(b"PK\x03\x04mock_defect_1")
    mock_defect_2 = tmp_path / "swg1_p01_INCOMING_1_CABLE_COMPARTMENT_01.docx"
    mock_defect_2.write_bytes(b"PK\x03\x04mock_defect_2")

    slicer = FakeDocumentSlicer(mock_cbm_defects=[mock_defect_1, mock_defect_2])

    # Test slice_cbm_defects
    cbm_slices = slicer.slice_cbm_defects(source_doc, target_dir)
    assert len(cbm_slices) == 2
    assert cbm_slices[0].name == "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx"
    assert cbm_slices[1].name == "swg1_p01_INCOMING_1_CABLE_COMPARTMENT_01.docx"
    assert (target_dir / "cbm_defects" / "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx").exists()

    # Test slice_sections includes cbm_defect_pages
    sections = slicer.slice_sections(source_doc, target_dir, station="CENDERAWASIH")
    assert len(sections.cbm_defect_pages) == 2
    assert sections.cbm_defect_pages[0].name == "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx"


def test_word_com_slicer_slices_cbm_defects_with_d37_names(tmp_path: Path):
    """Verify WordComDocumentSlicer identifies, slices, and names CBM defect pages per D37."""
    from unittest.mock import MagicMock
    from src.full_report.defect_parser import CbmDefectSliceMetadata

    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "CENDERAWASIH"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_new_doc = MagicMock()

    mock_word.Documents.Open.return_value = mock_source_doc
    mock_word.Documents.Add.return_value = mock_new_doc

    mock_source_doc.Content.Start = 0
    mock_source_doc.Content.End = 500
    mock_source_doc.ComputeStatistics.return_value = 5

    # Page navigation:
    # Page 2 start at 50 (CBM Summary)
    # Page 3 start at 100 (CBM Defect page)
    # Page 4 start at 200 (SUBSTATION CONDITION)
    # Page 5 start at 300 (NORMAL/DEFECT STICKER)
    p_goto_map = {
        2: MagicMock(Start=50),
        3: MagicMock(Start=100),
        4: MagicMock(Start=200),
        5: MagicMock(Start=300),
    }
    mock_source_doc.GoTo.side_effect = lambda what, which, count: p_goto_map.get(count, MagicMock(Start=500))

    # Paragraph boundaries:
    p_cbmsum = MagicMock()
    p_cbmsum.Range.Start = 50
    p_cbmsum.Range.End = 80
    p_cbmsum.Range.Text = "EXECUTIVE Summary\r"
    p_cbmsum.Range.Information.return_value = 2

    p_cond = MagicMock()
    p_cond.Range.Start = 200
    p_cond.Range.End = 230
    p_cond.Range.Text = "SUBSTATION CONDITION\r"
    p_cond.Range.Information.return_value = 4

    p_sticker = MagicMock()
    p_sticker.Range.Start = 300
    p_sticker.Range.End = 330
    p_sticker.Range.Text = "NORMAL/DEFECT STICKER\r"
    p_sticker.Range.Information.return_value = 5

    mock_source_doc.Paragraphs = [p_cbmsum, p_cond, p_sticker]
    mock_source_doc.Content.Find.Execute.return_value = False

    # Mock header parser to return D37 metadata for Page 3
    mock_parser = MagicMock()
    mock_parser.parse.return_value = CbmDefectSliceMetadata(
        equipment_category="swg",
        equipment_instance="swg1",
        sequence="p04",
        equipment_id="CKN01309",
        defect_area="FUSE_COMPARTMENT",
        index=1,
        filename="swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx",
    )

    slicer = WordComDocumentSlicer(word_app=mock_word, header_parser=mock_parser)

    # Test slice_cbm_defects directly
    cbm_slices = slicer.slice_cbm_defects(source_doc_path, target_dir)
    assert len(cbm_slices) == 1
    assert cbm_slices[0].name == "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx"

    # Test slice_sections includes CBM defect slices in SlicedSections
    sections = slicer.slice_sections(source_doc_path, target_dir, station="CENDERAWASIH")
    assert len(sections.cbm_defect_pages) == 1
    assert sections.cbm_defect_pages[0].name == "swg1_p04_CKN01309_FUSE_COMPARTMENT_01.docx"


def test_slice_cbm_defects_rejects_foreign_substation_metadata(tmp_path: Path):
    """Verify slice_cbm_defects unlinks and rejects candidate defect slices from foreign substations."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "PERPUSTAKAAN_AWAM"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_source_doc
    mock_source_doc.ComputeStatistics.return_value = 3
    mock_source_doc.GoTo.return_value = MagicMock(Start=50)

    p_cond = MagicMock()
    p_cond.Range.Start = 100
    p_cond.Range.Information.return_value = 3
    p_cond.Range.Text = "SUBSTATION CONDITION\r"
    mock_source_doc.Paragraphs = [p_cond]
    mock_source_doc.Content.Find.Execute.return_value = False

    # Mock parser returning foreign substation "TELEKOM TANAH PUTIH"
    mock_parser = MagicMock()
    mock_parser.parse.return_value = CbmDefectSliceMetadata(
        equipment_category="swg",
        equipment_instance="swg1",
        sequence="p01",
        equipment_id="VCB1",
        defect_area="CABLE_BOX",
        substation="TELEKOM TANAH PUTIH",
        filename="swg1_p01_VCB1_CABLE_BOX_01.docx",
    )

    slicer = WordComDocumentSlicer(word_app=mock_word, header_parser=mock_parser)
    cbm_slices = slicer.slice_cbm_defects(source_doc_path, target_dir, station="PERPUSTAKAAN AWAM")

    # Rejected because "TELEKOM TANAH PUTIH" conflicts with target "PERPUSTAKAAN AWAM"
    assert len(cbm_slices) == 0


def test_slice_sections_rejects_foreign_front_page_after_retries(tmp_path: Path):
    """Verify slice_sections retries 3 times and raises SlicingError if front page attribution fails."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "PERPUSTAKAAN_AWAM"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_source_doc
    mock_source_doc.ComputeStatistics.return_value = 3
    mock_source_doc.GoTo.return_value = MagicMock(Start=50)

    p_cond = MagicMock()
    p_cond.Range.Start = 100
    p_cond.Range.Information.return_value = 2
    p_cond.Range.Text = "SUBSTATION CONDITION\r"

    p_stk = MagicMock()
    p_stk.Range.Start = 200
    p_stk.Range.Information.return_value = 3
    p_stk.Range.Text = "NORMAL/DEFECT STICKER\r"

    mock_source_doc.Paragraphs = [p_cond, p_stk]
    mock_source_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)

    with patch("src.full_report.slicer._slice_range_to_doc") as mock_slice_range:
        # Create a real docx file on disk that contains foreign front page
        front_p = target_dir / "front_page.docx"
        def fake_slice(*args, **kwargs):
            doc = docx.Document()
            doc.add_table(rows=1, cols=3)
            t1 = doc.add_table(rows=4, cols=3)
            t1.rows[2].cells[0].text = "SUBSTATION NAME (ERMS)"
            t1.rows[2].cells[2].text = "TELEKOM TANAH PUTIH"
            t1.rows[3].cells[0].text = "SUBSTATION NAME (SITE)"
            t1.rows[3].cells[2].text = "TELEKOM TANAH PUTIH"
            front_p.parent.mkdir(parents=True, exist_ok=True)
            doc.save(front_p)
            return front_p

        mock_slice_range.side_effect = fake_slice

        with pytest.raises(SlicingError, match="failed attribution"):
            slicer.slice_sections(
                source_doc_path,
                target_dir,
                station="PERPUSTAKAAN AWAM",
                has_vi_summary=False,
                has_vi_defects=False,
            )
        # Verify it retried 3 times
        assert mock_slice_range.call_count == 3


def test_slice_sections_rejects_invalid_condition_structure_after_retries(tmp_path: Path):
    """Verify slice_sections retries 3 times and raises SlicingError if condition_pages has no images."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "PERPUSTAKAAN_AWAM"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_source_doc
    mock_source_doc.ComputeStatistics.return_value = 3
    mock_source_doc.GoTo.return_value = MagicMock(Start=50)

    p_cond = MagicMock()
    p_cond.Range.Start = 100
    p_cond.Range.Information.return_value = 2
    p_cond.Range.Text = "SUBSTATION CONDITION\r"

    p_stk = MagicMock()
    p_stk.Range.Start = 200
    p_stk.Range.Information.return_value = 3
    p_stk.Range.Text = "NORMAL/DEFECT STICKER\r"

    mock_source_doc.Paragraphs = [p_cond, p_stk]
    mock_source_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)

    with patch("src.full_report.slicer._slice_range_to_doc") as mock_slice_range:
        # Create valid front page but invalid condition page (no images)
        front_p = target_dir / "front_page.docx"
        cond_p = target_dir / "condition_pages.docx"

        def fake_slice(word_app, source_doc, start_pos, end_pos, output_path, is_front_page=False, **kwargs):
            doc = docx.Document()
            if is_front_page:
                doc.add_table(rows=1, cols=3)
                t1 = doc.add_table(rows=4, cols=3)
                t1.rows[2].cells[0].text = "SUBSTATION NAME (ERMS)"
                t1.rows[2].cells[2].text = "PERPUSTAKAAN AWAM"
                t1.rows[3].cells[0].text = "SUBSTATION NAME (SITE)"
                t1.rows[3].cells[2].text = "PERPUSTAKAAN AWAM"
            else:
                # condition_pages: has table but NO images
                doc.add_table(rows=2, cols=2)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(output_path)
            return output_path

        mock_slice_range.side_effect = fake_slice

        with pytest.raises(SlicingError, match="condition_pages failed attribution or structural guard after 3 attempts"):
            slicer.slice_sections(
                source_doc_path,
                target_dir,
                station="PERPUSTAKAAN AWAM",
                has_vi_summary=False,
                has_vi_defects=False,
            )


def test_slice_sections_recovers_on_retry(tmp_path: Path):
    """Verify slice_sections recovers if guard fails on first attempt and succeeds on second attempt."""
    source_doc_path = tmp_path / "source.docx"
    source_doc_path.touch()
    target_dir = tmp_path / "temp_parts" / "PERPUSTAKAAN_AWAM"

    mock_word = MagicMock()
    mock_source_doc = MagicMock()
    mock_word.Documents.Open.return_value = mock_source_doc
    mock_source_doc.ComputeStatistics.return_value = 3
    mock_source_doc.GoTo.return_value = MagicMock(Start=50)

    p_cond = MagicMock()
    p_cond.Range.Start = 100
    p_cond.Range.Information.return_value = 2
    p_cond.Range.Text = "SUBSTATION CONDITION\r"

    p_stk = MagicMock()
    p_stk.Range.Start = 200
    p_stk.Range.Information.return_value = 3
    p_stk.Range.Text = "NORMAL/DEFECT STICKER\r"

    mock_source_doc.Paragraphs = [p_cond, p_stk]
    mock_source_doc.Content.Find.Execute.return_value = False

    slicer = WordComDocumentSlicer(word_app=mock_word)

    attempt_count = [0]
    with patch("src.full_report.slicer._slice_range_to_doc") as mock_slice_range:
        def fake_slice(word_app, source_doc, start_pos, end_pos, output_path, is_front_page=False, **kwargs):
            doc = docx.Document()
            if is_front_page:
                attempt_count[0] += 1
                doc.add_table(rows=1, cols=3)
                t1 = doc.add_table(rows=4, cols=3)
                t1.rows[2].cells[0].text = "SUBSTATION NAME (ERMS)"
                t1.rows[3].cells[0].text = "SUBSTATION NAME (SITE)"
                if attempt_count[0] == 1:
                    # Stale clipboard on attempt 1
                    t1.rows[2].cells[2].text = "TELEKOM TANAH PUTIH"
                    t1.rows[3].cells[2].text = "TELEKOM TANAH PUTIH"
                else:
                    # Clean correct clipboard on attempt 2
                    t1.rows[2].cells[2].text = "PERPUSTAKAAN AWAM"
                    t1.rows[3].cells[2].text = "PERPUSTAKAAN AWAM"
            elif "condition_pages" in output_path.name:
                doc.add_table(rows=1, cols=1)
                p = doc.add_paragraph()
                r = p.add_run()
                from docx.oxml import parse_xml
                from docx.oxml.ns import nsdecls
                r._r.append(parse_xml(r'<w:drawing %s><w:inline><a:graphic %s><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic %s><pic:blipFill><a:blip r:embed="rId1"/></pic:blipFill></pic:pic></a:graphicData></a:graphic></w:inline></w:drawing>' % (nsdecls('w'), nsdecls('a'), nsdecls('pic', 'r'))))
            elif "sticker_page" in output_path.name:
                doc.add_paragraph("NORMAL STICKER")

            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(output_path)
            return output_path

        mock_slice_range.side_effect = fake_slice

        sections = slicer.slice_sections(
            source_doc_path,
            target_dir,
            station="PERPUSTAKAAN AWAM",
            has_vi_summary=False,
            has_vi_defects=False,
        )
        assert sections.station == "PERPUSTAKAAN AWAM"
        assert attempt_count[0] == 2




