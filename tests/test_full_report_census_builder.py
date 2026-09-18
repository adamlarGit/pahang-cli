"""Unit and integration tests for Full Report ExecutiveSummaryCensusBuilder (Ticket #28 / T3.2)."""

from pathlib import Path
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
import pytest

from src.full_report.census import (
    CensusRowItem,
    ExecutiveSummaryCensusBuilder,
    ExecutiveSummaryCensusContext,
    ExecutiveSummaryCensusResult,
    apply_column_vertical_merge,
    apply_post_render_dom,
    apply_severity_shading,
)
from src.quick_report.defects import CbmDefectRecord
from src.testsheet.models import (
    BatteryBankSpec,
    LVDBFeederSpec,
    LVDBSpec,
    SubstationEquipmentPackage,
    SwitchgearPanelSpec,
    SwitchgearSpec,
    TransformerSpec,
)


def _get_cell_shading(cell: _Cell):
    """Safely retrieve shading XML element from a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    return tcPr.find(qn("w:shd"))


def _get_cell_vmerge(cell: _Cell):
    """Safely retrieve vMerge XML element from a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    return tcPr.find(qn("w:vMerge"))


def _get_row_tc_vmerge(row: docx.table._Row, col_idx: int = 0):
    """Safely retrieve vMerge XML element from underlying row tc."""
    tc = row._tr.tc_lst[col_idx]
    tcPr = tc.get_or_add_tcPr()
    return tcPr.find(qn("w:vMerge"))


# ==============================================================================
# 1. Scaffolding & Dataclasses Tests
# ==============================================================================

def test_census_row_item_fields_and_dict():
    """Verify CensusRowItem initializes with standard defaults and serializes to dict."""
    item = CensusRowItem(
        no="1.",
        equipment="RMU SF6 - TAMCO",
        defect_area="OVERVIEW",
        ir_abs="-",
        us_dB="-",
        tev_dB="-",
        severity="-",
        group_no=1,
        is_overview=True,
    )
    assert item.no == "1."
    assert item.equipment == "RMU SF6 - TAMCO"
    assert item.defect_area == "OVERVIEW"
    assert item.is_overview is True
    assert item["equipment"] == "RMU SF6 - TAMCO"

    d = item.to_dict()
    assert d == {
        "no": "1.",
        "equipment": "RMU SF6 - TAMCO",
        "defect_area": "OVERVIEW",
        "ir_abs": "-",
        "us_dB": "-",
        "tev_dB": "-",
        "severity": "-",
    }


def test_census_context_and_result_dataclasses():
    """Verify ExecutiveSummaryCensusContext and ExecutiveSummaryCensusResult dataclasses."""
    item = CensusRowItem(no="1.", equipment="TX 1", defect_area="OVERVIEW", severity="-", group_no=1)
    context = ExecutiveSummaryCensusContext(
        census_items=[item],
        substation_number=5,
        station_name="TALAPIA",
    )
    assert context.substation_number == 5
    assert len(context.census_items) == 1
    d = context.to_dict()
    assert len(d["census_items"]) == 1
    assert d["station_name"] == "TALAPIA"

    result = ExecutiveSummaryCensusResult(
        docx_path=Path("dummy.docx"),
        items=[item],
        group_count=1,
        total_rows=1,
        has_defects=False,
    )
    assert result.group_count == 1
    assert result.has_defects is False


# ==============================================================================
# 2. Transformer 7-Point Rows (ADR 0004 & D27)
# ==============================================================================

def test_transformer_unconditional_7_point_rows():
    """Verify unconditional 7-point TX rows per ADR 0004 and D27."""
    tx = TransformerSpec(
        tx_id="Tx 1",
        rating_kva="1000",
        manufacturer="MTM",
        serial_no="SN-1234",
    )
    pkg = SubstationEquipmentPackage(transformers=(tx,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg)
    assert len(rows) == 7

    expected_defect_areas = [
        "OVERVIEW",
        "OVERVIEW TOP",
        "HV BUSHING",
        "HV CABLES",
        "CABLE SPLIT",
        "LV BUSHING",
        "LV CABLES",
    ]
    actual_areas = [r.defect_area for r in rows]
    assert actual_areas == expected_defect_areas

    # All rows belong to group 1
    for r in rows:
        assert r.group_no == 1
        assert r.no == "1."

    # Overview rows have severity "-" and is_overview=True
    assert rows[0].is_overview is True
    assert rows[0].severity == "-"
    assert rows[1].is_overview is True
    assert rows[1].severity == "-"

    # Component rows are NORMAL and is_overview=False
    for r in rows[2:]:
        assert r.is_overview is False
        assert r.severity == "NORMAL"
        assert r.ir_abs == "-"
        assert r.us_dB == "-"
        assert r.tev_dB == "-"


# ==============================================================================
# 3. Feeder Pillar Defect-Only Granularity (D22)
# ==============================================================================

def test_feeder_pillar_healthy_emits_single_overview_row():
    """Verify healthy Feeder Pillar emits strictly 1 OVERVIEW row per D22."""
    fp = LVDBSpec(
        name="FP TX1",
        label="FP",
        source="TX1",
        feeders=tuple(LVDBFeederSpec(channel=f"OT{i}") for i in range(1, 11)),
    )
    pkg = SubstationEquipmentPackage(lvdb_specs=(fp,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg, defects=())
    assert len(rows) == 1
    assert rows[0].defect_area == "OVERVIEW"
    assert rows[0].is_overview is True
    assert rows[0].severity == "-"


def test_feeder_pillar_with_defect_appends_active_feeder_rows():
    """Verify Feeder Pillar with active defect appends individual feeder rows per D22."""
    fp = LVDBSpec(
        name="FP TX1",
        label="FP",
        source="TX1",
        feeders=tuple(LVDBFeederSpec(channel=f"OT{i}") for i in range(1, 11)),
    )
    defect = CbmDefectRecord(
        equipment="FEEDER PILLAR",
        equipment_id="F2",
        defect_area="Fuse Contact",
        additional_remarks="Hotspot",
        technology="IR",
        ir_reading="78.4",
    )
    pkg = SubstationEquipmentPackage(lvdb_specs=(fp,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg, defects=(defect,))
    # 1 overview + 1 active defect = 2 rows
    assert len(rows) == 2
    assert rows[0].defect_area == "OVERVIEW"
    assert rows[0].is_overview is True
    assert rows[0].severity == "-"

    # Defect row
    defect_row = rows[1]
    assert defect_row.is_overview is False
    assert defect_row.severity == "DEFECT"
    assert "F2" in defect_row.equipment.upper()
    assert "FUSE CONTACT" in defect_row.defect_area.upper()
    assert defect_row.ir_abs == "78.4 °C"
    assert defect_row.us_dB == "-"
    assert defect_row.tev_dB == "-"


# ==============================================================================
# 4. Defect Cross-Referencing & Measurement Formatting (D23 & D28)
# ==============================================================================

def test_cross_reference_switchgear_defects():
    """Verify cross-referencing Quick Report CbmDefectRecord to populate readings and DEFECT severity."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="INDKOM",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="CKN01308", name="SPARE"),
            SwitchgearPanelSpec(panel_no=2, panel_feeder_no="CKN01310", name="TMN CENDRAWASIH 3"),
            SwitchgearPanelSpec(panel_no=3, panel_feeder_no="CKN01311", name="TAJ 33A LOT 11520"),
            SwitchgearPanelSpec(panel_no=4, panel_feeder_no="CKN01309", name="PANEL CKN01309 TX"),
        ),
    )
    defect_ir = CbmDefectRecord(
        equipment="RMU SF6",
        equipment_id="CKN01309",
        defect_area="FUSE COMPARTMENT",
        technology="IR",
        ir_reading="68.2",
    )
    defect_us = CbmDefectRecord(
        equipment="RMU SF6",
        equipment_id="CKN01309",
        defect_area="FUSE COMPARTMENT",
        technology="US",
        us_reading="15.0",
        us_char="CORONA",
    )
    pkg = SubstationEquipmentPackage(switchgears=(swg,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg, defects=(defect_ir, defect_us))
    # 1 overview + 4 panels (INDKOM: 1 comp each) = 5 rows
    assert len(rows) == 5
    assert rows[0].defect_area == "OVERVIEW"
    assert rows[0].severity == "-"

    # Panels 1 to 3 are NORMAL
    for r in rows[1:4]:
        assert r.severity == "NORMAL"
        assert r.ir_abs == "-"
        assert r.us_dB == "-"
        assert r.tev_dB == "-"

    # Panel 4 is DEFECT with paired IR & US readings
    p4_row = rows[4]
    assert p4_row.defect_area == "FUSE COMPARTMENT"
    assert p4_row.severity == "DEFECT"
    assert p4_row.ir_abs == "68.2 °C"
    assert p4_row.us_dB == "15dB"
    assert p4_row.tev_dB == "-"


def test_cross_reference_transformer_defect():
    """Verify cross-referencing defect on transformer component."""
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    defect = CbmDefectRecord(
        equipment="TRANSFORMER",
        equipment_id="TX 1",
        defect_area="HV Bushing",
        technology="US",
        us_reading="18.0",
        us_char="TRACKING",
    )
    pkg = SubstationEquipmentPackage(transformers=(tx,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg, defects=(defect,))
    assert len(rows) == 7

    bushing_row = next(r for r in rows if r.defect_area == "HV BUSHING")
    assert bushing_row.severity == "DEFECT"
    assert bushing_row.us_dB == "18dB"
    assert bushing_row.ir_abs == "-"

    split_row = next(r for r in rows if r.defect_area in ("CABLE SPLIT", "HV CABLE SPLIT"))
    assert split_row.severity == "NORMAL"


# ==============================================================================
# 5. Canonical Benchmark Substations (TALAPIA, CENDERAWASIH, TELEKOM)
# ==============================================================================

def test_canonical_benchmark_talapia_census():
    """Benchmark TALAPIA: TAMCO RMU (10 rows), 1 TX (7 rows), FP with F2 defect (2 rows), Battery (1 row) = 20 rows."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        model="GR1",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="CRB00679", name="PE KG ASLI BATU BALONG"),
            SwitchgearPanelSpec(panel_no=2, panel_feeder_no="CRB00679", name="TX B"),
            SwitchgearPanelSpec(panel_no=3, panel_feeder_no="CRB00678", name="TX A"),
            SwitchgearPanelSpec(panel_no=4, panel_feeder_no="CRB00677", name="CS LDG BILUT"),
        ),
    )
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    fp = LVDBSpec(
        name="FP TX1",
        label="FP",
        source="TX1",
        feeders=tuple(LVDBFeederSpec(channel=f"OT{i}") for i in range(1, 11)),
    )
    bb = BatteryBankSpec(name="BATTERY 1", manufacturer="Sunpower")

    pkg = SubstationEquipmentPackage(
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(fp,),
        battery_banks=(bb,),
    )

    defect_fp = CbmDefectRecord(
        equipment="FEEDER PILLAR",
        equipment_id="F2",
        defect_area="Fuse Contact",
        technology="IR",
        ir_reading="78.4",
    )

    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=(defect_fp,))

    # TAMCO: 2 overview + 4 * 2 panels = 10 rows (Group 1)
    # TX1: 7 rows (Group 2)
    # FP: 1 overview + 1 defect = 2 rows (Group 3)
    # Battery: 1 overview = 1 row (Group 4)
    # Total = 20 rows across 4 groups
    assert len(rows) == 20

    group_1_rows = [r for r in rows if r.group_no == 1]
    assert len(group_1_rows) == 10
    assert group_1_rows[0].defect_area == "OVERVIEW"
    assert group_1_rows[1].defect_area == "OVERVIEW BOTTOM"

    group_2_rows = [r for r in rows if r.group_no == 2]
    assert len(group_2_rows) == 7

    group_3_rows = [r for r in rows if r.group_no == 3]
    assert len(group_3_rows) == 2
    assert group_3_rows[0].defect_area == "OVERVIEW"
    assert group_3_rows[1].severity == "DEFECT"

    group_4_rows = [r for r in rows if r.group_no == 4]
    assert len(group_4_rows) == 1
    assert group_4_rows[0].defect_area == "OVERVIEW"


def test_canonical_benchmark_cenderawasih_census():
    """Benchmark CENDERAWASIH NO.1: INDKOM RMU (5 rows with panel 4 defect), 1 TX (7 rows), FP (1 row) = 13 rows."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="INDKOM",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="CKN01308", name="SPARE"),
            SwitchgearPanelSpec(panel_no=2, panel_feeder_no="CKN01310", name="TMN CENDRAWASIH 3"),
            SwitchgearPanelSpec(panel_no=3, panel_feeder_no="CKN01311", name="TAJ 33A LOT 11520"),
            SwitchgearPanelSpec(panel_no=4, panel_feeder_no="CKN01309", name="PANEL CKN01309 TX"),
        ),
    )
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="EWT")
    fp = LVDBSpec(name="FP TX1", label="FP", source="TX1")

    pkg = SubstationEquipmentPackage(
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(fp,),
        battery_banks=(),
    )

    defect_swg = CbmDefectRecord(
        equipment="RMU SF6",
        equipment_id="CKN01309",
        defect_area="FUSE COMPARTMENT",
        technology="IR",
        ir_reading="65.4",
    )

    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=(defect_swg,))

    # INDKOM: 1 overview + 4 panels = 5 rows (Group 1)
    # TX: 7 rows (Group 2)
    # FP: 1 row (Group 3)
    # Total = 13 rows across 3 groups
    assert len(rows) == 13
    assert rows[4].defect_area == "FUSE COMPARTMENT"
    assert rows[4].severity == "DEFECT"
    assert rows[4].ir_abs == "65.4 °C"


def test_canonical_benchmark_telekom_tanah_putih_census():
    """Benchmark TELEKOM TANAH PUTIH: INDKOM RMU (5 rows with TEV on all panels), 1 TX (7 rows), LVDB (1 row), Battery (1 row) = 14 rows."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="INDKOM",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="F1", name="INCOMING 1"),
            SwitchgearPanelSpec(panel_no=2, panel_feeder_no="F2", name="INCOMING 2"),
            SwitchgearPanelSpec(panel_no=3, panel_feeder_no="F3", name="BUS COUPLER"),
            SwitchgearPanelSpec(panel_no=4, panel_feeder_no="F4", name="TX 1 FEEDER"),
        ),
    )
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    lvdb = LVDBSpec(name="LVDB 1")
    bb = BatteryBankSpec(name="BATTERY 1")

    pkg = SubstationEquipmentPackage(
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(lvdb,),
        battery_banks=(bb,),
    )

    # TEV defect across all panels
    defects = [
        CbmDefectRecord(equipment="RMU SF6", equipment_id=f"F{i}", defect_area="CABLE COMPARTMENT" if i != 4 else "FUSE COMPARTMENT", technology="TEV", tev_reading="24.0")
        for i in range(1, 5)
    ]

    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=defects)

    # INDKOM: 1 overview + 4 panels = 5 rows (Group 1)
    # TX: 7 rows (Group 2)
    # LVDB: 1 row (Group 3)
    # Battery: 1 row (Group 4)
    # Total = 14 rows across 4 groups
    assert len(rows) == 14
    for i in range(1, 5):
        assert rows[i].severity == "DEFECT"
        assert rows[i].tev_dB == "24dB"


# ==============================================================================
# 6. Post-Render DOM Pass: Column 0 <w:vMerge> and Dynamic Shading (D21 & D24)
# ==============================================================================

def test_post_render_dom_vmerge_and_shading(tmp_path: Path):
    """Verify post-render DOM pass setting <w:vMerge> on Col 0 and green/red shading on Col 6."""
    builder = ExecutiveSummaryCensusBuilder()

    # Create mock items with 2 groups:
    # Group 1: 3 rows (1 overview, 1 normal, 1 defect)
    # Group 2: 2 rows (1 overview, 1 normal)
    items = [
        CensusRowItem(no="1.", equipment="SWG", defect_area="OVERVIEW", severity="-", group_no=1, is_overview=True),
        CensusRowItem(no="1.", equipment="SWG", defect_area="CABLE COMPARTMENT", severity="NORMAL", group_no=1, is_overview=False),
        CensusRowItem(no="1.", equipment="SWG", defect_area="FUSE COMPARTMENT", ir_abs="65.0 °C", severity="DEFECT", group_no=1, is_overview=False, is_defect=True),
        CensusRowItem(no="2.", equipment="TX 1", defect_area="OVERVIEW", severity="-", group_no=2, is_overview=True),
        CensusRowItem(no="2.", equipment="TX 1", defect_area="HV BUSHING", severity="NORMAL", group_no=2, is_overview=False),
    ]

    out_path = tmp_path / "census_rendered.docx"
    result = builder.render(
        ExecutiveSummaryCensusContext(census_items=items),
        output_path=out_path,
    )
    assert out_path.exists()
    assert result.total_rows == 5
    assert result.group_count == 2
    assert result.has_defects is True

    # Re-open docx and inspect DOM structure
    doc = docx.Document(str(out_path))
    table = doc.tables[0]
    assert len(table.rows) == 6  # 1 header + 5 data rows

    # Group 1: Table rows 1, 2, 3
    # Row 1 (first row of group 1): Col 0 must have vMerge restart and text '1.'
    vm1_0 = _get_row_tc_vmerge(table.rows[1], 0)
    assert vm1_0 is not None
    assert vm1_0.get(qn("w:val")) == "restart"
    assert table.rows[1].cells[0].text.strip() == "1."

    # Row 2 (continuing row of group 1): Col 0 must have vMerge continue (val is None) and underlying text cleared
    vm2_0 = _get_row_tc_vmerge(table.rows[2], 0)
    assert vm2_0 is not None
    assert vm2_0.get(qn("w:val")) is None or vm2_0.get(qn("w:val")) == "continue"
    assert _Cell(table.rows[2]._tr.tc_lst[0], table).text.strip() == ""

    # Row 3 (continuing row of group 1): Col 0 must have vMerge continue (val is None) and underlying text cleared
    vm3_0 = _get_row_tc_vmerge(table.rows[3], 0)
    assert vm3_0 is not None
    assert vm3_0.get(qn("w:val")) is None or vm3_0.get(qn("w:val")) == "continue"
    assert _Cell(table.rows[3]._tr.tc_lst[0], table).text.strip() == ""

    # Group 2: Table rows 4, 5
    # Row 4 (first row of group 2): Col 0 must have vMerge restart and text '2.'
    vm4_0 = _get_row_tc_vmerge(table.rows[4], 0)
    assert vm4_0 is not None
    assert vm4_0.get(qn("w:val")) == "restart"
    assert table.rows[4].cells[0].text.strip() == "2."

    # Row 5 (continuing row of group 2): Col 0 must have vMerge continue (val is None) and underlying text cleared
    vm5_0 = _get_row_tc_vmerge(table.rows[5], 0)
    assert vm5_0 is not None
    assert vm5_0.get(qn("w:val")) is None or vm5_0.get(qn("w:val")) == "continue"
    assert _Cell(table.rows[5]._tr.tc_lst[0], table).text.strip() == ""

    # Severity cell shading checks (Col 6):
    # Row 1: OVERVIEW -> unshaded, text is '-'
    c1_6 = table.rows[1].cells[6]
    shd1_6 = _get_cell_shading(c1_6)
    assert shd1_6 is None or shd1_6.get(qn("w:fill")) in (None, "auto")
    assert c1_6.text.strip() == "-"

    # Row 2: NORMAL -> shaded 00B050 (Green), text cleared
    c2_6 = table.rows[2].cells[6]
    shd2_6 = _get_cell_shading(c2_6)
    assert shd2_6 is not None
    assert shd2_6.get(qn("w:fill")) == "00B050"
    assert c2_6.text.strip() == ""

    # Row 3: DEFECT -> shaded EE0000 (Red), text cleared
    c3_6 = table.rows[3].cells[6]
    shd3_6 = _get_cell_shading(c3_6)
    assert shd3_6 is not None
    assert shd3_6.get(qn("w:fill")) == "EE0000"
    assert c3_6.text.strip() == ""

    # Row 4: OVERVIEW -> unshaded, text is '-'
    c4_6 = table.rows[4].cells[6]
    shd4_6 = _get_cell_shading(c4_6)
    assert shd4_6 is None or shd4_6.get(qn("w:fill")) in (None, "auto")
    assert c4_6.text.strip() == "-"

    # Row 5: NORMAL -> shaded 00B050 (Green), text cleared
    c5_6 = table.rows[5].cells[6]
    shd5_6 = _get_cell_shading(c5_6)
    assert shd5_6 is not None
    assert shd5_6.get(qn("w:fill")) == "00B050"
    assert c5_6.text.strip() == ""

    # Measurement columns (Cols 3, 4, 5) remain unshaded
    for row_idx in range(1, 6):
        for col_idx in (3, 4, 5):
            meas_cell = table.rows[row_idx].cells[col_idx]
            shd_meas = _get_cell_shading(meas_cell)
            assert shd_meas is None or shd_meas.get(qn("w:fill")) in (None, "auto")


def test_builder_end_to_end_from_package(tmp_path: Path):
    """Verify builder end-to-end rendering directly from SubstationEquipmentPackage."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        panels=(SwitchgearPanelSpec(panel_no=1, panel_feeder_no="P1", name="TX A"),),
    )
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    fp = LVDBSpec(name="FP TX1")
    pkg = SubstationEquipmentPackage(
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(fp,),
    )
    builder = ExecutiveSummaryCensusBuilder()
    out_path = tmp_path / "talapia_census.docx"
    result = builder.render(
        package_or_context=pkg,
        output_path=out_path,
        substation_number=5,
        station_name="TALAPIA",
    )
    assert result.docx_path == out_path
    assert out_path.exists()
    # TAMCO: 2 overview + 1 panel (2 comps) = 4 rows
    # TX: 7 rows
    # FP: 1 row
    # Total = 12 rows
    assert result.total_rows == 12
    assert result.group_count == 3


def test_tx_overview_defect_reporting():
    """Verify transformer tank/top defect is reported with DEFECT severity."""
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    pkg = SubstationEquipmentPackage(transformers=(tx,))
    defect = CbmDefectRecord(
        equipment="TRANSFORMER",
        equipment_id="TX 1",
        defect_area="BODY",
        additional_remarks="HOTSPOT ON MAIN TANK",
        technology="IR",
        raw_measurement="65.4",
        ir_reading="65.4",
    )
    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=[defect])
    assert len(rows) == 7
    ov_row = next(r for r in rows if r.defect_area == "OVERVIEW")
    assert ov_row.severity == "DEFECT"
    assert ov_row.is_defect is True
    assert "65.4" in ov_row.ir_abs


def test_battery_bank_defect_reporting():
    """Verify battery bank defect is reported with DEFECT severity."""
    bb = BatteryBankSpec(name="BATTERY BANK 1", manufacturer="CHLORIDE")
    pkg = SubstationEquipmentPackage(battery_banks=(bb,))
    defect = CbmDefectRecord(
        equipment="BATTERY",
        equipment_id="BATTERY BANK 1",
        defect_area="CELL TERMINAL",
        additional_remarks="CORROSION AND HIGH TEMP",
        technology="IR",
        raw_measurement="48.2",
        ir_reading="48.2",
    )
    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=[defect])
    assert len(rows) == 1
    assert rows[0].severity == "DEFECT"
    assert rows[0].is_defect is True
    assert "48.2" in rows[0].ir_abs


def test_multi_fp_defect_isolation():
    """Verify defects on FP 1 do not leak to FP 2 in a multi-FP station."""
    fp1 = LVDBSpec(name="FP 1")
    fp2 = LVDBSpec(name="FP 2")
    pkg = SubstationEquipmentPackage(lvdb_specs=(fp1, fp2))
    defect = CbmDefectRecord(
        equipment="FP 1",
        equipment_id="WAY 2",
        defect_area="HOTSPOT ON FUSE CONTACT",
        technology="IR",
        raw_measurement="55.0",
        ir_reading="55.0",
    )
    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg, defects=[defect])
    # FP 1: 1 overview + 1 defect row = 2 rows
    # FP 2: 1 overview = 1 row
    # Total = 3 rows
    fp1_rows = [r for r in rows if "FP 1" in r.equipment]
    fp2_rows = [r for r in rows if "FP 2" in r.equipment]
    assert len(fp1_rows) == 2
    assert len(fp2_rows) == 1
    assert fp2_rows[0].severity == "-"


def test_tx_has_hv_cable_split_predicate(monkeypatch):
    """Verify transformer component generation respects has_hv_cable_split."""
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    pkg = SubstationEquipmentPackage(transformers=(tx,))
    builder = ExecutiveSummaryCensusBuilder()

    # Default predicate: True -> 7 rows
    rows_default = builder.build_census_rows(pkg)
    assert len(rows_default) == 7
    assert any(r.defect_area in ("CABLE SPLIT", "HV CABLE SPLIT") for r in rows_default)

    # Monkeypatch to False -> 6 rows
    import src.full_report.census
    monkeypatch.setattr(src.full_report.census, "has_hv_cable_split", lambda _tx: False)
    rows_no_split = builder.build_census_rows(pkg)
    assert len(rows_no_split) == 6
    assert not any(r.defect_area in ("CABLE SPLIT", "HV CABLE SPLIT") for r in rows_no_split)


def test_tx_rating_omitted_when_not_accessible():
    """Verify that TX rating is omitted from equipment title if 'NOT ACCESSIBLE', '-', or 'N/A'."""
    tx1 = TransformerSpec(tx_id="Tx 1", rating_kva="NOT ACCESSIBLE", manufacturer="MTM")
    tx2 = TransformerSpec(tx_id="Tx 2", rating_kva="1000kVA", manufacturer="MTM")
    pkg = SubstationEquipmentPackage(transformers=(tx1, tx2))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg)
    # Tx 1 should format as 'TX1 - MTM'
    tx1_rows = [r for r in rows if r.group_no == 1]
    assert len(tx1_rows) == 7
    assert tx1_rows[0].equipment == "TX1 - MTM"

    # Tx 2 should format as 'TX2 - MTM 1000kVA'
    tx2_rows = [r for r in rows if r.group_no == 2]
    assert len(tx2_rows) == 7
    assert tx2_rows[0].equipment == "TX2 - MTM 1000kVA"


def test_column_vertical_merge_preserves_paragraph_centering_and_template_formatting():
    """Verify apply_column_vertical_merge sets w:jc=center, w:vAlign=center, and preserves run formatting on restart cell."""
    doc = docx.Document()
    table = doc.add_table(rows=4, cols=2)
    # Pre-style restart cell with template run properties (e.g. Calibri, 18pt)
    cell_p = table.rows[1].cells[0].paragraphs[0]
    run = cell_p.add_run("orig")
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rPr.append(rFonts)

    spans = [(1, 3)]  # rows 1-3 merged
    apply_column_vertical_merge(table, spans, group_numbers=[1])

    # Check restart cell at row 1 col 0
    cell = table.rows[1].cells[0]
    assert cell.text.strip() == "1."

    # Verify paragraph centering
    p = cell.paragraphs[0]
    pPr = p._p.find(qn("w:pPr"))
    assert pPr is not None
    jc = pPr.find(qn("w:jc"))
    assert jc is not None
    assert jc.get(qn("w:val")) == "center"

    # Verify run preserved template font Calibri
    assert len(p.runs) >= 1
    rPr = p.runs[0]._r.find(qn("w:rPr"))
    assert rPr is not None
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:ascii")) == "Calibri"

    # Verify cell vertical centering
    tcPr = cell._tc.find(qn("w:tcPr"))
    assert tcPr is not None
    v_align = tcPr.find(qn("w:vAlign"))
    assert v_align is not None
    assert v_align.get(qn("w:val")) == "center"


def test_feeder_pillar_defect_matching_quick_report_summary():
    """Verify FP defect rows reuse Quick Report summary rows matching ground truth exactly."""
    fp = LVDBSpec(name="FP TX1", label="FP", source="TX1")
    d1 = CbmDefectRecord(
        equipment="FP (J)",
        equipment_id="FP TX1 - INCOMING 1",
        defect_area="INCOMING LINK CONNECTION",
        additional_remarks="RED PHASE",
        technology="IR",
        ir_reading="64.9",
    )
    d2 = CbmDefectRecord(
        equipment="FP (J)",
        equipment_id="FP TX1 - OUTGOING F1",
        defect_area="INCOMING FUSE CONNECTION",
        additional_remarks="BLUE PHASE",
        technology="IR",
        ir_reading="64.9",
    )
    pkg = SubstationEquipmentPackage(lvdb_specs=(fp,))
    builder = ExecutiveSummaryCensusBuilder()

    rows = builder.build_census_rows(pkg, defects=(d1, d2))
    # 1 overview + 2 defects = 3 rows
    assert len(rows) == 3
    assert rows[0].equipment == "FP TX1"
    assert rows[0].defect_area == "OVERVIEW"
    assert rows[0].severity == "-"

    # Defect row 1 matches Quick Report Table 2 Row 1
    assert "INCOMING 1" in rows[1].equipment
    assert rows[1].defect_area == "INCOMING LINK CONNECTION/ RED PHASE"
    assert rows[1].ir_abs == "64.9 °C"
    assert rows[1].severity == "DEFECT"

    # Defect row 2 matches Quick Report Table 2 Row 2
    assert "OUTGOING F1" in rows[2].equipment
    assert rows[2].defect_area == "INCOMING FUSE CONNECTION/ BLUE PHASE"
    assert rows[2].ir_abs == "64.9 °C"
    assert rows[2].severity == "DEFECT"


def test_switchgear_overview_and_panel_equipment_formatting():
    """Verify Q4: SWG overview formatted as 'type, mfg' and panels as 'PANEL <feeder>\n<name>' without swg_base prefix."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="CRB00679", name="PE KG ASLI BATU BALONG"),
            SwitchgearPanelSpec(panel_no=2, panel_feeder_no="CRB00679", name="TX B"),
        ),
    )
    pkg = SubstationEquipmentPackage(switchgears=(swg,))
    builder = ExecutiveSummaryCensusBuilder()
    rows = builder.build_census_rows(pkg)

    # TAMCO: 2 overview rows + 2 panels * 2 comps = 6 rows
    assert len(rows) == 6
    assert rows[0].equipment == "RMU SF6, TAMCO"
    assert rows[1].equipment == "RMU SF6, TAMCO"
    assert rows[2].equipment == "PANEL CRB00679\nPE KG ASLI BATU BALONG"
    assert rows[3].equipment == "PANEL CRB00679\nPE KG ASLI BATU BALONG"
    assert rows[4].equipment == "PANEL CRB00679\nTX B"
    assert rows[5].equipment == "PANEL CRB00679\nTX B"


def test_column_1_dynamic_vertical_merging(tmp_path: Path):
    """Verify Q3 & Q4: Column 1 dynamic vertical merging for SWG overview, panels, and transformers."""
    swg = SwitchgearSpec(
        switchgear_type="RMU SF6",
        manufacturer="TAMCO",
        panels=(
            SwitchgearPanelSpec(panel_no=1, panel_feeder_no="CRB00679", name="PE KG ASLI BATU BALONG"),
        ),
    )
    tx = TransformerSpec(tx_id="Tx 1", rating_kva="1000", manufacturer="MTM")
    fp = LVDBSpec(name="FP TX1")
    pkg = SubstationEquipmentPackage(
        switchgears=(swg,),
        transformers=(tx,),
        lvdb_specs=(fp,),
    )
    builder = ExecutiveSummaryCensusBuilder()
    out_path = tmp_path / "test_merge.docx"
    result = builder.render(pkg, output_path=out_path)

    doc = docx.Document(str(out_path))
    table = doc.tables[0]

    # Row 1-2: SWG overview (2 rows)
    # Row 1: restart
    vm_swg_ov1 = _get_row_tc_vmerge(table.rows[1], 1)
    assert vm_swg_ov1 is not None and vm_swg_ov1.get(qn("w:val")) == "restart"
    assert table.rows[1].cells[1].text.strip() == "RMU SF6, TAMCO"
    # Row 2: continue, cleared text
    vm_swg_ov2 = _get_row_tc_vmerge(table.rows[2], 1)
    assert vm_swg_ov2 is not None and (vm_swg_ov2.get(qn("w:val")) is None or vm_swg_ov2.get(qn("w:val")) == "continue")
    assert _Cell(table.rows[2]._tr.tc_lst[1], table).text.strip() == ""

    # Row 3-4: Panel 1 compartments (2 rows: CABLE COMPARTMENT, CABLE ENTRY)
    # Row 3: restart
    vm_p1 = _get_row_tc_vmerge(table.rows[3], 1)
    assert vm_p1 is not None and vm_p1.get(qn("w:val")) == "restart"
    assert "PANEL CRB00679" in table.rows[3].cells[1].text
    # Row 4: continue, cleared text
    vm_p2 = _get_row_tc_vmerge(table.rows[4], 1)
    assert vm_p2 is not None and (vm_p2.get(qn("w:val")) is None or vm_p2.get(qn("w:val")) == "continue")
    assert _Cell(table.rows[4]._tr.tc_lst[1], table).text.strip() == ""

    # Rows 5-11: TX1 (7 rows)
    # Row 5: restart with TX1 title
    vm_tx_start = _get_row_tc_vmerge(table.rows[5], 1)
    assert vm_tx_start is not None and vm_tx_start.get(qn("w:val")) == "restart"
    assert table.rows[5].cells[1].text.strip() == "TX1 - MTM 1000kVA"
    # Rows 6-11: continue, cleared text
    for r_idx in range(6, 12):
        vm_tx_cont = _get_row_tc_vmerge(table.rows[r_idx], 1)
        assert vm_tx_cont is not None and (vm_tx_cont.get(qn("w:val")) is None or vm_tx_cont.get(qn("w:val")) == "continue")
        assert _Cell(table.rows[r_idx]._tr.tc_lst[1], table).text.strip() == ""

    # Row 12: FP TX1 (1 row, no Col 1 merge)
    vm_fp = _get_row_tc_vmerge(table.rows[12], 1)
    assert vm_fp is None or vm_fp.get(qn("w:val")) not in ("restart", "continue")
    assert table.rows[12].cells[1].text.strip() == "FP TX1"


def test_executive_summary_template_title_formatting(tmp_path: Path):
    """Verify Q2: Executive summary title in rendered docx carries template formatting ('EXECUTIVE SUMMARY')."""
    builder = ExecutiveSummaryCensusBuilder()
    out_path = tmp_path / "census_title_test.docx"
    builder.render(
        ExecutiveSummaryCensusContext(
            census_items=[CensusRowItem(no="1.", equipment="SWG", defect_area="OVERVIEW", severity="-", group_no=1)],
        ),
        output_path=out_path,
    )

    doc = docx.Document(str(out_path))
    p0 = doc.paragraphs[0]
    assert p0.text.strip() == "EXECUTIVE SUMMARY"
    # Check centered
    pPr = p0._p.find(qn("w:pPr"))
    assert pPr is not None
    jc = pPr.find(qn("w:jc"))
    assert jc is not None and jc.get(qn("w:val")) == "center"
    # Check run font
    assert len(p0.runs) >= 1
    rPr = p0.runs[0]._r.find(qn("w:rPr"))
    assert rPr is not None
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:asciiTheme")) == "minorHAnsi"
    color = rPr.find(qn("w:color"))
    assert color is not None and color.get(qn("w:val")) == "auto"


def test_census_builder_vcb_transition_panel_defect_matching():
    """Verify census builder maps BACK COMPARTMENT defect to REAR COMPARTMENT row on transition panel."""
    from src.full_report.census import _match_defects_for_swg_panel

    panel_trans = SwitchgearPanelSpec(
        panel_no=2,
        name="TRANSITION PANEL",
        panel_feeder_no="CKN02",
        panel_type="VCB",
    )
    swg = SwitchgearSpec(
        switchgear_type="VCB",
        panels=(
            SwitchgearPanelSpec(panel_no=1, name="FEEDER 1", panel_feeder_no="CKN01", panel_type="VCB"),
            panel_trans,
        ),
    )

    defect_back = CbmDefectRecord(
        equipment="VCB",
        equipment_id="CKN02",
        defect_area="BACK COMPARTMENT",
        additional_remarks="Partial discharge detected on rear barrier",
    )

    # REAR COMPARTMENT matches
    m_rear = _match_defects_for_swg_panel(swg, panel_trans, "REAR COMPARTMENT", [defect_back])
    assert len(m_rear) == 1
    assert m_rear[0].defect_area == "BACK COMPARTMENT"

    # FRONT COMPARTMENT, BUSBAR COMPARTMENT do not match
    m_front = _match_defects_for_swg_panel(swg, panel_trans, "FRONT COMPARTMENT", [defect_back])
    assert len(m_front) == 0
    m_busbar = _match_defects_for_swg_panel(swg, panel_trans, "BUSBAR COMPARTMENT", [defect_back])
    assert len(m_busbar) == 0



