"""Unit tests for PpuTestsheetParser (Ticket #092)."""

from __future__ import annotations

from pathlib import Path
import pytest

from src.testsheet.models import (
    SwitchgearPanelSpec,
    TestsheetData,
    VisualInspectionFinding,
)
from src.testsheet.parsers.ppu import PpuTestsheetParser
from tests.fixtures.ppu_workbook_factory import create_ppu_testsheet_workbook


def test_ppu_can_handle_sheets() -> None:
    parser = PpuTestsheetParser()
    assert parser.can_handle_sheets(["33kV (Main)", "Visual Inspec."]) is True
    assert parser.can_handle_sheets(["33kV (Extra)"]) is True
    assert parser.can_handle_sheets(["11kV (Main)"]) is True
    assert parser.can_handle_sheets(["11kV (Extra)"]) is True
    assert parser.can_handle_sheets(["11kV (Extra) (2)"]) is True
    assert parser.can_handle_sheets(["PCE Testsheet"]) is False
    assert parser.can_handle_sheets(["RAW DATA", "Sheet1"]) is False


def test_ppu_substation_number_resolution_hierarchy(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()

    # 1. Cell V2 on 33kV (Main) takes precedence
    wb1_path = tmp_path / "99. UNRELATED_FILENAME.xlsx"
    create_ppu_testsheet_workbook(
        path=wb1_path,
        v2_substation_num_33kv=5,
        v2_substation_num_11kv=10,
    )
    data1 = parser.extract_testsheet_data(wb1_path)
    assert data1.substation_number == 5

    # 2. Cell V2 on 11kV (Main) used when 33kV V2 is absent
    wb2_path = tmp_path / "99. UNRELATED_FILENAME.xlsx"
    create_ppu_testsheet_workbook(
        path=wb2_path,
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=12,
    )
    data2 = parser.extract_testsheet_data(wb2_path)
    assert data2.substation_number == 12

    # 3. Filename prefix used when both V2 cells are missing/blank
    wb3_path = tmp_path / "07. PPU APMM 33-11kV.xlsx"
    create_ppu_testsheet_workbook(
        path=wb3_path,
        v2_substation_num_33kv=None,
        v2_substation_num_11kv=None,
    )
    data3 = parser.extract_testsheet_data(wb3_path)
    assert data3.substation_number == 7

    # 4. ValueError raised when all resolution attempts fail
    wb4_path = tmp_path / "NO_NUMBER_PPU.xlsx"
    create_ppu_testsheet_workbook(
        path=wb4_path,
        v2_substation_num_33kv="",
        v2_substation_num_11kv=None,
    )
    with pytest.raises(ValueError, match="Unable to determine substation number for NO_NUMBER_PPU.xlsx"):
        parser.extract_testsheet_data(wb4_path)


def test_ppu_metadata_header_extraction_and_fallback(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()

    # Case A: 33kV present -> extracts header metadata from 33kV (Main)
    wb_33_path = tmp_path / "01. PPU DUAL.xlsx"
    create_ppu_testsheet_workbook(
        path=wb_33_path,
        substation_name="PPU GEBENG 33kV",
        fl_number="CGBG-PPU-GBNG",
        date_val="2026-10-15",
        time_val="11:45 AM",
        ambient_temp="31.2",
        humidity="60%",
        background_tev="4",
        gps_coordinate="3.987, 103.456",
    )
    data = parser.extract_testsheet_data(wb_33_path)
    assert data.substation_name_site == "PPU GEBENG 33kV"
    assert data.fl_site == "CGBG-PPU-GBNG"
    assert data.date_str == "15-10-2026"
    assert data.time == "11:45 AM"
    assert "31.2" in data.ambient
    assert data.humidity == "60%"
    assert data.tev_background == "4"
    assert data.gps_coordinate == "3.987, 103.456"

    # Case B: 11kV-only station (33kV absent) -> falls back to 11kV (Main)
    wb_11_path = tmp_path / "02. SSU 11KV ONLY.xlsx"
    create_ppu_testsheet_workbook(
        path=wb_11_path,
        has_33kv=False,
        has_11kv=True,
        substation_name="SSU INDUSTRIAL 11kV",
        fl_number="CGBG-SSU-IND",
        date_val="2026-10-16",
        ambient_temp="29.8",
        humidity="70%",
    )
    data_11 = parser.extract_testsheet_data(wb_11_path)
    assert data_11.substation_name_site == "SSU INDUSTRIAL 11kV"
    assert data_11.fl_site == "CGBG-SSU-IND"
    assert data_11.date_str == "16-10-2026"
    assert "29.8" in data_11.ambient
    assert data_11.humidity == "70%"


def test_ppu_dual_voltage_switchgear_assembly(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()

    # 1. Dual-voltage station: both 33kV and 11kV lineups present
    wb_dual = tmp_path / "01. PPU DUAL.xlsx"
    create_ppu_testsheet_workbook(
        path=wb_dual,
        has_33kv=True,
        has_11kv=True,
        panels_33kv=2,
        panels_11kv=3,
        extra_panels_33kv=([[{"feeder_no": "3P3", "name": "EXTRA 33kV"}]]),
    )
    data_dual = parser.extract_testsheet_data(wb_dual)
    assert len(data_dual.equipment.switchgears) == 2
    swg_33, swg_11 = data_dual.equipment.switchgears

    assert swg_33.switchgear_type == "GIS"
    # 2 main panels + 1 extra panel = 3 panels
    panel_names_33 = [p.name for p in swg_33.panels]
    assert "EXTRA 33kV" in panel_names_33
    assert len(panel_names_33) == 3

    assert swg_11.switchgear_type == "VCB"
    assert len(swg_11.panels) == 3

    # 2. 33kV only
    wb_33_only = tmp_path / "02. 33KV ONLY.xlsx"
    create_ppu_testsheet_workbook(path=wb_33_only, has_33kv=True, has_11kv=False, panels_33kv=2)
    data_33 = parser.extract_testsheet_data(wb_33_only)
    assert len(data_33.equipment.switchgears) == 1
    assert data_33.equipment.switchgears[0].switchgear_type == "GIS"

    # 3. 11kV only
    wb_11_only = tmp_path / "03. 11KV ONLY.xlsx"
    create_ppu_testsheet_workbook(path=wb_11_only, has_33kv=False, has_11kv=True, panels_11kv=2)
    data_11 = parser.extract_testsheet_data(wb_11_only)
    assert len(data_11.equipment.switchgears) == 1
    assert data_11.equipment.switchgears[0].switchgear_type == "VCB"

    # 4. Neither populated -> raises ValueError
    wb_empty = tmp_path / "04. EMPTY.xlsx"
    create_ppu_testsheet_workbook(path=wb_empty, has_33kv=True, has_11kv=True, panels_33kv=0, panels_11kv=0)
    with pytest.raises(ValueError, match="No populated 33kV or 11kV switchgear lineups found"):
        parser.extract_testsheet_data(wb_empty)


def test_ppu_panel_stride_and_compartment_mapping(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU STRIDE.xlsx"

    panel_custom = [
        {
            "feeder_no": "3P1",
            "name": "INCOMING 33kV",
            "status": "CLOSE",
            "load_amp": "250A",
            "cable_type": "XLPE 1C 500mm2",
            "heater_amp": "0.5A",
            "panel_type": "GIS",
            "breaker_serial": "CB-SN-3301",
            "breaker_photo": 11,
            "cable_photo": 12,
            "busbar_photo": 13,
            "pt_photo": 14,
            "cable_entry_photo": 15,
            "link_box_photo": 16,
            "secondary_photo": 17,
            "has_pt_measurement": True,
        }
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=panel_custom,
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert len(swg.panels) == 1
    panel = swg.panels[0]

    assert panel.panel_feeder_no == "3P1"
    assert panel.name == "INCOMING 33kV"
    assert panel.status == "CLOSE"
    assert panel.load_amp == "250"
    assert panel.cable_type == "XLPE 1C 500mm2"
    assert panel.heater_amp == "0.5A"
    assert panel.panel_type == "GIS"
    assert panel.serial_no == "CB-SN-3301"

    # Verify 6 physical compartments + secondary panel
    assert panel.breaker_photo == 11
    assert panel.cable_photo == 12
    assert panel.busbar_photo == 13
    assert panel.pt_photo == 14
    assert panel.cable_entry_photo == 15
    assert panel.link_box_photo == 16
    assert panel.secondary_photo == 17
    assert panel.has_pt_measurement is True

    # photo_numbers contains all distinct photo integers
    assert panel.photo_numbers == (11, 12, 13, 14, 15, 16, 17)


def test_ppu_overview_bays_and_photo_numbers(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU OVERVIEW.xlsx"

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=1,
        overview_rhs_photos_33kv=[201, 202, 203],
        overview_lhs_photos_33kv=[204, 205],
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]

    # Overview photos aggregated in SwitchgearSpec.photo_numbers
    assert 201 in swg.photo_numbers
    assert 202 in swg.photo_numbers
    assert 203 in swg.photo_numbers
    assert 204 in swg.photo_numbers
    assert 205 in swg.photo_numbers

    # Overview compartments stored on SwitchgearSpec without injecting fake panels
    assert len(swg.overview_compartments) > 0
    assert any("OVERVIEW FRONT" in c for c in swg.overview_compartments)
    panel_names = [p.name for p in swg.panels]
    assert "OVERVIEW RHS" not in panel_names
    assert panel_names[0] == "FEEDER 33kV 1"


def test_ppu_visual_inspection_findings(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU VI.xlsx"

    checklist_items = [
        {"category": "1. CIVIL", "item_no": "1.1", "description": "Door locked", "status": "YES", "remarks": ""},
        {"category": "1. CIVIL", "item_no": "1.2", "description": "Roof leakage", "status": "NO", "remarks": "Active leak above TX"},
        {"category": "2. SWITCHGEAR", "item_no": "2.1", "description": "Labeling missing", "status": "YES", "remarks": "Repasted label"},
        {"category": "2. SWITCHGEAR", "item_no": "2.2", "description": "Earthing intact", "status": "YES", "remarks": ""},
    ]
    other_findings = [
        "Oil spillage observed near cable trench",
        "Signs of rodent nesting behind battery rack",
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        visual_checklist_items=checklist_items,
        other_findings=other_findings,
    )

    data = parser.extract_testsheet_data(wb_path)
    findings = data.visual_findings

    # Checklist: 1.2 (marked NO) and 2.1 (non-empty remark) recorded. 1.1 and 2.2 omitted.
    item_nos = [f.item_no for f in findings]
    assert "1.2" in item_nos
    assert "2.1" in item_nos
    assert "1.1" not in item_nos
    assert "2.2" not in item_nos

    # Other findings recorded under Category "Other Findings"
    other_f = [f for f in findings if f.category == "Other Findings"]
    assert len(other_f) == 2
    assert any("Oil spillage" in f.remarks for f in other_f)
    assert any("rodent nesting" in f.remarks for f in other_f)


def test_ppu_raw_data_photo_ranges(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU RAW.xlsx"

    create_ppu_testsheet_workbook(
        path=wb_path,
        raw_photo_bounds={"ir_start": 101, "ir_end": 145, "dg_start": 2001, "dg_end": 2045},
    )

    ranges = parser.extract_photo_ranges(wb_path)
    assert ranges.ir.start_num == 101
    assert ranges.ir.end_num == 145
    assert ranges.dg.start_num == 2001
    assert ranges.dg.end_num == 2045


def test_ppu_fast_metadata_extraction(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "05. PPU META.xlsx"

    create_ppu_testsheet_workbook(
        path=wb_path,
        substation_number=5,
        substation_name="PPU META TEST",
        fl_number="FL-META-005",
        date_val="2026-10-18",
    )

    meta = parser.extract_testsheet_metadata(wb_path)
    assert isinstance(meta, TestsheetData)
    assert meta.substation_number == 5
    assert meta.substation_name_site == "PPU META TEST"
    assert meta.fl_site == "FL-META-005"
    assert meta.date_str == "18-10-2026"
    assert len(meta.equipment.switchgears) == 1
    assert meta.equipment.switchgears[0].switchgear_type == "GIS"


def test_ppu_factory_transformer_ratings(tmp_path: Path) -> None:
    wb_path = tmp_path / "06. PPU TX TEST.xlsx"
    wb = create_ppu_testsheet_workbook(
        path=wb_path,
        transformer_ratings=["15MVA", "30MVA"],
    )
    assert "TX" in wb.sheetnames
    ws_tx = wb["TX"]
    assert ws_tx["G8"].value == "15MVA"
    assert ws_tx["K8"].value == "30MVA"


def test_ppu_overview_only_without_panels_raises_value_error(tmp_path: Path) -> None:
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "07. PPU OVERVIEW ONLY.xlsx"
    create_ppu_testsheet_workbook(
        path=wb_path,
        panels_33kv=0,
        panels_11kv=0,
        overview_rhs_photos_33kv=[501, 502],
    )
    with pytest.raises(ValueError, match="No populated 33kV or 11kV switchgear lineups found"):
        parser.extract_testsheet_data(wb_path)


def test_ppu_cable_compartment_readings_extraction(tmp_path: Path) -> None:
    """Verify Row 16 (r+1) Cable Compartment TEV, PPC, and US readings are parsed into cable_* fields."""
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "08. PPU CABLE READINGS.xlsx"

    panel_custom = [
        {
            "feeder_no": "3P1",
            "name": "FEEDER 1 33kV",
            "status": "CLOSE",
            "tev_reading": "12",
            "tev_ppc": "5",
            "us_reading": "3",
            "cable_tev_reading": "24",
            "cable_tev_ppc": "18",
            "cable_us_reading": "7",
        }
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=panel_custom,
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    panel = swg.panels[0]

    # Breaker readings from Row 15
    assert panel.tev_reading == "12"
    assert panel.tev_ppc == "5"
    assert panel.us_reading == "3"

    # Cable Compartment readings from Row 16 (r+1)
    assert panel.cable_tev_reading == "24"
    assert panel.cable_tev_ppc == "18"
    assert panel.cable_us_reading == "7"


def test_ppu_breaker_readings_isolated_from_cable_sub_row(tmp_path: Path) -> None:
    """Verify that when Row r has no TEV/US readings but Row r+1 has Cable TEV/US readings,

    panel.tev_reading and panel.us_reading remain empty while cable_* fields are populated.
    """
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "09. PPU CABLE ISOLATION.xlsx"

    panel_custom = [
        {
            "feeder_no": "3P1",
            "name": "ISOLATION TEST 33kV",
            "status": "CLOSE",
            "tev_reading": "",
            "tev_ppc": "",
            "us_reading": "",
            "cable_tev_reading": "28.0",
            "cable_tev_ppc": "15",
            "cable_us_reading": "9.5",
        }
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=panel_custom,
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    panel = swg.panels[0]

    # Breaker readings must remain empty (no leakage from Row r+1)
    assert panel.tev_reading == ""
    assert panel.tev_ppc == ""
    assert panel.us_reading == ""

    # Cable readings are populated
    assert panel.cable_tev_reading == "28.0"
    assert panel.cable_tev_ppc == "15"
    assert panel.cable_us_reading == "9.5"


def test_ppu_canonical_panel_columns_and_compound_breaker_status(tmp_path: Path) -> None:
    """Verify canonical 33kV testsheet column mapping:
    - Col C: Feeder No
    - Col D: Feeder Name
    - Col E: Cable Type
    - Col F: CB Position (e.g. SERVICE)
    - Col G: Breaker Status (e.g. CLOSE) -> compound 'SERVICE/CLOSE'
    - Col H: Busbar Position (e.g. MAIN)
    - Col I: Load (A)
    - Col J: Heater (A)
    """
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "10. PPU CANONICAL COLUMNS.xlsx"

    panel_custom = [
        {
            "feeder_no": "32",
            "name": "TX 33/11kV 5MVA",
            "breaker_position": "SERVICE",
            "status": "CLOSE",
            "cable_type": "NOT ACCESSIBLE",
            "busbar_position": "MAIN",
            "load_amp": "92/600",
            "heater_amp": "0.63",
        },
        {
            "feeder_no": "2",
            "name": "SG. BANGKONG",
            "breaker_position": "SERVICE",
            "status": "TRIP",
            "cable_type": "NOT ACCESSIBLE",
            "busbar_position": "MAIN",
            "load_amp": "0/300",
            "heater_amp": "",
        },
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=panel_custom,
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert len(swg.panels) == 2

    p1, p2 = swg.panels
    assert p1.panel_feeder_no == "32"
    assert p1.name == "TX 33/11kV 5MVA"
    assert p1.cable_type == "NOT ACCESSIBLE"
    assert p1.status == "SERVICE/CLOSE"
    assert p1.busbar_position == "MAIN"
    assert p1.load_amp == "92"
    assert p1.heater_amp == "0.63"

    assert p2.panel_feeder_no == "2"
    assert p2.name == "SG. BANGKONG"
    assert p2.status == "SERVICE/TRIP"
    assert p2.busbar_position == "MAIN"
    assert p2.load_amp == "0"



def test_ppu_gps_coordinate_extraction(tmp_path: Path) -> None:
    import openpyxl
    parser = PpuTestsheetParser()

    # Case 1: Standard layout (L3="GPS:", M3=coords, D3=substation_name)
    wb1 = tmp_path / "01. PPU STANDARD GPS.xlsx"
    create_ppu_testsheet_workbook(
        path=wb1,
        substation_name="PPU STANDARD 33-11kV",
        gps_coordinate="4.123456, 103.654321",
    )
    data1 = parser.extract_testsheet_data(wb1)
    meta1 = parser.extract_testsheet_metadata(wb1)
    assert data1.gps_coordinate == "4.123456, 103.654321"
    assert meta1.gps_coordinate == "4.123456, 103.654321"

    # Case 2: Adjacent cell to "GPS:" label in different column
    wb2 = tmp_path / "02. PPU ADJACENT GPS.xlsx"
    create_ppu_testsheet_workbook(
        path=wb2,
        substation_name="PPU ADJACENT 33kV",
    )
    wb_obj = openpyxl.load_workbook(wb2)
    ws_vi = wb_obj["Visual Inspec."]
    ws_vi["L3"] = None
    ws_vi["M3"] = None
    ws_vi["K3"] = "GPS:"
    ws_vi["L3"] = "3.999111, 102.888222"
    wb_obj.save(wb2)
    wb_obj.close()
    data2 = parser.extract_testsheet_data(wb2)
    meta2 = parser.extract_testsheet_metadata(wb2)
    assert data2.gps_coordinate == "3.999111, 102.888222"
    assert meta2.gps_coordinate == "3.999111, 102.888222"

    # Case 3: Cell M3 directly without "GPS:" label
    wb3 = tmp_path / "03. PPU M3 DIRECT.xlsx"
    create_ppu_testsheet_workbook(
        path=wb3,
        substation_name="PPU M3 DIRECT",
    )
    wb_obj = openpyxl.load_workbook(wb3)
    ws_vi = wb_obj["Visual Inspec."]
    ws_vi["L3"] = None  # Remove "GPS:" label
    ws_vi["M3"] = "4.000123, 101.999456"
    wb_obj.save(wb3)
    wb_obj.close()
    data3 = parser.extract_testsheet_data(wb3)
    meta3 = parser.extract_testsheet_metadata(wb3)
    assert data3.gps_coordinate == "4.000123, 101.999456"
    assert meta3.gps_coordinate == "4.000123, 101.999456"

    # Case 4: Fallback to D3 only when D3 matches coordinate regex and does not match substation name
    wb4 = tmp_path / "04. PPU D3 FALLBACK.xlsx"
    create_ppu_testsheet_workbook(
        path=wb4,
        substation_name="PPU D3 FALLBACK",
    )
    wb_obj = openpyxl.load_workbook(wb4)
    ws_vi = wb_obj["Visual Inspec."]
    ws_vi["L3"] = None
    ws_vi["M3"] = None
    ws_vi["D3"] = "3.7501, 103.2205"  # coordinate in D3
    wb_obj.save(wb4)
    wb_obj.close()
    data4 = parser.extract_testsheet_data(wb4)
    meta4 = parser.extract_testsheet_metadata(wb4)
    assert data4.gps_coordinate == "3.7501, 103.2205"
    assert meta4.gps_coordinate == "3.7501, 103.2205"

    # Case 5: When D3 is substation name and M3 is empty, no coordinate extracted (no false positive from D3)
    wb5 = tmp_path / "05. PPU D3 SUBSTATION NAME.xlsx"
    create_ppu_testsheet_workbook(
        path=wb5,
        substation_name="PPU NAME ONLY 33kV",
    )
    wb_obj = openpyxl.load_workbook(wb5)
    ws_vi = wb_obj["Visual Inspec."]
    ws_vi["L3"] = None
    ws_vi["M3"] = None
    ws_vi["D3"] = "PPU NAME ONLY 33kV"
    wb_obj.save(wb5)
    wb_obj.close()
    data5 = parser.extract_testsheet_data(wb5)
    meta5 = parser.extract_testsheet_metadata(wb5)
    assert data5.gps_coordinate == ""
    assert meta5.gps_coordinate == ""


def test_ppu_extract_gps_production_kuala_tahan() -> None:
    prod_path = Path(
        r"C:\Users\ADAM\Documents\PO 42394550 - PAHANG - 33kV - THUNDER\TESTSHEET\JERANTUT\01. OCTOBER\22-10-2026\33. MINI PPU KUALA TAHAN 33-11kV (VI).xlsx"
    )
    if not prod_path.exists():
        pytest.skip(f"Production testsheet not found: {prod_path}")

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_data(prod_path)
    meta = parser.extract_testsheet_metadata(prod_path)
    assert data.gps_coordinate == "4.359481, 102.405296"
    assert meta.gps_coordinate == "4.359481, 102.405296"
    assert data.substation_type == "MPPU"
    assert meta.substation_type == "MPPU"


def test_ppu_resolve_substation_type_fallback_tokens() -> None:
    """Verify _resolve_substation_type resolves MINI PPU and MPU tokens to MPPU,

    and canonical tokens to PPU, SSU, PMU.
    """
    parser = PpuTestsheetParser()

    # 1. MINI PPU variants resolve to MPPU
    assert parser._resolve_substation_type("MINI PPU KUALA TAHAN 33/11kV", "dummy.xlsx", "") == "MPPU"
    assert parser._resolve_substation_type("MINI-PPU KUALA TAHAN", "dummy.xlsx", "") == "MPPU"
    assert parser._resolve_substation_type("", "33. MINI PPU KUALA TAHAN.xlsx", "") == "MPPU"

    # 2. \bMPU\b token in FL or filename resolves to MPPU
    assert parser._resolve_substation_type("", "", "CJRT/MPU/G00008") == "MPPU"
    assert parser._resolve_substation_type("", "", "CTRI/MPU/G00008") == "MPPU"
    assert parser._resolve_substation_type("", "08. MPU KUALA TAHAN.xlsx", "") == "MPPU"

    # 3. Explicit MPPU token resolves to MPPU
    assert parser._resolve_substation_type("MPPU BERA", "dummy.xlsx", "") == "MPPU"
    assert parser._resolve_substation_type("", "", "CJRT/MPPU/G00008") == "MPPU"

    # 4. Standard canonical tokens
    assert parser._resolve_substation_type("PPU APMM 33-11kV", "dummy.xlsx", "CGBG/PPU/G00001") == "PPU"
    assert parser._resolve_substation_type("SSU FLEXSYS", "dummy.xlsx", "CGBG/SSU/G00002") == "SSU"
    assert parser._resolve_substation_type("PMU JENGKA", "dummy.xlsx", "CJGK/PMU/G00003") == "PMU"


def test_ppu_pt_measurement_detection_excludes_column_k(tmp_path: Path) -> None:
    """Ticket #120 regression: Column K 'PT COMPARTMENT' label must not trigger has_pt_measurement.

    Only actual measurement cells in columns L, M, N, O, R, S, T, U, V should trigger
    has_pt_measurement = True.
    """
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU PT REGRESSION.xlsx"

    create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=[
            {
                "feeder_no": "3P1",
                "name": "BAY 1 - NO PT MEASUREMENT",
                "has_pt_measurement": False,
            },
            {
                "feeder_no": "3P2",
                "name": "BAY 2 - PT MEASUREMENT IN L",
                "has_pt_measurement": True,
            },
            {
                "feeder_no": "3P3",
                "name": "BAY 3 - PT MEASUREMENT IN O",
                "sub_rows": {
                    3: {"O": "1.2"},
                },
            },
            {
                "feeder_no": "3P4",
                "name": "BAY 4 - PT TEV READING IN R",
                "sub_rows": {
                    3: {"R": "8.0"},
                },
            },
        ],
    )

    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert len(swg.panels) == 4
    p1, p2, p3, p4 = swg.panels

    # Panel 1: has pre-printed "PT COMPARTMENT" in Column K, but no readings in L..V
    assert p1.has_pt_measurement is False

    # Panel 2: has measurement in Column L ("24.5")
    assert p2.has_pt_measurement is True

    # Panel 3: has measurement in Column O ("1.2")
    assert p3.has_pt_measurement is True

    # Panel 4: has TEV reading in Column R ("8.0")
    assert p4.has_pt_measurement is True


def test_ppu_non_rmu_does_not_extract_l9_into_serial_no(tmp_path: Path) -> None:
    """Ticket #135: PPU parser must not read L9 into serial_no for non-RMU switchgear."""
    wb_path = tmp_path / "01. SSU KIIP GIS.xlsx"
    wb = create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=[
            {
                "feeder_no": "3P1",
                "name": "INCOMING 33kV",
                "breaker_serial": "SN: GIS-BRK-001",
            },
            {
                "feeder_no": "3P2",
                "name": "OUTGOING 33kV",
                "breaker_serial": "",
            },
        ],
        swg_33kv_details={
            "manufacturer": "TAMCO",
            "switchgear_type": "GIS",
            "model": "GV3",
            "serial_no": "LHS:9.5.23, RHS:9.5.23",
        },
    )
    ws = wb["33kV (Main)"]
    ws["L9"] = "LHS:9.5.23, RHS:9.5.23"
    ws["V15"] = "SN: GIS-BRK-001"
    ws["V21"] = None
    wb.save(wb_path)

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert swg.switchgear_type == "GIS"
    # Board serial_no is empty for GIS (L9 ignored, V15 not used for board serial)
    assert swg.serial_no == ""
    # Panel 1 has its own breaker serial
    assert swg.panels[0].serial_no == "GIS-BRK-001"
    # Panel 2 remains empty (non-RMU panels do NOT inherit sheet V15)
    assert swg.panels[1].serial_no == ""

    meta = parser.extract_testsheet_metadata(wb_path)
    assert meta.equipment.switchgears[0].serial_no == ""


def test_ppu_rmu_extracts_v15_serial_and_propagates_to_blank_panels(tmp_path: Path) -> None:
    """Ticket #135: RMU switchgear extracts serial from V15 and propagates to blank panels."""
    wb_path = tmp_path / "22. MINI PPU KECHAU TUI.xlsx"
    wb = create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=[
            {
                "feeder_no": "3P1",
                "name": "INCOMING 1",
                "breaker_serial": "SN: 2023-RMU-001",
            },
            {
                "feeder_no": "3P2",
                "name": "TRANSFORMER 1",
                "breaker_serial": "",
            },
            {
                "feeder_no": "3P3",
                "name": "OUTGOING 1",
                "breaker_serial": "SN: EXPLICIT-999",
            },
        ],
        swg_33kv_details={
            "manufacturer": "TAMCO",
            "switchgear_type": "RMU SF6",
            "model": "VRN2A",
            "serial_no": "20.10.2023",
        },
    )
    ws = wb["33kV (Main)"]
    ws["L9"] = "20.10.2023"
    ws["V15"] = "SN: 2023-RMU-001"
    ws["V21"] = None
    ws["V27"] = "SN: EXPLICIT-999"
    wb.save(wb_path)

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert swg.switchgear_type == "RMU SF6"
    assert swg.serial_no == "2023-RMU-001"
    assert len(swg.panels) == 3
    assert swg.panels[0].serial_no == "2023-RMU-001"
    assert swg.panels[1].serial_no == "2023-RMU-001"
    assert swg.panels[2].serial_no == "EXPLICIT-999"

    meta = parser.extract_testsheet_metadata(wb_path)
    assert meta.equipment.switchgears[0].serial_no == "2023-RMU-001"


def test_ppu_rmu_bare_sn_prefix_does_not_propagate_to_panels(tmp_path: Path) -> None:
    """Ticket #135 Review: Bare 'SN:' or 'SN:   ' in V15 normalizes to '' and does not propagate to panels."""
    wb_path = tmp_path / "23. MINI PPU BARE SN.xlsx"
    wb = create_ppu_testsheet_workbook(
        path=wb_path,
        has_33kv=True,
        has_11kv=False,
        panels_33kv=[
            {
                "feeder_no": "3P1",
                "name": "INCOMING 1",
                "breaker_serial": "",
            },
            {
                "feeder_no": "3P2",
                "name": "TRANSFORMER 1",
                "breaker_serial": "",
            },
        ],
        swg_33kv_details={
            "manufacturer": "TAMCO",
            "switchgear_type": "RMU SF6",
            "model": "VRN2A",
            "serial_no": "20.10.2023",
        },
    )
    ws = wb["33kV (Main)"]
    ws["L9"] = "20.10.2023"
    ws["V15"] = "SN:   "
    ws["V21"] = "SN:"
    wb.save(wb_path)

    parser = PpuTestsheetParser()
    data = parser.extract_testsheet_data(wb_path)
    swg = data.equipment.switchgears[0]
    assert swg.switchgear_type == "RMU SF6"
    assert swg.serial_no == ""
    assert len(swg.panels) == 2
    assert swg.panels[0].serial_no == ""
    assert swg.panels[1].serial_no == ""

    meta = parser.extract_testsheet_metadata(wb_path)
    assert meta.equipment.switchgears[0].serial_no == ""


def test_ppu_multiline_whitespace_and_control_chars_in_panels_and_headers(tmp_path: Path) -> None:
    """Verify that multiline and control character strings in substation headers,
    panel feeder numbers, and feeder names are collapsed to single-line strings.
    """
    parser = PpuTestsheetParser()
    wb_path = tmp_path / "01. PPU MULTILINE.xlsx"

    panels_33kv = [
        {
            "feeder_no": "3P1\n(INC)",
            "name": "33kV TRANSITION\n(BUS - SECTION)",
            "status": "CLOSE",
            "breaker_serial": "SN-001\r\n",
        },
        {
            "feeder_no": "3P2\r\n(OUT)",
            "name": "FEEDER 2\r\nLINE",
            "status": "CLOSE",
            "breaker_serial": "SN-002",
        },
        {
            "feeder_no": "3P3\t",
            "name": "AUX TX\r\n\t 33kV",
            "status": "CLOSE",
            "breaker_serial": "SN-003",
        },
    ]

    create_ppu_testsheet_workbook(
        path=wb_path,
        substation_name="33kV PPU\nTRANSITION SUBSTATION",
        fl_number="CGBG-PPU-TRNS.0\r\n",
        panels_33kv=panels_33kv,
        panels_11kv=(),
        has_11kv=False,
    )

    data = parser.extract_testsheet_data(wb_path)

    # Substation header metadata should be clean single-line strings
    assert data.substation_name_site == "33kV PPU TRANSITION SUBSTATION"
    assert data.substation_name_erms == "33kV PPU TRANSITION SUBSTATION"
    assert data.fl_site == "CGBG-PPU-TRNS.0"
    assert data.fl_erms == "CGBG-PPU-TRNS"

    # Switchgear panels should have clean single-line feeder numbers and names
    swg = data.equipment.switchgears[0]
    assert len(swg.panels) == 3

    p1 = swg.panels[0]
    assert p1.panel_feeder_no == "3P1 (INC)"
    assert p1.name == "33kV TRANSITION (BUS - SECTION)"

    p2 = swg.panels[1]
    assert p2.panel_feeder_no == "3P2 (OUT)"
    assert p2.name == "FEEDER 2 LINE"

    p3 = swg.panels[2]
    assert p3.panel_feeder_no == "3P3"
    assert p3.name == "AUX TX 33kV"

