"""Smoke and unit tests for Full Report Executive Summary Census Jinja2 prototype template."""

from pathlib import Path
from typing import Any
import docx
from docx.oxml.ns import qn
from docx.table import _Cell
import docxtpl
import pytest


TEMPLATE_PATH = Path("templates/FULL REPORT/executive_summary_census.docx")
EXPECTED_COL_WIDTHS_DXA = [630, 2742, 2719, 1019, 720, 720, 1980]


def _get_cell_shading(cell: _Cell):
    """Safely retrieve shading XML element from a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    return tcPr.find(qn("w:shd"))


@pytest.fixture
def census_doc() -> docx.Document:
    """Load the prototype census docx document."""
    return docx.Document(str(TEMPLATE_PATH))


def test_census_template_file_exists():
    """Verify that the prototype census docx template artifact exists on disk."""
    assert TEMPLATE_PATH.exists(), f"Template not found at {TEMPLATE_PATH}"
    assert TEMPLATE_PATH.is_file()
    assert TEMPLATE_PATH.stat().st_size > 0


def test_census_template_heading_and_styling(census_doc: docx.Document):
    """Verify acceptance criteria: Template contains title heading '2.0 EXECUTIVE SUMMARY (EQUIPMENT CENSUS)'."""
    assert len(census_doc.paragraphs) >= 1

    heading_paragraph = census_doc.paragraphs[0]
    heading_text = heading_paragraph.text.strip()
    assert heading_text == "2.0 EXECUTIVE SUMMARY (EQUIPMENT CENSUS)"
    assert heading_paragraph.style.name == "Subtitle SVT"
    assert heading_paragraph.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    # Heading runs must be bold
    assert any(run.font.bold for run in heading_paragraph.runs)


def test_census_template_table_structure_and_widths(census_doc: docx.Document):
    """Verify acceptance criteria: Table structure defines exactly 7 columns with correct standard widths and header styling."""
    assert len(census_doc.tables) == 1, "Template must contain exactly one census table"

    table = census_doc.tables[0]
    assert len(table.columns) == 7

    # Validate header row text
    header_cells = table.rows[0].cells
    assert header_cells[0].text.strip() == "NO."
    assert header_cells[1].text.strip() == "EQUIPMENT"
    assert header_cells[2].text.strip() == "DEFECT AREA"
    assert "IR" in header_cells[3].text
    assert "U/S" in header_cells[4].text
    assert "TEV" in header_cells[5].text
    assert header_cells[6].text.strip() == "SEVERITY"

    # Validate header row styling: all header cells must have shading
    for cell_index, cell in enumerate(header_cells):
        shading = _get_cell_shading(cell)
        assert shading is not None, f"Header cell {cell_index} must have shading XML element"
        assert shading.get(qn("w:val")) == "pct15" or shading.get(qn("w:fill")) is not None

    # Validate column widths in tblGrid
    tblGrid = table._tbl.tblGrid
    assert tblGrid is not None
    grid_cols = tblGrid.findall(qn("w:gridCol"))
    assert len(grid_cols) == 7
    for col_index, (col_element, expected_width) in enumerate(zip(grid_cols, EXPECTED_COL_WIDTHS_DXA)):
        col_width = int(col_element.get(qn("w:w")))
        assert col_width == expected_width, f"Column {col_index} width dxa mismatch: {col_width} != {expected_width}"


def test_census_template_jinja_loop_and_field_bindings(census_doc: docx.Document):
    """Verify acceptance criteria: Jinja loop {% tr for item in census_items %} binds all 7 fields."""
    table = census_doc.tables[0]
    assert len(table.rows) >= 4, "Template table must have header, loop start, data, and loop end rows"

    # Row 1: Jinja loop start
    row1_text = table.rows[1].cells[0].text.strip()
    assert "item in census_items" in row1_text
    assert "{%tr" in row1_text or "{% tr" in row1_text

    # Row 2: 7 field placeholders
    expected_placeholders = [
        "{{ item.no }}",
        "{{ item.equipment }}",
        "{{ item.defect_area }}",
        "{{ item.ir_abs }}",
        "{{ item.us_dB }}",
        "{{ item.tev_dB }}",
        "{{ item.severity }}",
    ]
    for cell_index, (cell, expected_placeholder) in enumerate(zip(table.rows[2].cells, expected_placeholders)):
        cell_text = cell.text.strip()
        assert cell_text == expected_placeholder, f"Cell 2,{cell_index} expected {expected_placeholder}, got {cell_text}"

    # Row 2 Cell 6 must not have static red shading hardcoded
    severity_shading = _get_cell_shading(table.rows[2].cells[6])
    assert severity_shading is None, "Data row severity cell must not have static XML shading"

    # Row 3: Jinja loop end
    row3_text = table.rows[3].cells[0].text.strip()
    assert "endfor" in row3_text


def test_render_census_mock_items_cleanly(tmp_path: Path):
    """Verify acceptance criteria: Verifiable by rendering mock items through docxtpl."""
    template = docxtpl.DocxTemplate(str(TEMPLATE_PATH))

    mock_items = [
        {
            "no": "1.",
            "equipment": "RMU 01 - TAMCO",
            "defect_area": "OVERVIEW",
            "ir_abs": "-",
            "us_dB": "-",
            "tev_dB": "-",
            "severity": "NORMAL",
        },
        {
            "no": "1.",
            "equipment": "RMU 01 - TAMCO",
            "defect_area": "CABLE COMPARTMENT",
            "ir_abs": "48.5 °C",
            "us_dB": "10dB",
            "tev_dB": "15dB",
            "severity": "NORMAL",
        },
        {
            "no": "2.",
            "equipment": "TX 1 - 1000kVA",
            "defect_area": "HV CABLE SPLIT",
            "ir_abs": "51.0 °C",
            "us_dB": "-",
            "tev_dB": "-",
            "severity": "NORMAL",
        },
        {
            "no": "3.",
            "equipment": "FEEDER PILLAR 10-WAY",
            "defect_area": "OVERVIEW",
            "ir_abs": "-",
            "us_dB": "-",
            "tev_dB": "-",
            "severity": "NORMAL",
        },
        {
            "no": "3.",
            "equipment": "FEEDER PILLAR 10-WAY",
            "defect_area": "WAY 2 - FUSE CONTACT",
            "ir_abs": "78.4 °C",
            "us_dB": "-",
            "tev_dB": "-",
            "severity": "DEFECT",
        },
    ]

    context: dict[str, Any] = {"census_items": mock_items}
    template.render(context)

    output_path = tmp_path / "rendered_census.docx"
    template.save(str(output_path))
    assert output_path.exists()

    # Load and inspect rendered document
    rendered_doc = docx.Document(str(output_path))
    assert len(rendered_doc.tables) == 1
    rendered_table = rendered_doc.tables[0]

    # 1 header row + 5 data rows
    assert len(rendered_table.rows) == 6

    # Verify first data row
    assert [cell.text.strip() for cell in rendered_table.rows[1].cells] == [
        "1.",
        "RMU 01 - TAMCO",
        "OVERVIEW",
        "-",
        "-",
        "-",
        "NORMAL",
    ]

    # Verify fifth data row (Way 2 defect)
    assert [cell.text.strip() for cell in rendered_table.rows[5].cells] == [
        "3.",
        "FEEDER PILLAR 10-WAY",
        "WAY 2 - FUSE CONTACT",
        "78.4 °C",
        "-",
        "-",
        "DEFECT",
    ]

    # Verify no raw Jinja syntax remains anywhere in table
    for row_index, row in enumerate(rendered_table.rows):
        for col_index, cell in enumerate(row.cells):
            assert "{%" not in cell.text, f"Row {row_index} Col {col_index} contains unrendered Jinja tag: {cell.text}"
            assert "{{" not in cell.text, f"Row {row_index} Col {col_index} contains unrendered Jinja var: {cell.text}"


def test_render_empty_census_items(tmp_path: Path):
    """Verify that an empty list of census items cleanly renders only the header row."""
    template = docxtpl.DocxTemplate(str(TEMPLATE_PATH))
    template.render({"census_items": []})

    output_path = tmp_path / "empty_census.docx"
    template.save(str(output_path))

    rendered_doc = docx.Document(str(output_path))
    rendered_table = rendered_doc.tables[0]
    assert len(rendered_table.rows) == 1
    assert rendered_table.rows[0].cells[0].text.strip() == "NO."
    assert rendered_table.rows[0].cells[6].text.strip() == "SEVERITY"
