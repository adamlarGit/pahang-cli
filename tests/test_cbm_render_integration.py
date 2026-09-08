"""Integration tests for CBM defect rendering context and template compilation.

Verifies:
- Injection of analysis & recommendation into context dictionaries
- Overview cell post-processing and shading
- Clean Jinja2 substitution in docx templates without raw placeholder leaks
"""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from src.quick_report.cbm_render import (
    _build_fp_lvdb_render_context,
    _build_swg_render_context,
    _build_tx_render_context,
    _post_process_overview_cell,
    _render_docx_template,
)
from src.quick_report.defects import CbmDefectRecord


def test_fp_lvdb_render_context_includes_analysis_and_recommendation() -> None:
    record = CbmDefectRecord(
        equipment="FP (D)",
        technology="IR",
        defect_area="FUSE COMPARTMENT",
        additional_remarks="RED PHASE",
    )
    ctx = _build_fp_lvdb_render_context(record, overview=False)
    assert "analysis" in ctx
    assert "recommendation" in ctx
    assert "fp" in ctx
    assert "analysis" in ctx["fp"]
    assert "recommendation" in ctx["fp"]
    assert "FUSE COMPARTMENT" in ctx["analysis"]


def test_swg_render_context_includes_analysis_and_recommendation() -> None:
    record = CbmDefectRecord(
        equipment="RMU SF6",
        technology="IR",
        defect_area="CABLE TERMINATION/CABLE XLPE",
        additional_remarks="YELLOW PHASE",
    )
    ctx = _build_swg_render_context(record, overview=False)
    assert "analysis" in ctx
    assert "recommendation" in ctx
    assert "panel" in ctx
    assert "analysis" in ctx["panel"]
    assert "recommendation" in ctx["panel"]
    assert "CABLE TERMINATION/CABLE XLPE" in ctx["analysis"]


def test_tx_render_context_includes_analysis_and_recommendation() -> None:
    record = CbmDefectRecord(
        equipment="LTX/DTX",
        technology="IR",
        defect_area="LV BUSHING",
        additional_remarks="RED PHASE",
    )
    ctx = _build_tx_render_context(record, overview=False)
    assert "analysis" in ctx
    assert "recommendation" in ctx
    assert "tx" in ctx
    assert "analysis" in ctx["tx"]
    assert "recommendation" in ctx["tx"]
    assert "RED PHASE BUSHING CONNECTION" in ctx["analysis"]
    assert "internal rod" in ctx["recommendation"]


def test_overview_cell_shading(tmp_path: Path) -> None:
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    defective_cell = table.cell(0, 0)
    defective_cell.text = "Analysis: Please refer to the following page for details defect."
    normal_cell = table.cell(1, 0)
    normal_cell.text = "Analysis: No Anomaly."

    _post_process_overview_cell(defective_cell)
    _post_process_overview_cell(normal_cell)

    assert "EE0000" in defective_cell._tc.xml
    assert "00B050" in normal_cell._tc.xml


def test_template_compilation_clean_substitution(tmp_path: Path) -> None:
    template_path = Path("templates/QUICK REPORT/DEFECT IR/fp-individual-defect.docx")
    if not template_path.exists():
        pytest.skip("Template file not found at expected path")

    record = CbmDefectRecord(
        equipment="FP (D)",
        technology="IR",
        defect_area="FUSE COMPARTMENT",
        additional_remarks="RED PHASE",
        ir_reading="55.2",
    )
    ctx = _build_fp_lvdb_render_context(record, overview=False)
    ctx["substation"] = {
        "name_erms": "PE TEST",
        "date": "08-09-2026",
        "time": "10:00 AM",
        "ambient": "30.0",
        "humidity": "65%",
    }
    ctx["visual"] = {"image": "-"}

    output_file = tmp_path / "compiled_fp_defect.docx"
    _render_docx_template(template_path, output_file, ctx)

    assert output_file.exists()
    out_doc = docx.Document(output_file)
    full_text = "\n".join(
        [p.text for p in out_doc.paragraphs]
        + [c.text for t in out_doc.tables for r in t.rows for c in r.cells]
    )

    # Assert clean substitution
    assert "{{ analysis }}" not in full_text
    assert "{{ recommendation }}" not in full_text
    assert "Thermal image above indicates hot spot detected at FUSE COMPARTMENT." in full_text
    assert "To inspect, clean the contact surface and re-tighten the connection point." in full_text
